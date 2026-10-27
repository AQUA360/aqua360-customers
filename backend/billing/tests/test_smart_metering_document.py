import shutil
import tempfile
from datetime import date
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

import openpyxl
from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings
from rest_framework.test import APIRequestFactory, force_authenticate

from billing.models import Reading, ReadingBatch, ReadingDocument
from billing.utils.smart_metering_service import (
    build_smart_metering_preview,
    import_smart_metering_readings,
    smart_metering_document_is_replaceable,
    smart_metering_scope_options,
)
from billing.views.reading_document_view import ReadingDocumentViewSet
from contract.models import Contract, ContractStatus
from coredata.models import ConfigProject
from service.models import (
    Company, Meter, MeterStatus, Property, Route, RoutePosition, RouteZone, SupplyPoint, SupplyPointStatus,
)

MEDIA_ROOT = tempfile.mkdtemp()
READING_DATE = date(2026, 10, 1)


@override_settings(MEDIA_ROOT=MEDIA_ROOT)
class TestSmartMeteringDocument(TestCase):
    """
    The Smart Metering load lives in "Introduir lectures": it creates a reading
    document (with an Excel file) that is not tied to any reading batch.
    """

    @classmethod
    def setUpTestData(cls):
        for token, value in (
            ('contract_active_token', 'ACTIVE'),
            ('contract_terminated_status', 'TERMINATED'),
            ('contract_termination_completed_token', 'COMPLETED'),
            ('token_meter_status_no_meter', 'NO_METER'),
        ):
            ConfigProject.objects.get_or_create(token=token, defaults={'value': value})
        cls.contract_status, _ = ContractStatus.objects.get_or_create(token='ACTIVE', defaults={'name': 'Alta'})
        cls.sp_status, _ = SupplyPointStatus.objects.get_or_create(token='SP_ACTIVE', defaults={'name': 'Actiu'})
        cls.meter_status, _ = MeterStatus.objects.get_or_create(token='1', defaults={'name': 'Actiu'})
        cls.company = Company.objects.create(name='Company smart metering')
        cls.user = get_user_model().objects.create_superuser(username='smart-metering', password='x', email='')

    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(MEDIA_ROOT, ignore_errors=True)

    def setUp(self):
        self.meter = Meter.objects.create(
            code='SM-0001', status=self.meter_status,
            has_remote_reading=True, remote_reading_type='SMART_METERING',
        )
        self.sp = SupplyPoint.objects.create(token='SM-SP-1', meter=self.meter, status=self.sp_status)
        self.contract = Contract.objects.create(
            company=self.company, token='SM-C-1', status=self.contract_status,
            is_active=True, supply_point_default=self.sp,
        )
        self.contract.supply_points.add(self.sp)
        self.previous = Reading.objects.create(
            meter=self.meter, contract=self.contract, supply_point=self.sp,
            reading_date=date(2026, 7, 1), reading_value=100, calculated_value=0,
        )
        # A manual meter must not be part of the Smart Metering load.
        manual_meter = Meter.objects.create(code='MAN-0001', status=self.meter_status)
        manual_sp = SupplyPoint.objects.create(token='MAN-SP-1', meter=manual_meter, status=self.sp_status)
        manual_contract = Contract.objects.create(
            company=self.company, token='MAN-C-1', status=self.contract_status,
            is_active=True, supply_point_default=manual_sp,
        )
        manual_contract.supply_points.add(manual_sp)

    def _api_payload(self, reading_value=130):
        return [{'meter': 'SM-0001', 'reading': reading_value, 'timestamp': READING_DATE.isoformat()}]

    def _new_document(self):
        return ReadingDocument.objects.create(
            token='SM-DOC',
            source=ReadingDocument.SOURCE_SMART_METERING,
            reading_date=READING_DATE,
            status=ReadingDocument.STATUS_PROCESSING,
        )

    def _import(self, document, reading_value=130):
        with patch(
            'billing.utils.smart_metering_service.SmartMeteringClient.fetch_readings',
            return_value=self._api_payload(reading_value),
        ):
            return import_smart_metering_readings(document)

    def test_preview_covers_smart_metering_meters_without_batch(self):
        with patch(
            'billing.utils.smart_metering_service.SmartMeteringClient.fetch_readings',
            return_value=self._api_payload(),
        ):
            preview = build_smart_metering_preview(READING_DATE.isoformat())

        self.assertEqual([row['meter_code'] for row in preview['readings']], ['SM-0001'])
        self.assertTrue(preview['readings'][0]['matched'])
        self.assertEqual(preview['readings'][0]['last_previous_reading_value'], 100.0)

    def test_import_creates_document_readings_and_excel(self):
        document = self._new_document()

        result = self._import(document)

        document.refresh_from_db()
        self.assertEqual(result['saved_readings'], 1)
        self.assertEqual(document.status, ReadingDocument.STATUS_PROCESSED)
        reading = document.readings.get(is_control=False)
        self.assertIsNone(reading.batch_id)
        self.assertEqual(reading.reading_value, Decimal('130'))
        self.assertEqual(reading.previous_reading_id, self.previous.id)
        self.assertEqual(reading.calculated_value, Decimal('30'))

        self.assertIn('smart_metering_20261001', document.file.name)
        self.assertTrue(document.file.name.endswith('.xlsx'))
        with document.file.open('rb') as fh:
            wb = openpyxl.load_workbook(BytesIO(fh.read()))
        rows = list(wb['readings'].iter_rows(min_row=2, values_only=True))
        self.assertEqual(rows[0][0], 'SM-0001')
        self.assertEqual(rows[0][1], 130.0)

    def test_reload_same_document_does_not_duplicate_readings(self):
        document = self._new_document()
        self._import(document)

        self._import(document)

        self.assertEqual(document.readings.filter(is_control=False).count(), 1)
        self.assertEqual(
            Reading.objects.filter(meter=self.meter, reading_date=READING_DATE, is_control=False).count(), 1
        )

    def test_reading_used_by_a_batch_is_not_modified(self):
        batch = ReadingBatch.objects.create(name='Lot')
        batch_reading = Reading.objects.create(
            meter=self.meter, contract=self.contract, supply_point=self.sp, batch=batch,
            previous_reading=self.previous, reading_date=READING_DATE, reading_value=120, calculated_value=20,
        )
        document = self._new_document()

        result = self._import(document, reading_value=130)

        batch_reading.refresh_from_db()
        self.assertEqual(result['saved_readings'], 0)
        self.assertEqual(result['not_saved_readings'], 1)
        self.assertFalse(batch_reading.is_control)
        self.assertEqual(batch_reading.batch_id, batch.id)
        self.assertIsNone(batch_reading.document_id)
        self.assertFalse(document.readings.exists())

    def test_document_is_replaceable_until_used_by_a_batch(self):
        document = self._new_document()
        self._import(document)
        self.assertTrue(smart_metering_document_is_replaceable(document))

        document.batch = ReadingBatch.objects.create(name='Lot')
        document.save(update_fields=['batch'])

        self.assertFalse(smart_metering_document_is_replaceable(document))

    def _post_import(self, replace=False, scope=None):
        view = ReadingDocumentViewSet.as_view({'post': 'smart_metering_import'})
        request = APIRequestFactory().post(
            '/billing/reading-document/smart-metering/import/',
            {'reading_date': READING_DATE.isoformat(), 'preview': None, 'replace': replace, 'scope': scope},
            format='json',
        )
        force_authenticate(request, user=self.user)
        with patch('billing.views.reading_document_view.import_smart_metering_document_task.delay') as delay:
            delay.return_value.id = 'task-id'
            return view(request)

    def test_import_endpoint_blocks_repeated_date_unless_replace(self):
        first = self._post_import()
        self.assertEqual(first.status_code, 202, first.data)
        document = ReadingDocument.objects.get(id=first.data['document_id'])
        document.status = ReadingDocument.STATUS_PROCESSED
        document.save(update_fields=['status'])

        repeated = self._post_import()
        self.assertEqual(repeated.status_code, 409)
        self.assertTrue(repeated.data['existing_document']['replaceable'])

        replaced = self._post_import(replace=True)
        self.assertEqual(replaced.status_code, 202, replaced.data)
        self.assertEqual(replaced.data['document_id'], document.id)
        self.assertEqual(
            ReadingDocument.objects.filter(source=ReadingDocument.SOURCE_SMART_METERING).count(), 1
        )

    def test_import_endpoint_blocks_replace_when_document_is_in_a_batch(self):
        first = self._post_import()
        document = ReadingDocument.objects.get(id=first.data['document_id'])
        document.status = ReadingDocument.STATUS_PROCESSED
        document.batch = ReadingBatch.objects.create(name='Lot')
        document.save(update_fields=['status', 'batch'])

        response = self._post_import(replace=True)

        self.assertEqual(response.status_code, 409)
        self.assertFalse(response.data['existing_document']['replaceable'])


    def _smart_metering_sp_on_route(self, code, route):
        meter = Meter.objects.create(
            code=code, status=self.meter_status,
            has_remote_reading=True, remote_reading_type='SMART_METERING',
        )
        position = RoutePosition.objects.create(route=route)
        prop = Property.objects.create(route_position=position)
        sp = SupplyPoint.objects.create(token=f'SP-{code}', meter=meter, status=self.sp_status, property=prop)
        contract = Contract.objects.create(
            company=self.company, token=f'C-{code}', status=self.contract_status,
            is_active=True, supply_point_default=sp,
        )
        contract.supply_points.add(sp)
        return sp

    def _preview_codes(self, scope):
        with patch(
            'billing.utils.smart_metering_service.SmartMeteringClient.fetch_readings',
            return_value=[],
        ):
            preview = build_smart_metering_preview(READING_DATE.isoformat(), scope=scope)
        return sorted(row['meter_code'] for row in preview['readings'])

    def test_preview_only_covers_selected_routes_zones_or_batches(self):
        zone = RouteZone.objects.create(name='Zona')
        route_a = Route.objects.create(token='RA', route_zone=zone)
        route_b = Route.objects.create(token='RB')
        self._smart_metering_sp_on_route('SM-A', route_a)
        self._smart_metering_sp_on_route('SM-B', route_b)
        batch = ReadingBatch.objects.create(name='Lot', include_telecontrol=True)
        batch.routes.add(route_b)
        batch.fix_meters.add(self.meter)
        manuals_batch = ReadingBatch.objects.create(name='Manuals', include_manual=True)
        manuals_batch.routes.add(route_a)

        self.assertEqual(self._preview_codes(None), ['SM-0001', 'SM-A', 'SM-B'])
        self.assertEqual(self._preview_codes({'route_ids': [route_a.id]}), ['SM-A'])
        self.assertEqual(self._preview_codes({'route_zone_ids': [zone.id]}), ['SM-A'])
        self.assertEqual(self._preview_codes({'batch_ids': [batch.id]}), ['SM-0001', 'SM-B'])
        self.assertEqual(self._preview_codes({'batch_ids': [manuals_batch.id]}), [])
        self.assertEqual(
            self._preview_codes({'route_ids': [route_a.id], 'batch_ids': [batch.id]}),
            ['SM-0001', 'SM-A', 'SM-B'],
        )

        options = smart_metering_scope_options()
        self.assertEqual([item['id'] for item in options['routes']], [route_a.id, route_b.id])
        self.assertEqual([item['id'] for item in options['route_zones']], [zone.id])
        self.assertEqual([item['id'] for item in options['batches']], [batch.id])

    def test_import_endpoint_allows_other_routes_on_the_same_date(self):
        route_a = Route.objects.create(token='RA')
        route_b = Route.objects.create(token='RB')
        self._smart_metering_sp_on_route('SM-A', route_a)
        self._smart_metering_sp_on_route('SM-B', route_b)

        first = self._post_import(scope={'route_ids': [route_a.id]})
        self.assertEqual(first.status_code, 202, first.data)
        document = ReadingDocument.objects.get(id=first.data['document_id'])
        self.assertEqual(document.smart_metering_scope, {'route_ids': [route_a.id]})

        other_route = self._post_import(scope={'route_ids': [route_b.id]})
        self.assertEqual(other_route.status_code, 202, other_route.data)

        # Loading every meter overlaps both documents: it can not replace them.
        all_meters = self._post_import(replace=True)
        self.assertEqual(all_meters.status_code, 409)
        self.assertFalse(all_meters.data['existing_document']['replaceable'])

        same_route = self._post_import(scope={'route_ids': [route_a.id]})
        self.assertEqual(same_route.status_code, 409)

    def test_import_endpoint_rejects_invalid_scope_ids(self):
        for scope in ({'route_ids': ['abc']}, {'route_ids': ['\u00b2']}, {'batch_ids': [0]}):
            response = self._post_import(scope=scope)
            self.assertEqual(response.status_code, 400, scope)
        self.assertFalse(ReadingDocument.objects.exists())

    def test_import_endpoint_allows_replace_only_without_losing_meters(self):
        route_a = Route.objects.create(token='RA')
        route_b = Route.objects.create(token='RB')
        self._smart_metering_sp_on_route('SM-A', route_a)
        self._smart_metering_sp_on_route('SM-B', route_b)
        first = self._post_import(scope={'route_ids': [route_a.id, route_b.id]})
        document = ReadingDocument.objects.get(id=first.data['document_id'])
        document.status = ReadingDocument.STATUS_PROCESSED
        document.save(update_fields=['status'])

        reduced = self._post_import(replace=True, scope={'route_ids': [route_a.id]})
        self.assertEqual(reduced.status_code, 409)
        self.assertTrue(reduced.data['existing_document']['scope_reduced'])
        self.assertFalse(reduced.data['existing_document']['replaceable'])

        widened = self._post_import(replace=True)
        self.assertEqual(widened.status_code, 202, widened.data)
        self.assertEqual(widened.data['document_id'], document.id)
        document.refresh_from_db()
        self.assertIsNone(document.smart_metering_scope)

    def test_import_ignores_client_preview_rows_outside_the_scope(self):
        route_a = Route.objects.create(token='RA')
        route_b = Route.objects.create(token='RB')
        sp_a = self._smart_metering_sp_on_route('SM-A', route_a)
        sp_b = self._smart_metering_sp_on_route('SM-B', route_b)
        scope = {'route_ids': [route_a.id]}
        document = self._new_document()
        document.smart_metering_scope = scope
        document.save(update_fields=['smart_metering_scope'])
        row = {
            'matched': True, 'reading_value': 130, 'reading_date': READING_DATE.isoformat(),
            'meter_code': '', 'contracts': [{}],
        }
        preview = {
            'reading_date': READING_DATE.isoformat(),
            'scope': scope,
            'readings': [
                {**row, 'id': sp_a.id, 'contract_id': sp_a.contracts.get().id},
                {**row, 'id': sp_b.id, 'contract_id': sp_b.contracts.get().id},
            ],
        }

        result = import_smart_metering_readings(document, preview=preview)

        self.assertEqual(result['saved_readings'], 1)
        self.assertEqual(list(document.readings.values_list('supply_point_id', flat=True)), [sp_a.id])

    def test_document_keeps_supply_points_that_changed_route(self):
        route_b = Route.objects.create(token='RB')
        route_c = Route.objects.create(token='RC')
        sp_b = self._smart_metering_sp_on_route('SM-B', route_b)
        first = self._post_import(scope={'route_ids': [route_b.id]})
        document = ReadingDocument.objects.get(id=first.data['document_id'])
        document.status = ReadingDocument.STATUS_PROCESSED
        document.save(update_fields=['status'])
        Reading.objects.create(
            meter=sp_b.meter, contract=sp_b.contracts.get(), supply_point=sp_b, document=document,
            reading_date=READING_DATE, reading_value=130, calculated_value=0,
        )
        # The supply point moves to another route after the load.
        sp_b.property.route_position.route = route_c
        sp_b.property.route_position.save()

        load_c = self._post_import(scope={'route_ids': [route_c.id]})

        self.assertEqual(load_c.status_code, 409)
        self.assertEqual(load_c.data['existing_document']['id'], document.id)
        self.assertTrue(load_c.data['existing_document']['replaceable'])
