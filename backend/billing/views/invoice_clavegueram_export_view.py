import openpyxl
from io import BytesIO
from datetime import datetime
from openpyxl.styles import Font, PatternFill

from django.http import HttpResponse
from django.utils.translation import gettext as _
from rest_framework.views import APIView
from rest_framework.permissions import IsAuthenticated
from rest_framework import status

from billing.models import Invoice
from coredata.utils.language_utils import use_default_language
from coredata.models import ConfigProject

class InvoiceClavegueramExcelExportView(APIView):
    permission_classes = [IsAuthenticated]

    @use_default_language
    def get(self, request, *args, **kwargs):
        # Obtenim els paràmetres de filtre des de la petició GET
        start_date = request.query_params.get('start_date')
        end_date = request.query_params.get('end_date')
        periode = request.query_params.get('periode')
        batch_id = request.query_params.get('batch_id')
        
        # Filtres base obligatoris per l'informe de Clavegueram
        filters = {
            'line_items__description__icontains': 'Clavegueram',
            'is_confirmed': True,
        }

        # Només factures (No pressupostos)
        try:
            invoice_type_token = ConfigProject.objects.get(token='invoice_type_invoice_token').value
            filters['type__token'] = invoice_type_token
        except ConfigProject.DoesNotExist:
            pass

        # Només cobrades
        try:
            paid_status_token = ConfigProject.objects.get(token='invoice_status_paid_token').value
            filters['status__token'] = paid_status_token
        except ConfigProject.DoesNotExist:
            pass

        # Apliquem filtres opcionals
        if start_date:
            filters['issue_date__gte'] = start_date
        if end_date:
            filters['issue_date__lte'] = end_date
            
        if batch_id:
            # Factures poden tenir relació 'batch' o 'billing', depenent del vostre model assumim 'batch_id'
            filters['batch_id'] = batch_id

        if periode:
            if periode.upper() == '1T':
                filters['issue_date__month__in'] = [1, 2, 3]
            elif periode.upper() == '2T':
                filters['issue_date__month__in'] = [4, 5, 6]
            elif periode.upper() == '3T':
                filters['issue_date__month__in'] = [7, 8, 9]
            elif periode.upper() == '4T':
                filters['issue_date__month__in'] = [10, 11, 12]

        invoices = Invoice.objects.filter(**filters).exclude(status__token='1').distinct().select_related(
            'contract', 'contract__use_type'
        ).prefetch_related('contract__variables')

        # Moviments de cartera (ajustaments, retorns, anul·lacions d'altres períodes)
        movements = []
        if start_date and end_date:
            from billing.models import PaymentMovement
            # Busquem moviments negatius produïts en aquest període que afecten factures anteriors
            movements = PaymentMovement.objects.filter(
                movement_date__range=(start_date, end_date),
                is_positive=False,
                payment__invoice__line_items__description__icontains='Clavegueram',
                payment__invoice__issue_date__lt=start_date
            ).select_related(
                'payment__invoice', 'payment__invoice__contract', 'payment__invoice__contract__use_type'
            ).prefetch_related('payment__invoice__contract__variables').distinct()

        # Creem l'espai de memòria i el llibre Excel
        output = BytesIO()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = _("Informe Clavegueram")

        # Capçaleres
        headers = [
            _("Contracte"),
            _("Titular"),
            _("NIF"),
            _("Adreça Subministrament"),
            _("Número Factura"),
            _("Data Factura"),
            _("Període Liquidat"),
            _("Consum anual"),
            _("Tipologia de client"),
            _("Quota Fixa €"),
            _("Quota Variable €"),
            _("Total Clavegueram €"),
            _("Observacions")
        ]

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
        
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header)
            cell.font = header_font
            cell.fill = header_fill

        row_num = 2
        # Combinem dades
        data_to_export = []
        for inv in invoices:
            data_to_export.append({'invoice': inv, 'is_adjustment': False})
        for mov in movements:
            data_to_export.append({'invoice': mov.payment.invoice, 'is_adjustment': True, 'movement': mov})

        for item in data_to_export:
            invoice = item['invoice']
            is_adj = item['is_adjustment']
            multiplier = -1.0 if is_adj else 1.0
            
            clavegueram_lines = invoice.line_items.filter(description__icontains='Clavegueram')
            
            quota_fixa = 0.0
            quota_variable = 0.0
            
            for line in clavegueram_lines:
                desc = str(line.description).lower() if line.description else ""
                val = float(line.total or 0) * multiplier
                
                # Heurística bàsica de validació Fixa vs Variable
                if 'conservació' in desc or 'quota' in desc or 'fix' in desc:
                    quota_fixa += val
                else:
                    quota_variable += val

            # Període dinàmic
            periode_str = ""
            if invoice.issue_date:
                m = invoice.issue_date.month
                if m in [1, 2, 3]: periode_str = "1T"
                elif m in [4, 5, 6]: periode_str = "2T"
                elif m in [7, 8, 9]: periode_str = "3T"
                elif m in [10, 11, 12]: periode_str = "4T"
                periode_str = f"{periode_str} {invoice.issue_date.year}"

            # Tipologia de client
            client_type = "Desconegut"
            if invoice.contract and invoice.contract.use_type:
                client_type = invoice.contract.use_type.name

            # Consum anual variable
            consum_anual = ""
            if invoice.contract:
                ca_vars = [v for v in invoice.contract.variables.all() if 'Consum anual' in (v.name or '')]
                if ca_vars:
                    ca_vars.sort(key=lambda x: x.name or '', reverse=True)
                    consum_anual = ca_vars[0].value or ""

            # Dades bàsiques
            contract_token = invoice.contract.token if invoice.contract else ""
            titular = invoice.customer_final or ""
            nif = invoice.customer_token_final or invoice.payer_token_final or ""
            
            address_parts = filter(bool, [invoice.address_final, invoice.postal_code_final, invoice.city_final])
            address = ", ".join(address_parts)
            
            num_factura = invoice.serie_final or invoice.number or invoice.token or ""
            data_factura = invoice.issue_date.strftime("%Y-%m-%d") if invoice.issue_date else ""

            # Observacions d'ajustament
            obs = ""
            if is_adj:
                mov = item['movement']
                obs = f"AJUSTAMENT: {mov.current_status.name if mov.current_status else 'Retorn'} en data {mov.movement_date}"

            # Escrivim fila a Excel
            ws.cell(row=row_num, column=1, value=contract_token)
            ws.cell(row=row_num, column=2, value=titular)
            ws.cell(row=row_num, column=3, value=nif)
            ws.cell(row=row_num, column=4, value=address)
            ws.cell(row=row_num, column=5, value=num_factura)
            ws.cell(row=row_num, column=6, value=data_factura)
            ws.cell(row=row_num, column=7, value=periode_str)
            ws.cell(row=row_num, column=8, value=consum_anual)
            ws.cell(row=row_num, column=9, value=client_type)
            ws.cell(row=row_num, column=10, value=quota_fixa).number_format = '0.00 €'
            ws.cell(row=row_num, column=11, value=quota_variable).number_format = '0.00 €'
            ws.cell(row=row_num, column=12, value=(quota_fixa + quota_variable)).number_format = '0.00 €'
            ws.cell(row=row_num, column=13, value=obs)
            
            row_num += 1

        # Ajustem l'amplada de les columnes
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            ws.column_dimensions[column].width = max_length + 2

        # Desem el file object al buffer dictat prèviament
        wb.save(output)
        output.seek(0)
        
        filename = f"informe_clavegueram_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"

        # Resposta de l'API emetent fitxer adjunt
        response = HttpResponse(
            output,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename={filename}'
        response['Access-Control-Expose-Headers'] = 'Content-Disposition'

        return response
