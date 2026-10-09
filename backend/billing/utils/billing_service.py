import calendar
import datetime
import re
from django.db.models import Max, Q
from django.http import HttpResponse
from django.utils import timezone
from django.utils.translation import gettext as _
import openpyxl
from openpyxl.styles import PatternFill
from billing.models import Billing, Invoice, Reading
from contract.models import Contract, ContractTerminationRequest
from coredata.models import ConfigProject
from service.models import Exploitation
from statistics.views.reports_views import add_row, adjust_column_widths

PERIOD_TYPE_INTERVAL = {
    'mensual': 1,
    'bimestral': 2,
    'trimestral': 3,
    'quadrimestral': 4,
    'semestral': 6,
    'anual': 12,
}


def get_billing_period_range(billing):
    interval = PERIOD_TYPE_INTERVAL.get(billing.biller.period_type if billing.biller else None, 3)
    today = timezone.now().date()

    latest_own_reading_date = Reading.objects.filter(
        billing=billing, is_control=False
    ).aggregate(max_date=Max('reading_date'))['max_date']

    if latest_own_reading_date:
        end_month = latest_own_reading_date.month
        end_year = latest_own_reading_date.year
    else:
        period_month_year_match = re.search(r'-\s*(\d{1,2})(\d{4})\s*$', billing.token or '') or \
            re.search(r'-\s*(\d{1,2})(\d{4})\s*$', billing.name or '')
        end_month = int(period_month_year_match.group(1)) if period_month_year_match else None
        end_year = int(period_month_year_match.group(2)) if period_month_year_match else None
        if not end_month or not (1 <= end_month <= 12):
            period_reference = billing.created_at.date() if billing.created_at else today
            end_month = period_reference.month
            end_year = period_reference.year

    start_month = end_month - interval + 1
    start_year = end_year
    if start_month <= 0:
        start_month += 12
        start_year -= 1

    first_start_reading = datetime.date(start_year, start_month, 1)
    last_day_of_end_month = calendar.monthrange(end_year, end_month)[1]
    # No es poden generar factures amb dates futures.
    last_end_reading = min(datetime.date(end_year, end_month, last_day_of_end_month), today)

    return first_start_reading, last_end_reading


def get_billing_scope_reading_filter(billing):
    scope_filter = Q(supply_point__property__route_position__route__biller=billing.biller)

    exploitations = list(billing.invoices.values_list('exploitation__id', flat=True).distinct())
    if len(exploitations) == 1 and exploitations[0] and Exploitation.objects.count() > 1:
        scope_filter |= Q(supply_point__connection__exploitation__id=exploitations[0])

    return scope_filter


def generate_billing_missing_contracts_report(contracts_data, billing, first_start_reading, last_end_reading):
    
    contract_ids = [item['contract_id'] for item in contracts_data]
    contracts = Contract.objects.filter(id__in=contract_ids)
    readings = Reading.objects.filter(
        contract__in=contracts,
        reading_date__gte=first_start_reading,
        reading_date__lte=last_end_reading,
        is_control=False,
        is_initial=False,
        is_active=True,
    ).order_by('reading_date')
    max_previous_reading_date = first_start_reading - datetime.timedelta(days=90)
    previous_readings = Reading.objects.filter(
        contract__in=contracts,
        reading_date__lte=first_start_reading,
        reading_date__gte=max_previous_reading_date,
        is_control=False,
        is_initial=False,
        is_active=True,
    ).order_by('reading_date')
    terminated_status = ConfigProject.objects.get(token='contract_termination_completed_token').value
    contract_termination_requests = ContractTerminationRequest.objects.filter(
        contract__in=contracts,
        approved_at__isnull=False,
        status__token=terminated_status,
    ).order_by('-created_at')
    
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = _("REMESA AQUA")
    
    row = 0
    no_reading_text = _("No reading")
    no_price_rates_text = _("No price rates")
    titles = [
            _("Contract"), _("Creation date"),
            _("Status"), _("Termination date"),
            _("CIF Client"), _("Client Name"),
            _("Billable"), no_price_rates_text,
            _("Exploitation"), _("Route"),
            _("Reading date"), _("Reading value"), _("Consumption"),
            _("Meter"), _("Last invoice"), _("Last reading date"), _("Last reading consumption")
            ]
    

    row = add_row(sheet, row, titles, fill=PatternFill(start_color="b8cce4", end_color="b8cce4", fill_type="solid"))
    for item in contracts_data:
        reading = readings.filter(contract__id=item['contract_id']).order_by('-reading_date').first()
        prev_reading = previous_readings.filter(contract__id=item['contract_id']).order_by('-reading_date').first()
        contract = contracts.get(id=item['contract_id'])
        print(contract.token)
        termination = contract_termination_requests.filter(contract__id=item['contract_id']).order_by('-created_at').first()
        line_row = [
            contract.token, contract.created_at.strftime('%d/%m/%Y'), 
            contract.status.name.upper(), termination.approved_at.strftime('%d/%m/%Y') if termination else '',
            contract.holder.token, f"{contract.holder.name}{' ' + contract.holder.surname if contract.holder.surname else ''}",
            "No" if contract.block_billing else "", no_price_rates_text if 'block_billing' in item['missing_reasons'] else '',
            contract.supply_point_default.connection.exploitation.name, contract.supply_point_default.property.route_position.route.name,
            
            no_reading_text if 'no_reading' in item['missing_reasons'] else reading.reading_date.strftime('%d/%m/%Y') if reading else '', 
            no_reading_text if 'no_reading' in item['missing_reasons'] else reading.reading_value if reading else '',
            no_reading_text if 'no_reading' in item['missing_reasons'] else reading.calculated_value if reading else '',
            
            contract.supply_point_default.meter.code, 
            prev_reading.reading_date.strftime('%d/%m/%Y') if prev_reading else no_reading_text,
            prev_reading.reading_value if prev_reading else no_reading_text,
            prev_reading.calculated_value if prev_reading else no_reading_text,
        ]
        row = add_row(sheet, row, line_row)
        
    adjust_column_widths(sheet)
        
    clean_token = re.sub(r'[^A-Z0-9]', '', str(billing.token).upper())
    filename = f"{clean_token}.xlsx"
    
    wb.save(filename)

    with open(filename, "rb") as f:
        response = HttpResponse(f.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f"attachment; filename={filename}"
    
    return response


def get_billed_invoices_without_billing(billing_id, detailed=False, billed_type=None):
    billing = Billing.objects.get(id=billing_id)
    invoice_type_invoice_token = ConfigProject.objects.get(token='invoice_type_invoice_token').value
    invoice_status_pending_token = ConfigProject.objects.get(token='invoice_status_pending_token').value
    invoice_status_payoff_token = ConfigProject.objects.get(token='invoice_status_payoff_token').value

    first_start_reading, last_end_reading = get_billing_period_range(billing)

    billing_batch_ids = Reading.objects.filter(
        billing=billing,
        batch__isnull=False,
    ).values_list('batch__id', flat=True).distinct()

    scope_filter = get_billing_scope_reading_filter(billing)
    reading_source_filter = (
        Q(batch__id__in=billing_batch_ids) |
        (scope_filter & Q(reading_date__range=(first_start_reading, last_end_reading)))
    )

    relevant_readings = Reading.objects.filter(
        reading_source_filter,
        is_control=False,
        is_initial=False,
        is_active=True,
    )

    billed_without_billing = Invoice.objects.filter(
        readings__in=relevant_readings,
        is_active=True,
        billing__isnull=True,
    ).exclude(status__token=invoice_status_payoff_token).distinct()

    billing_invoices = Invoice.objects.filter(billing=billing, type_final=invoice_type_invoice_token, is_active=True).exclude(status__token=invoice_status_payoff_token).distinct()
    # Invoices store the original budget's token on budget_token; exclude those budgets by token.
    billing_invoice_budget_tokens = list(
        billing_invoices.filter(budget_token__isnull=False).values_list('budget_token', flat=True).distinct()
    )
    invoices = billed_without_billing.filter(type_final=invoice_type_invoice_token)
    invoice_budgets_token = list(
        invoices.filter(budget_token__isnull=False).values_list('budget_token', flat=True).distinct()
    )
    invoices_other_mngs = invoices.filter(Q(contract_request__isnull=False) | Q(contract_termination__isnull=False)).distinct()
    budgets = billed_without_billing.exclude(type_final=invoice_type_invoice_token).exclude(
        token__in=invoice_budgets_token
    ).exclude(
        token__in=billing_invoice_budget_tokens
    ).distinct()
    budgets_other_mngs = budgets.filter(Q(contract_request__isnull=False) | Q(contract_termination__isnull=False)).distinct()
    in_other_billing = Invoice.objects.filter(
        readings__in=relevant_readings,
        is_active=True,
        billing__isnull=False,
    ).exclude(billing=billing).exclude(status__token=invoice_status_payoff_token).distinct()

    total_invoices = invoices.count()
    total_budgets = budgets.count()
    
    readings = relevant_readings.filter(batch__id__in=billing_batch_ids, invoices__in=invoices | budgets).distinct()
    pending_invoices = invoices.filter(status__token=invoice_status_pending_token)
    reading_within_period = relevant_readings.filter(invoices__in=invoices | budgets).exclude(id__in=readings.values_list('id', flat=True)).distinct()

    pre_invoice_matches = _get_pre_invoice_matches(
        billing,
        reading_source_filter,
        set(list(budgets.values_list('id', flat=True)) + list(invoices.values_list('id', flat=True))),
        invoice_status_pending_token,
        invoice_type_invoice_token,
        invoice_status_payoff_token
    )

    return_data = {
        'total_billed': total_invoices + total_budgets,
        'total_invoices': total_invoices,
        'total_budgets': total_budgets,
        'total_in_other_billing': in_other_billing.count(),
        'total_readings_in_batch': readings.count(),
        'total_pending_invoices': pending_invoices.count(),
        'total_found_match': len(pre_invoice_matches),
        'total_other_mngs': invoices_other_mngs.count() + budgets_other_mngs.count(),
    }

    if not detailed:
        return return_data

    return_data['results'] = _get_billed_invoices_detail(
        billed_without_billing,
        billing,
        billing_batch_ids,
        reading_source_filter,
        invoice_type_invoice_token,
        first_start_reading,
        last_end_reading,
        invoice_status_pending_token,
        billed_type,
        invoice_status_payoff_token
    )

    return return_data


def _get_pre_invoice_matches(billing, reading_source_filter, budget_ids, invoice_status_pending_token, invoice_type_invoice_token, invoice_status_payoff_token):
    """PRE INVOICE FOUND IN BILLING WITH SAME READINGS AND CONTRACT, per budget id."""
    if not budget_ids:
        return {}

    contract_by_budget = dict(
        Invoice.objects.filter(id__in=budget_ids).exclude(status__token=invoice_status_payoff_token).values_list('id', 'contract__id')
    )

    budget_ids_by_reading = {}
    budget_reading_rows = Reading.objects.filter(
        reading_source_filter,
        is_control=False,
        is_initial=False,
        is_active=True,
        invoices__id__in=budget_ids,
    ).values_list('invoices__id', 'id').distinct()
    for budget_id, reading_id in budget_reading_rows:
        budget_ids_by_reading.setdefault(reading_id, set()).add(budget_id)

    if not budget_ids_by_reading:
        return {}

    match_id_by_budget = {}
    candidate_rows = Reading.objects.filter(
        id__in=list(budget_ids_by_reading.keys()),
        invoices__billing=billing,
        invoices__is_active=True,
    ).filter(
        Q(invoices__status__token=invoice_status_pending_token) |
        Q(invoices__type_final=invoice_type_invoice_token)
        ).values_list('id', 'invoices__id', 'invoices__contract__id', 'invoices__status__token', 'invoices__billing__id').order_by('-invoices__id').distinct()
    
    
    
    for reading_id, candidate_id, candidate_contract_id, candidate_status_token, candidate_billing_id in candidate_rows:
        for budget_id in budget_ids_by_reading.get(reading_id, ()):
            if budget_id in match_id_by_budget:
                continue
            if contract_by_budget.get(budget_id) != candidate_contract_id:
                continue
            if candidate_status_token == invoice_status_payoff_token:
                continue
            if budget_id == candidate_id:
                continue
            if not candidate_billing_id:
                continue
            match_id_by_budget[budget_id] = candidate_id
    if not match_id_by_budget:
        return {}

    candidate_by_id = {
        candidate.id: candidate
        for candidate in Invoice.objects.filter(
            id__in=set(match_id_by_budget.values()),
        ).select_related('contract', 'contract__holder')
    }
    
    matches = {}
    for budget_id, candidate_id in match_id_by_budget.items():
        # try:
        #     budget = Invoice.objects.get(id=budget_id)
        #     billed_invoice = Invoice.objects.filter(budget_token=budget.token, type_final=invoice_type_invoice_token).first().id or budget_id
        # except Exception as e:
        #     billed_invoice = budget_id
        if budget_id == candidate_id:
            continue
        candidate = candidate_by_id[candidate_id]
        matches[budget_id] = {
            'id': candidate.id,
            'serie_final': candidate.serie_final,
            'contract_token': candidate.contract.token if candidate.contract else None,
            'contract_holder_name': str(candidate.contract.holder) if candidate.contract and candidate.contract.holder else None,
            'total_final': candidate.total_final,
            'is_preinvoice': candidate.status.token == invoice_status_pending_token,
        }
    return matches


def _get_billed_invoices_detail(
    billed_without_billing,
    billing,
    billing_batch_ids,
    reading_source_filter,
    invoice_type_invoice_token,
    first_start_reading,
    last_end_reading,
    invoice_status_pending_token,
    billed_type,
    invoice_status_payoff_token
):
    invoices_qs = billed_without_billing.filter(type_final=invoice_type_invoice_token)
    invoice_budgets_token = list(
        invoices_qs.filter(budget_token__isnull=False).values_list('budget_token', flat=True).distinct()
    )
    billing_invoice_contract_ids = billing.invoices.filter(contract__isnull=False, is_active=True).values_list('contract__id', flat=True).distinct()
    billing_invoice_budget_tokens = list(
        Invoice.objects.filter(
            billing=billing,
            type_final=invoice_type_invoice_token,
            is_active=True,
            budget_token__isnull=False,
        ).exclude(status__token=invoice_status_payoff_token).values_list('budget_token', flat=True).distinct()
    )
    budgets_qs = billed_without_billing.exclude(type_final=invoice_type_invoice_token).exclude(
        token__in=invoice_budgets_token
    ).exclude(
        token__in=billing_invoice_budget_tokens
    ).distinct()

    if billed_type == 'confirmed_invoices':
        billed_without_billing = invoices_qs.distinct()
    elif billed_type == 'pending_budgets':
        billed_without_billing = budgets_qs
    else:
        billed_without_billing = (invoices_qs | budgets_qs).distinct()
    billed_invoice_ids = set(billed_without_billing.values_list('id', flat=True))
    if not billed_invoice_ids:
        return []

    billing_batch_id_set = set(billing_batch_ids)

    pre_invoice_matches = _get_pre_invoice_matches(
        billing,
        reading_source_filter,
        set(billed_without_billing.values_list('id', flat=True)),
        invoice_status_pending_token,
        invoice_type_invoice_token,
        invoice_status_payoff_token
    )

    reading_rows = Reading.objects.filter(
        reading_source_filter,
        is_control=False,
        is_initial=False,
        is_active=True,
        invoices__id__in=billed_invoice_ids,
    ).values(
        'invoices__id', 'id', 'reading_date', 'batch__id', 'billing__id'
    ).distinct()

    readings_by_invoice = {}
    for row in reading_rows:
        invoice_id = row['invoices__id']
        if invoice_id not in billed_invoice_ids:
            continue
        readings_by_invoice.setdefault(invoice_id, []).append(row)

    results = []

    invoices = list(billed_without_billing.select_related(
        'contract', 'contract__holder', 'status', 'contract_termination'
    ).order_by('-contract','-issue_date', '-id'))
    
    contract_set = {}
    invoice_contract = {}
    for invoice in invoices:
        contract = None
        if invoice.contract:
            contract = invoice.contract
        elif invoice.contract_request:
            try:
                contract = Contract.objects.get(contract_request=invoice.contract_request)
            except Exception as e:
                print(e)
        elif invoice.contract_termination:
            contract = invoice.contract_termination.contract
        invoice_contract[invoice.id] = contract
        if not contract:
            continue
        if contract.id not in contract_set:
            contract_set[contract.id] = 0
        contract_set[contract.id] += 1

    for invoice in invoices:
        reading_rows_for_invoice = readings_by_invoice.get(invoice.id, [])
        reasons = set()
        readings_detail = []
        contract = invoice_contract.get(invoice.id)
        contract_already_added = bool(contract and contract_set.get(contract.id, 0) > 1)
            

        for row in reading_rows_for_invoice:
            in_billing_batch = row['batch__id'] in billing_batch_id_set
            in_period = bool(
                row['reading_date'] and first_start_reading <= row['reading_date'] <= last_end_reading
            )
            if in_billing_batch:
                reasons.add('reading_in_billing_batch')
            if in_period:
                reasons.add('reading_in_period')
            if not row['batch__id']:
                reasons.add('no_batch')
            if not row['billing__id']:
                reasons.add('no_billing')
            elif row['billing__id'] == billing.id:
                reasons.add('reading_in_current_billing')

            readings_detail.append({
                'reading_id': row['id'],
                'reading_date': row['reading_date'],
                'batch_id': row['batch__id'],
                'reading_billing_id': row['billing__id'],
                'in_billing_batch': in_billing_batch,
                'in_period': in_period,
            })

        readings_detail.sort(key=lambda item: (item['reading_date'] is None, item['reading_date']))
        invoice_is_budget = invoice.type_final != invoice_type_invoice_token
        # found_match = pre_invoice_matches.get(invoice.id) if invoice_is_budget else None
        found_match = pre_invoice_matches.get(invoice.id)
        if billed_type == 'found_match':
            if not found_match:
                continue
        results.append({
            'invoice_id': invoice.id,
            'invoice_token': invoice.token,
            'invoice_serie_final': invoice.serie_final,
            'is_budget': invoice_is_budget,
            'issue_date': invoice.issue_date,
            'total': invoice.total_final,
            'invoice_status_color': invoice.status.color if invoice.status else None,
            'invoice_status_name': invoice.status.name if invoice.status else None,
            'invoice_status_id': invoice.status.id if invoice.status else None,
            'is_confirmed': invoice.is_confirmed,
            'contract_id': invoice.contract.id if invoice.contract else None,
            'contract_token': invoice.contract.token if invoice.contract else None,
            'contract_holder_name': str(invoice.contract.holder) if invoice.contract and invoice.contract.holder else None,
            'contract_status_name': invoice.contract.status.name if invoice.contract and invoice.contract.status else None,
            'contract_status_color': invoice.contract.status.color if invoice.contract and invoice.contract.status else None,
            'contract_status_id': invoice.contract.status.id if invoice.contract and invoice.contract.status else None,
            'readings': readings_detail,
            'reasons': sorted(reasons),
            'found_match': found_match,
            'multiple_budgets': contract_already_added,
            'other_managements': "contract_termination" if invoice.contract_termination else "contract_request" if invoice.contract_request else None,
            'contract_in_billing': not contract or (contract.id in billing_invoice_contract_ids),
        })

    return results


def move_budget_to_billing_pre_invoice(billing_id, budget_ids):
    try:
        billing = Billing.objects.get(id=billing_id)
        invoice_type_invoice_token = ConfigProject.objects.get(token='invoice_type_invoice_token').value
        budgets = Invoice.objects.filter(id__in=budget_ids)
        readings = Reading.objects.filter(invoices__in=budgets, billing__isnull=True)
        readings.update(billing=billing)
        for budget in budgets:
            if budget.type_final == invoice_type_invoice_token:
                print("budget is already a pre-invoice")
                continue
            budget.billing = billing
            budget.type_final = invoice_type_invoice_token
        Invoice.objects.bulk_update(budgets, ['billing', 'type_final'])
        return {
            'success': True,
            'message': 'Budgets moved to billing pre-invoice',
        }
    except Exception as e:
        print(e)
        return None

def generate_invoice_budgets_from_billing(billing_id, budget_ids):
    try:
        from billing.views.generate_invoice_budget_view import pass_budget_to_invoice
        billing = Billing.objects.get(id=billing_id)
        billing_batch_processed = ConfigProject.objects.get(token='billing_batch_processed').value
        budgets = Invoice.objects.filter(id__in=budget_ids)
        invoice_ids = []
        for budget in budgets:
            new_invoice = pass_budget_to_invoice(budget)
            if new_invoice:
                invoice_ids.append(new_invoice.id)
        if billing.status.token == billing_batch_processed and invoice_ids:
            Invoice.objects.filter(id__in=invoice_ids).update(billing=billing)
        return {
            'success': True,
            'message': 'Budgets moved to billing pre-invoice',
        }
    except Exception as e:
        print(e)
        return None

def move_invoice_to_billing(billing_id, invoice_ids):
    try:
        billing = Billing.objects.get(id=billing_id)
        invoices = Invoice.objects.filter(id__in=invoice_ids)
        invoices.update(billing=billing)
        return {
            'success': True,
            'message': 'Invoices moved to billing',
        }
    except Exception as e:
        print(e)
        return None