import csv
import io

from openpyxl import load_workbook
from django.db.models import Q

from billing.models import Reading
from billing.utils.reading_filters import PENDING_READING_FILTER
from coredata.models import ConfigProject
from service.models import Meter, SupplyPoint


def get_batch_supply_point_ids(reading_batch):
    """
    Supply points que HAURIEN de formar part del lot: els de les seves rutes
    (filtrats pels flags telecontrol/manual del lot) més els dels seus fix_meters,
    que s'inclouen sempre.
    """
    sp_query = SupplyPoint.objects.filter(
        property__route_position__route__reading_batches=reading_batch
    )

    if not reading_batch.include_telecontrol:
        sp_query = sp_query.filter(
            Q(meter__has_remote_reading=False) | Q(meter__force_manual_reading=True)
        )
    elif not reading_batch.include_manual:
        sp_query = sp_query.filter(
            meter__has_remote_reading=True, meter__force_manual_reading=False
        )

    supply_point_ids = set(sp_query.values_list('id', flat=True))
    supply_point_ids.update(
        SupplyPoint.objects.filter(
            meter__reading_batches=reading_batch
        ).values_list('id', flat=True)
    )
    return supply_point_ids


def get_missing_supply_points_queryset(reading_batch, active_contract_token=None):
    """
    Supply points del lot amb contracte actiu i sense cap lectura dins del lot.
    Mateix criteri que el llistat "Subministraments sense lectura" del setup.
    """
    if active_contract_token is None:
        active_contract_token = ConfigProject.objects.get(
            token='contract_active_token'
        ).value

    read_supply_points = set(
        reading_batch.readings.filter(is_control=False).values_list(
            'supply_point_id', flat=True
        )
    )
    missing_supply_point_ids = get_batch_supply_point_ids(reading_batch) - read_supply_points

    if not missing_supply_point_ids:
        return SupplyPoint.objects.none()

    return SupplyPoint.objects.select_related(
        'type', 'status', 'property__route_position__route'
    ).filter(
        id__in=missing_supply_point_ids,
        contracts__status__token=active_contract_token,
    ).order_by('id').distinct()


def get_extra_supply_point_ids(reading_batch):
    """
    Supply points que apareixen als fitxers de lectures del lot però que no
    formen part del lot (lectures sobrants).
    """
    files_supply_point_ids = set(
        Reading.objects.filter(
            document__batch=reading_batch, is_control=False
        ).values_list('supply_point_id', flat=True)
    )
    files_supply_point_ids.discard(None)
    return files_supply_point_ids - get_batch_supply_point_ids(reading_batch)


# Apartats del pas de lectures pendents del lot (paràmetre ?readings= del setup)
SETUP_FILTERS = (
    'missing',
    'extra',
    'reader_alert',
    'remote_alert',
    'no_reading_value',
    'remote_reader',
    'inactive_sp',
)


def get_setup_filter_supply_point_ids(reading_batch, filter_key, active_contract_token=None):
    """
    Supply points de cada apartat del pas de lectures pendents del lot.
    Font única per al comptador (ReadingBatchSummaryView) i per al llistat
    (ReadingBatchSetupViewSet), perquè el número i el detall sempre coincideixin.
    Els apartats basats en lectures només tenen en compte les pendents de facturar.
    """
    if filter_key == 'missing':
        return set(
            get_missing_supply_points_queryset(
                reading_batch, active_contract_token
            ).values_list('id', flat=True)
        )
    if filter_key == 'extra':
        return get_extra_supply_point_ids(reading_batch)

    readings = reading_batch.readings.filter(is_control=False).filter(PENDING_READING_FILTER)
    if filter_key == 'reader_alert':
        readings = readings.filter(reader_alert__isnull=False)
    elif filter_key == 'remote_alert':
        readings = readings.filter(remote_alert__isnull=False)
    elif filter_key == 'no_reading_value':
        readings = readings.filter(reading_value__isnull=True)
    elif filter_key == 'remote_reader':
        readings = readings.filter(
            meter__has_remote_reading=True, meter__force_manual_reading=False
        )
    elif filter_key == 'inactive_sp':
        supply_active_token = ConfigProject.objects.get(
            token='supply_point_status_activate_token'
        ).value
        supply_cut_token = ConfigProject.objects.get(
            token='supply_point_status_cut_token'
        ).value
        readings = readings.exclude(
            supply_point__status__token__in=[supply_active_token, supply_cut_token]
        )
    else:
        raise ValueError(f'Filtre de lot desconegut: {filter_key}')

    supply_point_ids = set(readings.values_list('supply_point_id', flat=True))
    supply_point_ids.discard(None)
    return supply_point_ids


def get_meters_from_file(file):
    if not file:
        return []

    filename = (getattr(file, "name", "") or "").lower()
    content_type = (getattr(file, "content_type", "") or "").lower()
    file_bytes = _read_file_bytes(file)
    
    if not file_bytes:
        return []

    is_spreadsheet = _is_spreadsheet(filename, content_type, file_bytes)
    
    rows = (
        _read_spreadsheet_rows(file_bytes)
        if is_spreadsheet
        else _read_csv_rows(file_bytes)
    )

    meter_reading_map = {}
    meter_codes = []
    reading_ids = []
    for row in rows:
        normalized_row = {
            (key or "").strip().lower(): (row_value if row_value is not None else "")
            for key, row_value in row.items()
            if key
        }
        meter_code = str(normalized_row.get("meter", "")).strip()
        reading_id = str(normalized_row.get("reading_id", "")).strip()
        meter_codes.append(meter_code)
        reading_ids.append(reading_id)
        if not meter_code or not reading_id:
            continue

        meter_reading_map[meter_code] = reading_id

    if not meter_reading_map:
        return []

    existing_meters = Meter.objects.filter(code__in=meter_reading_map.keys())
    existing_meters_codes = existing_meters.values_list('code', flat=True)
    not_found_meters = []

    for meter_code in meter_codes:
        if meter_code not in existing_meters_codes:
            not_found_meters.append(meter_code)

    return existing_meters_codes, reading_ids, not_found_meters


def _read_file_bytes(file):
    if hasattr(file, "seek"):
        file.seek(0)

    if hasattr(file, "read"):
        content = file.read()
    else:
        content = file

    if isinstance(content, str):
        return content.encode("utf-8")
    return content


def _is_spreadsheet(filename, content_type, file_bytes):
    if filename.endswith((".xlsx", ".xlsm")):
        return True
    if "spreadsheetml" in content_type:
        return True
    # XLSX files are zipped; detect PK header
    return isinstance(file_bytes, (bytes, bytearray)) and file_bytes[:2] == b"PK"


def _read_csv_rows(file_bytes):
    text = file_bytes.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text), delimiter=";")
    for row in reader:
        yield row


def _read_spreadsheet_rows(file_bytes):
    workbook = load_workbook(filename=io.BytesIO(file_bytes), read_only=True, data_only=True)
    sheet = workbook.active
    rows_iter = sheet.iter_rows(values_only=True)
    headers = [
        (str(cell).strip() if cell is not None else "")
        for cell in next(rows_iter, [])
    ]

    for values in rows_iter:
        row_dict = {}
        for index, header in enumerate(headers):
            if not header:
                continue
            value = values[index] if index < len(values) else None
            row_dict[header.lower()] = value
        yield row_dict