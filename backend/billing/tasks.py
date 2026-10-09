import csv
import re
from decimal import Decimal
from html import escape
from io import BytesIO, StringIO
import itertools
import time
import uuid
import base64
import io
import zipfile
from urllib.parse import urljoin
from celery import shared_task, chord, group
from billing.utils.barcode_service import generate_barcode
from statistics.utils.report_service import delete_file_later
from django.core.files.storage import default_storage
from celery_progress.backend import ProgressRecorder
from decouple import config
from django.conf import settings
from django.db import transaction, close_old_connections
from django.core.files.base import ContentFile
from billing.utils.confirm_invoice_service import confirm_invoice
from billing.utils.adjustment_service import billing_run_cache_scope
from billing.utils.payment_service import disconnect_payment_signals, get_status_map, reconnect_payment_signals
from billing.utils.datetime_utils import convert_date_with_timezone
from django.template.loader import render_to_string
from xhtml2pdf import pisa
from billing.utils.reading_service import (
    calculate_estimated_bag,
    check_billing_period,
    check_unusual_consumption,
    classify_reading_processing_error,
    fanout_general_meter_readings,
    find_canonical_general_meter_reading,
    get_estimated_reading_minimal_object,
    get_estimated_reading,
    READING_ESTIMATE_PERIOD_CHOICES,
    READING_ESTIMATE_PERIOD_SAME_PERIOD_PREVIOUS_YEAR,
    READING_ESTIMATE_STATISTIC_CHOICES,
    READING_ESTIMATE_STATISTIC_MEAN,
    get_reading_document_target_contract,
    get_reading_minimal_object,
)
from billing.serializers.reading_serializer import ReadingMinimalSerializer
from billing.views.invoice_pdf_view import generate_report_invoice_pdf
from claimrequest.models import ClaimRequest
from communication.models import Communication, CommunicationProcessStatus, CommunicationStatus
from coredata.models import ConfigProject
from contract.models import Contract, ContractStatus, PaymentType
from coredata.serializers import PersonSerializer
from documentmanager.utils.main_utils import upload_document
from documentmanager.utils.sign_certificate_service import sign_pdf
from logger.models import LogInvoiceChangeStatus
from notification.models import CalendarTask, Notification
from service.models import CompanyBank, Exploitation, Meter, Route, SupplyPoint
from service.serializers.company_serializer import CompanySerializer
from verifactu.utils.save_verifactu_info import notify_verifactu_unsent_batches
from .models import ( Biller, Billing, BillingBatch, BillingBatchStatus, BillingStatus, EstimatedBagMovement, Invoice, InvoiceStatus, InvoiceWarning, JoinedPayment, Payment, PaymentStatus, ReadingAlert, ReadingBatch, Reading, ReadingBatchStatus, ReadingDocument, ReadingDocumentNotFound, RemoteReadingAlert, EstimatedBag,)
from statistics.models import BillingConsumption, ReadingBatchExportColumn
from statistics.utils.billing_amount_average import TotalAmountAvgCache, is_high_amount_for_period
from billing.utils.invoice_service import generate_consumption_invoice_multiple, get_invoice_status, log_invoice_data_change
from datetime import date, datetime, timedelta
from django.utils import timezone
from django.utils import translation
from django.utils.translation import gettext as _

from django.db.models import Count, Sum, Q, Min, Max, F, ExpressionWrapper, fields, Avg, Prefetch
from django.db.models.functions import Abs
from statistics.tasks import update_billing_consumption
from billing.utils.smart_metering_service import assign_smart_metering_readings, build_smart_metering_preview, _json_safe, SMART_METERING_PREVIEW_API_TIMEOUT
from billing.utils.reading_filters import PENDING_READING_FILTER

class ExportReadingBatchTaskError(ValueError):
    """Raised so Celery records the task as FAILURE (AsyncResult.state == 'FAILURE')."""


def _reading_batch_export_headers():
    return [
        str(_("Comptador")),
        str(_("Adreça subm.")),
        str(_("Route")),
        str(_("Posició")),
        str(_("Contracte")),
    ]


def _export_meter_row(meter, supply_address, route_label, position_label, contract_token):
    code = (meter.code or "").strip() if meter else ""
    return [
        code,
        (supply_address or "").strip(),
        (route_label or "").strip(),
        (position_label or "").strip(),
        (contract_token or "").strip(),
    ]


_READING_BATCH_EXPORT_PLACEHOLDER_RE = re.compile(r"%([a-zA-Z_][a-zA-Z0-9_.]*)")
_READING_BATCH_EXPORT_ESCAPE_SENTINEL = "\x1fRBEXP\x1f"


def _format_reading_batch_export_value(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, Decimal):
        s = format(value, "f")
        if "." in s:
            s = s.rstrip("0").rstrip(".")
        return s or "0"
    if isinstance(value, (int, float)):
        return str(value)
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value)


def _resolve_reading_batch_export_attr_chain(obj, dotted_tail):
    """Walk Django attributes; formats scalars at the end."""
    if obj is None:
        return ""
    if not dotted_tail:
        return _format_reading_batch_export_value(obj)
    cur = obj
    for part in dotted_tail.split("."):
        if cur is None:
            return ""
        cur = getattr(cur, part, None)
    return _format_reading_batch_export_value(cur)


def _resolve_reading_batch_export_placeholder(path, ctx, row_index):
    """
    Column expression placeholder without leading %.
    Roots: contract, route, supply, meter, contact, position (RoutePosition); also increment (row numbers starting at 1).
    """
    key = path.strip()
    if key.lower() == "increment":
        return str(row_index)

    parts = key.split(".", 1)
    root_name = parts[0].lower()
    tail = parts[1] if len(parts) > 1 else None

    root_map = {
        "contract": ctx.get("contract"),
        "route": ctx.get("route"),
        "supply": ctx.get("supply_point"),
        "meter": ctx.get("meter"),
        "contact": ctx.get("contact"),
        "position": ctx.get("route_position"),
    }
    if root_name not in root_map:
        return ""
    root_obj = root_map[root_name]
    if root_obj is None:
        return ""
    if tail is None:
        return _format_reading_batch_export_value(root_obj)
    return _resolve_reading_batch_export_attr_chain(root_obj, tail)


def _reading_batch_export_expression(value_expr, ctx, row_index):
    """
    Configured column value: plain text is copied as-is; %path resolves DB fields.
    Use %% for a literal percent sign.
    """
    if value_expr is None:
        return ""
    s = str(value_expr)
    if "%" not in s:
        return s
    s = s.replace("%%", _READING_BATCH_EXPORT_ESCAPE_SENTINEL)

    def _repl(match):
        return _resolve_reading_batch_export_placeholder(match.group(1), ctx, row_index)

    out = _READING_BATCH_EXPORT_PLACEHOLDER_RE.sub(_repl, s)
    return out.replace(_READING_BATCH_EXPORT_ESCAPE_SENTINEL, "%")


def _primary_contract_for_supply_point(supply_point, active_contract_token):
    if not supply_point:
        return None
    contracts = getattr(supply_point, "contracts", None)
    if contracts is None:
        return None
    qs = contracts.filter(is_active=True)
    if active_contract_token:
        c = qs.filter(status__token=active_contract_token).first()
        if c:
            return c
    return qs.first()


def _primary_contract_token_for_supply_point(supply_point, active_contract_token):
    c = _primary_contract_for_supply_point(supply_point, active_contract_token)
    return (c.token or "").strip() if c and c.token else ""


def _primary_contact_for_contract(contract):
    if not contract:
        return None
    return (
        contract.contacts.filter(is_active=True).order_by("-is_default", "id").first()
    )


def _route_model_from_supply_point(supply_point):
    if not supply_point or not supply_point.property_id or not supply_point.property:
        return None
    rp = supply_point.property.route_position
    if not rp or not rp.route_id:
        return None
    return rp.route


def _route_position_from_supply_point(supply_point):
    if not supply_point or not supply_point.property_id or not supply_point.property:
        return None
    return supply_point.property.route_position


def _reading_batch_export_context_dict(
    supply_point,
    meter,
    supply_address,
    route_label,
    position_label,
    route_model,
    route_position,
    active_contract_token,
):
    contract = _primary_contract_for_supply_point(supply_point, active_contract_token)
    contract_token = (contract.token or "").strip() if contract and contract.token else ""
    contact = _primary_contact_for_contract(contract)
    return {
        "supply_point": supply_point,
        "meter": meter,
        "supply_address": supply_address,
        "route_label": route_label,
        "position_label": position_label,
        "contract_token": contract_token,
        "contract": contract,
        "route": route_model,
        "route_position": route_position,
        "contact": contact,
    }


def _supply_point_address_str(supply_point):
    if not supply_point or not supply_point.address_id or not supply_point.address:
        return ""
    return str(supply_point.address).strip()


def _route_position_labels_from_sp(supply_point):
    """Route/position from property.route_position (not necessarily same as batch route)."""
    if not supply_point or not supply_point.property_id or not supply_point.property:
        return "", ""
    rp = supply_point.property.route_position
    if not rp:
        return "", ""
    route_label = ""
    if rp.route_id and rp.route:
        r = rp.route
        route_label = (r.token or r.name or str(r.id)).strip()
    if rp.position is not None:
        position_label = str(rp.position)
    elif rp.token:
        position_label = str(rp.token).strip()
    else:
        position_label = ""
    return route_label, position_label


def _meter_matches_reading_batch_filters(reading_batch, meter):
    """
    Telecontrol / manual inclusion: same rules as ReadingBatchSetupViewSet route supply points
    (not include_telecontrol → manual only; elif not include_manual → telecontrol only; else both).
    """
    if not meter:
        return False
    if not reading_batch.include_telecontrol:
        return not meter.has_remote_reading
    if not reading_batch.include_manual:
        return meter.has_remote_reading
    return True


def _positions_qs_for_reading_batch_route(route, reading_batch):
    if not reading_batch.include_telecontrol:
        qs = route.positions.filter(
            Q(properties__supply_points__meter__has_remote_reading=False) | Q(properties__supply_points__meter__force_manual_reading=True)
        ).distinct()
    elif not reading_batch.include_manual:
        qs = route.positions.filter(
            Q(properties__supply_points__meter__has_remote_reading=True, properties__supply_points__meter__force_manual_reading=False)
        ).distinct()
    else:
        qs = route.positions.all()
    return qs.order_by("position", "id")


def _supply_points_qs_for_reading_batch_property(prop, reading_batch):
    qs = prop.supply_points.select_related("cluster_nozzle")
    if not reading_batch.include_telecontrol:
        qs = qs.filter(Q(meter__has_remote_reading=False) | Q(meter__force_manual_reading=True)).distinct()
    elif not reading_batch.include_manual:
        qs = qs.filter(meter__has_remote_reading=True, meter__force_manual_reading=False).distinct()
    return qs.order_by(F("cluster_nozzle__position").asc(nulls_first=True), "id")


def _sortable_int(value, nulls_last=True):
    """Numeric sort with None first or last (so 1, 20, 123 stay in numeric order)."""
    if value is None:
        return (nulls_last, 0)
    return (not nulls_last, value)


def _reading_batch_export_context_sort_key(ctx):
    """
    Walk order: route, then supply.property.route_position.position,
    then cluster_nozzle.position (supplies without a nozzle first).
    """
    route = ctx.get("route")
    rp = ctx.get("route_position")
    sp = ctx.get("supply_point")
    cn = getattr(sp, "cluster_nozzle", None) if sp is not None else None
    cn_pos = getattr(cn, "position", None) if cn is not None else None
    return (
        _sortable_int(getattr(route, "position", None) if route is not None else None),
        getattr(route, "id", 0) or 0 if route is not None else 0,
        _sortable_int(getattr(rp, "position", None) if rp is not None else None),
        _sortable_int(cn_pos, nulls_last=False),
        getattr(sp, "id", 0) or 0 if sp is not None else 0,
    )


def _reading_batch_export_row_contexts(reading_batch):
    """
    Same scope as the default export: routes -> fix_meters -> all active supply points.
    Supply points/meters respect include_telecontrol / include_manual like ReadingBatchSetupViewSet.
    Each context has supply_point, meter, supply_address, route_label, position_label,
    contract_token, contract, route (Route model when available), route_position (RoutePosition),
    contact (primary PersonContact). Rows are ordered by route, then
    supply.property.route_position.position, then cluster_nozzle.position.
    """
    active_contract_token = ConfigProject.objects.get(token="contract_active_token").value
    token_meter_status_no_meter = ConfigProject.objects.get(token="token_meter_status_no_meter").value
    activate_token = ConfigProject.objects.get(token="supply_point_status_activate_token").value

    contexts = []

    if reading_batch.routes.exists():
        routes = (
            reading_batch.routes.all()
            .prefetch_related(
                Prefetch(
                    "positions__properties__supply_points",
                    queryset=SupplyPoint.objects.all()
                    .select_related(
                        "meter",
                        "meter__status",
                        "address",
                        "address__street",
                        "address__street_number",
                        "address__city",
                        "cluster_nozzle",
                        "property__route_position",
                        "property__route_position__route",
                    )
                    .prefetch_related(
                        "contracts",
                        "contracts__status",
                        "contracts__contacts",
                        "contracts__contacts__person",
                    ),
                ),
            )
            .select_related("route_zone")
        )

        def _position_has_any_listable_meter(position):
            """At least one supply point meter listable for export (status + tele/manual flags)."""
            for prop in position.properties.all():
                for sp in _supply_points_qs_for_reading_batch_property(prop, reading_batch):
                    if not sp.meter or not sp.meter.status:
                        continue
                    if sp.meter.status.token == token_meter_status_no_meter:
                        continue
                    if _meter_matches_reading_batch_filters(reading_batch, sp.meter):
                        return True
            return False

        for route in routes:
            route_label = (route.token or route.name or str(route.id)).strip()
            positions_qs = _positions_qs_for_reading_batch_route(route, reading_batch)
            positions_list = list(
                positions_qs.prefetch_related(
                    "properties__supply_points",
                    "properties__supply_points__meter",
                    "properties__supply_points__meter__status",
                    "properties__supply_points__cluster_nozzle",
                )
            )
            for position in positions_list:
                if not _position_has_any_listable_meter(position):
                    continue
                if position.position is not None:
                    position_label = str(position.position)
                elif position.token:
                    position_label = str(position.token).strip()
                else:
                    position_label = ""
                for prop in position.properties.all():
                    sps = _supply_points_qs_for_reading_batch_property(prop, reading_batch)
                    for sp in sps:
                        if (
                            not sp.meter
                            or not sp.meter.status
                            or sp.meter.status.token == token_meter_status_no_meter
                            or not _meter_matches_reading_batch_filters(reading_batch, sp.meter)
                        ):
                            continue
                        supply_address = _supply_point_address_str(sp)
                        contexts.append(
                            _reading_batch_export_context_dict(
                                sp,
                                sp.meter,
                                supply_address,
                                route_label,
                                position_label,
                                route,
                                position,
                                active_contract_token,
                            )
                        )

    elif reading_batch.fix_meters.exists():
        fix_meters = reading_batch.fix_meters.prefetch_related(
            Prefetch(
                "supply_points",
                queryset=SupplyPoint.objects.filter(is_active=True)
                .select_related(
                    "address",
                    "address__street",
                    "address__street_number",
                    "address__city",
                    "cluster_nozzle",
                    "property__route_position",
                    "property__route_position__route",
                    "meter",
                    "meter__status",
                )
                .prefetch_related(
                    "contracts",
                    "contracts__status",
                    "contracts__contacts",
                    "contracts__contacts__person",
                ),
            ),
        )
        for meter in fix_meters:
            if not _meter_matches_reading_batch_filters(reading_batch, meter):
                continue
            sp = meter.supply_points.first()
            supply_address = _supply_point_address_str(sp) if sp else ""
            route_label, position_label = _route_position_labels_from_sp(sp) if sp else ("", "")
            route_model = _route_model_from_supply_point(sp) if sp else None
            route_position = _route_position_from_supply_point(sp) if sp else None
            contexts.append(
                _reading_batch_export_context_dict(
                    sp,
                    meter,
                    supply_address,
                    route_label,
                    position_label,
                    route_model,
                    route_position,
                    active_contract_token,
                )
            )

    else:
        sp_base = SupplyPoint.objects.filter(is_active=True, status__token=activate_token)
        if not reading_batch.include_telecontrol:
            sp_base = sp_base.filter(Q(meter__has_remote_reading=False) | Q(meter__force_manual_reading=True))
        elif not reading_batch.include_manual:
            sp_base = sp_base.filter(meter__has_remote_reading=True, meter__force_manual_reading=False)
        for sp in (
            sp_base.select_related(
                "meter",
                "meter__status",
                "address",
                "address__street",
                "address__street_number",
                "address__city",
                "cluster_nozzle",
                "property__route_position",
                "property__route_position__route",
            )
            .prefetch_related(
                "contracts",
                "contracts__status",
                "contracts__contacts",
                "contracts__contacts__person",
            )
            .order_by("id")
            .iterator(chunk_size=2000)
        ):
            if (
                not sp.meter
                or not sp.meter.status
                or sp.meter.status.token == token_meter_status_no_meter
                or not _meter_matches_reading_batch_filters(reading_batch, sp.meter)
            ):
                continue
            supply_address = _supply_point_address_str(sp)
            route_label, position_label = _route_position_labels_from_sp(sp)
            route_model = _route_model_from_supply_point(sp)
            route_position = _route_position_from_supply_point(sp)
            contexts.append(
                _reading_batch_export_context_dict(
                    sp,
                    sp.meter,
                    supply_address,
                    route_label,
                    position_label,
                    route_model,
                    route_position,
                    active_contract_token,
                )
            )

    contexts.sort(key=_reading_batch_export_context_sort_key)
    return contexts


def _reading_batch_export_rows(reading_batch):
    """
    Default table: meter code, supply address, route, position, contract token.
    """
    headers = _reading_batch_export_headers()
    contexts = _reading_batch_export_row_contexts(reading_batch)
    rows = [
        _export_meter_row(
            c["meter"],
            c["supply_address"],
            c["route_label"],
            c["position_label"],
            c["contract_token"],
        )
        for c in contexts
    ]
    return headers, rows


def _reading_batch_export_rows_configured(reading_batch):
    columns = list(ReadingBatchExportColumn.objects.order_by("position", "id"))
    headers = [(c.name or "").strip() for c in columns]
    contexts = _reading_batch_export_row_contexts(reading_batch)
    rows = [
        [_reading_batch_export_expression(col.value, ctx, idx) for col in columns]
        for idx, ctx in enumerate(contexts, start=1)
    ]
    return headers, rows


@shared_task(bind=True)
def export_reading_batch_task(self, reading_batch_id, file_format="csv"):
    reading_batch = ReadingBatch.objects.filter(id=reading_batch_id, is_active=True).first()
    if not reading_batch:
        raise ExportReadingBatchTaskError("Reading batch not found")

    fmt = (file_format or "csv").strip().lower()
    if fmt not in ("csv", "xls", "xlsx"):
        raise ExportReadingBatchTaskError("Invalid format; use csv, xls, or xlsx")

    if ReadingBatchExportColumn.objects.exists():
        headers, rows = _reading_batch_export_rows_configured(reading_batch)
    else:
        headers, rows = _reading_batch_export_rows(reading_batch)
    if not rows:
        raise ExportReadingBatchTaskError(
            "No rows to export (empty routes / fix meters filtered out, or no active supply points)."
        )
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

    if fmt == "csv":
        buffer = io.StringIO()
        writer = csv.writer(buffer, delimiter=";")
        writer.writerow(headers)
        writer.writerows(rows)
        csv_text = "\ufeff" + buffer.getvalue()
        file_bytes = csv_text.encode("utf-8")
        ext = "csv"
    elif fmt == "xlsx":
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = str(_("readings"))
        ws.append(headers)
        for row in rows:
            ws.append(row)
        buf = BytesIO()
        wb.save(buf)
        file_bytes = buf.getvalue()
        ext = "xlsx"
    else:
        # Legacy .xls: HTML table (Excel opens this without xlwt).
        parts = [
            "<html><head><meta charset=\"utf-8\"></head><body><table>",
            "<tr>" + "".join(f"<td>{escape(str(h))}</td>" for h in headers) + "</tr>",
        ]
        for row in rows:
            parts.append(
                "<tr>" + "".join(f"<td>{escape(str(c))}</td>" for c in row) + "</tr>"
            )
        parts.append("</table></body></html>")
        file_bytes = ("\ufeff" + "\n".join(parts)).encode("utf-8")
        ext = "xls"

    filename = f"readings_batch_{reading_batch_id}_{timestamp}.{ext}"
    service = settings.DOCUMENT_MANAGER_SERVICES.get("billing", "hdd")
    document = upload_document(
        file=ContentFile(file_bytes, name=filename),
        entity="READING_BATCH",
        field="EXPORT",
        entity_id=reading_batch_id,
        entity_token="READING_BATCH_EXPORT",
        folder="",
        service=service,
        document_name=filename,
        date=datetime.now(),
    )

    return {
        "status": "ok",
        "message": "Reading batch export generated successfully",
        "format": fmt,
        "document_id": document.id,
        "document_name": filename,
    }


READING_EXPORT_COLUMNS = {
    "meter": ("meter", lambda r: r.meter.code if r.meter else ""),
    "reading_date": ("reading_date", lambda r: r.reading_date),
    "reading_value": ("reading_value", lambda r: r.reading_value),
    "calculated_value": ("calculated_value", lambda r: int(r.calculated_value) if r.calculated_value is not None else None),
    "origin": ("origin", lambda r: r.origin),
    "reading_id": ("reading_id", lambda r: r.id),
}
READING_EXPORT_DEFAULT_COLUMNS = [
    "meter", "reading_date", "reading_value", "calculated_value", "origin", "reading_id",
]


@shared_task
def export_readings_task(reading_ids, columns=None):
    try:
        from billing.models import Reading
        from documentmanager.utils.generic_export_service import resolve_columns
        import openpyxl

        # select_related("meter"): la columna `meter` feia una consulta per lectura.
        readings = Reading.objects.filter(id__in=reading_ids or []).select_related("meter")
        resolved_columns = resolve_columns(READING_EXPORT_COLUMNS, READING_EXPORT_DEFAULT_COLUMNS, columns)

        # write_only: no manté totes les cel·les a memòria (exportacions de desenes de milers de files).
        wb = openpyxl.Workbook(write_only=True)
        sheet = wb.create_sheet(title=str(_("Readings")))
        sheet.append([header for _, header, _ in resolved_columns])
        for reading in readings.iterator(chunk_size=2000):
            sheet.append([getter(reading) for _, _, getter in resolved_columns])

        buffer = BytesIO()
        wb.save(buffer)
        file_bytes = buffer.getvalue()

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"readings_{timestamp}.xlsx"
        service = settings.DOCUMENT_MANAGER_SERVICES.get("billing", "hdd")
        document = upload_document(
            file=ContentFile(file_bytes, name=filename),
            entity="READING",
            field="EXPORT",
            entity_id=0,
            entity_token="READING_EXPORT",
            folder="",
            service=service,
            document_name=filename,
            date=datetime.now(),
        )

        return {
            "status": "ok",
            "message": "Readings export generated successfully",
            "document_id": document.id,
            "document_name": filename,
        }
    except Exception as e:
        return {
            "status": "error",
            "message": f"Error generating readings export: {str(e)}",
        }


def _reading_document_cell_to_str(cell_value):
    if cell_value is None:
        return ""
    if isinstance(cell_value, datetime):
        # Try multiple date formats for string output. If all fail, fallback to default ISO.
        for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y"):
            try:
                return cell_value.strftime(fmt)
            except Exception:
                continue
        # Fallback: use ISO format
        return cell_value.isoformat()
    if isinstance(cell_value, date):
        for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y"):
            try:
                return cell_value.strftime(fmt)
            except Exception:
                continue
        # Fallback: use ISO format
        return cell_value.isoformat()
    return str(cell_value)


def _detect_csv_delimiter(sample):
    try:
        sniffer = csv.Sniffer()
        return sniffer.sniff(sample).delimiter
    except Exception:
        pass

    first_line = sample.splitlines()[0] if sample else ""
    best_delimiter = ","
    best_count = 0
    for delimiter in (";", ",", "\t"):
        count = first_line.count(delimiter)
        if count > best_count:
            best_count = count
            best_delimiter = delimiter
    return best_delimiter


def _detect_reading_document_csv_encoding(raw_bytes):
    """Prefer UTF-8 (with optional BOM); fall back to Windows/ISO Latin encodings."""
    for encoding in ("utf-8-sig", "cp1252", "iso-8859-1"):
        try:
            raw_bytes.decode(encoding)
            return encoding
        except UnicodeDecodeError:
            continue
    return "iso-8859-1"


def _iter_reading_document_csv_rows(file_obj):
    raw_bytes = file_obj.read()
    encoding = _detect_reading_document_csv_encoding(raw_bytes)
    text = raw_bytes.decode(encoding)
    delimiter = _detect_csv_delimiter(text[:4096])
    yield from csv.DictReader(StringIO(text), delimiter=delimiter)


def _iter_reading_document_xlsx_rows(file_obj):
    import openpyxl

    workbook = openpyxl.load_workbook(file_obj, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        headers = None
        for row in worksheet.iter_rows(values_only=True):
            if headers is None:
                headers = [_reading_document_cell_to_str(header) for header in row]
                continue
            if not any(cell is not None for cell in row):
                continue
            row_dict = {}
            for index, cell_value in enumerate(row):
                if index < len(headers) and headers[index]:
                    row_dict[headers[index]] = _reading_document_cell_to_str(cell_value)
            yield row_dict
    finally:
        workbook.close()


def _is_xlsx_upload(file_obj, filename):
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        return True
    if name.endswith(".csv"):
        return False
    position = file_obj.tell()
    signature = file_obj.read(4)
    file_obj.seek(position)
    return signature[:2] == b"PK"


def iter_reading_document_rows(uploaded_file):
    with uploaded_file.open("rb") as file_obj:
        if _is_xlsx_upload(file_obj, uploaded_file.name):
            yield from _iter_reading_document_xlsx_rows(file_obj)
        else:
            yield from _iter_reading_document_csv_rows(file_obj)


_READING_DOCUMENT_DATE_ONLY_FORMATS = (
    "%d-%m-%Y",
    "%d/%m/%Y",
    "%Y-%m-%d",
    "%d-%m-%y",
    "%d/%m/%y",
)

_READING_DOCUMENT_DATETIME_FORMATS = (
    "%d/%m/%y %H:%M",
    "%d/%m/%Y %H:%M",
    "%d-%m-%y %H:%M",
    "%d-%m-%Y %H:%M",
    "%d/%m/%y %H:%M:%S",
    "%d/%m/%Y %H:%M:%S",
    "%d-%m-%y %H:%M:%S",
    "%d-%m-%Y %H:%M:%S",
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%dT%H:%M:%S",
    "%Y-%m-%d %H:%M",
)


def parse_reading_document_date(value):
    """Parse reading dates from import files (many formats; day-first)."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())

    reading_date_str = str(value).strip()
    if not reading_date_str:
        return None

    for fmt in _READING_DOCUMENT_DATETIME_FORMATS + _READING_DOCUMENT_DATE_ONLY_FORMATS:
        try:
            return datetime.strptime(reading_date_str, fmt)
        except ValueError:
            continue

    # Legacy behaviour: ISO / long strings often worked with the date part only.
    if len(reading_date_str) > 10:
        short = reading_date_str[:10]
        for fmt in _READING_DOCUMENT_DATE_ONLY_FORMATS:
            try:
                return datetime.strptime(short, fmt)
            except ValueError:
                continue

    try:
        from dateutil import parser as dateutil_parser

        return dateutil_parser.parse(reading_date_str, dayfirst=True)
    except (ValueError, TypeError, OverflowError):
        return None


def _parse_reading_document_numeric_value(raw_value):
    return int(float(str(raw_value or "0").replace(".", "").replace(",", "")))


def _reading_value_from_liters(raw_value):
    """Liters template: ignore last 3 digits (e.g. 11999 -> 11, 77 -> 0)."""
    return _parse_reading_document_numeric_value(raw_value) // 1000


def _clear_stale_reading_document_processing_lock(document):
    """If status is processing but Celery already finished/failed, unlock the document."""
    if document.status != ReadingDocument.STATUS_PROCESSING:
        return document
    if not document.task_id:
        document.status = ReadingDocument.STATUS_PENDING
        document.save(update_fields=['status'])
        return document

    from celery.result import AsyncResult

    task_result = AsyncResult(document.task_id)
    if task_result.state in ['PENDING', 'STARTED', 'PROGRESS', 'RETRY']:
        return document

    if task_result.state == 'SUCCESS':
        document.status = ReadingDocument.STATUS_PROCESSED
        if not document.processed_at:
            document.processed_at = timezone.now()
        document.save(update_fields=['status', 'processed_at'])
    else:
        document.status = (
            ReadingDocument.STATUS_FAILED
            if task_result.state == 'FAILURE'
            else ReadingDocument.STATUS_PENDING
        )
        document.save(update_fields=['status'])
    return document


@shared_task
def validate_reading_document(reading_document_id, params):
    """
    Build/slice reading-document preview. Return value is the validate response payload
    (available via AsyncResult once the task succeeds).
    """
    from billing.utils.reading_document_import_service import (
        ReadingDocumentPreviewValidationError,
        get_or_build_reading_document_preview,
        slice_reading_document_preview,
    )

    try:
        document = ReadingDocument.objects.get(id=reading_document_id)
    except ReadingDocument.DoesNotExist:
        return {"error": "Document no trobat."}

    if not document.file:
        return {"error": "El document no té fitxer."}

    params = params or {}
    try:
        full_preview, built = get_or_build_reading_document_preview(
            document,
            refresh=bool(params.get("refresh")),
        )
        response_payload = slice_reading_document_preview(
            full_preview,
            action=params.get("action"),
            reason=params.get("reason"),
            offset=params.get("offset", 0),
            max_rows=params.get("max_rows", 25),
            search=params.get("search"),
            ordering=params.get("ordering") or "row_index",
        )
    except ReadingDocumentPreviewValidationError as exc:
        return {"error": str(exc)}
    except ValueError as exc:
        return {"error": str(exc)}

    if built:
        document.last_preview = full_preview
        update_fields = ['last_preview']
        document = _clear_stale_reading_document_processing_lock(document)
        if document.status != ReadingDocument.STATUS_PROCESSING:
            document.status = ReadingDocument.STATUS_PENDING
            update_fields.append('status')
        document.save(update_fields=update_fields)

    return response_payload


@shared_task
def process_reading_document_file(reading_document_id, reprocess=False):
    try:
        document = ReadingDocument.objects.get(id=reading_document_id)
    except ReadingDocument.DoesNotExist:
        return

    uploaded_file = document.file
    if not uploaded_file:
        return

    if reprocess:
        document.not_found_readings.all().delete()

    document.status = ReadingDocument.STATUS_PROCESSING
    document.save(update_fields=['status'])

    try:
        return _process_reading_document_file(document, uploaded_file, reading_document_id)
    except Exception:
        document.status = ReadingDocument.STATUS_FAILED
        document.save(update_fields=['status'])
        raise


def _mapped_reading_document_cell(row, labels, label_key, default=None):
    """Ignore empty template mappings: CSV DictReader uses '' for untitled columns."""
    label = labels.get(label_key)
    if label is None or not str(label).strip():
        return default
    value = row.get(label, default)
    if value is None:
        return default
    if isinstance(value, str) and value.strip() == "" and default is not None:
        return default
    return value


def _process_reading_document_file(document, uploaded_file, reading_document_id):

    active_contract_status_token = ConfigProject.objects.get(token="contract_active_token").value
    reading_alert_unusual_consumption_token = ConfigProject.objects.get(token="reading_alert_unusual_consumption").value
    type_final = ConfigProject.objects.get(token='invoice_type_invoice_token').value
    invoice_status_cancelled = ConfigProject.objects.get(token='invoice_status_cancelled_token').value
    invoice_status_pending = ConfigProject.objects.get(token='invoice_status_pending_token').value
    reading_alert_unusual_consumption = ReadingAlert.objects.get(token=reading_alert_unusual_consumption_token)
    not_found_meter_tokens = []
    not_found_readings = []
    last_reading_close = []
    errors = []

    labels = {
      "meter_label" : "meter",
      "observation_label" : "observation",
      "reading_value_label" : "reading_value",
      "reading_label" : "reading",
      "reading_date_label" : "reading_date",
      "origin_label" : "origin",
      "leak_value_label" : "leak_value",
      "is_control_label" : "is_control",
      "contract_label" : "contract",
      "comm_module_label" : "comm_module"
    }
    if document.template:
      for column in document.template.columns.all():
        labels[column.original_name+"_label"] = (column.mapped_name or "").strip()

    reading_alert_negative_token = ConfigProject.objects.get(token='reading_alert_negative').value
    reading_alert_negative = ReadingAlert.objects.get(token=reading_alert_negative_token)
    template_is_liters = bool(document.template and document.template.is_liters)

    row_index = 0
    stats = {
        "rows": 0,
        "skipped_no_date": 0,
        "not_found_meter": 0,
        "no_supply_points": 0,
        "no_contracts": 0,
        "skipped_existing": 0,
        "created": 0,
        "updated": 0,
    }

    print(
        f"[reading-doc] start document_id={reading_document_id} "
        f"file={uploaded_file.name} template={document.template.name if document.template else None} "
        f"is_liters={template_is_liters} labels={labels}"
    )

    for row in iter_reading_document_rows(uploaded_file):
        row_index += 1
        stats["rows"] += 1
        meter_token = _mapped_reading_document_cell(
            row, labels, "meter_label",
            row.get("\ufeffmeter", row.get("meter_id", row.get("\ufeffmeter_id"))),
        )
        
        observation = _mapped_reading_document_cell(row, labels, "observation_label", "")

        reading_date_str = _mapped_reading_document_cell(
            row, labels, "reading_date_label", row.get("timestamp")
        )
        reading_date = parse_reading_document_date(reading_date_str)
        if not reading_date:
            stats["skipped_no_date"] += 1
            print(
                f"[reading-doc] row {row_index}: SKIP no reading_date "
                f"(raw={reading_date_str!r}, label={labels['reading_date_label']!r}, "
                f"row_keys={list(row.keys())[:6]})"
            )
            continue

        meter = None
        if meter_token:
            meter = Meter.objects.filter(code=meter_token).first()
            if not meter:
                meter = Meter.objects.filter(token=meter_token).first()
            if not meter:
                meter = Meter.objects.filter(code=meter_token[:-1]).first()
            if not meter:
                meter = Meter.objects.filter(code=meter_token[1:]).first()
            if not meter:
                meter = Meter.objects.filter(code=meter_token[1:-1]).first()
            if not meter:
                for length in [8, 9, 10]:
                    padded_token = meter_token.zfill(length)
                    meter = Meter.objects.filter(code=padded_token).first()
                    # ??
                    """ if meter:
                        break """

        comm_module = _mapped_reading_document_cell(row, labels, "comm_module_label")
        if not meter:
            if comm_module and str(comm_module).strip():
                meter = (
                    Meter.objects.filter(comm_module=str(comm_module).strip())
                    .exclude(Q(comm_module__isnull=True) | Q(comm_module=""))
                    .first()
                )

        # print(f"meter: {meter.code if meter else 'None'}")
        contract_mapped = bool(str(labels.get("contract_label") or "").strip())
        if not meter and contract_mapped and labels["contract_label"] in row:
            contract_token = row[labels["contract_label"]]
            # print(f"contract_token: {contract_token}")
            contract = Contract.objects.filter(token=contract_token).first()
            if contract:
                meter = contract.supply_point_default.meter if contract.supply_point_default else None

        # print(f"meter: {meter.code if meter else 'None'}")

        if not meter:
            meter_identifier = meter_token or comm_module or ""
            stats["not_found_meter"] += 1
            print(
                f"[reading-doc] row {row_index}: NOT_FOUND meter "
                f"(meter_token={meter_token!r}, comm_module={comm_module!r}, "
                f"label={labels['comm_module_label']!r})"
            )
            not_found_meter_tokens.append(meter_identifier)
            try:
                raw_not_found_value = _mapped_reading_document_cell(
                    row, labels, "reading_value_label", "0"
                )
                if template_is_liters:
                    not_found_reading_value = _reading_value_from_liters(raw_not_found_value)
                else:
                    not_found_reading_value = _parse_reading_document_numeric_value(raw_not_found_value)
            except (ValueError, TypeError):
                not_found_reading_value = 0
            not_found_readings.append(ReadingDocumentNotFound(
                token=f"{meter_identifier}/DOC{reading_date.strftime('%y%m%d')}",
                meter_code=meter_identifier,
                reading_value=not_found_reading_value,
                reading_date=reading_date.date(),
                origin=_mapped_reading_document_cell(row, labels, "origin_label"),
                is_control=_coerce_to_bool(_mapped_reading_document_cell(row, labels, "is_control_label", "")),
                leak_value=int(float(_mapped_reading_document_cell(row, labels, "leak_value_label", 0) or 0)),
                observation=observation,
                reading_document=document,
                ))
            continue

        is_control_value = _mapped_reading_document_cell(row, labels, "is_control_label", "")
        is_control = _coerce_to_bool(is_control_value)

        try:
            raw_reading_value = _mapped_reading_document_cell(row, labels, "reading_value_label", "0")
            raw_reading = _mapped_reading_document_cell(
                row, labels, "reading_label", row.get("reading", "0")
            )
            if template_is_liters:
                reading_value_str = _reading_value_from_liters(raw_reading_value)
                reading_str = _reading_value_from_liters(raw_reading)
            else:
                reading_value_str = _parse_reading_document_numeric_value(raw_reading_value)
                reading_str = _parse_reading_document_numeric_value(raw_reading)
            reading_value = max(reading_value_str, reading_str)
        except (ValueError, TypeError):
            print(f"Error converting reading_value: {_mapped_reading_document_cell(row, labels, 'reading_value_label')}")
            reading_value = 0
        origin = str(_mapped_reading_document_cell(row, labels, "origin_label", "TELECONTROL") or "TELECONTROL").upper()
        if origin.lower() == "estimada":
          origin = "MANUAL"

        try:
            leak_value = int(float(_mapped_reading_document_cell(row, labels, "leak_value_label", 0) or 0))
        except (ValueError, TypeError):
            leak_value = 0

        print(
            f"[reading-doc] row {row_index}: meter={meter.code} "
            f"comm_module={comm_module!r} date={reading_date.date()} "
            f"value={_mapped_reading_document_cell(row, labels, 'reading_value_label')!r}"
        )

        supply_points = list(meter.supply_points.all())
        if not supply_points:
            stats["no_supply_points"] += 1
            print(f"[reading-doc] row {row_index}: meter {meter.code} has no supply_points")
            continue

        # One reading per CSV row: only the active contract (ConfigProject contract_active_token).
        # Prefer SP where the meter contract has supply_point_default; never attach to terminated.
        target = None
        for supply_point in supply_points:
            contract = get_reading_document_target_contract(
                supply_point, active_contract_status_token
            )
            if contract:
                target = (supply_point, contract)
                if contract.supply_point_default_id == supply_point.id:
                    break

        if not target:
            stats["no_contracts"] += 1
            print(
                f"[reading-doc] row {row_index}: meter {meter.code} "
                f"no active contract (token={active_contract_status_token!r})"
            )
            continue

        supply_point, contract = target
        tele_obs = None
        if observation:
            try:
                tele_obs = RemoteReadingAlert.objects.get(token=observation)
            except RemoteReadingAlert.DoesNotExist:
                tele_obs = None

        base_date = reading_date.date()
        previous_reading = (
            Reading.objects.filter(
                supply_point=supply_point,
                reading_date__lt=base_date,
                meter=meter,
                contract=contract,
                is_control=False,
            )
            .order_by("-reading_date", "-reading_value")
            .first()
        )

        existing_reading = Reading.objects.filter(
                supply_point=supply_point,
                reading_date=base_date,
                meter=meter,
                contract=contract,
                is_control=is_control,
                is_initial=False,
            ).first()

        allow_by_period = check_billing_period(base_date, supply_point, meter, contract, previous_reading)
        prev_reading_invoice = False
        if previous_reading:
          prev_reading_invoice = previous_reading.invoices.filter(type_final=type_final).exclude(status__token__in=[invoice_status_cancelled, invoice_status_pending]).exists()

        if existing_reading:
            stats["skipped_existing"] += 1
            print(
                f"[reading-doc] row {row_index}: SKIP existing reading "
                f"contract={contract.token} date={base_date}"
            )
            continue

        try:
          calculated_value = (
              reading_value - previous_reading.reading_value if previous_reading else reading_value
          )
        except Exception as e:
            calculated_value = 0
            errors.append(f"Error calculating calculated_value for contract: {contract.token} - {e}")
        is_negative_reading = calculated_value < 0

        consumption_days = (
            (base_date - previous_reading.reading_date).days if previous_reading else 0
        )
        last_month = timezone.now().month
        if previous_reading and previous_reading.reading_date.month >= last_month:
            last_reading_close.append(previous_reading)

        if allow_by_period or prev_reading_invoice:
            stats["created"] += 1
            print(
                f"[reading-doc] row {row_index}: CREATE reading "
                f"contract={contract.token} status={getattr(contract.status, 'token', None)} value={reading_value}"
            )
            new_reading = Reading.objects.create(
                token=f"{meter.code}/DOC{reading_date.strftime('%y%m%d')}",
                document=document,
                supply_point=supply_point,
                reading_date=reading_date,
                reading_value=reading_value,
                meter=meter,
                contract=contract,
                origin=origin,
                is_control=is_control,
                calculated_value=calculated_value if previous_reading else 0,
                previous_reading=previous_reading if not is_control else None,
                consumption_days=consumption_days if not is_control else 0,
                leak_value=leak_value,
                remote_alert=tele_obs,
            )
            if is_negative_reading:
                new_reading.alert = reading_alert_negative
                new_reading.alert_notes = reading_alert_negative.name
                new_reading.save(update_fields=['alert', 'alert_notes'])
        else:
            if not previous_reading or previous_reading.reading_date > base_date:
                continue

            stats["updated"] += 1
            print(
                f"[reading-doc] row {row_index}: UPDATE reading "
                f"contract={contract.token} status={getattr(contract.status, 'token', None)} value={reading_value}"
            )
            new_reading = previous_reading
            new_reading.reading_value = reading_value
            new_reading.reading_date = base_date
            new_reading.contract = contract
            new_reading.calculated_value = new_reading.calculated_value + calculated_value if not is_control else 0
            new_reading.real_consumption = new_reading.calculated_value + calculated_value if not is_control else 0
            new_reading.consumption_days = new_reading.consumption_days + consumption_days if new_reading.consumption_days else consumption_days if not is_control else 0
            new_reading.leak_value = leak_value
            new_reading.remote_alert = tele_obs
            new_reading.document = document
            new_reading.save()
            if is_negative_reading:
                new_reading.alert = reading_alert_negative
                new_reading.alert_notes = reading_alert_negative.name
                new_reading.save()

        if not new_reading.is_control:
            try:
                estimated_bag = EstimatedBag.objects.get(
                    supply_point=new_reading.supply_point, contract=new_reading.contract
                )
            except EstimatedBag.DoesNotExist:
                estimated_bag = None

            if estimated_bag:
                calculate_estimated_bag(
                    new_reading.calculated_value,
                    new_reading.is_estimated,
                    new_reading,
                    estimated_bag,
                )
                if new_reading.estimated_used:
                    new_reading.real_consumption = new_reading.calculated_value - new_reading.estimated_used
                    new_reading.save()

        warning = check_unusual_consumption(new_reading, new_reading.real_consumption if new_reading.real_consumption else calculated_value)
        if warning:
            alert = reading_alert_unusual_consumption
            new_reading.alert = alert
            new_reading.alert_notes = alert.name
            new_reading.save()

                            
    print(f"[reading-doc] summary document_id={reading_document_id} stats={stats}")
    print(f"\nNot found meter tokens: {not_found_meter_tokens}")
    for meter_token in not_found_meter_tokens:
        print(f"Meter not found for token: {meter_token}")
    try:
      ReadingDocumentNotFound.objects.bulk_create(not_found_readings)
    except Exception as e:
      print(f"Error creating reading document not found: {e}")
        
    print("Last reading close: ", last_reading_close)
    for reading in last_reading_close:
        print(f"Reading: {reading.reading_date} - {reading.reading_value} - {reading.contract.token}")
      
    for error in errors:  
        print(f"Error: {error}")

    document.status = ReadingDocument.STATUS_PROCESSED
    document.processed_at = timezone.now()
    document.save(update_fields=['status', 'processed_at'])
    return stats

def _coerce_to_bool(value):
  if isinstance(value, bool):
    return value
  if isinstance(value, str):
    return value.strip().lower() in ['true', '1', 'yes', 'on']
  return bool(value)


@shared_task(bind=True)
def process_reading_batch(self, reading_batch_id):
  progress_recorder = ProgressRecorder(self)

  batch = ReadingBatch.objects.get(id = reading_batch_id)
  reading_files = []
  readings = []
  
  total_supplypoints = batch.routes.annotate(num_supplypoints=Count('positions__properties__supply_points')).aggregate(total=Sum('num_supplypoints'))['total']
  print("PROCESSING BATCH: "+str(reading_batch_id))
  print("TOTAL SUPPLYPOINTS: "+str(total_supplypoints))

  # reading_alert_estimated
  reading_alert_estimated_token = ConfigProject.objects.get(token='reading_alert_estimated').value
  reading_alert_estimated = ReadingAlert.objects.get(token=reading_alert_estimated_token)
  # reading_alert_negative
  reading_alert_negative_token = ConfigProject.objects.get(token='reading_alert_negative').value
  reading_alert_negative = ReadingAlert.objects.get(token=reading_alert_negative_token)
  # reading_alert_zero
  reading_alert_zero_token = ConfigProject.objects.get(token='reading_alert_zero').value
  reading_alert_zero = ReadingAlert.objects.get(token=reading_alert_zero_token)
  # reading_alert_meter_cycle
  reading_alert_meter_cycle_token = ConfigProject.objects.get(token='reading_alert_meter_cycle').value
  reading_alert_meter_cycle = ReadingAlert.objects.get(token=reading_alert_meter_cycle_token)
  # reading_alert_unusual_consumption
  reading_alert_unusual_consumption_token = ConfigProject.objects.get(token='reading_alert_unusual_consumption').value
  reading_alert_unusual_consumption = ReadingAlert.objects.get(token=reading_alert_unusual_consumption_token)
  
  # reading_alert_low_consumption
  reading_alert_low_consumption_config = ConfigProject.objects.filter(token='reading_alert_low_consumption').first()
  reading_alert_low_consumption = ReadingAlert.objects.filter(token=reading_alert_low_consumption_config.value).first() if reading_alert_low_consumption_config and reading_alert_low_consumption_config.value else None
  
  min_consumption = ConfigProject.objects.filter(token='minimum_consumption').first()
  
  i = 0
  task_errors = []

  for reading in batch.readings.filter(
    reading_value__isnull=False,
    is_control=False,
    is_initial=False,
    is_active=True,
    copied_from__isnull=True
  ).filter(
    PENDING_READING_FILTER
  ).distinct():
    try:
      r = get_reading_minimal_object(reading, reading.origin, reading_alert_negative, reading_alert_zero, reading_alert_meter_cycle, reading_alert_unusual_consumption, reading_alert_estimated, batch, min_consumption, reading_alert_low_consumption)
      readings.append(r)
    except Exception as e:
      print(f'Error processing reading={reading.id}: {type(e).__name__}: {str(e)}')
      classification = classify_reading_processing_error(e, reading=reading)
      task_errors.append({
          'reading_id': reading.id,
          'error': str(e),
          'type': type(e).__name__,
          **classification,
      })

    progress_recorder.set_progress(i + 1, total_supplypoints or 0)
    i+=1

  counters = { }
  counters['total'] = batch.readings.filter(reading_value__isnull=False, copied_from__isnull=True).filter(PENDING_READING_FILTER).distinct().count()
  
  null_readings = batch.readings.filter(reading_value__isnull=True, copied_from__isnull=True).filter(PENDING_READING_FILTER).distinct()
  alert_readings = batch.readings.filter(alert__isnull=False, copied_from__isnull=True).filter(PENDING_READING_FILTER).distinct()
  alert_readings_list = sorted(alert_readings, key=lambda x: x.alert.name)
  alerts = itertools.groupby(alert_readings_list, key=lambda x: x.alert.name)
  
  if batch.include_manual:
    meters = Meter.objects.filter(
      supply_points__property__route_position__route__reading_batches=batch
    ).filter(force_manual_reading=True)
    meters.update(force_manual_reading=False)
  
  batch.processed_at = timezone.now()
  batch.save()
  
  for name, group in alerts:
    counters[name] = len(list(group))
    
  if null_readings.exists():
    counters['readings_null'] = null_readings.count()
      
  batch.last_task_status = 'partial' if task_errors else 'ok'
  batch.last_task_errors = task_errors
  batch.save(update_fields=['last_task_status', 'last_task_errors'])

  response = {
    'id': reading_batch_id,
    'counters': counters,
    'status': 'partial' if task_errors else 'ok',
    'errors': task_errors,
  }

  return response


@shared_task(bind=True)
def assign_readings_task(self, reading_batch_id, payload=None):
  payload = payload or {}
  progress_recorder = ProgressRecorder(self)
 
  reading_batch = ReadingBatch.objects.get(id=reading_batch_id)
  active_contract_token = ConfigProject.objects.get(token='contract_active_token').value
  terminated_contract_token = ConfigProject.objects.get(token='contract_terminated_status').value
  contract_termination_completed_token = ConfigProject.objects.get(token="contract_termination_completed_token").value

 
  reading_files_ids = payload.get('reading_files') or []
  reading_date = payload.get('reading_date')
  not_billed = _coerce_to_bool(payload.get('not_billed'))
  minimum_date = payload.get('minimum_date')
  if minimum_date and isinstance(minimum_date, str):
    minimum_date = datetime.strptime(minimum_date, "%Y-%m-%d").date()
 
  date_obj = None
  start_date = None
  end_date = None
  if reading_date:
    date_obj = datetime.strptime(reading_date, "%Y-%m-%d")
    start_date = date_obj - timedelta(days=5)
    end_date = date_obj + timedelta(days=5)
 
  total_contracts_routes = reading_batch.routes.aggregate(
      num_contracts=Count(
          'positions__properties__supply_points__contracts__id',
          filter=(
              Q(positions__properties__supply_points__contracts__status__token=active_contract_token) |
              (
                  Q(positions__properties__supply_points__contracts__status__token=terminated_contract_token) &
                  Q(positions__properties__supply_points__contracts__contractterminationrequest__status__token=contract_termination_completed_token) &
                  (
                      Q(positions__properties__supply_points__contracts__contractterminationrequest__approved_at__date__gte=date_obj.date() if date_obj else timezone.now().date()) |
                      Q(positions__properties__supply_points__contracts__contractterminationrequest__approved_at__isnull=True, positions__properties__supply_points__contracts__contractterminationrequest__created_at__date__gte=date_obj.date() if date_obj else timezone.now().date())
                  )
              )
          ),
          distinct=True
      )
  )
  fix_supplies_contracts_data = reading_batch.fix_meters.aggregate(
      num_contracts=Count(
          'supply_points__contracts__id',
          filter=(
              Q(supply_points__contracts__status__token=active_contract_token) |
              (
                  Q(supply_points__contracts__status__token=terminated_contract_token) &
                  Q(supply_points__contracts__contractterminationrequest__status__token=contract_termination_completed_token) &
                  (
                      Q(supply_points__contracts__contractterminationrequest__approved_at__date__gte=date_obj.date() if date_obj else timezone.now().date()) |
                      Q(supply_points__contracts__contractterminationrequest__approved_at__isnull=True, supply_points__contracts__contractterminationrequest__created_at__date__gte=date_obj.date() if date_obj else timezone.now().date())
                  )
              )
          ),
          distinct=True
      )
  )
  total_contracts = (total_contracts_routes['num_contracts'] or 0) + (fix_supplies_contracts_data['num_contracts'] or 0)
 
  total_route_supplypoints = reading_batch.routes.aggregate(
      total=Count('positions__properties__supply_points__id', distinct=True)
  )['total'] or 0
  total_fix_supplypoints = reading_batch.fix_meters.aggregate(
      total=Count('supply_points__id', distinct=True)
  )['total'] or 0
  total_supplypoints = total_route_supplypoints + total_fix_supplypoints or 1
 
  reading_files = ReadingDocument.objects.filter(id__in=reading_files_ids) if reading_files_ids else ReadingDocument.objects.none()
  reading_files_present = reading_files.exists()
  num_readings_in_files = 0
  if reading_files_present:
    for reading_file in reading_files:
      num_readings_in_files += reading_file.readings.filter(copied_from__isnull=True).count()
      reading_file.batch = reading_batch
      reading_file.save(update_fields=['batch'])
 
  exclude_telecontrol = not reading_batch.include_telecontrol
  exclude_manual = not reading_batch.include_manual
 
  readings_to_assign = []
  reading_ids = set()
  total_billed_readings = 0
  processed_supplypoints = 0
 
  def add_reading(reading_obj):
    if reading_obj and reading_obj.id not in reading_ids:
      reading_ids.add(reading_obj.id)
      readings_to_assign.append(reading_obj)
 
  def update_progress():
    progress_recorder.set_progress(
        processed_supplypoints if processed_supplypoints <= total_supplypoints else total_supplypoints,
        total_supplypoints
    )
 
  def get_supply_points(property_obj):
    if exclude_telecontrol:
      return property_obj.supply_points.filter(
          Q(meter__has_remote_reading=False) | Q(meter__sub_meters__has_remote_reading=False)
      | Q(meter__force_manual_reading=True) | Q(meter__sub_meters__force_manual_reading=True)
      ).distinct()
    if exclude_manual:
      return property_obj.supply_points.filter(
          Q(meter__has_remote_reading=True, meter__force_manual_reading=False) | Q(meter__sub_meters__has_remote_reading=True, meter__sub_meters__force_manual_reading=False)
      ).distinct()
    return property_obj.supply_points.all()

  def resolve_general_meter_canonical(meter, for_reading_date=False, for_not_billed=False, reading_file=None):
    """Fallback: physical general-meter reading (often contract/SP null) before billing fan-out."""
    if not meter or not meter.is_general:
      return None
    if reading_file is not None:
      canonical = reading_file.readings.filter(
          meter=meter,
          copied_from__isnull=True,
          is_control=False,
          is_close=False,
          is_initial=False,
      ).filter(
            Q(previous_reading__is_close=False) |
            Q(previous_reading__isnull=True)
          ).order_by('contract_id', '-reading_date', '-id').first()
      if canonical:
        return canonical
    if for_reading_date:
      return find_canonical_general_meter_reading(
          meter, start_date=start_date, end_date=end_date, batch=reading_batch
      )
    if for_not_billed:
      qs = Reading.objects.filter(
          meter=meter,
          copied_from__isnull=True,
          is_control=False,
          is_initial=False,
          is_close=False,
          reading_date__gte=minimum_date,
      ).filter(
          Q(previous_reading__is_close=False) |
          Q(previous_reading__isnull=True)
        ).filter(PENDING_READING_FILTER).distinct().order_by('contract_id', '-reading_date', '-id')
      return qs.first()
    return find_canonical_general_meter_reading(meter, batch=reading_batch)

  def try_add_reading(contract, meter, reading_db):
    if not reading_db or not meter:
      return
    key = (contract.token, meter.code)
    if key in set_added_readings:
      return
    set_added_readings.add(key)
    add_reading(reading_db)
 
  filters_present = reading_files_present or bool(reading_date) or not_billed
  # SET OF CONTRACT_TOKEN + METER_CODE TO AVOID ADDING MORE THAN 1 READING FOR THE SAME CONTRACT AND METER
  set_added_readings = set()
  task_errors = []

  if filters_present:
    for route in reading_batch.routes.all():
      for position in route.positions.all():
        for property_obj in position.properties.all():
          supply_points = get_supply_points(property_obj)
          for supply_point in supply_points:
            processed_supplypoints += 1
            update_progress()
            try:
              contracts = supply_point.contracts.filter(
                  Q(status__token=active_contract_token) |
                  (
                      Q(status__token=terminated_contract_token) &
                      Q(contractterminationrequest__status__token=contract_termination_completed_token) &
                      (
                          Q(contractterminationrequest__approved_at__date__gte=date_obj.date() if date_obj else (minimum_date if minimum_date else timezone.now().date())) |
                          Q(contractterminationrequest__approved_at__isnull=True, contractterminationrequest__created_at__date__gte=date_obj.date() if date_obj else (minimum_date if minimum_date else timezone.now().date()))
                      )
                  )
              ).distinct()
              if reading_files_present:
                for reading_file in reading_files:
                  for contract in contracts:
                    reading_db = reading_file.readings.filter(supply_point=supply_point, contract=contract, is_control=False, is_close=False, is_initial=False).first()
                    if reading_file.readings.filter(supply_point=supply_point, contract=contract, invoices__type_final='F', is_close=False, is_initial=False).exists():
                      total_billed_readings += 1
                    if not reading_db:
                      reading_db = resolve_general_meter_canonical(supply_point.meter, reading_file=reading_file)
                    try_add_reading(contract, supply_point.meter, reading_db)
              elif reading_date:
                for contract in contracts:
                  reading_db = Reading.objects.filter(
                      reading_date__range=(start_date, end_date),
                      supply_point=supply_point,
                      contract=contract,
                      is_control=False,
                      is_initial=False,
                      is_close=False,
                  ).filter(
                      Q(previous_reading__is_close=False) |
                      Q(previous_reading__isnull=True)
                    ).annotate(
                      date_diff=ExpressionWrapper(
                          F('reading_date') - date_obj.date(),
                          output_field=fields.DurationField()
                      )
                  ).order_by('date_diff').first()
                  if Reading.objects.filter(
                      reading_date__range=(start_date, end_date),
                      supply_point=supply_point,
                      contract=contract,
                      invoices__type_final='F',
                      is_initial=False,
                      is_close=False,
                  ).filter(
                      Q(previous_reading__is_close=False) |
                      Q(previous_reading__isnull=True)
                    ).exists():
                    total_billed_readings += 1
                  if not reading_db:
                    reading_db = resolve_general_meter_canonical(supply_point.meter, for_reading_date=True)
                  try_add_reading(contract, supply_point.meter, reading_db)
              elif not_billed:
                for contract in contracts:
                  reading_db = Reading.objects.filter(
                      supply_point=supply_point,
                      contract=contract,
                      is_control=False,
                      is_initial=False,
                      is_close=False,
                      reading_date__gte=minimum_date
                  ).filter(
                      Q(previous_reading__is_close=False) |
                      Q(previous_reading__isnull=True)
                    ).filter(
                      PENDING_READING_FILTER
                      ).distinct().order_by('-reading_date').first()
                  if not reading_db:
                    reading_db = resolve_general_meter_canonical(supply_point.meter, for_not_billed=True)
                  try_add_reading(contract, supply_point.meter, reading_db)
            except Exception as e:
              print(f'Error processing supply_point={supply_point.id}: {type(e).__name__}: {str(e)}')
              classification = classify_reading_processing_error(e)
              task_errors.append({
                  'supply_point_id': supply_point.id,
                  'error': str(e),
                  'type': type(e).__name__,
                  **classification,
              })
              continue

    for meter in reading_batch.fix_meters.all():
      for supply_point in meter.supply_points.all():
        processed_supplypoints += 1
        update_progress()
        try:
          contracts = supply_point.contracts.filter(
              Q(status__token=active_contract_token) |
              (
                  Q(status__token=terminated_contract_token) &
                  Q(contractterminationrequest__status__token=contract_termination_completed_token) &
                  (
                      Q(contractterminationrequest__approved_at__date__gte=date_obj.date() if date_obj else (minimum_date if minimum_date else timezone.now().date())) |
                      Q(contractterminationrequest__approved_at__isnull=True, contractterminationrequest__created_at__date__gte=date_obj.date() if date_obj else (minimum_date if minimum_date else timezone.now().date()))
                  )
              )
          ).distinct()
          if reading_files_present:
            for reading_file in reading_files:
              for contract in contracts:
                reading_db = reading_file.readings.filter(supply_point=supply_point, contract=contract).first()
                if reading_file.readings.filter(supply_point=supply_point, contract=contract, invoices__type_final='F', is_close=False, is_initial=False).exists():
                  total_billed_readings += 1
                if not reading_db:
                  reading_db = resolve_general_meter_canonical(supply_point.meter, reading_file=reading_file)
                try_add_reading(contract, supply_point.meter, reading_db)
          elif reading_date:
            for contract in contracts:
              reading_db = Reading.objects.filter(
                  reading_date__range=(start_date, end_date),
                  supply_point=supply_point,
                  is_control=False,
                  is_initial=False,
                  is_close=False,
                  contract=contract
              ).filter(
                  Q(previous_reading__is_close=False) |
                  Q(previous_reading__isnull=True)
                ).annotate(
                  date_diff=ExpressionWrapper(
                      F('reading_date') - date_obj.date(),
                      output_field=fields.DurationField()
                  )
              ).order_by('date_diff').first()
              if Reading.objects.filter(
                  reading_date__range=(start_date, end_date),
                  supply_point=supply_point,
                  contract=contract,
                  invoices__type_final='F',
                  is_close=False,
                  is_initial=False,
              ).filter(
                    Q(previous_reading__is_close=False) |
                    Q(previous_reading__isnull=True)
                  ).exists():
                total_billed_readings += 1
              if not reading_db:
                reading_db = resolve_general_meter_canonical(supply_point.meter, for_reading_date=True)
              try_add_reading(contract, supply_point.meter, reading_db)
          elif not_billed:
            for contract in contracts:
              reading_db = Reading.objects.filter(
                  supply_point=supply_point,
                  contract=contract,
                  is_control=False,
                  is_initial=False,
                  is_close=False,
                  reading_date__gte=minimum_date
              ).filter(
                  Q(previous_reading__is_close=False) |
                  Q(previous_reading__isnull=True)
                ).filter(PENDING_READING_FILTER).distinct().order_by('-reading_date').first()
              if not reading_db:
                reading_db = resolve_general_meter_canonical(supply_point.meter, for_not_billed=True)
              try_add_reading(contract, supply_point.meter, reading_db)
        except Exception as e:
          print(f'Error processing fix_meter supply_point={supply_point.id}: {type(e).__name__}: {str(e)}')
          classification = classify_reading_processing_error(e)
          task_errors.append({
              'supply_point_id': supply_point.id,
              'meter_id': meter.id,
              'error': str(e),
              'type': type(e).__name__,
              **classification,
          })
          continue

    if readings_to_assign:
      reading_batch.readings.set(readings_to_assign)
  else:
    readings_to_assign = list(reading_batch.readings.filter(PENDING_READING_FILTER).distinct())
    total_supplypoints = len(readings_to_assign) or 1
    processed_supplypoints = total_supplypoints
    update_progress()
 
  status_token = ConfigProject.objects.get(token='reading_batch_pending_processing_token').value
  reading_batch.status = ReadingBatchStatus.objects.get(token=status_token)
  reading_batch.assign_readings_task_id = None
  reading_batch.last_task_status = 'partial' if task_errors else 'ok'
  reading_batch.last_task_errors = task_errors
  reading_batch.save(update_fields=['status', 'assign_readings_task_id', 'last_task_status', 'last_task_errors'])

  supply_active_token = ConfigProject.objects.get(token='supply_point_status_activate_token').value
  supply_cut_token = ConfigProject.objects.get(token='supply_point_status_cut_token').value
 
  readings_with_reader_alerts = reading_batch.readings.filter(reader_alert__isnull=False).filter(PENDING_READING_FILTER).distinct().count()
  readings_with_remote_alerts = reading_batch.readings.filter(remote_alert__isnull=False).filter(PENDING_READING_FILTER).distinct().count()
  readings_without_reading_value = reading_batch.readings.filter(reading_value__isnull=True).filter(PENDING_READING_FILTER).distinct().count()
  readings_remote_reading = reading_batch.readings.filter(meter__has_remote_reading=True, meter__force_manual_reading=False).filter(PENDING_READING_FILTER).distinct().count()
  readings_inactive_supplypoints = reading_batch.readings.exclude(
      supply_point__status__token__in=[supply_active_token, supply_cut_token]
  ).filter(PENDING_READING_FILTER).distinct().count()
 
  assigned_readings_count = len(readings_to_assign)
 
  result = {
      "assigned_readings": assigned_readings_count - total_billed_readings,
      "readings_with_reader_alerts": readings_with_reader_alerts,
      "readings_with_remote_alerts": readings_with_remote_alerts,
      "missing_readings": total_contracts - assigned_readings_count if total_contracts - assigned_readings_count > 0 else 0,
      "discarded_readings": num_readings_in_files - total_contracts if num_readings_in_files - total_contracts > 0 else 0,
      "readings_without_reading_value": readings_without_reading_value,
      "readings_remote_reading": readings_remote_reading,
      "readings_inactive_supplypoints": readings_inactive_supplypoints,
      "status": 'partial' if task_errors else 'ok',
      "errors": task_errors,
  }

  progress_recorder.set_progress(total_supplypoints, total_supplypoints)
  return result


@shared_task(bind=True)
def preview_smart_metering_task(self, reading_batch_id, reading_date):
  progress_recorder = ProgressRecorder(self)
  reading_batch = ReadingBatch.objects.get(id=reading_batch_id)

  def on_progress(current, total, description=""):
    progress_recorder.set_progress(current, total, description)

  result = build_smart_metering_preview(
      reading_batch,
      reading_date,
      api_timeout=SMART_METERING_PREVIEW_API_TIMEOUT,
      progress_callback=on_progress,
  )
  return _json_safe(result)


@shared_task(bind=True)
def assign_smart_metering_task(self, reading_batch_id, reading_date, preview=None):

  progress_recorder = ProgressRecorder(self)
  reading_batch = ReadingBatch.objects.get(id=reading_batch_id)
  progress_recorder.set_progress(0, 1)
  try:
    result = assign_smart_metering_readings(reading_batch, reading_date, preview=preview)
  except Exception:
    reading_batch.assign_readings_task_id = None
    reading_batch.save(update_fields=['assign_readings_task_id'])
    raise
  progress_recorder.set_progress(1, 1)
  return result


@shared_task(bind=True)
def estimate_readings_task(self, reading_batch_id, payload=None):
  payload = payload or {}
  progress_recorder = ProgressRecorder(self)

  reading_batch = ReadingBatch.objects.select_related('status').prefetch_related(
      'routes__positions__properties__supply_points'
  ).get(id=reading_batch_id)
  active_contract_token = ConfigProject.objects.get(token='contract_active_token').value
  PERIOD_MONTHS_MAP = {
      'trimestral': 90,
      'semestral': 180,
      'bimestral': 60,
      'quadrimestral': 120,
      'anual': 360,
      'mensual': 30,
  }

  extra_filter = payload.get('extra_filter')
  requested_reading_date = payload.get('reading_date')
  estimation_type = payload.get('estimation_type')
  days_from_last = payload.get('days_from_last')
  reading_estimate_period = payload.get('period') if payload.get('period') in READING_ESTIMATE_PERIOD_CHOICES else None
  reading_estimate_statistic = payload.get('statistic') if payload.get('statistic') in READING_ESTIMATE_STATISTIC_CHOICES else None

  batch_supply_points = set()
  exclude_telecontrol = not reading_batch.include_telecontrol
  exclude_manual = not reading_batch.include_manual
  
  batch_billers = {route.biller_id for route in reading_batch.routes.all() if route.biller_id}
  biller = reading_batch.routes.first().biller if len(batch_billers) == 1 else None
  biller_days = PERIOD_MONTHS_MAP[biller.period_type] if biller else 90
  terminated_contract_token = ConfigProject.objects.get(token='contract_terminated_status').value
  contract_termination_completed_token = ConfigProject.objects.get(token="contract_termination_completed_token").value

  reading_dates = list(reading_batch.readings.exclude(reading_date__isnull=True).values_list('reading_date', flat=True))
  avg_batch_date = None
  target_date = None
  if reading_dates:
    epoch = datetime(1970, 1, 1).date()
    total_days = sum((reading_date - epoch).days for reading_date in reading_dates)
    avg_days = total_days // len(reading_dates)
    avg_batch_date = epoch + timedelta(days=avg_days)

  if requested_reading_date:
    # Una data demanada explicitament pel front sempre te prioritat sobre la
    # mitjana del lot (avg_batch_date), que es nomes un fallback quan no es
    # demana cap data concreta.
    if isinstance(requested_reading_date, str):
      try:
        target_date = datetime.strptime(requested_reading_date, "%Y-%m-%d").date()
      except ValueError:
        target_date = avg_batch_date or timezone.now().date()
    else:
      target_date = requested_reading_date
  elif avg_batch_date:
    target_date = avg_batch_date
  else:
    target_date = timezone.now().date()
    
  
  for route in reading_batch.routes.all():
    for position in route.positions.all():
      for property_obj in position.properties.all():
        if exclude_telecontrol:
          supply_points = property_obj.supply_points.filter(
              Q(meter__has_remote_reading=False) | Q(meter__sub_meters__has_remote_reading=False)
              | Q(meter__force_manual_reading=True) | Q(meter__sub_meters__force_manual_reading=True)
          )
        elif exclude_manual:
          supply_points = property_obj.supply_points.filter(
              Q(meter__has_remote_reading=True, meter__force_manual_reading=False) | Q(meter__sub_meters__has_remote_reading=True, meter__sub_meters__force_manual_reading=False)
          )
        else:
          supply_points = property_obj.supply_points.all()
        supply_points = supply_points.filter(
            Q(contracts__status__token=active_contract_token) |
            (
                Q(contracts__status__token=terminated_contract_token) &
                Q(contracts__contractterminationrequest__status__token=contract_termination_completed_token) &
                (
                    Q(contracts__contractterminationrequest__approved_at__date__gte=target_date) |
                    Q(contracts__contractterminationrequest__approved_at__isnull=True, contracts__contractterminationrequest__created_at__date__gte=target_date)
                )
            )
        ).distinct()
        for supply_point in supply_points:
          batch_supply_points.add(supply_point.id)

  for meter in reading_batch.fix_meters.all():
    for supply_point in meter.supply_points.all():
      batch_supply_points.add(supply_point.id)

  readings_supply_points = set(reading_batch.readings.values_list('supply_point_id', flat=True))
  readings_without_value = set(
      reading_batch.readings.filter(
          reading_value__isnull=True,
          supply_point__contracts__is_active=True,
          supply_point__contracts__status__token__in=[active_contract_token, terminated_contract_token]
      ).values_list('supply_point_id', flat=True)
  )

  if extra_filter == 'no_reading_value':
    missing_supply_points = readings_without_value
  else:
    missing_supply_points = (batch_supply_points - readings_supply_points) | readings_without_value

  if missing_supply_points:
    supply_points_qs = SupplyPoint.objects.select_related(
        'type', 'status', 'property__route_position__route'
    ).filter(id__in=missing_supply_points).order_by('id')
    supply_points = list(supply_points_qs)
  else:
    supply_points = []

  

  readings = []
  total_supplypoints = len(supply_points) or 1

  # Add safeguard: maximum iterations to prevent infinite processing
  MAX_ITERATIONS = 100000  # Safety limit
  processed_count = 0
  task_errors = []
  today = timezone.now().date()
  terminated_contract_token = ConfigProject.objects.get(token='contract_terminated_status').value
  contract_termination_completed_token = ConfigProject.objects.get(token="contract_termination_completed_token").value
  try:
    for index, supply_point in enumerate(supply_points, start=1):
      # Safety check to prevent infinite loops
      if processed_count >= MAX_ITERATIONS:
        print(f'WARNING: Reached maximum iteration limit ({MAX_ITERATIONS}), stopping processing')
        break
      
      progress_recorder.set_progress(index, total_supplypoints)
      sp_biller_days = biller_days
      sp_biller = supply_point.property.route_position.route.biller if supply_point.property and supply_point.property.route_position and supply_point.property.route_position.route else None
      if sp_biller:
        sp_biller_days = PERIOD_MONTHS_MAP.get(sp_biller.period_type, biller_days)
      if supply_point.contracts.exists():
        contracts = supply_point.contracts.filter(
            Q(status__token=active_contract_token) |
            (
                Q(status__token=terminated_contract_token) &
                Q(contractterminationrequest__status__token=contract_termination_completed_token) &
                (
                    Q(contractterminationrequest__approved_at__date__gte=target_date) |
                    Q(contractterminationrequest__approved_at__isnull=True, contractterminationrequest__created_at__date__gte=target_date)
                )
            )
        ).distinct()
        for contract in contracts:
          try:
            if not Reading.objects.filter(supply_point=supply_point, contract=contract, batch=reading_batch, is_estimated=True).exists():
              
              # Lògica de càlcul de data específica per punt de subministrament
              sp_target_date = target_date
              if avg_batch_date and abs((contract.created_at.date() - avg_batch_date).days) < 30:
                  parsed_requested_date = requested_reading_date
                  if isinstance(parsed_requested_date, str):
                      try:
                          parsed_requested_date = datetime.strptime(parsed_requested_date, "%Y-%m-%d").date()
                      except ValueError:
                          parsed_requested_date = None
                  sp_target_date = parsed_requested_date if parsed_requested_date else today
              elif estimation_type == 'days_from_last' and days_from_last:
                  last_reading = Reading.objects.filter(
                      supply_point=supply_point,
                      contract=contract,
                      reading_value__isnull=False
                  ).order_by('-reading_date').first()

                  if last_reading:
                      sp_target_date = last_reading.reading_date + timedelta(days=int(days_from_last))
              elif not requested_reading_date:
                  # avg_batch_date es pot distorsionar per lectures disperses/antigues
                  # arrossegades al lot (p.ex. no facturades de cicles anteriors); cada
                  # punt fa servir la seva propia ultima lectura real + el seu propi
                  # periode com a referencia, en lloc de la mitjana de tot el lot.
                  last_reading = Reading.objects.filter(
                      supply_point=supply_point,
                      contract=contract,
                      reading_value__isnull=False
                  ).order_by('-reading_date').first()
                  if last_reading:
                      sp_target_date = last_reading.reading_date + timedelta(days=sp_biller_days)

              # For terminated contracts, only estimate if target_date is before or equal to termination
              is_billable = True
              if contract.status.token == terminated_contract_token:
                  termination = contract.contractterminationrequest_set.filter(status__token=contract_termination_completed_token).first()
                  term_date = (termination.approved_at.date() if termination.approved_at else termination.created_at.date()) if termination else None
                  if term_date and sp_target_date > term_date:
                      is_billable = False
              
              if is_billable:
                if reading_estimate_period or reading_estimate_statistic:
                  # Comportament configurable: nomes s'activa si el front ha triat
                  # explicitament periode/estadistic. Sense aquests parametres es
                  # mante el comportament original (get_estimated_reading_minimal_object).
                  reading = get_estimated_reading(
                      supply_point, contract, sp_target_date, batch=reading_batch,
                      offset_days=sp_biller_days,
                      period=reading_estimate_period or READING_ESTIMATE_PERIOD_SAME_PERIOD_PREVIOUS_YEAR,
                      statistic=reading_estimate_statistic or READING_ESTIMATE_STATISTIC_MEAN,
                      max_day_limit=today,
                  )
                else:
                  reading = get_estimated_reading_minimal_object(supply_point, contract, sp_target_date, reading_batch, sp_biller_days, max_day_limit = today)
                if reading:
                  readings.append(reading)
            processed_count += 1
          except Exception as e:
            print(f'Error processing supply_point={supply_point.id}, contract={contract.id}: {type(e).__name__}: {str(e)}')
            import traceback
            traceback.print_exc()
            classification = classify_reading_processing_error(e)
            task_errors.append({
                'supply_point_id': supply_point.id,
                'contract_id': contract.id,
                'error': str(e),
                'type': type(e).__name__,
                **classification,
            })
            # Continue processing other contracts instead of failing completely
            continue

    serialized_readings = ReadingMinimalSerializer(readings, many=True).data
    result = {
        'readings': serialized_readings,
        'status': 'partial' if task_errors else 'ok',
        'errors': task_errors,
    }
    progress_recorder.set_progress(total_supplypoints, total_supplypoints)
    return result
  finally:
    reading_batch.estimating_task_id = None
    reading_batch.last_task_status = 'partial' if task_errors else 'ok'
    reading_batch.last_task_errors = task_errors
    reading_batch.save(update_fields=['estimating_task_id', 'last_task_status', 'last_task_errors'])

@shared_task(bind=True)
@billing_run_cache_scope()
def process_billing_batch(self, billing_id, queue_item_id=None):
  progress_recorder = ProgressRecorder(self)
  
  print("Start processing invoices")
  
  billing = Billing.objects.select_related('biller', 'excluded_from', 'excluded_from__biller').get(id=billing_id)

  # Un lot d'excloses es crea a partir d'un altre i ha de facturar amb el mateix
  # facturador. Si es va crear sense, el recuperem de l'origen.
  if billing.biller_id is None and billing.excluded_from_id and billing.excluded_from.biller_id:
    billing.biller = billing.excluded_from.biller
    billing.save(update_fields=['biller'])
    if not billing.routes.exists():
      billing.routes.set(billing.excluded_from.routes.all())

  if billing.biller_id is None:
    raise ValueError(f"Billing {billing.id} has no biller")

  # Expand 1 physical general-meter reading into N contract-scoped copies (copied_from)
  fanout_general_meter_readings(billing)
  
  period_type = billing.biller.period_type
  PERIOD_MONTHS_MAP = {
      'trimestral': 90,
      'semestral': 180,
      'bimestral': 60,
      'quadrimestral': 120,
      'anual': 360,
      'mensual': 30,
  }
  period_days = PERIOD_MONTHS_MAP[period_type]
  period_months = period_days / 30
  
  contract_active_token = ConfigProject.objects.get(token='contract_active_token').value
  contract_terminated_token = ConfigProject.objects.get(token='contract_terminated_status').value
  contract_termination_completed_token = ConfigProject.objects.get(token="contract_termination_completed_token").value

  readings_qs = (
    billing.readings.filter(is_control=False, is_active=True, is_close=False)
    .filter(contract__block_billing=False)
    .filter(
        Q(contract__status__token=contract_active_token) |
        (
            Q(contract__status__token=contract_terminated_token) &
            Q(contract__contractterminationrequest__status__token=contract_termination_completed_token) &
            (
                Q(contract__contractterminationrequest__approved_at__date__gte=F('reading_date')) |
                Q(contract__contractterminationrequest__approved_at__isnull=True, contract__contractterminationrequest__created_at__date__gte=F('reading_date'))
            )
        )
    ).exclude(
      Q(invoices__type_final='F') | Q(invoices__type_final='P'),
    ).values_list('contract_id', flat=True).distinct()
  )
  
  contract_ids = list(readings_qs)
  
  name = billing.name

  if not contract_ids:
    from django.utils import timezone
    if queue_item_id:
      _close_billing_queue_item(queue_item_id, self.request.id, status='completed', completed_at=timezone.now())
      
    return {
      'id': billing_id,
      'counters': {'total': 0}
    }



  close_old_connections()
  invoices_count = 0
  try:
    invoice_warning_negative_token = ConfigProject.objects.get(token='invoice_warning_negative').value
    invoice_warning_negative = InvoiceWarning.objects.get(token=invoice_warning_negative_token)
    
    invoice_warning_zero_token = ConfigProject.objects.get(token='invoice_warning_zero').value
    invoice_warning_zero = InvoiceWarning.objects.get(token=invoice_warning_zero_token)
    
    invoice_warning_bank_missing_token = ConfigProject.objects.get(token='invoice_warning_bank_missing').value
    invoice_warning_bank_missing = InvoiceWarning.objects.get(token=invoice_warning_bank_missing_token)
    
    invoice_warning_payment_missing_token = ConfigProject.objects.get(token='invoice_warning_payment_missing').value
    invoice_warning_payment_missing = InvoiceWarning.objects.get(token=invoice_warning_payment_missing_token)
    
    invoice_warning_simplified_over_400_token = ConfigProject.objects.get(token='invoice_warning_simplified_over_400').value
    invoice_warning_simplified_over_400 = InvoiceWarning.objects.get(token=invoice_warning_simplified_over_400_token)
    
    invoice_warning_high_amount_token = ConfigProject.objects.get(token='invoice_warning_high_amount').value
    invoice_warning_high_amount = InvoiceWarning.objects.get(token=invoice_warning_high_amount_token)

    contracts = Contract.objects.filter(id__in=contract_ids).select_related(
        'general_invoice', 'holder', 'address_billing__address__city', 
        'address_billing__address__province', 'address_billing__address__country'
    )
    
    readings_qs = (
      Reading.objects.filter(billing_id=billing_id, contract_id__in=contract_ids, is_control=False, is_active=True, is_close=False)
      .select_related('contract', 'contract__status', 'supply_point', 'meter', 'batch', 'billing')
      .order_by('-reading_date')
    )
    
    contract_readings = {}
    set_added_readings = set()
    for reading in readings_qs:
      contract = reading.contract
      if (contract.token, reading.meter.code) in set_added_readings:
        continue
      if contract not in contract_readings:
        contract_readings[contract] = []
      set_added_readings.add((contract.token, reading.meter.code))
      contract_readings[contract].append(reading)

    invoices = []
    
    avg_total_cache = TotalAmountAvgCache(contract_ids)

    total_contracts = len(contracts)
    processed_contracts = 0
    progress_recorder.set_progress(0, total_contracts)

    for contract in contracts:
        try:
          warning = None
          readings_list = contract_readings.get(contract, [])
          if not readings_list:
            continue
          invoice = generate_consumption_invoice_multiple(contract, readings_list, name, billing=billing, period_months=period_months)
          if not invoice:
            continue
          if contract.general_invoice:
            general_invoice = contract.general_invoice
            general_contracts = [ctr for ctr in contracts if ctr.general_invoice == general_invoice]
            invoice.general_contracts.set(general_contracts)

          invoices.append(invoice)
          if invoice.total_final <= 0:
            warning = invoice_warning_negative
          elif invoice.total_final == 0:
            warning = invoice_warning_zero
          elif invoice.payment_type is None:
            warning = invoice_warning_payment_missing
          elif invoice.payment_type.token == 'DIRECT_DEBIT' and invoice.payment_bank is None:
            warning = invoice_warning_bank_missing
          elif invoice.simplified and invoice.total_final >= 400:
            warning = invoice_warning_simplified_over_400
          elif invoice.total_final:
            if avg_total_cache.is_high_amount(invoice.total_final, contract.id, invoice.billing_period_year, invoice.billing_period_month):
              warning = invoice_warning_high_amount

          if warning:
            invoice.warning = warning
            invoice.save(update_fields=['warning'])
        finally:
          processed_contracts += 1
          progress_recorder.set_progress(processed_contracts, total_contracts)

    if invoices:
      invoices_count = len(invoices)
      invoice_ids = [invoice.id for invoice in invoices]
      Invoice.objects.filter(id__in=invoice_ids).update(billing=billing)

  finally:
    close_old_connections()

  # Set Billing pending status
  pending_status_token = ConfigProject.objects.get(token='billing_batch_pending').value
  pending_status = BillingStatus.objects.get(token=pending_status_token)
  billing.status = pending_status
  billing.save()

  if queue_item_id:
    from django.utils import timezone
    _close_billing_queue_item(queue_item_id, self.request.id, status='completed', completed_at=timezone.now())

  invoices_qs = Invoice.objects.filter(billing=billing)
  counters = {}
  counters['total'] = invoices_qs.count()
  warning_counts = (
    invoices_qs.filter(warning__isnull=False)
    .values('warning__name')
    .annotate(count=Count('id'))
  )
  for wc in warning_counts:
    counters[wc['warning__name']] = wc['count']

  try:
      high_amount_token = ConfigProject.objects.get(token='invoice_warning_high_amount').value
      counters['high_billing_amount'] = invoices_qs.filter(warning__token=high_amount_token).count()
  except ConfigProject.DoesNotExist:
      counters['high_billing_amount'] = 0

  return {
    'id': billing_id,
    'counters': counters
  }


def expire_invoices_with_returned_payments(now=None):
  """Passa a Vençuda les factures amb el venciment passat i el rebut retornat.

  Quan es processa una devolució (`billing/views/payment_view.py`) l'estat de la
  factura es decideix comparant el venciment amb la data d'aquell moment: si el
  venciment encara era futur la factura torna a Confirmada. Aquesta decisió no es
  tornava a revisar mai, de manera que en arribar el venciment la factura es
  quedava Confirmada tot i tenir el rebut retornat i el deute viu.

  L'estat del Payment NO es toca: la devolució (i el seu motiu de rebuig) s'ha de
  conservar tal com la va retornar el banc. Només es mou l'estat de la factura.
  """
  if now is None:
    now = timezone.now().date()

  payment_status_returned = get_status_map()["payment_status_returned_token"]
  invoice_status_expired = get_invoice_status('invoice_status_expired_token')
  open_status_tokens = [
    ConfigProject.objects.get(token='invoice_status_confirmed_token').value,
    ConfigProject.objects.get(token='invoice_status_sent_token').value,
  ]

  invoices = Invoice.objects.filter(
    is_active=True,
    is_excluded=False,
    due_date__lte=now,
    status__token__in=open_status_tokens,
    payments__is_active=True,
    payments__is_duplicate=False,
    payments__status=payment_status_returned,
  ).select_related('status').distinct()

  logs = []
  invoice_ids = []
  for invoice in invoices:
    logs.append(LogInvoiceChangeStatus(
      object=invoice,
      previous_status=invoice.status,
      current_status=invoice_status_expired,
    ))
    invoice_ids.append(invoice.id)

  if not invoice_ids:
    return 0

  LogInvoiceChangeStatus.objects.bulk_create(logs, batch_size=500)
  # update() en lloc de save(): és un procés massiu i no ha de disparar els
  # senyals de Invoice/Payment (pagaments, verifactu, moneder...).
  Invoice.objects.filter(id__in=invoice_ids).update(
    status=invoice_status_expired,
    updated_at=timezone.now(),
  )

  # Mateixa actualització que fa el senyal `update_payment` quan un pagament venç:
  # deixar constància al contracte de la data del deute més recent.
  contract_debt_dates = (
    Invoice.objects.filter(id__in=invoice_ids, contract__isnull=False)
    .values('contract_id')
    .annotate(last_due_date=Max('due_date'))
  )
  for row in contract_debt_dates:
    Contract.objects.filter(id=row['contract_id']).filter(
      Q(last_debt_data__isnull=True) | Q(last_debt_data__lt=row['last_due_date'])
    ).update(last_debt_data=row['last_due_date'])

  print(f"Marked {len(invoice_ids)} invoices with returned payments as expired")
  return len(invoice_ids)


@shared_task
def check_payments_due_date():
  print("check due date")
  now = timezone.now().date()
  payment_status_map = get_status_map()
  payment_status_pending = payment_status_map["payment_status_pending_token"]
  payment_status_expired = payment_status_map["payment_status_expired_token"]
  passed_due_date_payments = Payment.objects.filter(due_date__lte=now, status=payment_status_pending)
  print(f"Found {passed_due_date_payments.count()} payments due date")

  # Els rebuts retornats abans del venciment deixen la factura Confirmada i cap
  # procés la tornava a revisar en arribar la data límit.
  expired_returned_invoices = expire_invoices_with_returned_payments(now)

  if len(passed_due_date_payments) > 0 or expired_returned_invoices > 0:
    for payment in passed_due_date_payments:
        # El senyal `update_payment` ja arrossega la factura a Vençuda.
        payment.status = payment_status_expired
        payment.reject_date = payment.due_date
        payment.save()
    with translation.override(settings.LANGUAGE_CODE):
      notification_save = {
        'token': uuid.uuid4(),
        'name': _("Expired invoices"),
        'description': _("Invoices declared as expired on %(date)s") % {"date": now.strftime('%d/%m/%Y')},
        'module': 'billing',
        'entity': 'invoice',
        'object_id': None,
        'is_active': True
      }
    notification = Notification.objects.create(**notification_save)


@shared_task
def update_missing_periods_for_invoices():
  # no need to disconnect signals for bulk update
  origin_reading_token = ConfigProject.objects.get(token='origin_reading_token').value
  invoices = Invoice.objects.filter(
      origin__token=origin_reading_token,
      issue_date__gte=datetime.date(2025,10,1),
      billing_period_days__isnull=True,
  )
  print(f"Found {invoices.count()} reading invoices without periods to update")
  monthly_invoices = invoices.filter(
      Q(billing__biller__period_type='mensual'),
      Q(contract__supply_point_default__property__route_position__route__biller__period_type='mensual'),
  )
  for invoice in invoices:
      invoice.billing_period_month = invoice.issue_date.month
      invoice.billing_period_year = invoice.issue_date.year
      if invoice in monthly_invoices:
          invoice.billing_period_days = 30
      else:
          invoice.billing_period_days = 90
  Invoice.objects.bulk_update(invoices, fields=['billing_period_month', 'billing_period_year', 'billing_period_days'])
  

@shared_task
def create_billings():
  print("check if any billing must start")
  active_billers = Biller.objects.filter(is_active=True)
  if active_billers.exists():
    current_year = timezone.now().date().year
    current_month = timezone.now().date().month
    
    months = ['jan','feb','mar','apr','may','jun','jul','aug','sep','oct','nov','dec']
    period_type = ['mensual','bimestral','trimestral','semestral','anual']
    period_jumps = [1,2,3,6,12]
    
    for biller in active_billers:
      
      print(f"Check process for {biller.name}")
      
      first_index = months.index(biller.initial_month)+1
      interval = period_jumps[period_type.index(biller.period_type)]
      
      if (first_index > interval):
        first_index-= interval
        
      diff = interval - first_index
      
      # print(f"Current month {current_month}")
      # print(f"First index {first_index}, diff {diff}, interval {interval}")
      
      i = 1
      process = False
      for month in months:
        if (i+diff) % interval == 0:
          if (current_month == i):
            process = True
        i+=1
    
      if process:
        print(f"Check if billig of this biller is still active and send an alert")
        
        billing = Billing.objects.filter(
          biller=biller, 
          is_active=True,
          created_at__year=current_year,
          created_at__month=current_month
        )
        if billing.exists():
          print(f"Billing of this biller is still active. SEND ALERT")
          notification_save = {
            'token': uuid.uuid4(),
            'name': f"Fact. activa",
            'description': f"Fact. {billing.first().name} activa",
            'module': 'billing',
            'entity': 'billing',
            'object_id': billing.first().id,
            'is_active': True
          }
          notification = Notification.objects.create(**notification_save)
        else:
          token = f"{biller.token}-{current_month}{current_year}"
          name = f"Fact. {biller.name} - {current_month}{current_year}"
          
          default_status = BillingStatus.objects.get(is_default=True)
          
          billing = Billing.objects.create(
            status = default_status,
            name = name,
            token = token,
            biller=biller,
            is_active=True)
          
          pending_status_token = ConfigProject.objects.get(token='billing_batch_pending').value
          pending_status = BillingBatchStatus.objects.get(token=pending_status_token)
          if biller.routes.exists():
            routes = biller.routes.all()
            billing.routes.set(routes)
          else:
            routes = Route.objects.filter(is_active = True).all()
            billing.routes.set(routes)
          billing.save()
          
@shared_task(bind=True)
def process_invoice_documents(self, billing_id, invoice_ids, context, issue_date, end_date, send_at, queue_item_id=None):
  progress_recorder = ProgressRecorder(self)

  # Pre-fetch all config project tokens to avoid database roundtrips in the loop
  tokens_to_fetch = [
    'invoice_status_confirmed_token',
    'billing_batch_processed',
    'reading_batch_billed_token',
    'status_payoff',
    'invoice_status_cancelled_token',
    'invoice_status_pending_token',
    'invoice_status_paid_token',
    'direct_debit_token',
    'payment_status_pending_token',
    'payment_status_paid_token',
    'payment_status_cancelled_token',
    'token_ordinary_invoice_serie',
    'token_returned_invoice_serie',
    'token_refactored_invoice_serie',
    'invoice_type_invoice_token',
    'invoice_status_payoff_token',
    'payment_status_payoff_token',
    'has_preprinted_template',
    'tertiary_color',
    'token_simplified_invoice_serie',
    'invoice_type_budget_token',
    'invoice_status_sent_token',
    'invoice_status_expired_token',
    'invoice_status_endowment_token',
    'invoice_status_commitment_token',
    'invoice_main_color',
    'invoice_secondary_color',
    'token_meter_status_no_meter',
  ]
  config_cache = {
    cfg.token: cfg.value
    for cfg in ConfigProject.objects.filter(token__in=tokens_to_fetch)
  }

  invoice_type_invoice_token = config_cache.get('invoice_type_invoice_token')
  invoice_status_pending_token = config_cache.get('invoice_status_pending_token')
  confirmed_status_token = config_cache.get('invoice_status_confirmed_token')
  confirmed_status = InvoiceStatus.objects.get(token=confirmed_status_token)

  print("--------------------------------")
  print(f"Generate invoices final documents")
  print(f"Number of invoices: {len(invoice_ids)}")
  print(f"Issue date: {issue_date}")
  print(f"End date: {end_date}")
  print(f"Send date: {send_at}")
  print("--------------------------------")
  
  # Convertir send_at a datetime amb timezone si existeix
  issue_date = convert_date_with_timezone(issue_date) if issue_date else None
  send_at_aware = convert_date_with_timezone(send_at) if send_at else None
  end_at_aware = convert_date_with_timezone(end_date) if end_date else None
  
  problematic_readings = Reading.objects.filter(
    billing__id=billing_id,
    invoices__billing__isnull=True,
    invoices__type_final=invoice_type_invoice_token
  )
  
  pre_invoices_to_delete = Invoice.objects.filter(
    id__in=invoice_ids,
    readings__in=problematic_readings
  )
  
  pre_invoices_to_delete.update(is_active=False)
    
  
  invoices_qs = Invoice.objects.filter(
      id__in=invoice_ids,
      status__token=invoice_status_pending_token,
      is_active=True,
      is_excluded=False,
  ).select_related(
      'contract',
      'contract__payment',
      'contract__payment__type',
      'contract__payment__IBAN',
      'contract__payment__company_iban',
      'contract__general_invoice',
      'exploitation',
      'company'
  )

  # Check how many invoices will need to be confirmed and generate a PDF
  invoices_needing_pdf_count = invoices_qs.exclude(status=confirmed_status).count()
  # El total real de la cua és el nombre de factures a generar, no la suma de les
  # dues fases (confirmació + generació de PDF), que és el que abans doblava el comptador.
  total_invoices = len(invoice_ids)

  if queue_item_id:
    from billing.models import BillingQueue
    try:
      q_item = BillingQueue.objects.get(id=queue_item_id)
      q_item.total_items = total_invoices
      q_item.processed_items = 0
      q_item.invoices_processed = 0
      q_item.documents_generated = 0
      q_item.save(update_fields=['total_items', 'processed_items', 'invoices_processed', 'documents_generated'])
    except BillingQueue.DoesNotExist:
      pass

  def _report_progress(invoices_processed, documents_generated):
    # Combina les dues fases (confirmació de factures i generació de PDF) en un
    # únic percentatge sobre `total_invoices`, ponderant cadascuna al 50%.
    phase1 = (invoices_processed / total_invoices) if total_invoices else 0.0
    phase2 = (documents_generated / invoices_needing_pdf_count) if invoices_needing_pdf_count else 1.0
    percent = round(((phase1 + phase2) / 2) * 100, 2)
    current = round(((phase1 + phase2) / 2) * total_invoices)
    progress_recorder.set_progress(current, total_invoices)
    if queue_item_id:
      BillingQueue.objects.filter(id=queue_item_id).update(
        invoices_processed=invoices_processed,
        documents_generated=documents_generated,
        processed_items=current,
      )
    return percent

  invoices_processed = 0
  documents_generated = 0
  _report_progress(invoices_processed, documents_generated)

  invoices_needing_pdf = []
  confirmed_invoice_ids = []

  for invoice in invoices_qs:
    if invoice.status != confirmed_status:
      invoice.status = confirmed_status

      update_fields = ['status']
      if issue_date:
        invoice.issue_date = issue_date.date() if hasattr(issue_date, 'hour') else issue_date
        invoice.billing_period_year = invoice.issue_date.year
        invoice.billing_period_month = invoice.issue_date.month
        update_fields.extend(['issue_date', 'billing_period_year', 'billing_period_month'])
      if end_at_aware:
        invoice.due_date = end_at_aware.date() if hasattr(end_at_aware, 'hour') else end_at_aware
        update_fields.append('due_date')
      if send_at_aware:
        invoice.send_at = send_at_aware
        update_fields.append('send_at')
      invoice._generate_verifactu = True
      invoice._bulk_batch = True
      invoice.save(update_fields=update_fields)
      confirm_invoice(invoice, generate_verifactu=True, skip_gen_pdf=True, config_cache=config_cache)

      confirmed_invoice_ids.append(invoice.id)
      invoices_needing_pdf.append(invoice.id)

    invoices_processed += 1
    if invoices_processed % 10 == 0 or invoices_processed == total_invoices:
      _report_progress(invoices_processed, documents_generated)

  if confirmed_invoice_ids:
    from statistics.tasks import update_billing_consumption_batch
    chunk_size = 1000
    for idx in range(0, len(confirmed_invoice_ids), chunk_size):
      batch_ids = confirmed_invoice_ids[idx:idx + chunk_size]
      update_billing_consumption_batch.delay(batch_ids)

  if invoices_needing_pdf:
    from django.db import close_old_connections
    from billing.views.invoice_pdf_view import generate_report_invoice_pdf

    for invoice_id in invoices_needing_pdf:
      try:
        close_old_connections()
        invoice = Invoice.objects.get(id=invoice_id)
        generate_report_invoice_pdf(invoice, config_cache=config_cache)
      except Exception as e:
        print(f"Error generating PDF for invoice {invoice_id}: {e}")
      finally:
        close_old_connections()
        documents_generated += 1
        if documents_generated % 10 == 0 or documents_generated == invoices_needing_pdf_count:
          _report_progress(invoices_processed, documents_generated)

  if config('VERIFACTU_ENABLED', default=False, cast=bool):
    notify_verifactu_unsent_batches()
  billing = Billing.objects.get(id=billing_id)

  processed_status_token = config_cache.get('billing_batch_processed')
  processed_status = BillingStatus.objects.get(token=processed_status_token)
  reading_batch_billed_token = config_cache.get('reading_batch_billed_token')
  reading_batch_billed = ReadingBatchStatus.objects.get(token=reading_batch_billed_token)
  
  billing.status = processed_status
  billing.send_at = send_at_aware
  billing.billing_batches.update(
    issue_date=issue_date.date() if issue_date else None,
    due_date=end_at_aware.date() if end_at_aware else None,
    send_at=send_at_aware
  )
  billing.save()
  
  batches = ReadingBatch.objects.filter(
            readings__billing=billing,
            is_active=True
        ).distinct()
  batches.update(status=reading_batch_billed)
  
  with translation.override(settings.LANGUAGE_CODE):
    notification_save = {
      'token': uuid.uuid4(),
      'name': _("Billing finished"),
      'description': _("Invoices generated for billing %(name)s") % {"name": billing.name},
      'module': 'billing',
      'entity': 'billing',
      'object_id': billing.id,
      'is_active': True
    }
  notification = Notification.objects.create(**notification_save)
  
  if queue_item_id:
    from django.utils import timezone
    _close_billing_queue_item(
      queue_item_id, self.request.id,
      status='completed',
      completed_at=timezone.now(),
      processed_items=total_invoices,
      invoices_processed=invoices_processed,
      documents_generated=documents_generated,
    )

  invoices = Invoice.objects.filter(billing = billing)
  progress_recorder.set_progress(total_invoices, total_invoices)
  response = {
    'id': billing_id,
    'counters': {'total':invoices.count()}
  }
  
  return response

""" @shared_task(bind=True)
def set_next_step_claim_request(self, claim_request_id, step):
  
  claim_request = ClaimRequest.objects.get(id=claim_request_id)
  claim_request_accepted = ClaimRequestStatus.objects.get(token=ConfigProject.objects.get(token='claim_request_status_accepted_token').value)
  if claim_request.status == claim_request_accepted:
    claim_step = ClaimStep.objects.get(id=claim_request.step.id)
    today = timezone.now().date()
    duration = claim_step.duration
    
    if claim_step.duration_type == 'NATURAL':
      if claim_request.previous_step_date + timedelta(days=duration) >= today:
        result = process_claim_request_next_step(claim_request_id)
    
    if claim_step.duration_type == 'WORK':
      #check only work days (not weekends)
      next_step_date = add_workdays(claim_request.previous_step_date, duration)
      if next_step_date >= today:
        result = process_claim_request_next_step(claim_request_id)
    
    if 'error' in result:
      print(f"Error processing claim request {claim_request_id}: {result['error']}")
    else:
      print(
        f"Claim request {claim_request_id} processed successfully.  Data:"
        f" {result['data']}"
      ) """

def add_workdays(start_date, workdays):
  current_date = start_date
  while workdays > 0:
    current_date += timedelta(days=1)
    if current_date.weekday() < 5:  # Monday to Friday (0-4)
      workdays -= 1
  return current_date
  
def assign_batches_to_invoices(billing_batch_ids):
  with transaction.atomic():  # Ensures data integrity
    for batch in BillingBatch.objects.filter(id__in=billing_batch_ids).prefetch_related(
      'templates__routes__positions__properties__supply_points__contracts__invoices'
    ):
      for template in batch.templates.all():
        for route in template.routes.all():
          for position in route.positions.all():
            for property_obj in position.properties.all():  # Multiple properties
              for supply_point in property_obj.supply_points.all():  # Multiple supply points
                for contract in supply_point.contracts.all():  # Multiple contracts
                  for invoice in contract.invoices.all():  # Invoices linked to contract
                    print("ASSIGN BATCH TO INVOICE:")
                    print(invoice)
                    invoice.batch = batch  # Assign batch
                    invoice.save(update_fields=['batch'])  # Save efficiently

    print("Batch assignment completed!")
  

@shared_task(bind=True)
def generate_claim_document_pdf(self, claim_request_id, contract_id, joined_payment_id=None):
    """Generates a PDF for a specific contract within a claim request."""
    print("GOES IN HERE")
    claim_request = ClaimRequest.objects.get(id=claim_request_id)
    contract = Contract.objects.get(id=contract_id)
    claim_step = claim_request.current_step
    document_type = claim_step.document_type
    now_date = datetime.now().date()
    limit_date = now_date + timedelta(days=claim_step.duration)
    formatted_date = now_date.strftime("%d de %B de %Y")

    company = None
    company_obj = None
    sp_address = None

    if (
        contract.supply_point_default
        and contract.supply_point_default.connection
        and contract.supply_point_default.connection.exploitation
        and contract.supply_point_default.connection.exploitation.company
    ):
        company_obj = contract.supply_point_default.connection.exploitation.company
        company = CompanySerializer(company_obj).data

    if contract.supply_point_default and contract.supply_point_default.address:
        sp_address = contract.supply_point_default.address

    holder = PersonSerializer(contract.holder).data
    holder_address = next(
        (address for address in holder.get("addresses", []) if address.get("is_billing")),
        None,
    )

    payment_ids = claim_request.payments.filter(contract=contract).values_list("payment_id", flat=True)
    contract_invoices = Invoice.objects.filter(payments__id__in=payment_ids, contract=contract).distinct()
    invoice_total = sum(invoice.total_final for invoice in contract_invoices)
    invoice_total += Decimal(holder.get("debt_amount", 0) or 0)

    joined_payment = None
    expenses_total = Decimal(0)
    if claim_request.current_step.step_template.group_payments and joined_payment_id:
      try:
        joined_payment = JoinedPayment.objects.get(id=joined_payment_id)
        original_payments = joined_payment.payments.filter(claim_requests__claim_step__isnull=True)
        expense_payments = joined_payment.payments.filter(claim_requests__claim_step__isnull=False)
        contract_invoices = Invoice.objects.filter(payments__in=original_payments, contract=contract).distinct()
        invoice_total = original_payments.aggregate(total=Sum('amount'))['total'] or Decimal(0)
        expenses_total = expense_payments.aggregate(total=Sum('amount'))['total'] or Decimal(0)
      except JoinedPayment.DoesNotExist:
        joined_payment = None
        expenses_total = Decimal(0)

    joined_payment_barcode_base64 = None
    if joined_payment:
      ident = joined_payment.due_date.strftime("%d%m%y") if joined_payment.due_date else limit_date.strftime("%d%m%y")
      if limit_date:
          ident = limit_date.strftime("%d%m%y")
      token = joined_payment.token or ""
      barcode_data = {
          'reference': token if len(token) == 11 else token[:-2],
          'total_final': joined_payment.total_final,
          'company': company_obj,
          'ident': ident,
      }
      barcode = generate_barcode(barcode_data)
      if barcode:
          joined_payment_barcode_base64 = base64.b64encode(barcode.getvalue()).decode('utf-8')
      
    
    
    html_content = render_to_string(
        f"{document_type.token}_template.html",
        {
            "now_date": now_date,
            "limit_date": limit_date,
            "formatted_date": formatted_date,
            "company": company,
            "sp_address": sp_address,
            "holder": holder,
            "holder_address": holder_address.get("address") if holder_address else None,
            "contract": contract,
            "invoices": contract_invoices,
            "invoice_total": invoice_total,
            "return_fee_amount": expenses_total,
            "joined_payment": joined_payment,
            "joined_payment_barcode_base64": joined_payment_barcode_base64,
        },
    )
    pdf_buffer = BytesIO()
    pisa_status = pisa.CreatePDF(html_content, dest=pdf_buffer)

    if pisa_status.err:
        return {"status": "error", "message": "PDF generation failed"}

    pdf_buffer.seek(0)
    
    #SIGNATURE
    signed_pdf_buffer = sign_pdf(pdf_buffer, company, document_type.name, document_type.token)
    
    return {"status": "success", "pdf_content": signed_pdf_buffer.getvalue()}
  
  
@shared_task
def generate_and_upload_xml( payment_id, bank_id, request_data ):
    from billing.views.epayment_document_generate_view import generate_xml
  
    try:
        print("generate_and_upload_xml")
        payment = Payment.objects.get(pk=payment_id)
        invoice = payment.invoice
        try:
          xml_data = generate_xml(payment)  
        except Exception as e:
          raise Exception(f"Error generating XML for payment {payment_id}: {e}")
        xml_file_name = f"{payment.invoice.serie_final.replace('/', '_')}_{payment.invoice.id}.xml"
        
        xml_file = ContentFile(xml_data.encode("utf-8"), name=xml_file_name)
        document_file = upload_document(xml_file, 'REBUTS', 'EFACTURA', invoice.id, invoice.customer_token_final, '', settings.DOCUMENT_MANAGER_SERVICES.get("billing"), xml_file_name)
        
        return document_file.id
    except Exception as e:
        print(f"Error in generate_and_upload_xml task: {e}")
        return None
      
      
@shared_task
def generate_invoice_pdf(invoice_id, find_existing=False):
  from billing.views.invoice_pdf_view import generate_report_invoice_pdf
  
  try:
      #invoice = Invoice.objects.get(id=invoice_id)
      if isinstance(invoice_id, int):
          invoice = Invoice.objects.get(id=invoice_id)
      else:
          invoice = invoice_id
        
      if find_existing:
        if invoice.invoice_file and invoice.invoice_file.id:
          return invoice.invoice_file.id

      file_url, invoice_file_id, _ = generate_report_invoice_pdf(invoice)
      
      return invoice_file_id
      
  except Exception as e:
      print(f"Error generating PDF for invoice {invoice_id}: {e}")
      return None


      
      
      
@shared_task(bind=True)
def regenerate_billing_pdfs(self, billing_id, queue_item_id=None):
  from billing.views.invoice_pdf_view import generate_report_invoice_pdf
  from billing.models import BillingQueue
  from django.db import close_old_connections
  from django.utils import timezone

  progress_recorder = ProgressRecorder(self)

  q_item = None
  if queue_item_id:
    try:
      q_item = BillingQueue.objects.get(id=queue_item_id)
    except BillingQueue.DoesNotExist:
      pass

  try:
    invoice_ids = list(Invoice.objects.filter(billing_id=billing_id, is_active=True).values_list('id', flat=True))
    total = len(invoice_ids)

    if q_item:
      q_item.total_items = total
      q_item.save(update_fields=['total_items'])

    if total == 0:
      if q_item:
        _close_billing_queue_item(q_item.id, self.request.id, status='completed', completed_at=timezone.now(), processed_items=0)
      return {'billing_id': billing_id, 'total': 0, 'generated': 0}

    progress_recorder.set_progress(0, total)
    generated = 0

    for i, invoice_id in enumerate(invoice_ids):
      try:
        close_old_connections()
        invoice = Invoice.objects.get(id=invoice_id)
        generate_report_invoice_pdf(invoice)
        generated += 1
      except Exception as e:
        print(f"Error regenerating PDF for invoice {invoice_id}: {e}")
      finally:
        close_old_connections()
        progress_recorder.set_progress(i + 1, total)

    if q_item:
      _close_billing_queue_item(q_item.id, self.request.id, status='completed', completed_at=timezone.now(), processed_items=generated)

    return {'billing_id': billing_id, 'total': total, 'generated': generated}

  except Exception as e:
    if q_item:
      _close_billing_queue_item(q_item.id, self.request.id, status='failed', completed_at=timezone.now(), error_message=str(e))
    raise Exception(f"Error in regenerate_billing_pdfs: {e}")


@shared_task(bind=True)
def export_wincen_task(self, billing_id, queue_item_id=None):
  from billing.models import BillingQueue
  from billing.utils.wincen_service import generate_wincen_file
  from django.utils import timezone

  progress_recorder = ProgressRecorder(self)

  q_item = None
  if queue_item_id:
    try:
      q_item = BillingQueue.objects.get(id=queue_item_id)
    except BillingQueue.DoesNotExist:
      pass

  try:
    billing = Billing.objects.get(id=billing_id)
    total = Invoice.objects.filter(billing_id=billing_id, is_active=True, is_excluded=False).count()

    if q_item:
      q_item.total_items = total
      q_item.save(update_fields=['total_items'])

    if total == 0:
      if q_item:
        _close_billing_queue_item(q_item.id, self.request.id, status='completed', completed_at=timezone.now(), processed_items=0)
      return {'billing_id': billing_id, 'total': 0, 'file_url': None}

    progress_recorder.set_progress(0, total)

    processed = [0]

    def on_progress(current, t):
      processed[0] = current
      progress_recorder.set_progress(current, t)

    file_content = generate_wincen_file(
      billing_id,
      include_readings=True,
      include_payments=True,
      progress_callback=on_progress,
    )

    filename = f'wincen_{billing_id}_{timezone.now().strftime("%Y%m%d_%H%M%S")}.dat'

    from statistics.views.reports_views import save_report
    document_id = save_report(
      content=file_content,
      filename=filename,
      name=f'WinCen - {billing.name or billing_id}',
      type_id=None,
      start_date=None,
      end_date=None,
    )

    file_url = None
    if document_id:
      from documentmanager.models import Document
      try:
        doc = Document.objects.get(id=document_id)
        file_url = doc.location_url or (doc.file.url if doc.file else None)
      except Document.DoesNotExist:
        pass

    if q_item:
      _close_billing_queue_item(
        q_item.id, self.request.id,
        status='completed', completed_at=timezone.now(), processed_items=processed[0], file_url=file_url,
      )

    return {'billing_id': billing_id, 'total': total, 'file_url': file_url, 'filename': filename, 'document_id': document_id}

  except Exception as e:
    if q_item:
      _close_billing_queue_item(q_item.id, self.request.id, status='failed', completed_at=timezone.now(), error_message=str(e))
    raise Exception(f"Error in export_wincen_task: {e}")


@shared_task
def check_today_tasks():
  try:
      tasks = CalendarTask.objects.filter(task_done=False, set_date__lte=timezone.now().date()).distinct()

      for task in tasks:
        notification_save = {
          'token': uuid.uuid4(),
          'name': f"Tasca del dia pendent",
          'description': f"Tasca: '{task.name}' pendent de finalització",
          'module': '',
          'entity': '',
          'object_id': None,
          'is_active': True,
          'user': task.user,
        }
        notification = Notification.objects.create(**notification_save)
  except Exception as e:
      print(f"Error checking today's tasks: {e}")
      return None


@shared_task
def send_payment_remittances(payment_remittance_ids, sent_at):
  from billing.models import PaymentRemittance
  today = timezone.now().date()
  
  if sent_at is not None:
    if isinstance(sent_at, str):
      try:
        dt = datetime.fromisoformat(sent_at.replace('Z', '+00:00'))
        sent_at = dt.date()
      except (ValueError, AttributeError):
        parsed = convert_date_with_timezone(sent_at)
        sent_at = parsed.date() if parsed else None
    elif hasattr(sent_at, 'date') and callable(getattr(sent_at, 'date')):
      sent_at = sent_at.date()
    elif not isinstance(sent_at, date):
      sent_at = None
  payment_status_paid = PaymentStatus.objects.get(token=ConfigProject.objects.get(token='payment_status_paid_token').value)
  payment_remittances = PaymentRemittance.objects.filter(id__in=payment_remittance_ids)
  payment_direct_debit = PaymentType.objects.get(token="DIRECT_DEBIT")
  payment_bank_transfer = PaymentType.objects.get(token="BANK_TRANSFER")
  for remittance in payment_remittances:
    for payment in remittance.payments.all():
      if not remittance.is_return and payment.payment_type_token != "DIRECT_DEBIT":
        if payment.invoice:
          disconnect_payment_signals()
          log_invoice_data_change(None, payment.invoice, None, payment_direct_debit.name)
          payment.invoice.payment_type_final = payment_direct_debit.name
          payment.invoice.payment_type_token_final = payment_direct_debit.token
          payment.invoice.save(update_fields=['payment_type_final', 'payment_type_token_final'])
        payment.payment_type_token = payment_direct_debit.token
        payment.payment_type = payment_direct_debit.name
      if remittance.is_return and payment.payment_type_token != "BANK_TRANSFER":
        payment.payment_type_token = payment_bank_transfer.token
        payment.payment_type = payment_bank_transfer.name
      invoice_send_at = payment.invoice.send_at if payment.invoice and payment.invoice.send_at else None
      invoice_send_at_date = invoice_send_at.date() if invoice_send_at else None
      payment_send_date = sent_at if sent_at and not invoice_send_at_date else invoice_send_at_date if invoice_send_at_date else today
      """ if invoice_send_at_date and sent_at:
        if sent_at > invoice_send_at_date: """
      payment_send_date = sent_at if sent_at else today
      print("RECONNECTING")
      reconnect_payment_signals()
      payment_date = payment_send_date
      payment.status = payment_status_paid
      payment.paid_at = payment_date
      payment.payment_date = payment_date
      payment.save()


@shared_task
def rerun_invoices(invoice_ids, invoice_type, status_pending_token):
  try:
    start_time = time.time()
    disconnect_payment_signals()
    batch_size = 500
    invoices = Invoice.objects.filter(id__in=invoice_ids).order_by('issue_date')
    from django.db.models import Max
    max_serie = Invoice.objects.filter(
        issue_date__year=2025, 
        type_final=invoice_type, 
    ).exclude(
        status__token=status_pending_token
    ).aggregate(max_serie=Max('serie_final'))['max_serie']
    
    next_num = 1
    if max_serie:
        try:
            last_part = max_serie.split('/')[-1]
            current_num = int(last_part)
            next_num = current_num + 1
        except (ValueError, IndexError):
            next_num = 1
    
    for invoice in invoices:
      serie_final = f"{invoice.serie_final}{next_num:06d}" if invoice.serie_final else None
      
      invoice.serie_final = serie_final
      invoice.save(update_fields=['serie_final'])
      generate_report_invoice_pdf(invoice)
      next_num += 1
    end_time = time.time()
    print(f"Time taken: {end_time - start_time} seconds")
    reconnect_payment_signals()
    
  except Exception as e:
    reconnect_payment_signals()
    raise Exception(f"Error rerunning invoices: {e}")

@shared_task
def fix_readings_batch(reading_ids):
  try:
    readings = Reading.objects.filter(id__in=reading_ids)
    
    estimated_readings = readings.filter(is_estimated=True)
    for reading in estimated_readings:
      
      estimated_bag = EstimatedBag.objects.get(supply_point=reading.supply_point, contract=reading.contract)
      if estimated_bag:
        estimated_bag.total_consumption -= reading.calculated_value
        if estimated_bag.total_consumption < 0:
          estimated_bag.total_consumption = 0
        EstimatedBagMovement.objects.create(
          token=uuid.uuid4(),
          amount=reading.calculated_value,
          movement_date=timezone.now().date(),
          estimated_bag=estimated_bag,
          reading=reading,
          is_positive=False
        )
        estimated_bag.save()
    readings.update(is_control=True, batch=None, is_estimated=False)
    for reading in readings:
      same_meter_readings = Reading.objects.filter(meter=reading.meter, reading_date=reading.reading_date, contract=reading.contract, is_control=False)
      if same_meter_readings.count() > 0:
        print(f"Same meter readings found for reading meter {reading.meter.code}: {same_meter_readings.count()}")
        fix_readings_batch.delay(list(same_meter_readings.values_list('id', flat=True)))
    
  except Exception as e:
    raise Exception(f"Error fixing estimated readings batch: {e}")

@shared_task(bind=True)
def generate_electronic_invoices(self, communication_process_id, base_url, invoice_ids=None):
  progress_recorder = ProgressRecorder(self)
  
  sent_status = CommunicationStatus.objects.get(
    token=ConfigProject.objects.get(token='communication_status_sent_token').value
    )
  
  if communication_process_id and not invoice_ids:
    communications = Communication.objects.filter(process__id=communication_process_id, types__token='electronic_inv')
    invoices = Invoice.objects.filter(communications__in=communications)
  else:
    communications = None
    invoices = Invoice.objects.filter(id__in=invoice_ids)
  if not invoices.exists():
      return {
        'status': 'error',
        'message': 'No invoices found',
      }

  counter = 0
  buffer = io.BytesIO()
  from billing.views.epayment_document_generate_view import generate_xml
  with zipfile.ZipFile(buffer, mode='w', compression=zipfile.ZIP_DEFLATED) as zip_file:
      for invoice in invoices:
          xml_data = generate_xml(None, invoice)
          xml_file_name = f"{invoice.serie_final.replace('/', '_')}_{invoice.id}.xml"
          zip_file.writestr(xml_file_name, xml_data)
          progress_recorder.set_progress(counter, invoices.count() + 2)  # +2 because of the zip file, make a bit of margin
          counter += 1

  if communications:
    communications.update(
      status=sent_status
    )
    if communications[0].process:
      if all(communication.status == sent_status for communication in communications[0].process.communications.all()):
        process_finalized = CommunicationProcessStatus.objects.get(token=ConfigProject.objects.get(token='communication_process_status_finalized_token').value)
        communications[0].process.status = process_finalized
        communications[0].process.save()
  
  buffer.seek(0)
  if communications:
    zip_file_name = f"einvoices_{communications[0].process.token}.zip"
  else:
    now_date = timezone.now()
    zip_file_name = f"einvoices_{now_date.strftime('%Y%m%d%H%M%S')}.zip"
  zip_content = ContentFile(buffer.getvalue(), name=zip_file_name)

  temp_rel_path = f"tmp/einvoice/{zip_file_name}"
  saved_path = default_storage.save(temp_rel_path, zip_content)
  file_url = urljoin(base_url, default_storage.url(saved_path))
  delete_file_later(saved_path, delay_seconds=2400)
  progress_recorder.set_progress(100, 100)
  
  response = {
    'file_url': file_url,
    'status': 'ok',
  }
  return response

@shared_task(bind=True)
def generate_massive_invoice_download(self, invoice_ids, base_url, in_zip=True):
  from documentmanager.models import Document
  from documentmanager.utils.main_utils import download_document, download_multiple_documents

  progress_recorder = ProgressRecorder(self)
  invoices = Invoice.objects.filter(id__in=invoice_ids)
  total = invoices.count()
  if not total:
    return {
      'status': 'error',
      'message': 'No invoices found',
    }

  document_ids = []
  for counter, invoice in enumerate(invoices):
    if not invoice.invoice_file:
      _, doc_id, _ = generate_report_invoice_pdf(invoice)
      document_ids.append(doc_id)
    else:
      document_ids.append(invoice.invoice_file.id)
    progress_recorder.set_progress(counter, total + 1)

  documents = Document.objects.filter(id__in=document_ids)
  service = documents[0].service if documents and all(document.service in ['aws', 'azure'] for document in documents) else None

  def iter_document_contents():
    if service:
      for filename, file_content in download_multiple_documents(document_ids, service):
        yield filename, file_content
    else:
      for document in documents:
        document_content = download_document(document)
        try:
          file_content = document_content.getvalue()
        except AttributeError:
          file_content = document_content.content
        yield document.document_name.replace('/', '_'), file_content

  now_date = timezone.now()
  if in_zip:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, mode='w', compression=zipfile.ZIP_DEFLATED) as zip_file:
      for filename, file_content in iter_document_contents():
        zip_file.writestr(filename, file_content)
    file_name = f"invoices_{now_date.strftime('%Y%m%d%H%M%S')}.zip"
    file_content_wrapper = ContentFile(buffer.getvalue(), name=file_name)
  else:
    from PyPDF2 import PdfMerger
    merger = PdfMerger()
    for _filename, file_content in iter_document_contents():
      merger.append(BytesIO(file_content))
    merged_buffer = BytesIO()
    merger.write(merged_buffer)
    merger.close()
    file_name = f"invoices_{now_date.strftime('%Y%m%d%H%M%S')}.pdf"
    file_content_wrapper = ContentFile(merged_buffer.getvalue(), name=file_name)

  temp_rel_path = f"tmp/massive_invoice_download/{file_name}"
  saved_path = default_storage.save(temp_rel_path, file_content_wrapper)
  file_url = urljoin(base_url, default_storage.url(saved_path))
  delete_file_later(saved_path, delay_seconds=2400)
  progress_recorder.set_progress(total, total)

  return {
    'file_url': file_url,
    'status': 'ok',
  }


@shared_task(bind=True)
def generate_sepa_document(self, request_data):
  """
  Background SEPA document generation for `SEPAPaymentDocumentGenerateViewSet.put()`.

  It creates one `PaymentRemittance` per batch, sets status to "processing" while
  generating/uploading the XML, then switches it back to the default status.

  Returns: {status: "ok", document_ids: [...]} (or raises on error).
  """
  try:
    from billing.views.sepa_payment_document_generate_view import (
      RequestData,
      get_bank_distribution,
      get_sepa_payments_and_batches,
    )
    from billing.utils.sepa_file_service import generate_xml_payments
    from billing.models import PaymentRemittance, PaymentRemittanceStatus
    progress_recorder = ProgressRecorder(self)

    request_data = request_data or {}

    processing_token = ConfigProject.objects.get(token="payment_remittance_status_processing_token").value
    processing_status = PaymentRemittanceStatus.objects.get(token=processing_token)
    pending_status = PaymentRemittanceStatus.objects.get(is_default=True)

    # Cada banc seleccionat genera els seus propis fitxers, amb nomes els
    # pagaments que te assignats (veure `get_bank_distribution`). Es resolen tots
    # els lots abans de comencar per poder informar del progres total.
    plan = []
    for bank_id, bank_data in get_bank_distribution(request_data):
      bank = CompanyBank.objects.get(id=bank_id)
      _, payments_batches, is_return = get_sepa_payments_and_batches(RequestData(bank_data))
      for batch, batch_send_date in payments_batches:
        if not batch:
          continue
        plan.append((bank, batch, batch_send_date, is_return))

    if not plan:
      raise ValueError("No payments found to generate SEPA documents.")

    payments_counter = 0

    document_ids = []
    for bank, batch, batch_send_date, is_return in plan:
      document = None
      try:
        try:
          send_date_inst = datetime.strptime(batch_send_date, "%Y-%m-%d")
        except Exception:
          send_date_inst = datetime.now()

        progress_recorder.set_progress(payments_counter, len(plan) + 2)  # +2 because of doc save and object
        payments_counter += 1
        batch_qs = Payment.objects.filter(id__in=[p.id for p in batch]).order_by("-token")
        xml_data, ident_msg = generate_xml_payments(batch_qs, bank, send_date_inst, is_return)
        random_uuid = uuid.uuid4()
        xml_file_name = f"{'SEPA' if not is_return else 'TRF'}_{send_date_inst.year}{send_date_inst.month}{str(random_uuid.int)[:11]}.xml"

        service = settings.DOCUMENT_MANAGER_SERVICES.get("billing")
        xml_file = ContentFile(xml_data.encode("utf-8"), name=xml_file_name)

        document = PaymentRemittance.objects.create(
          token=ident_msg,
          company_bank=bank,
          desired_send_at=send_date_inst,
          status=processing_status,
          is_return=is_return,
        )
        document.task_id = self.request.id
        document.save(update_fields=["task_id"])

        document.payments.set(batch)
        document_file = upload_document(
          xml_file,
          "REBUTS",
          "REBUTS",
          document.id,
          document.token,
          "",
          service,
          xml_file_name,
        )
        document.document = document_file
        document.status = pending_status
        document.task_id = self.request.id
        document.save()

        batch_ids = [p.id for p in batch]
        Payment.objects.filter(id__in=batch_ids).update(document=document_file)

        document_ids.append(document.document.id if document.document else None)
      except Exception:
        # Ensure status doesn't stay stuck in "processing" if generation fails.
        if document:
          document.status = pending_status
          document.task_id = None
          document.save(update_fields=["status", "task_id"])
        raise

    return {"status": "ok", "document_ids": document_ids}
  except Exception as e:
    raise Exception(f"Error generating SEPA documents: {e}")


@shared_task(bind=True)
def regenerate_sepa_file(self, payment_remittance_id, payment_ids, send_date):
  """
  Regenerates the SEPA XML/document for an existing `PaymentRemittance`.

  This is the async counterpart to `PaymentRemittanceRegenerateSepaView.post()`.
  """
  try:
    from billing.models import PaymentRemittance, PaymentRemittanceStatus
    from billing.utils.sepa_file_service import generate_xml_payments

    # Parse send_date similarly to the sync view.
    try:
      send_date_inst = (
        datetime.strptime(send_date, "%Y-%m-%d")
        if isinstance(send_date, str) and send_date
        else (send_date if send_date else datetime.now())
      )
    except Exception:
      send_date_inst = datetime.now()

    processing_token = ConfigProject.objects.get(
      token="payment_remittance_status_processing_token"
    ).value
    processing_status = PaymentRemittanceStatus.objects.get(token=processing_token)
    pending_status = PaymentRemittanceStatus.objects.get(is_default=True)

    payment_remittance = PaymentRemittance.objects.get(id=payment_remittance_id)
    payments_qs = payment_remittance.payments.filter(id__in=payment_ids).order_by("-token")

    if not payments_qs.exists():
      raise ValueError("No payments found to regenerate SEPA file.")

    if not payment_remittance.company_bank:
      raise ValueError("PaymentRemittance has no company_bank.")

    # Mark it as processing for UI consistency.
    payment_remittance.status = processing_status
    payment_remittance.task_id = self.request.id
    payment_remittance.pending_changes = True
    payment_remittance.save(update_fields=["status", "task_id", "pending_changes"])

    bank = payment_remittance.company_bank
    xml_data, ident_msg = generate_xml_payments(payments_qs, bank, send_date_inst, is_return=payment_remittance.is_return)

    date_year = datetime.now().year
    date_month = datetime.now().month
    random_uuid = uuid.uuid4()
    xml_file_name = f"{'SEPA' if not payment_remittance.is_return else 'TRF'}_{date_year}{date_month}{str(random_uuid.int)[:11]}.xml"
    xml_file = ContentFile(xml_data.encode("utf-8"), name=xml_file_name)

    # Update token before uploading the new version of the document.
    payment_remittance.token = ident_msg

    service = settings.DOCUMENT_MANAGER_SERVICES.get("billing")
    document_file = upload_document(
      xml_file,
      "REBUTS",
      "REBUTS",
      payment_remittance.id,
      payment_remittance.token,
      "",
      service,
      xml_file_name,
    )

    payment_remittance.document = document_file
    payment_remittance.status = pending_status
    payment_remittance.pending_changes = False
    payment_remittance.save(update_fields=["token", "document", "status", "pending_changes"])

    payments_qs.update(document=document_file)

    anomalies = [p.id for p in payments_qs if p.amount <= 0]

    return {
      "status": "ok",
      "remittance_id": payment_remittance.id,
      "payment_ids": list(payment_ids),
      "token": ident_msg,
      "document_file_id": document_file.id if document_file else None,
      "anomalies": anomalies if anomalies else None,
    }
  except Exception as e:
    # Best-effort rollback to avoid leaving the remittance stuck in "processing".
    try:
      from billing.models import PaymentRemittance, PaymentRemittanceStatus

      pending_status = PaymentRemittanceStatus.objects.get(is_default=True)
      payment_remittance = PaymentRemittance.objects.get(id=payment_remittance_id)
      payment_remittance.status = pending_status
      payment_remittance.task_id = None
      payment_remittance.save(update_fields=["status", "task_id"])
    except Exception:
      pass
    raise e


@shared_task
def export_readings_by_batch_csv_task(batch_id):
    try:
        from billing.models import Reading
        
        # Filter readings for specified batch and sort by date
        readings = Reading.objects.filter(
            batch_id=batch_id, 
            is_active=True,
            is_control=False
        ).select_related(
            'contract',
            'contract__holder',
            'contract__use_type',
            'meter',
            'supply_point',
            'supply_point__address',
            'supply_point__property__route_position',
        ).order_by('-reading_date')
        
        if not readings.exists():
            return {
                "status": "error",
                "message": "No s'han trobat lectures per aquest batch"
            }
        
        # Create CSV in memory
        buffer = io.StringIO()
        # Add BOM for UTF-8
        buffer.write('\ufeff')
        
        writer = csv.writer(buffer, delimiter=';')
        
        # Headers
        headers = [
            str(_('Data de lectura')),
            str(_('Contracte')),
            str(_('Titular')),
            str(_("Tipus d'ús")),
            str(_('Punt de subministrament')),
            str(_('Codi comptador')),
            str(_('Codi posició ruta')),
            str(_('Lectura')),
            str(_('Consum')),
            str(_('Origen')),
            str(_('Consum anterior 1')),
            str(_('Consum anterior 2')),
            str(_('Consum anterior 3')),
            str(_('Consum anterior 4'))
        ]
        writer.writerow(headers)
        
        # Write data
        for reading in readings:
            prev_consumptions = []
            if reading.contract:
                prev_readings = Reading.objects.filter(
                    contract_id=reading.contract_id,
                    reading_date__lt=reading.reading_date,
                    is_active=True,
                    is_control=False
                ).order_by('-reading_date')[:4]
                for pr in prev_readings:
                    prev_consumptions.append(pr.calculated_value if pr.calculated_value is not None else '')
            
            while len(prev_consumptions) < 4:
                prev_consumptions.append('')

            route_position = None
            if reading.supply_point and reading.supply_point.property:
                route_position = reading.supply_point.property.route_position
            
            row = [
                reading.reading_date.strftime('%d/%m/%Y') if reading.reading_date else '',
                reading.contract.token if reading.contract else '',
                str(reading.contract.holder) if reading.contract and reading.contract.holder else '',
                reading.contract.use_type.name if reading.contract and reading.contract.use_type and reading.contract.use_type.name else '',
                str(reading.supply_point.address) if reading.supply_point and reading.supply_point.address else '',
                reading.meter.code if reading.meter and reading.meter.code else '',
                route_position.token if route_position and route_position.token else '',
                reading.reading_value if reading.reading_value is not None else '',
                reading.calculated_value if reading.calculated_value is not None else '',
                reading.origin if reading.origin else ''
            ] + prev_consumptions
            writer.writerow(row)
            
        csv_text = buffer.getvalue()
        file_bytes = csv_text.encode('utf-8')
        
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"readings_batch_{batch_id}_{timestamp}.csv"
        relative_path = f"tmp/EXTRA/{filename}"
        
        saved_path = default_storage.save(relative_path, ContentFile(file_bytes))
        delete_file_later(saved_path, delay_seconds=3600)
        
        return {
            "status": "ok",
            "file_url": default_storage.url(saved_path),
            "filename": filename,
        }
    except Exception as e:
        return {
            "status": "error",
            "message": f"Error intern del servidor: {str(e)}"
        }


@shared_task
def export_reading_batch_supply_points_csv_task(batch_id):
    """
    Exporta els SupplyPoints de les Routes d'un ReadingBatch amb la lectura del lot
    (si n'hi ha) i els contractes associats (columnes dinàmiques Contract[n][...]).
    """
    try:
        from service.utils.supply_points_route_csv_export import (
            build_reading_batch_supply_points_csv_bytes,
        )

        file_bytes, error_message = build_reading_batch_supply_points_csv_bytes(batch_id)
        if error_message:
            return {
                "status": "error",
                "message": error_message,
            }

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"supply_points_batch_{batch_id}_{timestamp}.csv"
        relative_path = f"tmp/EXTRA/{filename}"

        saved_path = default_storage.save(relative_path, ContentFile(file_bytes))
        delete_file_later(saved_path, delay_seconds=3600)

        return {
            "status": "ok",
            "file_url": default_storage.url(saved_path),
            "filename": filename,
        }
    except Exception as e:
        return {
            "status": "error",
            "message": f"Error intern del servidor: {str(e)}"
        }


@shared_task
def fix_invoices_persons_task():
  from billing.utils.invoice_service import (
      find_person_for_invoice_customer,
      person_candidates_for_customer_final,
      person_matches_customer_final,
  )
  invoice_type_invoice_token = ConfigProject.objects.get(token='invoice_type_invoice_token').value
  invoices = Invoice.objects.filter(
      type_final=invoice_type_invoice_token,
      person__isnull=True
  )
  none_valid = ['00000000T', '00000001R', '99999999R', 'X0000000T']
  
  updating_invoices = []
  counter = 0
  
  person_multiple_times = []
  not_found_person = []
  
  for invoice in invoices:
      if len(updating_invoices) > 500:
          Invoice.objects.bulk_update(updating_invoices, ['person'])
          updating_invoices = []
          print(f"UPDATED {counter} OUT OF {invoices.count()} INVOICES")
      counter += 1
      person = None
      if invoice.contract:
          person = invoice.contract.holder
      elif invoice.contract_request:
          person = invoice.contract_request.holder
      elif invoice.contract_termination:
          person = invoice.contract_termination.contract.holder
      elif invoice.connection_request:
          person = invoice.connection_request.person
      
      candidates = []
      
      if not person:
          person = find_person_for_invoice_customer(invoice)
          candidates = person_candidates_for_customer_final(
              invoice.customer_final,
              exclude_token=person.token if person else None,
          )
      
      if not person:
          not_found_person.append(invoice)
          continue
      
      multiple_persons_found = [
          candidate for candidate in candidates
          if person_matches_customer_final(candidate, invoice.customer_final)
      ]

      if multiple_persons_found:
          person_multiple_times.append({
              'invoice': invoice,
              'person': person,
              'others': multiple_persons_found,
          })
          continue

      invoice.person = person
      updating_invoices.append(invoice)
  print("\n\n")
  if updating_invoices:
      pass
      Invoice.objects.bulk_update(updating_invoices, ['person'])

  print(f"FIXED {counter - len(not_found_person) - len(person_multiple_times)} INVOICES")
  for invoice in invoices[:5]:
      print(invoice.serie_final)
  print(f"NOT FOUND PERSON: {len(not_found_person)}")
  for invoice in not_found_person:
      print(f"Invoice {invoice.id}: {invoice.serie_final}, customer final: {invoice.customer_final}, customer token final: {invoice.customer_token_final}")
  print(f"MULTIPLE PERSONS FOUND: {len(person_multiple_times)}")
  for invoice in person_multiple_times[:5]:
      print(invoice['invoice'].serie_final)
      print(invoice['person'].name)
      print(invoice['others'])


@shared_task
@billing_run_cache_scope()
def process_selected_contracts_task(billing_id, contract_ids):
    print(f"Start process_selected_contracts_task for billing={billing_id}, contracts={contract_ids}")
    try:
        billing = Billing.objects.get(id=billing_id)
    except Billing.DoesNotExist:
        print(f"Billing {billing_id} not found.")
        return

    fanout_general_meter_readings(billing)

    PERIOD_MONTHS_MAP = {
        'trimestral': 90,
        'semestral': 180,
        'bimestral': 60,
        'quadrimestral': 120,
        'anual': 360,
        'mensual': 30,
    }
    biller = billing.biller
    billing_period_days = PERIOD_MONTHS_MAP.get(biller.period_type, 90)
    period_months = billing_period_days / 30

    try: 
        limit = int(ConfigProject.objects.get(token='reading_estimate_limit_percentage').value)
    except Exception:
        limit = 33

    from django.db.models import Min, Max
    reading_dates = billing.readings.aggregate(
        min_reading=Min('reading_date'),
        min_previous=Min('previous_reading__reading_date'),
        max_reading=Max('reading_date')
    )

    # Find or create ReadingBatch for billing_missing
    ref_batch = ReadingBatch.objects.filter(billing_missing=billing).first()
    if not ref_batch:
        first_billing_reading = billing.readings.select_related('batch').first()
        if first_billing_reading and first_billing_reading.batch:
            ref_batch = first_billing_reading.batch
        else:
            try:
                default_status = ReadingBatchStatus.objects.get(is_default=True)
            except ReadingBatchStatus.DoesNotExist:
                default_status = ReadingBatchStatus.objects.first()

            start_date = reading_dates['min_previous'] or reading_dates['min_reading'] or (date.today() - timedelta(days=billing_period_days))
            end_date = reading_dates['max_reading'] or date.today()

            ref_batch = ReadingBatch.objects.create(
                token=billing.token,
                name=billing.name,
                include_telecontrol=True,
                include_manual=True,
                status=default_status,
                billing_missing=billing,
                missing_start=start_date,
                missing_end=end_date,
            )

    first_start_reading = ref_batch.missing_start if ref_batch and ref_batch.missing_start else (reading_dates['min_previous'] or reading_dates['min_reading'])
    last_end_reading = (ref_batch.missing_end - timedelta(days=(billing_period_days * (limit/100)))) if ref_batch and ref_batch.missing_end else reading_dates['max_reading']

    target_date = ref_batch.missing_end if ref_batch and ref_batch.missing_end else last_end_reading
    if not target_date:
        target_date = date.today()
    if isinstance(target_date, datetime):
        target_date = target_date.date()

    for contract_id in contract_ids:
        try:
            contract = Contract.objects.get(id=contract_id)
        except Contract.DoesNotExist:
            print(f"Contract {contract_id} not found.")
            continue

        # Check if contract has an active reading in this period
        reading = Reading.objects.filter(
            contract=contract,
            is_initial=False,
            is_control=False,
            is_active=True,
            invoices__isnull=True,
        ).filter(
            reading_date__gte=contract.created_at.date() if contract.created_at else contract.registration_date
        )

        if first_start_reading and last_end_reading:
            margin_days = timedelta(days=(billing_period_days * (limit/100)))
            reading = reading.filter(
                reading_date__range=(first_start_reading, last_end_reading + margin_days)
            )

        reading = reading.order_by('-reading_date', '-id').first()

        # If no reading, generate estimated reading
        if not reading:
            supply_point = contract.supply_point_default
            if not supply_point:
                supply_point = contract.supply_points.first()
            if not supply_point:
                print(f"No supply point for contract {contract_id}")
                continue

            from billing.utils.reading_service import get_estimated_reading_minimal_object
            reading = get_estimated_reading_minimal_object(
                supply_point=supply_point,
                contract=contract,
                date=target_date,
                batch=ref_batch,
                offset_days=billing_period_days,
                max_day_limit=date.today(),
                create_reading=True
            )

        if not reading:
            print(f"Could not find or generate reading for contract {contract_id}")
            continue

        # Ensure reading is associated to ref_batch and billing
        if reading.batch != ref_batch or reading.billing != billing:
            reading.batch = ref_batch
            reading.billing = billing
            reading.save()

        # Process invoice
        contract_readings = list(billing.readings.filter(
            contract=contract,
            is_control=False,
            is_active=True,
            is_close=False
        ).select_related('contract', 'contract__status', 'supply_point', 'meter', 'batch', 'billing').exclude(
            invoices__type_final='F'
        ).order_by('reading_date').distinct())

        if not contract_readings:
            print(f"No active readings for contract {contract_id} in billing {billing_id}")
            continue

        # Deduplicate readings
        set_added_readings = set()
        filtered_readings = []
        for r in contract_readings:
            if (contract.token, r.meter.code) in set_added_readings:
                continue
            set_added_readings.add((contract.token, r.meter.code))
            filtered_readings.append(r)

        from billing.utils.invoice_service import generate_consumption_invoice_multiple
        billing_batch = billing.billing_batches.first()
        invoice = generate_consumption_invoice_multiple(
            contract=contract,
            readings=filtered_readings,
            title=billing.name,
            billing=billing,
            billing_batch=billing_batch,
            period_months=period_months
        )

        if not invoice:
            print(f"Failed to generate invoice for contract {contract_id}")
            continue

        if contract.general_invoice:
            general_invoice = contract.general_invoice
            general_contracts = list(Contract.objects.filter(general_invoice=general_invoice))
            invoice.general_contracts.set(general_contracts)

        # Apply warnings
        from billing.models import InvoiceWarning
        try:
            invoice_warning_negative = InvoiceWarning.objects.filter(token=ConfigProject.objects.get(token='invoice_warning_negative').value).first()
            invoice_warning_zero = InvoiceWarning.objects.filter(token=ConfigProject.objects.get(token='invoice_warning_zero').value).first()
            invoice_warning_bank_missing = InvoiceWarning.objects.filter(token=ConfigProject.objects.get(token='invoice_warning_bank_missing').value).first()
            invoice_warning_payment_missing = InvoiceWarning.objects.filter(token=ConfigProject.objects.get(token='invoice_warning_payment_missing').value).first()
            invoice_warning_simplified_over_400 = InvoiceWarning.objects.filter(token=ConfigProject.objects.get(token='invoice_warning_simplified_over_400').value).first()
            invoice_warning_high_amount = InvoiceWarning.objects.filter(token=ConfigProject.objects.get(token='invoice_warning_high_amount').value).first()
        except Exception:
            invoice_warning_negative = None
            invoice_warning_zero = None
            invoice_warning_bank_missing = None
            invoice_warning_payment_missing = None
            invoice_warning_simplified_over_400 = None
            invoice_warning_high_amount = None

        warning = None
        if invoice.total_final <= 0:
            warning = invoice_warning_negative
        elif invoice.total_final == 0:
            warning = invoice_warning_zero
        elif invoice.payment_type is None:
            warning = invoice_warning_payment_missing
        elif invoice.payment_type.token == 'DIRECT_DEBIT' and invoice.payment_bank is None:
            warning = invoice_warning_bank_missing
        elif invoice.simplified and invoice.total_final >= 400:
            warning = invoice_warning_simplified_over_400
        elif invoice.total_final and invoice_warning_high_amount:
            if is_high_amount_for_period(invoice.total_final, [contract], invoice.billing_period_year, invoice.billing_period_month, exclude_invoice_id=invoice.id):
                warning = invoice_warning_high_amount

        if warning:
            invoice.warning = warning
            invoice.save(update_fields=['warning'])

        invoice.billing = billing
        invoice.save()

    print(f"Finished process_selected_contracts_task for billing={billing_id}")
def process_next_queue_item():
  from billing.models import BillingQueue
  from django.utils import timezone
  
  # Check if any task is currently running
  if BillingQueue.objects.filter(status='running').exists():
    return

  # Get next pending task
  next_item = BillingQueue.objects.filter(status='pending').order_by('created_at').first()
  if not next_item:
    return

  task_by_type = {
    'PRE_INVOICE': process_billing_batch,
    'DEFINITIVE_INVOICE': process_invoice_documents,
    'REGENERATE_PDFS': regenerate_billing_pdfs,
    'WINCEN_EXPORT': export_wincen_task,
  }
  task = task_by_type.get(next_item.task_type)

  # El task_id es genera abans d'encuar i es desa amb l'estat `running` en un sol pas:
  # així un kill fet just després d'engegar l'item sempre troba quina tasca ha de matar,
  # i la tasca pot comprovar en acabar que l'item continua sent seu.
  task_id = str(uuid.uuid4()) if task else None
  next_item.status = 'running'
  next_item.started_at = timezone.now()
  next_item.task_id = task_id
  next_item.save()

  # Trigger task
  try:
    if next_item.task_type == 'DEFINITIVE_INVOICE':
      payload = next_item.payload or {}
      args = [
        next_item.billing_id,
        payload.get('invoice_ids', []),
        payload.get('context', {}),
        payload.get('issue_date'),
        payload.get('end_date'),
        payload.get('send_at'),
      ]
    else:
      args = [next_item.billing_id]
    if task:
      Billing.objects.filter(id=next_item.billing_id).update(task_id=task_id)
      task.apply_async(args=args, kwargs={'queue_item_id': next_item.id}, task_id=task_id)
  except Exception as e:
    next_item.status = 'failed'
    next_item.completed_at = timezone.now()
    next_item.error_message = str(e)
    next_item.save()
    # Trigger next task recursively
    process_next_queue_item()


def _close_billing_queue_item(queue_item_id, task_id, **fields):
  """
  Tanca l'item de la cua amb `fields` (status, completed_at...) només si continua en
  `running` i assignat a aquesta tasca, i llavors avança la cua. Si mentre corria l'han
  matat, saltat o reiniciat des del frontal, l'item ja no és d'aquesta tasca: no se'n
  trepitja l'estat ni s'engega el següent, que ja ho ha fet `apply_billing_queue_action`.
  Retorna si l'ha tancat.
  """
  from billing.models import BillingQueue

  closed = BillingQueue.objects.filter(id=queue_item_id, status='running', task_id=task_id).update(**fields)
  if closed:
    process_next_queue_item()
  return bool(closed)


# Estat del Billing mentre corre cada tipus de tasca (el que hi posen les vistes en encuar-la)
# i on es torna quan s'atura a mitges, perquè es pugui revisar i tornar a llançar.
BILLING_QUEUE_STATUS_TOKENS = {
  'PRE_INVOICE': ('billing_batch_processing', 'billing_batch_pending'),
  'DEFINITIVE_INVOICE': ('billing_batch_processing_documents', 'billing_batch_pending'),
}


def _sync_billing_status_after_queue_action(queue_item, action, previous_status):
  tokens = BILLING_QUEUE_STATUS_TOKENS.get(queue_item.task_type)
  if not tokens:
    return
  running_config, stopped_config = tokens
  running_token = ConfigProject.objects.get(token=running_config).value
  billing = queue_item.billing

  if action == 'restart':
    target_token = running_token
  elif previous_status in ('pending', 'running') and billing.status and billing.status.token == running_token:
    target_token = ConfigProject.objects.get(token=stopped_config).value
  else:
    return

  Billing.objects.filter(id=billing.id).update(status=BillingStatus.objects.get(token=target_token))


def apply_billing_queue_action(queue_item, action):
  """
  Aplica una acció manual (kill/skip/restart) sobre un BillingQueue.
  - kill: mata la tasca de Celery (si està en running) i el marca com a failed.
  - skip: mata la tasca (si està en running) i el marca com a skipped, sense reintentar-lo.
  - restart: mata la tasca (si està en running) i el torna a encuar des de zero (status=pending),
    conservant el `payload`/`task_type` originals. Un cop marcat com a pending,
    `process_next_queue_item` el recollirà quan li toqui el torn (ordenat per `created_at`,
    i sempre que no hi hagi cap altre item en `running`).

  Si l'ordre no arriba a Celery es llança `QueueTaskRevokeError` sense tocar res: marcar
  l'item com a aturat engegaria el següent mentre aquest continua corrent.

  Matar no desfà la feina feta: les prefactures o factures ja desades es queden. Per això,
  en aturar un PRE_INVOICE o un DEFINITIVE_INVOICE el Billing torna a "Pendent de
  confirmació" (on es poden revisar i tornar a llançar) en lloc de quedar-se en "Processant".
  """
  from django.utils import timezone
  from customers.queue_utils import revoke_queue_task

  if action not in ('kill', 'skip', 'restart'):
    raise ValueError(f"Acció no vàlida: {action}")

  previous_status = queue_item.status
  if previous_status == 'running':
    revoke_queue_task(queue_item.task_id)

  now = timezone.now()
  if action == 'kill':
    queue_item.status = 'failed'
    queue_item.error_message = 'Tasca aturada manualment per un usuari.'
    queue_item.completed_at = now
  elif action == 'skip':
    queue_item.status = 'skipped'
    queue_item.error_message = queue_item.error_message or 'Tasca saltada manualment per un usuari.'
    queue_item.completed_at = now
  elif action == 'restart':
    queue_item.status = 'pending'
    queue_item.task_id = None
    queue_item.error_message = None
    queue_item.started_at = None
    queue_item.completed_at = None

  queue_item.save()
  _sync_billing_status_after_queue_action(queue_item, action, previous_status)

  # Allibera/avança la cua FIFO (només actua si no hi ha cap altre item en running).
  process_next_queue_item()
