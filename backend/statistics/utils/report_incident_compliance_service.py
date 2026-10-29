"""
Informe de compliment del servei d'atenció (incidències i trucades): volums,
terminis, temps de resposta i resolució en dies hàbils, reobertures i
temps d'atenció telefònica. Els objectius de qualitat els fixa
cada client; l'informe només en dona les dades.
"""
import datetime
import traceback

import openpyxl
from django.db.models import Q
from django.db.models.functions import Coalesce
from django.utils import timezone
from django.utils.translation import gettext as _

from coredata.models import CallRegister
from coredata.utils.working_days import count_working_days
from logger.models import LogIncidentStatusChange
from notification.models import Incident
from notification.utils.incident_utils import get_closed_status_token, incident_exploitation, is_overdue
from statistics.utils.report_advanced_billing_service import parse_date
from statistics.views.reports_views import add_row, adjust_column_widths, save_report

# Temps màxim d'espera telefònica de referència (art. 8.2 i 10.3: 95 % en menys de 3 minuts)
PHONE_WAIT_SECONDS = 180


def _pct(part, total):
    return round(part * 100 / total, 2) if total else 0.0


def _avg(values):
    return round(sum(values) / len(values), 2) if values else ''


def _local(value):
    return timezone.localtime(value).strftime('%d/%m/%Y %H:%M') if value else ''


def generate_incident_compliance_report(request, black_fill, white_bold_font, task=None):
    try:
        date_range = request.data.get('date_range')
        name = request.data.get('name', "Compliment del servei d'atenció")
        type_id = request.data.get('type_id')
        exploitation_id = request.data.get('exploitation_id')
        if not date_range:
            return None, None, None
        start_date = parse_date(date_range[0], make_aware=True)
        end_date = parse_date(date_range[1], make_aware=True)
        if not start_date or not end_date:
            return None, None, None
        # Límit exclusiu: l'inici del dia següent a la data final, tant si arriba
        # com a data sola com amb l'hora de final de dia
        end_limit = timezone.make_aware(datetime.datetime.combine(end_date.date() + datetime.timedelta(days=1), datetime.time.min))

        closed_token = get_closed_status_token()
        now = timezone.now()
        incidents = Incident.objects.filter(
            created_at__gte=start_date, created_at__lt=end_limit, is_active=True,
        ).select_related(
            'type', 'status', 'channel', 'assigned_to',
            'contract__supply_point_default__connection__exploitation',
            'supply_point__connection__exploitation', 'cluster__connection__exploitation',
        )
        # Trucades entrants de la centraleta (incoming_at) o registrades al formulari (call_start / time_call),
        # amb la data de creació com a darrera alternativa, igual que a la cronologia
        calls = CallRegister.objects.annotate(
            received_at=Coalesce('incoming_at', 'call_start', 'time_call', 'created_at'),
        ).filter(
            Q(direction=CallRegister.DIRECTION_INBOUND) | Q(contacted_by_client=True),
            received_at__gte=start_date, received_at__lt=end_limit,
        ).exclude(channel='EMAIL')
        if exploitation_id and exploitation_id != 'all':
            # Mateixes rutes que incident_exploitation: contracte, punt de subministrament o agrupació
            incidents = incidents.filter(
                Q(contract__supply_point_default__connection__exploitation_id=exploitation_id)
                | Q(supply_point__connection__exploitation_id=exploitation_id)
                | Q(cluster__connection__exploitation_id=exploitation_id)
            )
            calls = calls.filter(contract__supply_point_default__connection__exploitation_id=exploitation_id)

        reopened_ids = set(
            LogIncidentStatusChange.objects.filter(
                object__in=incidents, previous_status__token=closed_token,
            ).exclude(current_status__token=closed_token).values_list('object_id', flat=True)
        )

        groups = {'type': {}, 'channel': {}}
        rows = []
        total = incidents.count()
        for counter, incident in enumerate(incidents.iterator(chunk_size=500), start=1):
            if task and counter % 50 == 0:
                task.update_state(state='PROGRESS', meta={
                    'current': counter, 'total': total, 'percent': round(counter * 100 / total, 2) if total else 0.0})
            created = timezone.localtime(incident.created_at).date()
            exploitation = incident_exploitation(incident)
            response_days = count_working_days(created, timezone.localtime(incident.first_response_at).date(), exploitation) if incident.first_response_at else None
            resolution_days = count_working_days(created, timezone.localtime(incident.closed_at).date(), exploitation) if incident.closed_at else None
            on_time = incident.closed_on_time
            overdue = is_overdue(incident, closed_token, now)

            for key, label in (('type', (incident.type.name if incident.type else None) or _("Sense tipus")),
                               ('channel', (incident.channel.name if incident.channel else None) or _("Sense canal"))):
                data = groups[key].setdefault(label, {
                    'total': 0, 'closed': 0, 'on_time': 0, 'late': 0, 'overdue': 0, 'reopened': 0,
                    'response': [], 'resolution': [],
                })
                data['total'] += 1
                data['closed'] += 1 if incident.closed_at else 0
                data['on_time'] += 1 if on_time is True else 0
                data['late'] += 1 if on_time is False else 0
                data['overdue'] += 1 if overdue else 0
                data['reopened'] += 1 if incident.id in reopened_ids else 0
                if response_days is not None:
                    data['response'].append(response_days)
                if resolution_days is not None:
                    data['resolution'].append(resolution_days)

            rows.append([
                incident.token, incident.name or '', incident.type.name if incident.type else '',
                incident.channel.name if incident.channel else '', incident.status.name if incident.status else '',
                incident.status.get_phase_display() if incident.status and incident.status.phase else '',
                incident.assigned_to.username if incident.assigned_to else '',
                _local(incident.created_at), _local(incident.first_response_at), _local(incident.due_at),
                _local(incident.closed_at),
                '' if response_days is None else response_days,
                '' if resolution_days is None else resolution_days,
                '' if on_time is None else (_("Sí") if on_time else _("No")),
                _("Sí") if overdue else '', _("Sí") if incident.id in reopened_ids else '',
            ])

        wb = openpyxl.Workbook()
        headers = [_("Total"), _("Tancades"), _("Dins de termini"), _("Fora de termini"), _("Dins de termini (%)"),
                   _("Pendents vençudes"), _("Reobertes"), _("Mitjana 1a resposta (dies hàbils)"),
                   _("Mitjana resolució (dies hàbils)")]
        for index, (key, title, first_header) in enumerate((
            ('type', _("Per tipus"), _("Tipus")),
            ('channel', _("Per canal"), _("Canal")),
        )):
            sheet = wb.active if index == 0 else wb.create_sheet()
            sheet.title = title
            row = add_row(sheet, 0, [first_header] + headers, fill=black_fill, font=white_bold_font)
            for label, data in sorted(groups[key].items()):
                with_term = data['on_time'] + data['late']
                row = add_row(sheet, row, [
                    label, data['total'], data['closed'], data['on_time'], data['late'],
                    _pct(data['on_time'], with_term), data['overdue'], data['reopened'],
                    _avg(data['response']), _avg(data['resolution']),
                ])
            adjust_column_widths(sheet)

        detail = wb.create_sheet(title=_("Detall"))
        row = add_row(detail, 0, [
            _("Referència"), _("Títol"), _("Tipus"), _("Canal"), _("Estat"), _("Fase"), _("Assignada a"),
            _("Registre"), _("Primera resposta"), _("Data límit"), _("Tancament"),
            _("Dies hàbils fins a resposta"), _("Dies hàbils fins a resolució"), _("Dins de termini"),
            _("Vençuda"), _("Reoberta"),
        ], fill=black_fill, font=white_bold_font)
        for values in rows:
            row = add_row(detail, row, values)
        adjust_column_widths(detail)

        # Atenció telefònica: l'espera surt de les marques de la centraleta o del temps d'espera del formulari
        answered = calls.filter(Q(answered_at__isnull=False) | Q(answered=True)).exclude(abandoned=True).exclude(abandoned_call=True)
        waits = []
        for call in answered.only('incoming_at', 'answered_at', 'wait_time'):
            if call.incoming_at and call.answered_at:
                waits.append((call.answered_at - call.incoming_at).total_seconds())
            elif call.wait_time is not None:
                waits.append(call.wait_time * 60)
        total_calls = calls.count()
        within = len([wait for wait in waits if wait <= PHONE_WAIT_SECONDS])
        phone = wb.create_sheet(title=_("Atenció telefònica"))
        row = add_row(phone, 0, [_("Indicador"), _("Valor")], fill=black_fill, font=white_bold_font)
        for label, value in (
            (_("Trucades entrants"), total_calls),
            (_("Ateses"), answered.count()),
            (_("Abandonades"), calls.filter(Q(abandoned=True) | Q(abandoned_call=True)).count()),
            (_("Ateses en menys de 3 minuts"), within),
            # Les ateses sense espera coneguda no es poden comptar com a ateses a temps
            (_("% ateses en menys de 3 minuts"), _pct(within, answered.count())),
            (_("Espera mitjana (segons)"), _avg(waits)),
            (_("Transferides"), calls.filter(
                Q(transfered_to__isnull=False) | (Q(transfered_to_name__isnull=False) & ~Q(transfered_to_name=''))
                | (Q(transferred_to__isnull=False) & ~Q(transferred_to=''))
            ).count()),
            (_("Peticions de supervisor"), calls.filter(supervisor_requested=True).count()),
            (_("Peticions de trucada de retorn"), calls.filter(callback_requested=True).count()),
        ):
            row = add_row(phone, row, [label, value])
        adjust_column_widths(phone)

        filename = f"report_compliment_sac_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        document_id = save_report(content=wb, filename=filename, name=name, type_id=type_id,
                                  start_date=start_date, end_date=end_date)
        return None, document_id, filename
    except Exception as e:
        print(f"Error in generate_incident_compliance_report: {e}")
        traceback.print_exc()
        return None, None, None
