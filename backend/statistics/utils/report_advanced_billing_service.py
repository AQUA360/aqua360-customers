import datetime
from django.db.models import Sum, Q, Count, F, Min, Max
from billing.models import Invoice, InvoiceLineItem
from contract.models import Contract, ContractStatus
from service.models import Exploitation, SupplyPoint
from pricing.models import Product, ProductOrigin, LineItemType, BillingPeriod
from coredata.models import ConfigProject
from statistics.views.reports_views import add_row, adjust_column_widths, filter_pending_invoices, save_report, jump_row
import openpyxl
from openpyxl.styles import PatternFill, Font
from django.utils.translation import gettext as _
from django.db.models import OuterRef, Subquery
from dateutil.relativedelta import relativedelta
from django.db.models.functions import Coalesce
from contract.models import ContractTerminationRequest, ContractRequest, ContractSurrogation, ContractTenantChange, ContractDataChange, ContractObservation, Bonification, ContractLog, Variable
from claimrequest.models import ClaimRequest, ClaimRequestStatus, VulnerabilityRequest
from coredata.models import CallRegister, ConfigProject, PersonObservation
import traceback
from notification.models import Incident
from django.utils import timezone
from statistics.utils.report_filters import (contract_multi_filter_q, get_billing_ids, get_multi_ids, has_serie_final_range,
                                             invoice_multi_filter_q, serie_final_range_invoices, serie_final_range_period)

def parse_date(date_str, make_aware=False):
    if not date_str:
        return None
    for fmt in ('%Y-%m-%dT%H:%M:%S.%fZ', '%Y-%m-%dT%H:%M:%SZ', '%Y-%m-%d'):
        try:
            dt = datetime.datetime.strptime(date_str, fmt)
            if make_aware and timezone.is_naive(dt):
                return timezone.make_aware(dt)
            return dt
        except ValueError:
            continue
    return None

def generate_detailed_typology_periodicity_report(request, black_fill, white_bold_font, task=None):
    try:
        id = request.data.get('id')
        date_range = request.data.get('date_range')
        billing_ids = get_billing_ids(request.data)
        name = request.data.get('name', _('Ingrés per tipologia i periodicitat'))
        type_id = request.data.get('type_id')
        exploitation_id = request.data.get('exploitation_id')

        start_date = None
        end_date = None
        if billing_ids:
            invoices = Invoice.objects.filter(billing_id__in=billing_ids)
        elif date_range:
            start_date = parse_date(date_range[0])
            end_date = parse_date(date_range[1])
            if not start_date or not end_date:
                return None, None, None
            invoice_type = ConfigProject.objects.get(token="invoice_type_invoice_token").value
            invoices = Invoice.objects.filter(issue_date__range=(start_date.date(), end_date.date()), type_final=invoice_type).distinct()
        elif has_serie_final_range(request.data):
            # Seleccio nomes pel rang de numero de factura: el rang ja identifica les
            # factures, sense lot ni periode. filter_pending_invoices hi aplica el rang
            # (apply_serie_final_range) i el periode surt de les factures resultants.
            invoices = filter_pending_invoices(serie_final_range_invoices(), request)
            start_date, end_date = serie_final_range_period(invoices)
            if start_date is None:
                return None, None, None
        else:
            return None, None, None

        multi_q = invoice_multi_filter_q(**get_multi_ids(request.data))
        if multi_q:
            invoices = invoices.filter(multi_q).distinct()

        if exploitation_id and exploitation_id != 'all':
            invoices = invoices.filter(exploitation_id=exploitation_id)

        invoices = filter_pending_invoices(invoices, request)

        # Obtenir les dates de moviment de cartera per a les factures
        from billing.models import PaymentMovement
        pm_qs = PaymentMovement.objects.filter(
            payment__invoice__in=invoices,
            is_active=True
        ).values('payment__invoice_id', 'movement_date')
        
        invoice_movement_dates = {}
        for pm in pm_qs:
            inv_id = pm['payment__invoice_id']
            m_date = pm['movement_date']
            if m_date:
                m_date_str = m_date.strftime("%d/%m/%Y")
                if inv_id not in invoice_movement_dates:
                    invoice_movement_dates[inv_id] = []
                if m_date_str not in invoice_movement_dates[inv_id]:
                    invoice_movement_dates[inv_id].append(m_date_str)

        # Grup per tipologia de tarifa (product) i periodicitat (biller period_type)
        line_items = InvoiceLineItem.objects.filter(invoice__in=invoices, is_active=True)
        
        data = line_items.values(
            prod_name=Coalesce(F('product_name'), F('price_rate__product__name'), F('product__name')),
            rate_name=Coalesce(F('price_rate__name'), F('price_rate_name')),
            use_type_name=F('invoice__contract__use_type__name'),
            periodicity=F('invoice__contract__supply_point_default__property__route_position__route__biller__period_type')
        ).annotate(
            total_base=Sum('price'),
            total_tax=Sum('tax_price'),
            invoice_count=Count('invoice_id', distinct=True)
        ).order_by('prod_name', 'rate_name', 'use_type_name', 'periodicity')

        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.title = _("Tipologia i Periodicitat")

        row = 0
        titles = [_("Producte"), _("Tarifa"), _("Tipus d'ús"), _("Periodicitat"), _("Num. Factures"), _("Total (€)")]
        row = add_row(sheet, row, titles, fill=black_fill, font=white_bold_font)

        # Prepare detail sheet
        detail_sheet = wb.create_sheet(title=_("Detall Registres"))
        detail_titles = [_("Producte"), _("Tarifa"), _("Tipus d'ús"), _("Periodicitat"), _("Num. Factura"), _("Data"), _("Data Moviment Cartera"), _("Titular"), _("Import (€)")]
        add_row(detail_sheet, 0, detail_titles, fill=black_fill, font=white_bold_font)
        detail_row = 1

        tot_lines = 0.0
        for item in data:
            base = float(item['total_base'] or 0)
            tax = float(item['total_tax'] or 0)
            total = base + tax
            tot_lines += total
            
            row_data = [
                item['prod_name'] or _("Sense producte"),
                item['rate_name'] or _("Sense tarifa"),
                item['use_type_name'] or _("Sense tipus d'ús"),
                item['periodicity'] or _("Sense periodicitat"),
                item['invoice_count'],
                total
            ]
            row = add_row(sheet, row, row_data)
            sheet.cell(row=sheet.max_row, column=6).number_format = '#,##0.00'

        # Consolidar factures sense línies desglossades per quadrar amb el total d'explotació
        invoices_with_lines = line_items.values_list('invoice_id', flat=True).distinct()
        invoices_without_lines = invoices.exclude(id__in=invoices_with_lines)
        missing_count = invoices_without_lines.count()
        if missing_count > 0:
            missing_total = sum(float(inv.total_final or 0) for inv in invoices_without_lines)
            if missing_total != 0:
                tot_lines += missing_total
                row_data = [
                    _("Factura sense línies desglossades"),
                    _("Manual / Directa"),
                    _("Sense tipus d'ús"),
                    _("Puntual"),
                    missing_count,
                    missing_total
                ]
                row = add_row(sheet, row, row_data)
                sheet.cell(row=sheet.max_row, column=6).number_format = '#,##0.00'

        # Afegir fila de TOTALS al final de la taula
        total_row_data = [
            _("TOTAL FACTURACIÓ"),
            "", "", "",
            invoices.count(),
            tot_lines
        ]
        row = add_row(sheet, row, total_row_data, fill=black_fill, font=white_bold_font)
        sheet.cell(row=sheet.max_row, column=6).number_format = '#,##0.00'

        # Fill detail sheet
        # Agrupem per factura, producte, origen i tipus d'ús per al detall
        invoice_totals = line_items.values(
            'invoice_id',
            'invoice__serie_final', 'invoice__number', 'invoice__token', 'invoice__issue_date', 'invoice__customer_final',
            prod_name=Coalesce(F('product_name'), F('price_rate__product__name'), F('product__name')),
            rate_name=Coalesce(F('price_rate__name'), F('price_rate_name')),
            use_type_name=F('invoice__contract__use_type__name'),
            periodicity=F('invoice__contract__supply_point_default__property__route_position__route__biller__period_type')
        ).annotate(
            total_val=Sum(F('price') + F('tax_price'))
        ).order_by('invoice__issue_date')

        total_invoice_totals = invoice_totals.count()
        for detail_counter, inv in enumerate(invoice_totals, start=1):
            if task and detail_counter % 50 == 0:
                task.update_state(
                    state='PROGRESS',
                    meta={
                        'current': detail_counter,
                        'total': total_invoice_totals,
                        'percent': round((detail_counter / total_invoice_totals) * 100, 2) if total_invoice_totals else 0.0
                    }
                )
            inv_id = inv['invoice_id']
            m_dates = ", ".join(invoice_movement_dates.get(inv_id, []))
            add_row(detail_sheet, detail_row, [
                inv['prod_name'] or _("Sense producte"),
                inv['rate_name'] or _("Sense tarifa"),
                inv['use_type_name'] or _("Sense tipus d'ús"),
                inv['periodicity'] or _("Sense periodicitat"),
                inv['invoice__serie_final'] or inv['invoice__number'] or inv['invoice__token'],
                inv['invoice__issue_date'].strftime("%Y-%m-%d") if inv['invoice__issue_date'] else "",
                m_dates,
                inv['invoice__customer_final'] or "",
                round(inv['total_val'] or 0, 2)
            ])
            detail_row += 1

        # Afegir les factures sense línies al full de detall
        for inv in invoices_without_lines:
            m_dates = ", ".join(invoice_movement_dates.get(inv.id, []))
            add_row(detail_sheet, detail_row, [
                _("Sense producte"),
                _("Manual / Directa"),
                _("Sense tipus d'ús"),
                _("Puntual"),
                inv.serie_final or inv.number or inv.token,
                inv.issue_date.strftime("%Y-%m-%d") if inv.issue_date else "",
                m_dates,
                inv.customer_final or "",
                round(float(inv.total_final or 0), 2)
            ])
            detail_row += 1

        adjust_column_widths(sheet)
        adjust_column_widths(detail_sheet)

        # Create route sheet
        route_sheet = wb.create_sheet(title=_("Consum i Quotes per Ruta"))
        route_sheet.views.sheetView[0].showGridLines = True

        route_lines_data = line_items.values(
            'invoice_id',
            'invoice__contract__supply_point_default__property__route_position__route__token',
            'invoice__contract__supply_point_default__property__route_position__route__name',
            'product__token',
            'product_name',
            'price_rate__product__name',
            'product__name',
            'price_rate_name',
            'price_rate__name',
            'name',
            'line_item_type__name',
            'units',
            'price',
            'tax_price'
        )
        
        route_groups = {}
        for item in route_lines_data:
            r_token = item['invoice__contract__supply_point_default__property__route_position__route__token'] or _("Sense ruta")
            r_name = item['invoice__contract__supply_point_default__property__route_position__route__name'] or _("Sense nom de ruta")
            
            if r_token not in route_groups:
                route_groups[r_token] = {
                    'name': r_name,
                    'water_invoices': set(),
                    'water_m3': 0.0,
                    'water_quota': 0.0,
                    'water_consum': 0.0,
                    
                    'sewer_invoices': set(),
                    'sewer_m3': 0.0,
                    'sewer_quota': 0.0,
                    'sewer_consum': 0.0,
                    
                    'rec_invoices': set(),
                    'rec_m3': 0.0,
                    'rec_quota': 0.0,
                    'rec_consum': 0.0,
                    
                    'esco_invoices': set(),
                    'esco_val': 0.0,
                    
                    'comp_invoices': set(),
                    'comp_val': 0.0
                }
                
            p_token = item['product__token']
            p_name = item['product_name'] or item['price_rate__product__name'] or item['product__name']
            l_name = item['name']
            lt_name = item['line_item_type__name']
            
            units = float(item['units'] or 0)
            base = float(item['price'] or 0)
            tax = float(item['tax_price'] or 0)
            
            # Correct sign inconsistency if any
            if base < 0 and tax > 0:
                tax = -tax
            elif base > 0 and tax < 0:
                tax = -tax
                
            total = base + tax
            
            p_token_str = p_token or ""
            p_name_str = (p_name or "").upper()
            pr_name_str = (item['price_rate_name'] or item['price_rate__name'] or "").upper()
            l_name_str = (l_name or "").upper()
            lt_name_str = (lt_name or "").upper()
            
            is_comp = p_token_str == '1003' or 'CONSERVACIÓ COMPTADOR' in p_name_str or 'CONSERVACIÓN CONTADOR' in p_name_str
            is_esco = (p_token_str.startswith('ESCO') or 'ESCOMESA' in p_name_str) and 'BATERIA' not in p_name_str
            
            is_rec_concept = (
                p_token_str == 'REC' or 
                p_token_str.startswith('REG') or 
                p_token_str.startswith('RIEG') or 
                p_token_str.startswith('IRRIG') or 
                'RIEGO' in p_name_str or 
                'IRRIGATION' in p_name_str or 
                'RIEGO' in pr_name_str or
                'REG' in pr_name_str or
                ('REG' in p_name_str and 'REGISTRE' not in p_name_str and 'REGISTRO' not in p_name_str and 'BOCA-REG-COMPT' not in p_token_str and 'BOCA DE REG' not in p_name_str)
            )
            
            is_rec_q = False
            is_rec_c = False
            is_sewer_q = False
            is_sewer_c = False
            is_water_q = False
            is_water_c = False
            
            if is_comp or is_esco:
                pass
            elif is_rec_concept:
                is_rec_q = 'QUOTA' in p_name_str or 'CUOTA' in p_name_str or 'FIX' in p_name_str or 'FIJA' in p_name_str or 'SERVEI' in p_name_str or 'SERVICIO' in p_name_str or 'FIXA' in lt_name_str or 'FIXA' in l_name_str
                is_rec_c = not is_rec_q
            elif p_token_str == 'CLV' or 'CLAVEGUERAM' in p_name_str or 'ALCANTARILLADO' in p_name_str:
                is_sewer_q = (p_token_str == 'CLV' and (lt_name_str == 'QUOTA FIXA' or 'FIXA' in l_name_str)) or ('ALCANTARILLADO' in p_name_str and p_name_str.endswith('CUOTA'))
                is_sewer_c = not is_sewer_q
            elif p_token_str == '1000' or p_token_str == '1001' or 'AIGUA' in p_name_str or 'AGUA' in p_name_str:
                is_water_q = p_token_str == '1000' or 'QUOTA DE SERVEI' in p_name_str or 'CUOTA AGUA' in p_name_str or 'QUOTA AGUA' in p_name_str
                is_water_c = not is_water_q
            
            if is_water_q:
                route_groups[r_token]['water_invoices'].add(item['invoice_id'])
                route_groups[r_token]['water_quota'] += total
            elif is_water_c:
                route_groups[r_token]['water_invoices'].add(item['invoice_id'])
                route_groups[r_token]['water_m3'] += units
                route_groups[r_token]['water_consum'] += total
            elif is_sewer_q:
                route_groups[r_token]['sewer_invoices'].add(item['invoice_id'])
                route_groups[r_token]['sewer_quota'] += total
            elif is_sewer_c:
                route_groups[r_token]['sewer_invoices'].add(item['invoice_id'])
                route_groups[r_token]['sewer_m3'] += units
                route_groups[r_token]['sewer_consum'] += total
            elif is_rec_q:
                route_groups[r_token]['rec_invoices'].add(item['invoice_id'])
                route_groups[r_token]['rec_quota'] += total
            elif is_rec_c:
                route_groups[r_token]['rec_invoices'].add(item['invoice_id'])
                route_groups[r_token]['rec_m3'] += units
                route_groups[r_token]['rec_consum'] += total
            elif is_esco:
                route_groups[r_token]['esco_invoices'].add(item['invoice_id'])
                route_groups[r_token]['esco_val'] += total
            elif is_comp:
                route_groups[r_token]['comp_invoices'].add(item['invoice_id'])
                route_groups[r_token]['comp_val'] += total

        # Determine dynamic presence of dynamic column groups
        has_water = True
        has_sewer = True
        has_rec = any(len(rg['rec_invoices']) > 0 for rg in route_groups.values())
        has_esco = any(len(rg['esco_invoices']) > 0 for rg in route_groups.values())
        has_comp = any(len(rg['comp_invoices']) > 0 for rg in route_groups.values())

        active_groups = []
        if has_water:
            active_groups.append({
                'key': 'water',
                'title': _("Aigua"),
                'sub_headers': [_("nº de Factures"), _("m3 facturats"), _("Quota servei aigua"), _("Consum aigua (€)")],
                'formats': ['#,##0', '#,##0.00', '#,##0.00', '#,##0.00']
            })
        if has_sewer:
            active_groups.append({
                'key': 'sewer',
                'title': _("Clavegueram"),
                'sub_headers': [_("nº de Factures"), _("m3 facturats"), _("Quota servei clavegueram"), _("Consum clavegueram (€)")],
                'formats': ['#,##0', '#,##0.00', '#,##0.00', '#,##0.00']
            })
        if has_rec:
            active_groups.append({
                'key': 'rec',
                'title': _("Rec"),
                'sub_headers': [_("nº de Factures"), _("m3 facturats"), _("Quota servei rec"), _("Consum rec (€)")],
                'formats': ['#,##0', '#,##0.00', '#,##0.00', '#,##0.00']
            })
        if has_esco:
            active_groups.append({
                'key': 'esco',
                'title': _("Escomeses"),
                'sub_headers': [_("nº de Factures"), _("Conservació escomeses")],
                'formats': ['#,##0', '#,##0.00']
            })
        if has_comp:
            active_groups.append({
                'key': 'comp',
                'title': _("Comptadors"),
                'sub_headers': [_("nº de Factures"), _("Conservació comptadors")],
                'formats': ['#,##0', '#,##0.00']
            })

        group_headers = [_("Ruta"), _("Nom de la ruta")]
        sub_headers = ["", ""]
        
        for g in active_groups:
            group_headers.append(g['title'])
            group_headers.extend([""] * (len(g['sub_headers']) - 1))
            sub_headers.extend(g['sub_headers'])

        route_row = 0
        route_row = add_row(route_sheet, route_row, group_headers, fill=black_fill, font=white_bold_font)
        route_row = add_row(route_sheet, route_row, sub_headers, fill=black_fill, font=white_bold_font)

        from openpyxl.utils import get_column_letter

        # Merge Ruta and Nom de la ruta
        route_sheet.merge_cells('A1:A2')
        route_sheet.merge_cells('B1:B2')
        
        current_col = 3
        for g in active_groups:
            width = len(g['sub_headers'])
            if width > 1:
                start_letter = get_column_letter(current_col)
                end_letter = get_column_letter(current_col + width - 1)
                route_sheet.merge_cells(f"{start_letter}1:{end_letter}1")
            current_col += width

        total_cols = len(group_headers)
        # Apply header styling and alignment
        for r in [1, 2]:
            for c in range(1, total_cols + 1):
                cell = route_sheet.cell(row=r, column=c)
                cell.fill = black_fill
                cell.font = white_bold_font
                cell.alignment = openpyxl.styles.Alignment(horizontal="center", vertical="center", wrap_text=True)

        total_water_invs = set()
        total_water_m3 = 0.0
        total_water_quota = 0.0
        total_water_consum = 0.0
        
        total_sewer_invs = set()
        total_sewer_m3 = 0.0
        total_sewer_quota = 0.0
        total_sewer_consum = 0.0
        
        total_rec_invs = set()
        total_rec_m3 = 0.0
        total_rec_quota = 0.0
        total_rec_consum = 0.0
        
        total_esco_invs = set()
        total_esco_val = 0.0
        
        total_comp_invs = set()
        total_comp_val = 0.0
        
        for r_token, r_data in sorted(route_groups.items()):
            row_vals = [r_token, r_data['name']]
            
            for g in active_groups:
                if g['key'] == 'water':
                    row_vals.extend([
                        len(r_data['water_invoices']),
                        r_data['water_m3'],
                        r_data['water_quota'],
                        r_data['water_consum']
                    ])
                elif g['key'] == 'sewer':
                    row_vals.extend([
                        len(r_data['sewer_invoices']),
                        r_data['sewer_m3'],
                        r_data['sewer_quota'],
                        r_data['sewer_consum']
                    ])
                elif g['key'] == 'rec':
                    row_vals.extend([
                        len(r_data['rec_invoices']),
                        r_data['rec_m3'],
                        r_data['rec_quota'],
                        r_data['rec_consum']
                    ])
                elif g['key'] == 'esco':
                    row_vals.extend([
                        len(r_data['esco_invoices']),
                        r_data['esco_val']
                    ])
                elif g['key'] == 'comp':
                    row_vals.extend([
                        len(r_data['comp_invoices']),
                        r_data['comp_val']
                    ])
                    
            route_row = add_row(route_sheet, route_row, row_vals)
            
            # Apply formatting to data row
            current_col = 3
            for g in active_groups:
                for fmt in g['formats']:
                    route_sheet.cell(row=route_sheet.max_row, column=current_col).number_format = fmt
                    current_col += 1
                    
            if 'water' in [g['key'] for g in active_groups]:
                total_water_invs.update(r_data['water_invoices'])
                total_water_m3 += r_data['water_m3']
                total_water_quota += r_data['water_quota']
                total_water_consum += r_data['water_consum']
                
            if 'sewer' in [g['key'] for g in active_groups]:
                total_sewer_invs.update(r_data['sewer_invoices'])
                total_sewer_m3 += r_data['sewer_m3']
                total_sewer_quota += r_data['sewer_quota']
                total_sewer_consum += r_data['sewer_consum']
                
            if 'rec' in [g['key'] for g in active_groups]:
                total_rec_invs.update(r_data['rec_invoices'])
                total_rec_m3 += r_data['rec_m3']
                total_rec_quota += r_data['rec_quota']
                total_rec_consum += r_data['rec_consum']
                
            if 'esco' in [g['key'] for g in active_groups]:
                total_esco_invs.update(r_data['esco_invoices'])
                total_esco_val += r_data['esco_val']
                
            if 'comp' in [g['key'] for g in active_groups]:
                total_comp_invs.update(r_data['comp_invoices'])
                total_comp_val += r_data['comp_val']
            
        total_vals = [_("TOTALS"), ""]
        for g in active_groups:
            if g['key'] == 'water':
                total_vals.extend([
                    len(total_water_invs),
                    total_water_m3,
                    total_water_quota,
                    total_water_consum
                ])
            elif g['key'] == 'sewer':
                total_vals.extend([
                    len(total_sewer_invs),
                    total_sewer_m3,
                    total_sewer_quota,
                    total_sewer_consum
                ])
            elif g['key'] == 'rec':
                total_vals.extend([
                    len(total_rec_invs),
                    total_rec_m3,
                    total_rec_quota,
                    total_rec_consum
                ])
            elif g['key'] == 'esco':
                total_vals.extend([
                    len(total_esco_invs),
                    total_esco_val
                ])
            elif g['key'] == 'comp':
                total_vals.extend([
                    len(total_comp_invs),
                    total_comp_val
                ])
                
        route_row = add_row(route_sheet, route_row, total_vals, fill=black_fill, font=white_bold_font)
        
        # Apply formatting to totals row
        current_col = 3
        for g in active_groups:
            for fmt in g['formats']:
                route_sheet.cell(row=route_sheet.max_row, column=current_col).number_format = fmt
                current_col += 1
            
        adjust_column_widths(route_sheet)

        filename = f"report_tipologia_periodicitat_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        document_id = save_report(content=wb, filename=filename, name=name, type_id=type_id, start_date=start_date, end_date=end_date)
        return None, document_id, filename
    except Exception as e:
        print(f"Error in generate_detailed_typology_periodicity_report: {e}")
        traceback.print_exc()
        return None, None, None

def generate_detailed_concept_report(request, black_fill, white_bold_font, task=None):
    try:
        id = request.data.get('id')
        date_range = request.data.get('date_range')
        billing_ids = get_billing_ids(request.data)
        name = request.data.get('name', 'Ingrés detallat per concepte')
        type_id = request.data.get('type_id')
        exploitation_id = request.data.get('exploitation_id')

        start_date = None
        end_date = None
        if billing_ids:
            invoices = Invoice.objects.filter(billing_id__in=billing_ids)
        elif date_range:
            start_date = parse_date(date_range[0])
            end_date = parse_date(date_range[1])
            if not start_date or not end_date:
                return None, None, None
            invoices = Invoice.objects.filter(issue_date__range=(start_date.date(), end_date.date()))
        elif has_serie_final_range(request.data):
            # Seleccio nomes pel rang de numero de factura: el rang ja identifica les
            # factures, sense lot ni periode. filter_pending_invoices hi aplica el rang
            # (apply_serie_final_range) i el periode surt de les factures resultants.
            invoices = filter_pending_invoices(serie_final_range_invoices(), request)
            start_date, end_date = serie_final_range_period(invoices)
            if start_date is None:
                return None, None, None
        else:
            return None, None, None

        multi_q = invoice_multi_filter_q(**get_multi_ids(request.data))
        if multi_q:
            invoices = invoices.filter(multi_q).distinct()

        if exploitation_id and exploitation_id != 'all':
            invoices = invoices.filter(exploitation_id=exploitation_id)

        invoices = filter_pending_invoices(invoices, request)

        # Ingressos per concepte i subconcepte (InvoiceLineItem)
        # S'elimina el filtre is_active=True per evitar discrepàncies amb el total_final de la factura
        line_items = InvoiceLineItem.objects.filter(invoice__in=invoices)
        
        data = line_items.values(
            concept=F('product_name'),
            subconcept=F('name'),
            tax_pct=F('tax_percent')
        ).annotate(
            total_base=Sum('price'),
            total_tax=Sum('tax_price'),
            total_units=Sum('units'),
            count=Count('id')
        ).order_by('concept', 'subconcept')

        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.title = _("Conceptes")

        row = 0
        titles = [_("Concepte"), _("Subconcepte (Tram/Quota)"), _("IVA %"), _("Núm. Línies"), _("Unitats/m³"), _("Base (€)"), _("IVA (€)"), _("Total (€)")]
        row = add_row(sheet, row, titles, fill=black_fill, font=white_bold_font)

        total_base_all = 0
        total_tax_all = 0
        total_sum_all = 0

        # Prepare detail sheet
        detail_sheet = wb.create_sheet(title=_("Detall Registres"))
        detail_titles = [_("Concepte"), _("Subconcepte"), _("Num. Factura"), _("Data"), _("Titular"), _("Unitats/m³"), _("Base (€)"), _("IVA %"), _("IVA (€)"), _("Total (€)")]
        add_row(detail_sheet, 0, detail_titles, fill=black_fill, font=white_bold_font)
        detail_row = 1

        for item in data:
            base = item['total_base'] or 0
            tax = item['total_tax'] or 0
            if tax == 0 and base != 0 and item['tax_pct']:
                # Fallback in case tax_price was somehow null but there is a pct
                tax = base * (float(item['tax_pct']) / 100)
            
            total = base + tax
            units = item['total_units'] or 0
            
            total_base_all += base
            total_tax_all += tax
            total_sum_all += total

            row = add_row(sheet, row, [
                item['concept'] or "",
                item['subconcept'] or "",
                f"{item['tax_pct'] or 0} %",
                item['count'],
                round(units, 4),
                round(base, 2),
                round(tax, 2),
                round(total, 2)
            ])

        # Fill detail sheet
        total_line_items = line_items.count()
        for line_counter, line in enumerate(line_items.select_related('invoice'), start=1):
            if task and line_counter % 50 == 0:
                task.update_state(
                    state='PROGRESS',
                    meta={
                        'current': line_counter,
                        'total': total_line_items,
                        'percent': round((line_counter / total_line_items) * 100, 2) if total_line_items else 0.0
                    }
                )
            base = line.price or 0
            tax = line.tax_price or 0
            if tax == 0 and base != 0 and line.tax_percent:
                tax = base * (float(line.tax_percent) / 100)
            
            total = base + tax
            units = line.units or 0
            
            add_row(detail_sheet, detail_row, [
                line.product_name or "",
                line.name or "",
                line.invoice.serie_final or line.invoice.number or line.invoice.token,
                line.invoice.issue_date.strftime("%Y-%m-%d") if line.invoice.issue_date else "",
                line.invoice.customer_final or "",
                round(units, 4),
                round(base, 2),
                f"{line.tax_percent or 0} %",
                round(tax, 2),
                round(total, 2)
            ])
            detail_row += 1

        row = add_row(sheet, row, [
            _("TOTALS"), "", "", "", "",
            round(total_base_all, 2),
            round(total_tax_all, 2),
            round(total_sum_all, 2)
        ], fill=black_fill, font=white_bold_font)

        adjust_column_widths(sheet)
        adjust_column_widths(detail_sheet)
        filename = f"report_conceptes_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        document_id = save_report(content=wb, filename=filename, name=name, type_id=type_id, start_date=start_date, end_date=end_date)
        return None, document_id, filename
    except Exception as e:
        print(f"Error in generate_detailed_concept_report: {e}")
        traceback.print_exc()
        return None, None, None

def generate_non_tariff_income_report(request, black_fill, white_bold_font, task=None):
    try:
        id = request.data.get('id')
        date_range = request.data.get('date_range')
        billing_ids = get_billing_ids(request.data)
        name = request.data.get('name', 'Ingressos no tarifaris')
        type_id = request.data.get('type_id')
        exploitation_id = request.data.get('exploitation_id')

        start_date = None
        end_date = None
        if billing_ids:
            invoices = Invoice.objects.filter(billing_id__in=billing_ids)
        elif date_range:
            start_date = parse_date(date_range[0])
            end_date = parse_date(date_range[1])
            if not start_date or not end_date:
                return None, None, None
            invoices = Invoice.objects.filter(issue_date__range=(start_date.date(), end_date.date()))
        elif has_serie_final_range(request.data):
            # Seleccio nomes pel rang de numero de factura: el rang ja identifica les
            # factures, sense lot ni periode. filter_pending_invoices hi aplica el rang
            # (apply_serie_final_range) i el periode surt de les factures resultants.
            invoices = filter_pending_invoices(serie_final_range_invoices(), request)
            start_date, end_date = serie_final_range_period(invoices)
            if start_date is None:
                return None, None, None
        else:
            return None, None, None

        multi_q = invoice_multi_filter_q(**get_multi_ids(request.data))
        if multi_q:
            invoices = invoices.filter(multi_q).distinct()

        if exploitation_id and exploitation_id != 'all':
            invoices = invoices.filter(exploitation_id=exploitation_id)

        invoices = filter_pending_invoices(invoices, request)

        # Ingressos no tarifaris: origins 'contracte', 'escomesa', 'altres', 'subministrament'
        non_tariff_origins = ['contracte', 'escomesa', 'altres', 'subministrament']
        line_items = InvoiceLineItem.objects.filter(
            invoice__in=invoices, 
            is_active=True,
            product__origin__token__in=non_tariff_origins
        )

        data = line_items.values(
            concept=F('product_name'),
            origin=F('product__origin__name')
        ).annotate(
            total_base=Sum('price'),
            count=Count('id')
        ).order_by('origin', 'concept')

        if exploitation_id and exploitation_id != 'all':
            data = data.filter(invoice__exploitation_id=exploitation_id)

        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.title = _("No Tarifaris")

        row = 0
        titles = [_("Origen"), _("Concepte"), _("Línies"), _("Total Base")]
        row = add_row(sheet, row, titles, fill=black_fill, font=white_bold_font)

        # Prepare detail sheet
        detail_sheet = wb.create_sheet(title=_("Detall Registres"))
        detail_titles = [_("Origen"), _("Concepte"), _("Num. Factura"), _("Data"), _("Titular"), _("Base")]
        add_row(detail_sheet, 0, detail_titles, fill=black_fill, font=white_bold_font)
        detail_row = 1

        for item in data:
            row = add_row(sheet, row, [
                item['origin'],
                item['concept'],
                item['count'],
                f"{round(item['total_base'] or 0, 2)} €"
            ])

        # Fill detail sheet
        total_line_items = line_items.count()
        for line_counter, line in enumerate(line_items.select_related('invoice', 'product__origin'), start=1):
            if task and line_counter % 50 == 0:
                task.update_state(
                    state='PROGRESS',
                    meta={
                        'current': line_counter,
                        'total': total_line_items,
                        'percent': round((line_counter / total_line_items) * 100, 2) if total_line_items else 0.0
                    }
                )
            add_row(detail_sheet, detail_row, [
                line.product.origin.name if line.product and line.product.origin else "",
                line.product_name,
                line.invoice.serie_final or line.invoice.number or line.invoice.token,
                line.invoice.issue_date.strftime("%Y-%m-%d") if line.invoice.issue_date else "",
                line.invoice.customer_final or "",
                f"{round(line.price or 0, 2)} €"
            ])
            detail_row += 1

        adjust_column_widths(sheet)
        adjust_column_widths(detail_sheet)
        filename = f"report_no_tarifaris_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        document_id = save_report(content=wb, filename=filename, name=name, type_id=type_id, start_date=start_date, end_date=end_date)
        return None, document_id, filename
    except Exception as e:
        print(f"Error in generate_non_tariff_income_report: {e}")
        traceback.print_exc()
        return None, None, None

def generate_subscriber_evolution_report(request, black_fill, white_bold_font, task=None):
    try:
        date_range = request.data.get('date_range')
        name = request.data.get('name', 'Evolució abonats')
        type_id = request.data.get('type_id')
        exploitation_id = request.data.get('exploitation_id')

        if not date_range:
            return None, None, None

        start_date = parse_date(date_range[0])
        if not start_date:
            return None, None, None
        
        # Get configuration tokens
        try:
            completed_status_token = ConfigProject.objects.get(token='contract_termination_completed_token').value
        except:
            completed_status_token = '3' # Default per Finalitzat segons DB
        
        baixa_status_token = '-1' # Token estàndard per Baixa en ContractStatus

        multi_ids = get_multi_ids(request.data)
        multi_q = contract_multi_filter_q(person_ids=multi_ids['person_ids'], contract_ids=multi_ids['contract_ids'])
        ctr_multi_q = None
        if multi_ids['contract_ids']:
            ctr_multi_q = Q(contract_id__in=multi_ids['contract_ids'])
        if multi_ids['person_ids']:
            ctr_person_q = (
                Q(contract__owner_id__in=multi_ids['person_ids'])
                | Q(contract__tenant_id__in=multi_ids['person_ids'])
                | Q(contract__holder_id__in=multi_ids['person_ids'])
            )
            ctr_multi_q = ctr_person_q if ctr_multi_q is None else ctr_multi_q & ctr_person_q

        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.title = _("Evolució Abonats")

        row = 0
        titles = [_("Mes"), _("Total Abonats (actius)"), _("Altes"), _("Baixes")]
        row = add_row(sheet, row, titles, fill=black_fill, font=white_bold_font)

        # We need to iterate month by month, starting from the first of the month
        curr = start_date.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
        months = []
        # Get end_date if available
        end_date_dt = None
        if date_range and len(date_range) > 1 and date_range[1]:
            end_date_dt = parse_date(date_range[1])
        
        # If no end_date, show 6 months from start_date
        if not end_date_dt:
            end_date_dt = start_date + relativedelta(months=5)

        while curr <= end_date_dt:
            months.append((curr.year, curr.month))
            curr = curr + relativedelta(months=1)

        # Prepare detail sheet
        detail_sheet = wb.create_sheet(title=_("Detall Registres"))
        detail_titles = [_("Mes"), _("Tipus"), _("Token Contracte"), _("Titular"), _("Data Alta"), _("Data Baixa"), _("Explotació")]
        add_row(detail_sheet, 0, detail_titles, fill=black_fill, font=white_bold_font)
        detail_row = 1

        # 1. Calculate Initial Running Total (Active subscribers before start_date)
        # ---------------------------------------------------------------------
        month_start_initial = datetime.date(months[0][0], months[0][1], 1)
        
        # All contracts registered before the report start date
        initial_altes_qs = Contract.objects.filter(
            Q(registration_date__lt=month_start_initial) |
            Q(registration_date__isnull=True, created_at__date__lt=month_start_initial)
        )
        if multi_q:
            initial_altes_qs = initial_altes_qs.filter(multi_q).distinct()

        # All terminations before the report start date
        initial_baixes_req_qs = ContractTerminationRequest.objects.annotate(
            term_date=Coalesce('approved_at', 'requested_at', 'created_at')
        ).filter(
            Q(status__token=completed_status_token) | Q(contract__status__token=baixa_status_token),
            term_date__date__lt=month_start_initial
        )
        if ctr_multi_q:
            initial_baixes_req_qs = initial_baixes_req_qs.filter(ctr_multi_q).distinct()
        initial_baixes_req_ids = list(initial_baixes_req_qs.values_list('contract_id', flat=True))

        initial_baixes_directes_qs = Contract.objects.filter(
            status__token=baixa_status_token,
            updated_at__date__lt=month_start_initial
        ).exclude(id__in=initial_baixes_req_ids)
        if multi_q:
            initial_baixes_directes_qs = initial_baixes_directes_qs.filter(multi_q).distinct()
        initial_baixes_directes_ids = list(initial_baixes_directes_qs.values_list('id', flat=True))
        
        initial_baixes_total_ids = set(initial_baixes_req_ids + initial_baixes_directes_ids)
        
        if exploitation_id and exploitation_id != 'all':
            initial_altes_qs = initial_altes_qs.filter(supply_point_default__connection__exploitation_id=exploitation_id)
            # For baixes, we filter the contracts by exploitation
            initial_baixes_total_ids = set(Contract.objects.filter(
                id__in=initial_baixes_total_ids,
                supply_point_default__connection__exploitation_id=exploitation_id
            ).values_list('id', flat=True))

        running_total = initial_altes_qs.count() - len(initial_baixes_total_ids)

        # 2. Iterate months and calculate monthly changes
        # ---------------------------------------------------------------------
        total_months = len(months)
        for month_counter, (year, month) in enumerate(months, start=1):
            if task:
                task.update_state(
                    state='PROGRESS',
                    meta={
                        'current': month_counter,
                        'total': total_months,
                        'percent': round((month_counter / total_months) * 100, 2) if total_months else 0.0
                    }
                )
            month_start = datetime.date(year, month, 1)
            if month == 12:
                next_month_start = datetime.date(year + 1, 1, 1)
            else:
                next_month_start = datetime.date(year, month + 1, 1)
            end_of_month = next_month_start - datetime.timedelta(days=1)
            
            # Altes in this month (any contract registered this month, regardless of current status)
            altes_qs = Contract.objects.filter(
                Q(registration_date__range=(month_start, end_of_month)) |
                Q(registration_date__isnull=True, created_at__date__range=(month_start, end_of_month))
            )
            if multi_q:
                altes_qs = altes_qs.filter(multi_q).distinct()

            # Baixes in this month (termination requests completed this month)
            baixes_req_qs = ContractTerminationRequest.objects.annotate(
                term_date=Coalesce('approved_at', 'requested_at', 'created_at')
            ).filter(
                Q(status__token=completed_status_token) | Q(contract__status__token=baixa_status_token),
                term_date__date__range=(month_start, end_of_month)
            )
            if ctr_multi_q:
                baixes_req_qs = baixes_req_qs.filter(ctr_multi_q).distinct()

            # Plus direct baixes in this month (without request)
            baixes_req_ids = list(baixes_req_qs.values_list('contract_id', flat=True))
            baixes_directes_qs = Contract.objects.filter(
                status__token=baixa_status_token,
                updated_at__date__range=(month_start, end_of_month)
            ).exclude(id__in=baixes_req_ids)
            if multi_q:
                baixes_directes_qs = baixes_directes_qs.filter(multi_q).distinct()

            if exploitation_id and exploitation_id != 'all':
                altes_qs = altes_qs.filter(supply_point_default__connection__exploitation_id=exploitation_id)
                baixes_req_qs = baixes_req_qs.filter(contract__supply_point_default__connection__exploitation_id=exploitation_id)
                baixes_directes_qs = baixes_directes_qs.filter(supply_point_default__connection__exploitation_id=exploitation_id)

            num_altes = altes_qs.count()
            num_baixes = baixes_req_qs.count() + baixes_directes_qs.count()
            
            # Update running total
            running_total = running_total + num_altes - num_baixes
            
            row = add_row(sheet, row, [
                f"{month:02d}/{year}",
                running_total,
                num_altes,
                num_baixes
            ])

            # Add to detail sheet
            for c in altes_qs:
                add_row(detail_sheet, detail_row, [f"{month:02d}/{year}", "ALTA", c.token, str(c.holder), c.registration_date or c.created_at.date(), "", str(c.supply_point_default.connection.exploitation if c.supply_point_default and c.supply_point_default.connection else "")])
                detail_row += 1
            for b in baixes_req_qs:
                term_date_str = b.term_date.strftime("%Y-%m-%d") if b.term_date else ""
                add_row(detail_sheet, detail_row, [f"{month:02d}/{year}", "BAIXA", b.contract.token, str(b.contract.holder), b.contract.registration_date or b.contract.created_at.date(), term_date_str, str(b.contract.supply_point_default.connection.exploitation if b.contract.supply_point_default and b.contract.supply_point_default.connection else "")])
                detail_row += 1
            for c in baixes_directes_qs:
                term_date_str = c.updated_at.strftime("%Y-%m-%d")
                add_row(detail_sheet, detail_row, [f"{month:02d}/{year}", "BAIXA", c.token, str(c.holder), c.registration_date or c.created_at.date(), term_date_str, str(c.supply_point_default.connection.exploitation if c.supply_point_default and c.supply_point_default.connection else "")])
                detail_row += 1

            # Si és l'últim mes del filtre, mostrem també tots els actius actuals al detall
            if year == months[-1][0] and month == months[-1][1]:
                # Calculem els actius reals a final de període per al detall
                all_baixes_qs = ContractTerminationRequest.objects.annotate(
                    term_date=Coalesce('approved_at', 'requested_at', 'created_at')
                ).filter(
                    Q(status__token=completed_status_token) | Q(contract__status__token=baixa_status_token),
                    term_date__date__lte=end_of_month
                )
                if ctr_multi_q:
                    all_baixes_qs = all_baixes_qs.filter(ctr_multi_q).distinct()
                all_baixes_ids = list(all_baixes_qs.values_list('contract_id', flat=True))

                direct_baixes_qs = Contract.objects.filter(
                    status__token=baixa_status_token,
                    updated_at__date__lte=end_of_month
                ).exclude(id__in=all_baixes_ids)
                if multi_q:
                    direct_baixes_qs = direct_baixes_qs.filter(multi_q).distinct()
                direct_baixes_ids = list(direct_baixes_qs.values_list('id', flat=True))

                final_active_qs = Contract.objects.filter(
                    Q(registration_date__lte=end_of_month) |
                    Q(registration_date__isnull=True, created_at__date__lte=end_of_month)
                ).exclude(id__in=set(all_baixes_ids + direct_baixes_ids))
                if multi_q:
                    final_active_qs = final_active_qs.filter(multi_q).distinct()

                if exploitation_id and exploitation_id != 'all':
                    final_active_qs = final_active_qs.filter(supply_point_default__connection__exploitation_id=exploitation_id)
                
                for c in final_active_qs.distinct():
                    add_row(detail_sheet, detail_row, [f"{month:02d}/{year}", "ACTIU", c.token, str(c.holder), c.registration_date or c.created_at.date(), "", str(c.supply_point_default.connection.exploitation if c.supply_point_default and c.supply_point_default.connection else "")])
                    detail_row += 1

        adjust_column_widths(sheet)
        adjust_column_widths(detail_sheet)
        filename = f"report_evolucio_abonats_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        document_id = save_report(content=wb, filename=filename, name=name, type_id=type_id, start_date=start_date, end_date=None)
        return None, document_id, filename
    except Exception as e:
        print(f"Error in generate_subscriber_evolution_report: {e}")
        traceback.print_exc()
        return None, None, None

def generate_social_tariff_evolution_report(request, black_fill, white_bold_font, task=None):
    try:
        date_range = request.data.get('date_range')
        name = request.data.get('name', 'Evolució tarifa social')
        type_id = request.data.get('type_id')
        exploitation_id = request.data.get('exploitation_id')

        if not date_range:
            return None, None, None

        start_date = parse_date(date_range[0])
        end_date = parse_date(date_range[1])
        if not start_date or not end_date:
            return None, None, None

        # Month by month, starting from the first of the month
        curr = start_date.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
        months = []
        while curr <= end_date:
            months.append((curr.year, curr.month))
            curr = curr + relativedelta(months=1)

        try:
            tarifa_social_token = ConfigProject.objects.get(token='tarifa_social').value
        except:
            tarifa_social_token = 'tarifa_social'
        
        try:
            completed_status_token = ConfigProject.objects.get(token='contract_termination_completed_token').value
        except:
            completed_status_token = '3' # Default per Finalitzat segons DB
        
        baixa_status_token = '-1' # Token estàndard per Baixa en ContractStatus

        multi_ids = get_multi_ids(request.data)
        multi_q = contract_multi_filter_q(person_ids=multi_ids['person_ids'], contract_ids=multi_ids['contract_ids'])
        ctr_multi_q = None
        if multi_ids['contract_ids']:
            ctr_multi_q = Q(contract_id__in=multi_ids['contract_ids'])
        if multi_ids['person_ids']:
            ctr_person_q = (
                Q(contract__owner_id__in=multi_ids['person_ids'])
                | Q(contract__tenant_id__in=multi_ids['person_ids'])
                | Q(contract__holder_id__in=multi_ids['person_ids'])
            )
            ctr_multi_q = ctr_person_q if ctr_multi_q is None else ctr_multi_q & ctr_person_q

        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.title = _("Tarifa Social")

        row = 0
        titles = [_("Mes"), _("Contractes amb Tarifa Social")]
        row = add_row(sheet, row, titles, fill=black_fill, font=white_bold_font)

        # Prepare detail sheet
        detail_sheet = wb.create_sheet(title=_("Detall Registres"))
        detail_titles = [_("Mes"), _("Token Contracte"), _("Titular"), _("Bonificació"), _("Data Inici"), _("Data Fi"), _("Explotació")]
        add_row(detail_sheet, 0, detail_titles, fill=black_fill, font=white_bold_font)
        detail_row = 1

        total_months = len(months)
        for month_counter, (year, month) in enumerate(months, start=1):
            if task:
                task.update_state(
                    state='PROGRESS',
                    meta={
                        'current': month_counter,
                        'total': total_months,
                        'percent': round((month_counter / total_months) * 100, 2) if total_months else 0.0
                    }
                )
            month_start = datetime.date(year, month, 1)
            if month == 12:
                next_month_start = datetime.date(year + 1, 1, 1)
            else:
                next_month_start = datetime.date(year, month + 1, 1)
            end_of_month = next_month_start - datetime.timedelta(days=1)

            # Baixes detectades fins al final d'aquest mes
            baixes_req_qs = ContractTerminationRequest.objects.annotate(
                term_date=Coalesce('approved_at', 'requested_at', 'created_at')
            ).filter(
                Q(status__token=completed_status_token) | Q(contract__status__token=baixa_status_token),
                term_date__date__lte=end_of_month
            )
            if ctr_multi_q:
                baixes_req_qs = baixes_req_qs.filter(ctr_multi_q).distinct()
            baixes_req_ids = list(baixes_req_qs.values_list('contract_id', flat=True))

            # També excloem contractes que estiguin en estat Baixa actualment però no tinguin sol·licitud (ex. dades importades)
            baixes_directes_qs = Contract.objects.filter(
                status__token=baixa_status_token,
                updated_at__date__lte=end_of_month
            ).exclude(id__in=baixes_req_ids)
            if multi_q:
                baixes_directes_qs = baixes_directes_qs.filter(multi_q).distinct()
            baixes_directes_ids = list(baixes_directes_qs.values_list('id', flat=True))

            baixes_en_periode_ids = baixes_req_ids + baixes_directes_ids

            # Total with social tariff at the end of the month
            active_social_contracts = Contract.objects.filter(
                status__token='1',
                variables__type__token=tarifa_social_token,
                variables__is_active=True,
            ).filter(
                Q(variables__start_at__lte=end_of_month) | Q(variables__start_at__isnull=True)
            ).filter(
                Q(variables__end_at__isnull=True) | Q(variables__end_at__gte=month_start)
            ).exclude(
                id__in=baixes_en_periode_ids
            ).distinct()
            if multi_q:
                active_social_contracts = active_social_contracts.filter(multi_q).distinct()

            if exploitation_id and exploitation_id != 'all':
                active_social_contracts = active_social_contracts.filter(supply_point_default__connection__exploitation_id=exploitation_id)

            row = add_row(sheet, row, [
                f"{month:02d}/{year}",
                active_social_contracts.count()
            ])

            # Add to detail sheet
            for c in active_social_contracts:
                # Find the relevant variable for this contract in this month
                v = c.variables.filter(type__token=tarifa_social_token, start_at__lte=end_of_month, is_active=True).filter(Q(end_at__isnull=True) | Q(end_at__gte=month_start)).first()
                add_row(detail_sheet, detail_row, [f"{month:02d}/{year}", c.token, str(c.holder), v.type.name if v else "Tarifa Social", v.start_at if v else "", v.end_at if v else "", str(c.supply_point_default.connection.exploitation if c.supply_point_default and c.supply_point_default.connection else "")])
                detail_row += 1

        adjust_column_widths(sheet)
        adjust_column_widths(detail_sheet)
        filename = f"report_evolucio_social_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        document_id = save_report(content=wb, filename=filename, name=name, type_id=type_id, start_date=start_date, end_date=end_date)
        return None, document_id, filename
    except Exception as e:
        print(f"Error in generate_social_tariff_evolution_report: {e}")
        traceback.print_exc()
        return None, None, None

def generate_claim_response_time_report(request, black_fill, white_bold_font, task=None):
    try:
        from django.utils import timezone
        from coredata.models import ConfigProject
        
        date_range = request.data.get('date_range')
        name = request.data.get('name', 'Temps de resposta reclamacions')
        type_id = request.data.get('type_id')
        exploitation_id = request.data.get('exploitation_id')

        if not date_range:
            return None, None, None

        start_date = parse_date(date_range[0], make_aware=True)
        end_date = parse_date(date_range[1], make_aware=True)
        if not start_date or not end_date:
            return None, None, None

        try:
            closed_status_token = ConfigProject.objects.get(token='incident_status_closed_token').value
        except Exception:
            closed_status_token = '2'

        # Filter incidents in range
        incidents = Incident.objects.filter(created_at__range=(start_date, end_date))
        
        if exploitation_id and exploitation_id != 'all':
            incidents = incidents.filter(contract__supply_point_default__connection__exploitation_id=exploitation_id)

        data_by_type = {}

        total_incidents = incidents.count()
        for incident_counter, incident in enumerate(incidents, start=1):
            if task and incident_counter % 50 == 0:
                task.update_state(
                    state='PROGRESS',
                    meta={
                        'current': incident_counter,
                        'total': total_incidents,
                        'percent': round((incident_counter / total_incidents) * 100, 2) if total_incidents else 0.0
                    }
                )
            itype = incident.type.name if incident.type else _("Altres")
            status_name = incident.status.name if incident.status else _("Sense Estat")
            
            if incident.status and incident.status.token == closed_status_token:
                duration = (incident.updated_at - incident.created_at).days
            else:
                duration = (timezone.now() - incident.created_at).days
                
            key = (itype, status_name)
            if key not in data_by_type:
                data_by_type[key] = []
            data_by_type[key].append(duration)

        def get_stats(durations):
            if not durations:
                return [0] * 8
            durations.sort()
            n = len(durations)
            
            def percentile(p):
                idx = (p / 100) * (n - 1)
                if idx.is_integer():
                    return durations[int(idx)]
                else:
                    lower = durations[int(idx)]
                    upper = durations[int(idx) + 1]
                    return lower + (upper - lower) * (idx % 1)

            avg = sum(durations) / n
            p50 = percentile(50)
            mx = max(durations)
            mn = min(durations)
            p75 = percentile(75)
            p90 = percentile(90)
            p95 = percentile(95)
            p99 = percentile(99)
            
            return [round(avg, 2), p50, mx, mn, p75, p90, p95, p99]

        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.title = _("Temps Resposta")

        row = 0
        titles = [_("Tipus Reclamació"), _("Estat"), _("Promig (dies)"), _("Mediana (P50)"), _("Màxim"), _("Mínim"), _("P75"), _("P90"), _("P95"), _("P99")]
        row = add_row(sheet, row, titles, fill=black_fill, font=white_bold_font)

        # Prepare detail sheet
        detail_sheet = wb.create_sheet(title=_("Detall Registres"))
        detail_titles = [_("Tipus Reclamació"), _("Estat"), _("Token"), _("Nom"), _("Descripció"), _("Data Creació"), _("Data Resolució"), _("Dies"), _("Explotació")]
        add_row(detail_sheet, 0, detail_titles, fill=black_fill, font=white_bold_font)
        detail_row = 1

        # RE-FILTERING for detail sheet (to keep it simple and correct)
        for detail_incident_counter, incident in enumerate(incidents, start=1):
            if task and detail_incident_counter % 50 == 0:
                task.update_state(
                    state='PROGRESS',
                    meta={
                        'current': detail_incident_counter,
                        'total': total_incidents,
                        'percent': round((detail_incident_counter / total_incidents) * 100, 2) if total_incidents else 0.0
                    }
                )
            itype = incident.type.name if incident.type else _("Altres")
            status_name = incident.status.name if incident.status else _("Sense Estat")

            if incident.status and incident.status.token == closed_status_token:
                duration = (incident.updated_at - incident.created_at).days
                resolution_date = incident.updated_at.date()
            else:
                duration = (timezone.now() - incident.created_at).days
                resolution_date = ""

            exploitation = str(incident.contract.supply_point_default.connection.exploitation if incident.contract and incident.contract.supply_point_default and incident.contract.supply_point_default.connection else "")
            add_row(detail_sheet, detail_row, [itype, status_name, incident.token, incident.name, incident.description or "", incident.created_at.date(), resolution_date, duration, exploitation])
            detail_row += 1

        for (ctype, status_name), durations in data_by_type.items():
            stats = get_stats(durations)
            row = add_row(sheet, row, [ctype, status_name] + stats)

        adjust_column_widths(sheet)
        adjust_column_widths(detail_sheet)
        filename = f"report_temps_resposta_reclamacions_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        document_id = save_report(content=wb, filename=filename, name=name, type_id=type_id, start_date=start_date, end_date=end_date)
        return None, document_id, filename
    except Exception as e:
        print(f"Error in generate_claim_response_time_report: {e}")
        traceback.print_exc()
        return None, None, None

def generate_management_volume_report(request, black_fill, white_bold_font, task=None):
    try:
        date_range = request.data.get('date_range')
        name = request.data.get('name', 'Volum de gestions')
        type_id = request.data.get('type_id')
        exploitation_id = request.data.get('exploitation_id')

        if not date_range:
            return None, None, None

        start_date = parse_date(date_range[0])
        end_date = parse_date(date_range[1])
        if not start_date or not end_date:
            return None, None, None

        # Month by month
        curr = start_date.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
        months = []
        while curr <= end_date:
            months.append((curr.year, curr.month))
            curr = curr + relativedelta(months=1)

        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.title = _("Volum Gestions")

        row = 0
        month_titles = [f"{m:02d}/{y}" for y, m in months]
        titles = [_("Tipologia de Gestió")] + month_titles + [_("Total")]
        row = add_row(sheet, row, titles, fill=black_fill, font=white_bold_font)

        # Define categories and their querysets
        # Altes: Count new contracts as suggested by the user
        # Canvis de nom: combine several models
        # Actualitzacions: combine several models
        categories = [
            (_("Canvis de nom / Titular"), [
                ContractRequest.objects.filter(Q(type__token__in=['surrogacio', 'canvi_titular']) | Q(is_change_of_name=True)),
                ContractSurrogation.objects.all(),
                ContractTenantChange.objects.all()
            ]),
            (_("Altes"), [Contract.objects.filter(status__token='1')]),
            (_("Baixes"), [ContractTerminationRequest.objects.all()]),
            (_("Reclamacions"), [Incident.objects.filter(type__token='reclamacio')]),
            (_("Actualitzacions de dades"), [
                ContractDataChange.objects.all(),
                ContractObservation.objects.all()
            ]),
            (_("Ampliacions de tram"), [
                Bonification.objects.filter(bonification_type__token__in=['ACA-TRAM', 'ACA-TRAM-75']),
                ContractLog.objects.filter(field_name='total_persons'),
                ContractObservation.objects.filter(
                    Q(observation__icontains='tram') | 
                    Q(observation__icontains='membres') | 
                    Q(observation__icontains='convivència') |
                    Q(observation__icontains='familia nombrosa')
                ),
                Variable.objects.filter(type__token__in=['ACA-TRAM', 'ACA-TRAM-MEMBRES']),
                ContractRequest.objects.filter(type__token='ampliacio_tram')
            ]),
            (_("Sol·licitud tarifes social"), [
                VulnerabilityRequest.objects.all(),
                Variable.objects.filter(type__token__in=['tarifa-social', 'TARIFA-SOCIAL', 'vulnerability']),
                Bonification.objects.filter(bonification_type__token__icontains='EXCLUSIO')
            ]),
            (_("Consultes"), [CallRegister.objects.all()]),
        ]

        # Prepare detail sheet
        detail_sheet = wb.create_sheet(title=_("Detall Registres"))
        detail_titles = [_("Mes"), _("Categoria"), _("Token"), _("Data"), _("Explotació")]
        add_row(detail_sheet, 0, detail_titles, fill=black_fill, font=white_bold_font)
        detail_row = 1

        total_categories = len(categories)
        for cat_counter, (cat_name, qs_list) in enumerate(categories, start=1):
            if task:
                task.update_state(
                    state='PROGRESS',
                    meta={
                        'current': cat_counter,
                        'total': total_categories,
                        'percent': round((cat_counter / total_categories) * 100, 2) if total_categories else 0.0
                    }
                )
            row_data = [cat_name]
            total_cat = 0
            
            for y, m in months:
                m_start = datetime.date(year=y, month=m, day=1)
                if m == 12:
                    m_end = datetime.date(year=y+1, month=1, day=1) - datetime.timedelta(days=1)
                else:
                    m_end = datetime.date(year=y, month=m+1, day=1) - datetime.timedelta(days=1)
                
                count = 0
                for qs in qs_list:
                    # Apply exploitation filter if needed
                    active_qs = qs # Don't mutate original
                    if exploitation_id and exploitation_id != 'all':
                        try:
                            if hasattr(active_qs.model, 'contract'):
                                active_qs = active_qs.filter(contract__supply_point_default__connection__exploitation_id=exploitation_id)
                            elif hasattr(active_qs.model, 'supply_point_default'):
                                active_qs = active_qs.filter(supply_point_default__connection__exploitation_id=exploitation_id)
                            elif active_qs.model == Contract:
                                active_qs = active_qs.filter(supply_point_default__connection__exploitation_id=exploitation_id)
                            elif active_qs.model == Incident:
                                active_qs = active_qs.filter(contract__supply_point_default__connection__exploitation_id=exploitation_id)
                        except:
                            pass
                    
                    # Check which date field to use
                    date_field_base = 'created_at'
                    if hasattr(active_qs.model, 'requested_at') and active_qs.model != Contract:
                        date_field_base = 'requested_at'
                    elif hasattr(active_qs.model, 'request_at'):
                        date_field_base = 'request_at'
                    elif hasattr(active_qs.model, 'start_at'):
                        date_field_base = 'start_at'
                    elif hasattr(active_qs.model, 'time_call'):
                        date_field_base = 'time_call'
                    
                    try:
                        field_type = active_qs.model._meta.get_field(date_field_base).get_internal_type()
                        date_field = f"{date_field_base}__date" if field_type in ['DateTimeField'] else date_field_base
                    except Exception:
                        date_field = f"{date_field_base}__date"
                    
                    if active_qs.model == Contract:
                        # For Altes, use registration_date or created_at
                        filtered_qs = active_qs.filter(
                            Q(registration_date__range=(m_start, m_end)) |
                            Q(registration_date__isnull=True, created_at__date__range=(m_start, m_end))
                        )
                    else:
                        filtered_qs = active_qs.filter(**{f"{date_field}__range": (m_start, m_end)})
                    
                    m_count = filtered_qs.distinct().count()
                    count += m_count

                    # Add to detail sheet
                    for obj in filtered_qs.distinct():
                        obj_date = getattr(obj, date_field.split('__')[0], None)
                        if hasattr(obj, 'registration_date') and obj.registration_date:
                            obj_date = obj.registration_date
                        
                        if hasattr(obj_date, 'tzinfo') and obj_date.tzinfo is not None:
                            obj_date = obj_date.replace(tzinfo=None)
                        
                        exploitation = ""
                        try:
                            if hasattr(obj, 'contract') and obj.contract.supply_point_default:
                                exploitation = str(obj.contract.supply_point_default.connection.exploitation)
                            elif hasattr(obj, 'supply_point_default') and obj.supply_point_default:
                                exploitation = str(obj.supply_point_default.connection.exploitation)
                            elif isinstance(obj, Contract) and obj.supply_point_default:
                                exploitation = str(obj.supply_point_default.connection.exploitation)
                        except:
                            pass

                        # Per a les accions de gestió (observacions, canvis de dades, etc.),
                        # prioritzem el token del contracte si existeix, ja que el token 
                        # propi de l'acció (com CDC_...) pot ser menys reconeixible.
                        token_display = None
                        if not isinstance(obj, (Contract, ContractRequest)):
                            if hasattr(obj, 'contract') and obj.contract:
                                token_display = getattr(obj.contract, 'token', None)
                        
                        if not token_display:
                            token_display = getattr(obj, 'token', None)
                        
                        if not token_display:
                            token_display = str(obj.id)
                        
                        add_row(detail_sheet, detail_row, [f"{m:02d}/{y}", cat_name, token_display, obj_date, exploitation])
                        detail_row += 1

                row_data.append(count)
                total_cat += count
            
            row_data.append(total_cat)
            row = add_row(sheet, row, row_data)

        adjust_column_widths(sheet)
        adjust_column_widths(detail_sheet)
        filename = f"report_volum_gestions_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        document_id = save_report(content=wb, filename=filename, name=name, type_id=type_id, start_date=start_date, end_date=end_date)
        return None, document_id, filename
    except Exception as e:
        print(f"Error in generate_management_volume_report: {e}")
        traceback.print_exc()
        return None, None, None

def generate_complex_route_report(request, black_fill, white_bold_font):
    try:
        from openpyxl.styles import Alignment
        id = request.data.get('id')
        date_range = request.data.get('date_range')
        billing_ids = get_billing_ids(request.data)
        name = request.data.get('name', 'Informe de Facturació per Rutes')
        type_id = request.data.get('type_id')
        exploitation_id = request.data.get('exploitation_id')

        start_date = None
        end_date = None
        if billing_ids:
            invoices = Invoice.objects.filter(billing_id__in=billing_ids)
        elif date_range:
            start_date = parse_date(date_range[0])
            end_date = parse_date(date_range[1])
            if not start_date or not end_date:
                return None, None, None
            invoices = Invoice.objects.filter(issue_date__range=(start_date.date(), end_date.date()))
        elif has_serie_final_range(request.data):
            # Seleccio nomes pel rang de numero de factura: el rang ja identifica les
            # factures, sense lot ni periode. filter_pending_invoices hi aplica el rang
            # (apply_serie_final_range) i el periode surt de les factures resultants.
            invoices = filter_pending_invoices(serie_final_range_invoices(), request)
            start_date, end_date = serie_final_range_period(invoices)
            if start_date is None:
                return None, None, None
        else:
            return None, None, None

        multi_q = invoice_multi_filter_q(**get_multi_ids(request.data))
        if multi_q:
            invoices = invoices.filter(multi_q).distinct()

        if exploitation_id and exploitation_id != 'all':
            invoices = invoices.filter(exploitation_id=exploitation_id)

        invoices = filter_pending_invoices(invoices, request)

        # Optimize queries with select_related and prefetch_related
        invoices = invoices.select_related(
            'contract__supply_point_default__property__route_position__route',
            'contract_termination__contract__supply_point_default__property__route_position__route',
            'contract__use_type', 'contract__category',
            'origin'
        ).prefetch_related(
            'line_items__product__origin'
        )

        routes_data = {}
        detail_records = []

        # Helper to classify a single line item
        def classify_line(line):
            p_orig = line.product.origin.token if line.product and line.product.origin else ""
            desc = str(line.description or "").lower()
            n_str = str(line.name or "").lower()
            p_name = str(line.product_name or "").lower()
            full_t = f"{p_orig} {desc} {n_str} {p_name}"
            
            if p_orig in ['escomesa', 'escomeses'] or any(k in full_t for k in ['escomesa', 'acometida', 'escomeses', 'acometidas', 'connexio', 'connexió']):
                return 'escomeses'
            elif p_orig in ['subministrament', 'comptador', 'contador'] or any(k in full_t for k in ['comptador', 'contador', 'comptadors', 'contadores', 'subministrament', 'subministre', 'suministro']):
                return 'contadors'
            return None

        for inv in invoices:
            route_name = "Sense Ruta"
            contract = inv.contract
            if not contract and inv.contract_termination:
                contract = inv.contract_termination.contract

            route_obj = None
            if contract and contract.supply_point_default and contract.supply_point_default.property and contract.supply_point_default.property.route_position and contract.supply_point_default.property.route_position.route:
                route_obj = contract.supply_point_default.property.route_position.route
            elif inv.connection:
                # Fallback for invoices billed directly to a Connection
                sp = SupplyPoint.objects.filter(connection=inv.connection, property__route_position__route__isnull=False).select_related('property__route_position__route').first()
                if sp and sp.property and sp.property.route_position and sp.property.route_position.route:
                    route_obj = sp.property.route_position.route

            if route_obj:
                route_name = f"{route_obj.name or route_obj.token or 'Sense Ruta'}"

            if route_name not in routes_data:
                routes_data[route_name] = {
                    'rec_m3': 0.0,
                    'rec_quota': 0.0,
                    'rec_consum': 0.0,
                    'contadors_m3': 0.0,
                    'contadors_conservacio': 0.0,
                    'escomeses_m3': 0.0,
                    'escomeses_conservacio': 0.0,
                }

            active_lines = [line for line in inv.line_items.all() if line.is_active]

            # Determine invoice overall primary category
            inv_category = 'rec'
            inv_origin_token = inv.origin.token if inv.origin else ""
            
            if inv_origin_token in ['escomesa', 'escomeses'] or inv.connection:
                inv_category = 'escomeses'
            elif inv_origin_token in ['subministrament', 'comptador', 'contador']:
                inv_category = 'contadors'
            else:
                # Check contract fields if available
                if contract:
                    u_type = str(contract.use_type.name or "").lower() if contract.use_type else ""
                    c_cat = str(contract.category.name or "").lower() if contract.category else ""
                    combined_cat = f"{u_type} {c_cat}"
                    if any(k in combined_cat for k in ['escomesa', 'acometida']):
                        inv_category = 'escomeses'
                    elif any(k in combined_cat for k in ['comptador', 'contador', 'subministrament']):
                        inv_category = 'contadors'
                    else:
                        # Check if active line items override overall classification
                        line_cats = [classify_line(l) for l in active_lines]
                        valid_cats = [c for c in line_cats if c]
                        if valid_cats:
                            if 'escomeses' in valid_cats:
                                inv_category = 'escomeses'
                            else:
                                inv_category = 'contadors'

            consumption = float(inv.consumption or 0.0)
            if consumption > 0:
                if inv_category == 'escomeses':
                    routes_data[route_name]['escomeses_m3'] += consumption
                elif inv_category == 'contadors':
                    routes_data[route_name]['contadors_m3'] += consumption
                else:
                    routes_data[route_name]['rec_m3'] += consumption

            # Line items processing
            for line in active_lines:
                # Use line.price directly as standard across reports, fallback to line.total
                val = float(line.price if line.price is not None else (line.total or 0.0))
                if val == 0.0:
                    continue

                desc = str(line.description or "").lower()
                name_str = str(line.name or "").lower()
                prod_name = str(line.product_name or "").lower()
                full_desc = f"{desc} {name_str} {prod_name}"

                line_cat = classify_line(line)
                
                if line_cat == 'escomeses':
                    routes_data[route_name]['escomeses_conservacio'] += val
                elif line_cat == 'contadors':
                    routes_data[route_name]['contadors_conservacio'] += val
                else:
                    # Line doesn't explicitly match connection/meter keywords.
                    # Inherit invoice category if dedicated
                    if inv_category == 'escomeses':
                        routes_data[route_name]['escomeses_conservacio'] += val
                    elif inv_category == 'contadors':
                        routes_data[route_name]['contadors_conservacio'] += val
                    else:
                        # Invoice and line are both 'rec' or general. Check maintenance override
                        if any(k in full_desc for k in ['conservació', 'conservacio', 'conservación', 'conservacion', 'manteniment', 'mantenimiento']):
                            routes_data[route_name]['contadors_conservacio'] += val
                        elif any(k in full_desc for k in ['quota', 'cuota', 'servei', 'servicio', 'fix', 'fijo']):
                            routes_data[route_name]['rec_quota'] += val
                        else:
                            routes_data[route_name]['rec_consum'] += val

            # Save detail record
            num_factura = inv.serie_final or inv.number or inv.token or ""
            data_factura = inv.issue_date.strftime("%Y-%m-%d") if inv.issue_date else ""
            titular = inv.customer_final or ""
            total_factura = float(inv.total_final or 0.0)
            
            cat_display = inv_category.upper()
            detail_records.append([
                route_name, num_factura, data_factura, titular, cat_display, round(consumption, 4), round(total_factura, 2)
            ])

        # Prepare Workbook
        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.title = _("Resum per Rutes")

        row = 0
        titles_row1 = [_("Ruta"), _("REC"), "", "", _("CONTADORS"), "", _("ESCOMESES"), ""]
        row = add_row(sheet, row, titles_row1, fill=black_fill, font=white_bold_font)
        
        sheet.merge_cells(start_row=1, start_column=2, end_row=1, end_column=4)
        sheet.merge_cells(start_row=1, start_column=5, end_row=1, end_column=6)
        sheet.merge_cells(start_row=1, start_column=7, end_row=1, end_column=8)

        center_alignment = Alignment(horizontal="center", vertical="center")
        for col_idx in [2, 5, 7]:
            sheet.cell(row=1, column=col_idx).alignment = center_alignment

        titles_row2 = [
            "", 
            _("m³ Facturats"), _("Quota Servei (€)"), _("Consum (€)"), 
            _("m³ Facturats"), _("Conservació (€)"), 
            _("m³ Facturats"), _("Conservació (€)")
        ]
        row = add_row(sheet, row, titles_row2, fill=black_fill, font=white_bold_font)

        # Totals accumulators
        tot_rec_m3 = tot_rec_quota = tot_rec_consum = 0.0
        tot_cont_m3 = tot_cont_cons = 0.0
        tot_esco_m3 = tot_esco_cons = 0.0

        for r_name in sorted(routes_data.keys()):
            d = routes_data[r_name]
            tot_rec_m3 += d['rec_m3']
            tot_rec_quota += d['rec_quota']
            tot_rec_consum += d['rec_consum']
            tot_cont_m3 += d['contadors_m3']
            tot_cont_cons += d['contadors_conservacio']
            tot_esco_m3 += d['escomeses_m3']
            tot_esco_cons += d['escomeses_conservacio']

            row = add_row(sheet, row, [
                r_name,
                round(d['rec_m3'], 4),
                round(d['rec_quota'], 2),
                round(d['rec_consum'], 2),
                round(d['contadors_m3'], 4),
                round(d['contadors_conservacio'], 2),
                round(d['escomeses_m3'], 4),
                round(d['escomeses_conservacio'], 2),
            ])

        # Add Totals Row
        row = add_row(sheet, row, [
            _("TOTALS"),
            round(tot_rec_m3, 4),
            round(tot_rec_quota, 2),
            round(tot_rec_consum, 2),
            round(tot_cont_m3, 4),
            round(tot_cont_cons, 2),
            round(tot_esco_m3, 4),
            round(tot_esco_cons, 2),
        ], fill=black_fill, font=white_bold_font)

        # Prepare Detail Sheet
        detail_sheet = wb.create_sheet(title=_("Detall Factures"))
        detail_titles = [_("Ruta"), _("Num. Factura"), _("Data"), _("Titular"), _("Categoria"), _("m³ Facturats"), _("Total Factura (€)")]
        add_row(detail_sheet, 0, detail_titles, fill=black_fill, font=white_bold_font)
        detail_row = 1

        for record in detail_records:
            add_row(detail_sheet, detail_row, record)
            detail_row += 1

        adjust_column_widths(sheet)
        adjust_column_widths(detail_sheet)

        filename = f"report_rutes_creuat_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        document_id = save_report(content=wb, filename=filename, name=name, type_id=type_id, start_date=start_date, end_date=end_date)
        return None, document_id, filename
    except Exception as e:
        print(f"Error in generate_complex_route_report: {e}")
        import traceback
        traceback.print_exc()
        return None, None, None

