"""Dashboard de facturació.

Informe propi: una fila per factura, amb els imports agregats per producte i
els camps de contracte, lectura, tarifa social i ampliació de trams. No forma
part del padró de facturació.
"""
import datetime
from io import BytesIO

from django.db.models import Prefetch
from django.http import HttpResponse
from django.utils.translation import get_language, gettext as _
import openpyxl
from openpyxl.utils import get_column_letter
from rest_framework import status
from rest_framework.response import Response

from billing.models import Invoice, InvoiceLineItem, Reading
from billing.utils.invoice_range_utils import filter_invoices_by_serie_final_range
from contract.models import Contract, Variable
from coredata.models import ConfigProject
from coredata.utils.address_utils import get_address_complete_without_city
from coredata.utils.other_utils import round_ceil
from pricing.models import Product
from service.models import Exploitation
from statistics.utils.report_filters import (
    get_billing_ids, get_multi_ids, has_serie_final_range, invoice_multi_filter_q,
    serie_final_range_period,
)
from statistics.views.reports_views import filter_pending_invoices, save_report


_SOCIAL_TARIFF_FALSE = {'false', '0', 'n', 'no', 'fals'}
_ALERT_TRANSLATIONS = {
    'reading_alert_turn': {'ca': 'Volta de comptador', 'es': 'Vuelta de contador', 'en': 'Meter turn'},
    'reading_alert_zero': {'ca': 'Consum zero', 'es': 'Consumo cero', 'en': 'Zero consumption'},
    'reading_alert_negative': {'ca': 'Consum negatiu', 'es': 'Consumo negativo', 'en': 'Negative consumption'},
    'reading_alert_estimated': {'ca': 'Consum estimat', 'es': 'Consumo estimado', 'en': 'Estimated consumption'},
    'reading_alert_unusual': {'ca': 'Lectura inusual', 'es': 'Lectura inusual', 'en': 'Unusual reading'},
    'reading_alert_unusual_consumption': {'ca': 'Lectura inusual', 'es': 'Lectura inusual', 'en': 'Unusual reading'},
    'reading_alert_low_consumption': {'ca': 'Consum baix', 'es': 'Consumo bajo', 'en': 'Low consumption'},
}


def _address_relations(prefix):
    return [
        f'{prefix}__street__type',
        f'{prefix}__street_number__number_type',
        f'{prefix}__city',
        f'{prefix}__country',
    ]


def _translated_alert(alert_text):
    if not alert_text:
        return ""
    lang = (get_language() or 'ca').split('-')[0]
    stripped = alert_text.strip()
    translations = _ALERT_TRANSLATIONS.get(stripped)
    if not translations:
        return alert_text
    return translations.get(lang, translations.get('ca', alert_text))


def _sheet_title(base, used_titles):
    for forbidden in ':\\/?*[]':
        base = base.replace(forbidden, ' ')
    base = ' '.join(base.split())[:31] or str(_('Dashboard'))
    title = base
    suffix = 2
    while title in used_titles:
        title = f"{base[:28]} {suffix}"
        suffix += 1
    used_titles.add(title)
    return title


def _product_id_filter(raw_ids):
    if not raw_ids:
        return None
    try:
        return [int(item) for item in str(raw_ids).split(',') if item]
    except (TypeError, ValueError):
        return None


def _request_with_data(data):
    """Objecte mínim amb `.data` perquè `filter_pending_invoices` no muti el request original."""
    return type('ReportRequest', (), {'data': data})()


def _select_invoices(request):
    # Còpia local: no escrivim a request.data (QueryDict immutable / efectes laterals).
    data = request.data.copy() if hasattr(request.data, 'copy') else dict(request.data)
    if not data.get('generate_object', True):
        data['include_preinvoices'] = True
    filter_request = _request_with_data(data)

    date_range = data.get('date_range', None)
    billing_ids = get_billing_ids(data)
    serie_final_from = data.get('serie_final_from', None)
    serie_final_to = data.get('serie_final_to', None)
    prefix = data.get('prefix', None)
    name = data.get('name', '')
    type_id = data.get('type_id', None)
    exploitation_id = data.get('exploitation_id', None)
    if exploitation_id == 'all':
        exploitation_id = None

    has_serie_range = has_serie_final_range(data)
    if not billing_ids and not date_range and not has_serie_range:
        return None, None, Response(
            {"error": "billing id, date range or serie_final range is required."},
            status=status.HTTP_400_BAD_REQUEST,
        )

    start_date = None
    end_date = None
    if billing_ids:
        invoices = filter_pending_invoices(
            Invoice.objects.filter(billing_id__in=billing_ids), filter_request
        )
    elif date_range:
        start_date = datetime.datetime.strptime(date_range[0], '%Y-%m-%dT%H:%M:%S.%fZ')
        end_date = datetime.datetime.strptime(date_range[1], '%Y-%m-%dT%H:%M:%S.%fZ')
        invoice_type = ConfigProject.objects.get(token="invoice_type_invoice_token").value
        invoices = Invoice.objects.filter(
            issue_date__range=(start_date.date(), end_date.date()),
            type_final=invoice_type,
        ).distinct()
        invoices = filter_pending_invoices(invoices, filter_request)
    else:
        invoice_type = ConfigProject.objects.get(token="invoice_type_invoice_token").value
        invoices = filter_invoices_by_serie_final_range(
            Invoice.objects.filter(type_final=invoice_type),
            serie_final_from=serie_final_from,
            serie_final_to=serie_final_to,
            prefix=prefix,
        ).distinct()
        if data.get('exclude_reviewed', False):
            invoices = invoices.filter(reviewed=False)
        invoices = filter_pending_invoices(invoices, filter_request)
        start_date, end_date = serie_final_range_period(invoices)

    if exploitation_id:
        invoices = invoices.filter(exploitation_id=exploitation_id)

    multi_q = invoice_multi_filter_q(**get_multi_ids(data))
    if multi_q:
        invoices = invoices.filter(multi_q).distinct()
    return invoices, (start_date, end_date, name, type_id, exploitation_id), None


def _load_invoices(invoices):
    line_items_qs = InvoiceLineItem.objects.filter(is_active=True)
    readings_qs = Reading.objects.select_related('alert', 'remote_alert')
    variables_qs = Variable.objects.select_related('type')
    invoices = invoices.select_related(
        'exploitation',
        'contract',
        'contract__use_type',
        'contract__supply_point_default__address',
        'contract__supply_point_default__connection__dma',
        'contract_termination__contract',
        'contract_termination__contract__use_type',
        'contract_termination__contract__supply_point_default__address',
        'contract_termination__contract__supply_point_default__connection__dma',
        'serie',
        *_address_relations('contract__supply_point_default__address'),
        *_address_relations('contract_termination__contract__supply_point_default__address'),
    ).prefetch_related(
        Prefetch('line_items', queryset=line_items_qs, to_attr='prefetched_active_line_items'),
        Prefetch('readings', queryset=readings_qs),
        Prefetch('contract__variables', queryset=variables_qs),
        Prefetch('contract_termination__contract__variables', queryset=variables_qs),
    ).order_by('exploitation_id', 'id')
    invoice_list = list(invoices)

    request_ids = [
        invoice.contract_request_id
        for invoice in invoice_list
        if not invoice.contract_id
        and not (
            invoice.contract_termination_id
            and invoice.contract_termination
            and invoice.contract_termination.contract_id
        )
        and invoice.contract_request_id
    ]
    contracts_by_request = {}
    if request_ids:
        for contract in (
            Contract.objects.filter(contract_request_id__in=request_ids)
            .select_related(
                'use_type',
                'supply_point_default__address',
                'supply_point_default__connection__dma',
                *_address_relations('supply_point_default__address'),
            )
            .prefetch_related(Prefetch('variables', queryset=Variable.objects.select_related('type')))
        ):
            contracts_by_request[contract.contract_request_id] = contract
    return invoice_list, contracts_by_request


def _resolve_contract(invoice, contracts_by_request):
    if invoice.contract_id:
        return invoice.contract
    termination = invoice.contract_termination
    if termination and termination.contract_id:
        return termination.contract
    if invoice.contract_request_id:
        return contracts_by_request.get(invoice.contract_request_id)
    return None


def _line_items(invoice, product_ids):
    items = getattr(invoice, 'prefetched_active_line_items', None)
    if items is None:
        items = [line for line in invoice.line_items.all() if line.is_active]
    if product_ids:
        items = [line for line in items if line.product_id in product_ids]
    return items


def _product_columns(invoices, product_ids):
    present_ids = set()
    has_other = False
    for invoice in invoices:
        for line in _line_items(invoice, product_ids):
            if line.product_id:
                present_ids.add(line.product_id)
            else:
                has_other = True
    if product_ids:
        present_ids.update(product_ids)
        has_other = False

    products = list(
        Product.objects.filter(id__in=present_ids).select_related('exploitation')
    ) if present_ids else []
    products.sort(key=lambda item: (item.position if item.position is not None else 0, item.name or '', item.id))

    name_counts = {}
    for product in products:
        name_counts[product.name or ''] = name_counts.get(product.name or '', 0) + 1
    columns = []
    for product in products:
        label = product.name or ''
        if name_counts.get(label, 0) > 1:
            qualifier = (
                product.exploitation.name if product.exploitation and product.exploitation.name
                else product.token or str(product.id)
            )
            label = f"{label} ({qualifier})"
        columns.append((product.id, label))
    if has_other:
        columns.append((None, str(_("Altres"))))
    return columns


def _has_social_tariff(contract, issue_date, tarifa_social_token):
    if not contract or not tarifa_social_token:
        return False
    for variable in contract.variables.all():
        if not variable.is_active or not variable.type:
            continue
        if (variable.type.token or '') != tarifa_social_token:
            continue
        if (variable.value or '').strip().lower() in _SOCIAL_TARIFF_FALSE:
            continue
        if issue_date:
            if variable.start_at and variable.start_at > issue_date:
                continue
            if variable.end_at and variable.end_at < issue_date:
                continue
        return True
    return False


def _sector_name(contract):
    supply_point = contract.supply_point_default if contract else None
    connection = supply_point.connection if supply_point else None
    dma = connection.dma if connection else None
    return dma.name if dma and dma.name else ""


def _reading_flags(invoice):
    readings = list(invoice.readings.all())
    estimated_used = sum(reading.estimated_used or 0 for reading in readings)
    if not readings:
        return estimated_used, "", ""
    reading = readings[0]
    estimated = "S" if reading.is_estimated else "N"
    observation = _translated_alert(
        reading.alert_notes if reading.alert_notes
        else reading.remote_alert.name if reading.remote_alert
        else reading.alert.name if reading.alert else ""
    )
    return estimated_used, estimated, observation


def _tram_extension(invoice, contract):
    if invoice.persons_final is not None:
        return invoice.persons_final
    if contract and contract.total_persons is not None:
        return contract.total_persons
    return ""


def _invoice_row(invoice, contract, product_columns, product_ids, tarifa_social_token):
    estimated_used, reading_estimated, reading_observation = _reading_flags(invoice)
    address = (
        get_address_complete_without_city(contract.supply_point_default.address)
        if contract and contract.supply_point_default and contract.supply_point_default.address
        else invoice.address_final
    )
    row = [
        invoice.customer_final,
        invoice.customer_token_final,
        address,
        invoice.exploitation.name if invoice.exploitation else "",
        contract.token if contract else invoice.customer_token_final,
        invoice.issue_date.strftime('%d/%m/%Y') if invoice.issue_date else "",
        invoice.serie_final,
        invoice.serie.name if invoice.serie else "Factura Original",
        _sector_name(contract),
        invoice.payment_type_final or "",
        contract.use_type.name if contract and contract.use_type else "",
        contract.get_use_aca_display() if contract and contract.use_aca else "",
        reading_estimated,
        reading_observation,
        "S" if _has_social_tariff(contract, invoice.issue_date, tarifa_social_token) else "N",
        _tram_extension(invoice, contract),
        float(invoice.consumption) - float(estimated_used) if invoice.consumption else 0,
        invoice.subtotal_final,
        invoice.total_final - invoice.subtotal_final,
        invoice.total_final,
    ]

    product_id_set = {key for key, _label in product_columns if key is not None}
    amounts = {}
    for line in _line_items(invoice, product_ids):
        product_key = line.product_id if line.product_id in product_id_set else None
        subtotal, tax_value = amounts.get(product_key, (0, 0))
        amounts[product_key] = (
            subtotal + round_ceil(line.price),
            tax_value + round_ceil(line.tax_price),
        )
    for product_key, _label in product_columns:
        subtotal, tax_value = amounts.get(product_key, (0, 0))
        subtotal = round_ceil(subtotal)
        tax_value = round_ceil(tax_value)
        row.extend([subtotal, tax_value, round_ceil(subtotal + tax_value)])
    return row


def _headers():
    reading_prefix = "%s %s" % (_("Lectura"), 1)
    return [
        _("Abonat"), _("NIF/CIF"), _("Dom. del subministre"),
        _("Població"), _("Identificador"), _("Data emissió"),
        _("Núm. factura"), _("Tipus factura"),
        _("Sector"), _("Tipus pagament"), _("Tipus d'ús contracte"), _("Ús ACA"),
        f"{reading_prefix} · {_('Estimada')}",
        f"{reading_prefix} · {_('Observ. actual')}",
        _("Tarifa Social"), _("Ampliació de Trams"),
        _("Consum (m3)"),
        _("Total sense IVA"), _("Import IVA Total"), _("Total"),
    ]


def _render_sheet(sheet, invoices, contracts_by_request, product_ids, tarifa_social_token, task, progress_offset, total):
    product_columns = _product_columns(invoices, product_ids)
    titles = _headers()
    for _product_key, label in product_columns:
        titles.append(f"{label} · {_('Subtotal')}")
        titles.append(f"{label} · {_('IVA')}")
        titles.append(f"{label} · {_('Total')}")
    amount_start = titles.index(_("Consum (m3)"))

    for col_idx, title in enumerate(titles, start=1):
        width = min(max(len(str(title)) + 2, 10), 40)
        sheet.column_dimensions[get_column_letter(col_idx)].width = width
    sheet.append(titles)

    totals_by_column = {}
    for index, invoice in enumerate(invoices, start=1):
        if task and (progress_offset + index) % 200 == 0:
            task.update_state(
                state='PROGRESS',
                meta={
                    'current': progress_offset + index,
                    'total': total,
                    'percent': round(((progress_offset + index) / total) * 100, 2) if total else 100,
                }
            )
        row = _invoice_row(
            invoice,
            _resolve_contract(invoice, contracts_by_request),
            product_columns,
            product_ids,
            tarifa_social_token,
        )
        for column_idx in range(amount_start, len(row)):
            try:
                value = float(row[column_idx])
            except (TypeError, ValueError):
                continue
            totals_by_column[column_idx] = totals_by_column.get(column_idx, 0) + value
        sheet.append(row)

    if invoices:
        totals_row = [""] * len(titles)
        totals_row[0] = _("TOTAL")
        for column_idx, value in totals_by_column.items():
            totals_row[column_idx] = round_ceil(value)
        sheet.append(totals_row)


def _sheet_plan(invoices, exploitation_id):
    if exploitation_id:
        return [(None, invoices)]
    grouped = {}
    for invoice in invoices:
        grouped.setdefault(invoice.exploitation_id, []).append(invoice)
    if len(grouped) <= 1:
        return [(None, invoices)]
    exploitations = {
        item.id: item
        for item in Exploitation.objects.filter(id__in=[key for key in grouped if key is not None])
    }
    plan = []
    for exploitation_key, sheet_invoices in grouped.items():
        exploitation = exploitations.get(exploitation_key)
        label = exploitation.name if exploitation and exploitation.name else str(_("Sense explotació"))
        plan.append((label, sheet_invoices))
    plan.sort(key=lambda item: item[0])
    return plan


def generate_billing_product_dashboard(request, _black_fill=None, _white_bold_font=None, task=None):
    # _black_fill / _white_bold_font: signatura compartida amb run_report_task; aquest Excel no els fa servir.
    invoices, meta, error = _select_invoices(request)
    if error:
        return error, None, None
    start_date, end_date, name, type_id, exploitation_id = meta
    product_ids = _product_id_filter(request.data.get('product_ids', ''))
    if request.data.get('product_ids') and product_ids is None:
        return Response(
            {"error": "product_ids must be a comma-separated list of integers."},
            status=status.HTTP_400_BAD_REQUEST,
        ), None, None
    invoice_list, contracts_by_request = _load_invoices(invoices)
    total = len(invoice_list)

    tarifa_social_config = ConfigProject.objects.filter(token='tarifa_social').first()
    tarifa_social_token = (tarifa_social_config.value or '').strip() if tarifa_social_config else ''

    wb = openpyxl.Workbook()
    # Treiem el full buit per defecte: només es creen els fulls amb dades.
    default_sheet = wb.active
    wb.remove(default_sheet)
    used_titles = set()
    rendered = 0
    for label, sheet_invoices in _sheet_plan(invoice_list, exploitation_id):
        if label:
            title = _sheet_title(label, used_titles)
        else:
            title = _sheet_title(str(_("Dashboard de facturació")), used_titles)
        sheet = wb.create_sheet(title=title)
        _render_sheet(
            sheet, sheet_invoices, contracts_by_request, product_ids,
            tarifa_social_token, task, rendered, total,
        )
        rendered += len(sheet_invoices)

    filename = f"billing_dashboard_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    buffer = BytesIO()
    wb.save(buffer)
    content = buffer.getvalue()
    response = HttpResponse(
        content,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f"attachment; filename={filename}"
    document_id = save_report(
        content=content, filename=filename, name=name, type_id=type_id,
        start_date=start_date, end_date=end_date,
    )
    return response, document_id, None
