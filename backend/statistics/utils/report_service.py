# FOR ACCOUNTING REPORTS EXPORTATION
import datetime
from django.utils.dateparse import parse_date
import os
import threading
from django.core.files.base import ContentFile
from django.db import connection
from django.db.models import F, Q, Count, Max, Sum, OuterRef, Subquery, Prefetch
from django.http import HttpResponse, JsonResponse
from django.utils.translation import gettext as _
import openpyxl
import xml.etree.ElementTree as ET
from io import BytesIO
from openpyxl.styles import Alignment, PatternFill
from openpyxl.utils import get_column_letter
from rest_framework.response import Response
from django.core.files.storage import default_storage
from rest_framework import status

from billing.models import Biller, CommitmentDeposit, Invoice, InvoiceLineItem, Payment, PaymentRemittance, Reading
from billing.utils.invoice_service import get_invoice_status, get_totals
from billing.utils.payment_service import get_status_map
from billing.views.detailed_billing_document_generate_view import normalize_text, repeat_to_at_least_length
from contract.models import Bail, Contract, ContractTerminationRequest, PaymentType, Variable
from coredata.models import Bank, ConfigProject, StreetType
from coredata.utils.address_utils import get_address_complete_without_city
from coredata.utils.other_utils import round_ceil
from coredata.utils.validators_utils import validate_nif
from logger.models import LogCommitmentDepositMovement, LogPaymentStatusChange
from order.models import Operator, Order, OrderReport
from pricing.models import Adjustment, BillingRange, Product, Tax
from service.models import Company, Exploitation, Meter
from statistics.utils.incasol import resolve_incasol_num
from statistics.views.reports_views import add_row, adjust_column_widths, filter_pending_invoices, jump_row, save_report
from statistics.utils.report_filters import (contract_multi_filter_q, get_billing_ids, get_multi_ids, has_serie_final_range,
                                             invoice_multi_filter_q, serie_final_range_invoices, serie_final_range_period)
from billing.utils.invoice_range_utils import filter_invoices_by_serie_final_range


def delete_file_later(path: str, delay_seconds: int = 100) -> None:
    def _run():
        try:
            default_storage.delete(path)
        except Exception:
            pass
    t = threading.Timer(delay_seconds, _run)
    t.daemon = True
    t.start()
    
def generate_accounting_values_report(request, black_fill=None, white_bold_font=None, task=None):

    print("getting billing summary")
    id = request.data.get('id', None)
    date_range = request.data.get('date_range', None)
    billing_ids = get_billing_ids(request.data)
    name = request.data.get('name', '')
    type_id = request.data.get('type_id', None)
    exploitation_id = request.data.get('exploitation_id', None)
    accounting_type = request.data.get('accounting_type', None) or request.data.get('account_type', None)
    if not accounting_type:
        accounting_type = "INVOICE"
    if accounting_type == "INVOICES":
        accounting_type = "INVOICE"
    elif accounting_type == "COMMITMENTDEPOSIT":
        accounting_type = "COMMITMENT"
    print("exploitation_id: ", exploitation_id)
    print("type_id: ", type_id)
    print("name: ", name)
    print("id: ", id)
    print("billing_ids: ", billing_ids)
    print("date_range: ", date_range)
    has_serie_range = has_serie_final_range(request.data)
    if not billing_ids and not date_range and not has_serie_range:
        return Response({"error": "billing id, date range or serie_final range is required."}, status=status.HTTP_400_BAD_REQUEST), None, None
    print("got id or date range")
    billing = None
    billings = []
    start_date = None
    end_date = None
    
    exploitation = None
    if exploitation_id:
        exploitation = Exploitation.objects.get(id=exploitation_id)
    
    invoice_type = ConfigProject.objects.get(token="invoice_type_invoice_token").value
    commitment_cancelled_token = ConfigProject.objects.get(token="commitment_deposit_status_cancelled_token").value
    commitments = []
    if billing_ids:
        invoices = Invoice.objects.filter(billing_id__in=billing_ids)
        if exploitation:
            invoices = invoices.filter(exploitation=exploitation)
        commitments = CommitmentDeposit.objects.filter(invoices__in=invoices).distinct()
    elif date_range:
        start_date = datetime.datetime.strptime(date_range[0], '%Y-%m-%dT%H:%M:%S.%fZ')
        end_date = datetime.datetime.strptime(date_range[1], '%Y-%m-%dT%H:%M:%S.%fZ')
        invoices = Invoice.objects.filter(issue_date__range=(start_date.date(), end_date.date()), type_final=invoice_type).distinct()
        invoices = filter_pending_invoices(invoices, request)
        
        commitments = CommitmentDeposit.objects.filter(created_at__range=(start_date.date(), end_date.date())).exclude(status__token=commitment_cancelled_token).distinct()
        if exploitation:
            invoices = invoices.filter(exploitation=exploitation)   
            commitments = commitments.filter(invoices__exploitation=exploitation)
        
    elif has_serie_range:
        # Seleccio nomes pel rang de numero de factura: el rang ja identifica les
        # factures, sense lot ni periode. Els dipositos son els lligats a aquestes
        # factures (com a la seleccio per lot) i el periode surt de les factures
        # resultants, que es el que fan servir els titols i el nom del fitxer.
        invoices = filter_pending_invoices(serie_final_range_invoices(), request)
        start_date, end_date = serie_final_range_period(invoices)
        if start_date is None:
            return Response({"error": "the serie_final range has no invoices."}, status=status.HTTP_400_BAD_REQUEST), None, None
        commitments = CommitmentDeposit.objects.filter(invoices__in=invoices).distinct()
        if exploitation:
            invoices = invoices.filter(exploitation=exploitation)
            commitments = commitments.filter(invoices__exploitation=exploitation)
    else:
        raise Exception("Error obtaining invoices for report billing summary")

    multi_q = invoice_multi_filter_q(**get_multi_ids(request.data))
    if multi_q:
        invoices = invoices.filter(multi_q).distinct()

    wb = openpyxl.Workbook()
    sheet = wb.active
    payment_status_map = get_status_map()

    invoice_status_payoff_token = ConfigProject.objects.get(token="invoice_status_payoff_token").value
    main_company = Company.objects.get(vat=ConfigProject.objects.get(token="main_company_token").value)
    main_company_bank = main_company.company_banks.filter(is_sepa=True).first()

    payment_method_values = {
        "DIRECT_DEBIT": "02",
        "BANK_TRANSFER": "03",
        "BANK_PAYMENT": "06",
        "CASH": "01",
        "CARD": "05",
        "BALANCE": "07",
        "TPV_ONLINE": "05"
    }

    row = 0
    if accounting_type == "PAYMENT":
        
        sheet.title = _("Cobraments")
        # payments = Payment.objects.filter(invoice__in=invoices)
        payments_inv = Payment.objects.filter(invoice__in=invoices)
        #paid_status_token = ConfigProject.objects.get(token='payment_status_paid_token').value
        paid_status_token = payment_status_map['payment_status_paid_token'].token
        piggy_status_token = payment_status_map['payment_status_piggy_token'].token
        returned_status_token = payment_status_map['payment_status_returned_token'].token
        payment_cancelled_token = payment_status_map['payment_status_cancelled_token'].token
        
        log_payments_inv = []
        log_payments_comm = []
        payments_without_log = []
        returned_invoices = []
        
        if billing_ids:
            log_payments_inv = LogPaymentStatusChange.objects.filter(
                object__in=payments_inv
                ).filter(
                    Q(current_status__token__in=[paid_status_token, piggy_status_token]) | Q(previous_status__token__in=[paid_status_token])
                ).order_by('timestamp').distinct()
            log_payments_comm = LogCommitmentDepositMovement.objects.filter(
                object__in=commitments,
                new_wallet__in=Payment.objects.filter(commitment_deposit__in=commitments),
            ).exclude(previous_remaining=F('current_remaining')).order_by('timestamp').distinct()
            payments_without_log = payments_inv.filter(
                status__token__in=[paid_status_token, piggy_status_token],
                ).exclude(id__in=log_payments_inv.values_list('object__id', flat=True)).order_by('payment_date').distinct()
            returned_invoices = invoices.filter(status__token=invoice_status_payoff_token)
        elif date_range:
            start_date_date = start_date.date() if hasattr(start_date, 'date') else start_date
            end_date_date = end_date.date() if hasattr(end_date, 'date') else end_date
            log_payments_by_timestamp = LogPaymentStatusChange.objects.filter(
                timestamp__range=(start_date, end_date),
                object__invoice__isnull=False,
                object__is_active=True
            ).filter(
                Q(current_status__token__in=[paid_status_token, piggy_status_token]) | Q(previous_status__token__in=[paid_status_token])
            )
            log_payments_by_sent_at = LogPaymentStatusChange.objects.filter(
                object__remittances__sent_at__isnull=False,
                object__remittances__sent_at__range=(start_date_date, end_date_date),
                object__invoice__isnull=False,
                object__is_active=True
            ).filter(
                Q(current_status__token__in=[paid_status_token, piggy_status_token]) | Q(previous_status__token__in=[paid_status_token])
            )
            log_payments_inv = (log_payments_by_timestamp | log_payments_by_sent_at).order_by('timestamp').distinct()
            log_payments_comm = LogCommitmentDepositMovement.objects.filter(
                timestamp__range=(start_date, end_date),
                new_wallet__isnull=False,
            ).exclude(previous_remaining=F('current_remaining')
                      ).exclude(object__status__token=commitment_cancelled_token).order_by('timestamp').distinct()
            inv_comm_payment_ids = list(log_payments_inv.values_list('object__id', flat=True)) + list(log_payments_comm.values_list('object__id', flat=True))
            payments_without_log = Payment.objects.exclude(
                id__in=inv_comm_payment_ids).filter(
                payment_date__range=(start_date, end_date),
                status__token__in=[paid_status_token, piggy_status_token],
                is_active=True,
            ).order_by('payment_date').distinct()
            returned_invoices = invoices.filter(status__token=invoice_status_payoff_token)
        if exploitation:
            log_payments_inv = log_payments_inv.filter(
                object__invoice__exploitation=exploitation
            ).distinct()
            log_payments_comm = log_payments_comm.filter(
                object__invoices__exploitation=exploitation
            ).distinct()
            payments_comm_without_log = payments_without_log.filter(
                invoice__exploitation=exploitation,
                payment_date__isnull=False
                ).filter(
                    status__token__in=[paid_status_token, piggy_status_token]
                )
            payments_inv_without_log = payments_without_log.filter(
                commitment_deposit__invoices__exploitation=exploitation,
                payment_date__isnull=False
                ).exclude(
                    commitment_deposit__status__token=commitment_cancelled_token
                ).distinct()
            payments_no_without_log = payments_without_log.filter(
                Q(invoice__isnull=True, commitment_deposit__isnull=True),
                payment_date__isnull=False
                )
            payments_without_log = payments_comm_without_log | payments_inv_without_log | payments_no_without_log
            returned_invoices = returned_invoices.filter(exploitation=exploitation)
        
        print("total payments logs: ", log_payments_inv.count())
        print("total payments commitments logs: ", log_payments_comm.count())
        print("total payments without log: ", payments_without_log.count())
        

        row = 0
        titles = [ 
                    _("TipusCobrament"), _("AnyConveni"),
                    _("NumeroConveni"), _("FormaPagament"),
                    _("TerminiConveni"), _("IdRemesa"),
                    _("DataRemesa"), _("IBANEmpresaRemesa"),
                    _("DataCobrament"), _("Import"), _("IBANClient"),
                    _("FacturaClient"), _("Abonament")
                    ]

        row = add_row(sheet, row, titles)
        counter_log_payments_inv = 0
        log_payments_inv_list = list(log_payments_inv)
        start_date_date = start_date.date() if date_range and start_date and hasattr(start_date, 'date') else None
        end_date_date = end_date.date() if date_range and end_date and hasattr(end_date, 'date') else None

        def _get_used_date_str(item_inv):
            ts_date = item_inv.timestamp.date() if isinstance(item_inv.timestamp, datetime.datetime) else item_inv.timestamp
            payment_date_near_timestamp = item_inv.object.payment_date and abs((item_inv.object.payment_date - ts_date).days) <= 2
            if item_inv.object.remittances and item_inv.object.remittances.filter(sent_at__isnull=False).count() > 0 and item_inv.current_status.token in [paid_status_token, piggy_status_token]:
                return item_inv.object.remittances.filter(sent_at__isnull=False).order_by('-created_at').first().sent_at.strftime('%Y%m%d')
            return item_inv.timestamp.strftime('%Y%m%d')

        def _get_used_date_str_for_payment(item_log):
            if (
                item_log.remittances
                and item_log.remittances.filter(sent_at__isnull=False).count() > 0
                and item_log.status.token in [paid_status_token, piggy_status_token]
            ):
                return (
                    item_log.remittances.filter(sent_at__isnull=False)
                    .order_by('-created_at')
                    .first()
                    .sent_at.strftime('%Y%m%d')
                )
            if item_log.payment_date and item_log.status.token in [paid_status_token, piggy_status_token]:
                return item_log.payment_date.strftime('%Y%m%d')
            return item_log.timestamp.strftime('%Y%m%d')

        def _process_log_payment_item(item_inv):
            nonlocal counter_log_payments_inv
            if item_inv.object.status.token == payment_cancelled_token and item_inv.current_status.token == payment_cancelled_token:
                return
            
            used_date_obj = None
    
            if date_range and start_date_date and end_date_date:
                payment_date = item_inv.object.payment_date
                used_date_str = _get_used_date_str(item_inv)
                used_date_obj = datetime.datetime.strptime(used_date_str, '%Y%m%d').date()
                if (
                    payment_date
                    and item_inv.current_status.token in [paid_status_token, piggy_status_token]
                ):
                    if abs(item_inv.timestamp.date() - payment_date) < datetime.timedelta(days=5):
                        used_date_obj = (
                            payment_date.date()
                            if isinstance(payment_date, datetime.datetime)
                            else payment_date
                        ) 
                
                if not (start_date_date <= used_date_obj <= end_date_date):
                    return
            counter_log_payments_inv += 1
            if counter_log_payments_inv % 1000 == 0:
                print(f"total payments logs processed: {counter_log_payments_inv} out of {len(log_payments_inv_list)}")
            if task and counter_log_payments_inv % 50 == 0:
                total_log_payments_inv = len(log_payments_inv_list)
                task.update_state(
                    state='PROGRESS',
                    meta={
                        'current': counter_log_payments_inv,
                        'total': total_log_payments_inv,
                        'percent': round((counter_log_payments_inv / total_log_payments_inv) * 100, 2) if total_log_payments_inv else 0.0
                    }
                )
            bail = None
            try:
                bail = Bail.objects.filter(invoice=item_inv.object.invoice).first()
            except Exception:
                pass
            try:
                payment_method_value = payment_method_values[item_inv.object.payment_type_token]
            except Exception:
                payment_method_value = "01"
            ts_date = item_inv.timestamp.date() if isinstance(item_inv.timestamp, datetime.datetime) else item_inv.timestamp
            payment_date_near_timestamp = item_inv.object.payment_date and abs((item_inv.object.payment_date - ts_date).days) <= 2
            used_date_str = used_date_obj.strftime('%Y%m%d') if used_date_obj else ""

            remittances_qs = item_inv.object.remittances.filter(sent_at__isnull=False)
            remittance = None
            if remittances_qs.exists():
                if used_date_obj:
                    remittances_list = list(remittances_qs)

                    def _remittance_date(rem):
                        sent = rem.sent_at
                        return sent.date() if isinstance(sent, datetime.datetime) else sent

                    remittance = min(
                        remittances_list,
                        key=lambda r: abs((_remittance_date(r) - used_date_obj).days),
                    )
                else:
                    remittance = remittances_qs.order_by('-created_at').first()
            line_row = [
                "04" if bail else "01", "",
                "", payment_method_value,
                "",
                remittance.token if remittance and payment_method_value == "02" else "",
                remittance.sent_at.strftime('%Y%m%d') if remittance and payment_method_value == "02" else "",
                main_company_bank.iban if payment_method_value in ['02', '03'] else "",
                used_date_str,
                item_inv.object.amount if item_inv.current_status.token in [paid_status_token, piggy_status_token] else -item_inv.object.amount,
                item_inv.object.payment_bank, item_inv.object.invoice.serie_final, "N"
            ]
            return line_row

        for invoice in returned_invoices:
            try:
                payment_method_value = payment_method_values[invoice.payment_type_token_final]
            except Exception:
                payment_method_value = "01"
            inv_line_row = [
                "01",
                "", "",
                payment_method_value,
                "", "", "", "",
                invoice.issue_date.strftime('%Y%m%d'), invoice.total_final,
                invoice.payment_bank_final if payment_method_value == "02" else "", 
                invoice.serie_final, "S"
            ]
            row = add_row(sheet, row, inv_line_row)
        
        for item_inv in log_payments_inv_list:
            
            has_remittances_sent = item_inv.object.remittances and item_inv.object.remittances.filter(sent_at__isnull=False).count() > 0
            ts_date = item_inv.timestamp.date() if isinstance(item_inv.timestamp, datetime.datetime) else item_inv.timestamp
            payment_date_near_timestamp = item_inv.object.payment_date and abs((item_inv.object.payment_date - ts_date).days) <= 2
            if item_inv.current_status.token in [paid_status_token, piggy_status_token]:
                if has_remittances_sent and item_inv.current_status.token in [paid_status_token, piggy_status_token]:
                    line_row = _process_log_payment_item(item_inv)
                    if line_row:
                        row = add_row(sheet, row, line_row)
                elif not has_remittances_sent and item_inv.object.payment_date and payment_date_near_timestamp:
                    line_row = _process_log_payment_item(item_inv)
                    if line_row:
                        row = add_row(sheet, row, line_row)
            if not (has_remittances_sent and item_inv.current_status.token in [paid_status_token, piggy_status_token]) and not (item_inv.object.payment_date and item_inv.current_status.token in [paid_status_token, piggy_status_token] and payment_date_near_timestamp):
                line_row = _process_log_payment_item(item_inv)
                if line_row:
                    row = add_row(sheet, row, line_row)

        counter_log_payments_comm = 0
        for item_com in log_payments_comm:
            counter_log_payments_comm += 1
            if counter_log_payments_comm % 1000 == 0:
                print(f"total payments commitments logs processed: {counter_log_payments_comm} out of {log_payments_comm.count()}")
            try:
                payment_method_value = payment_method_values[item_com.object.payment_type_token]
            except:
                payment_method_value = "01"
            line_row = [
                "90", item_com.object.created_at.strftime('%Y') if item_com.object else "", 
                item_com.object.token if item_com.object else "",  payment_method_value,
                item_com.object.due_date.strftime('%Y%m%d') if item_com.object.commitment_deposit and item_com.object.commitment_deposit.due_date else "", 
                item_com.object.remittances.filter(sent_at__isnull=False).order_by('created_at').first().token if item_com.object.remittances and item_com.object.remittances.filter(sent_at__isnull=False).count() > 0 and payment_method_value == "02" else "",
                item_com.new_wallet.remittances.filter(sent_at__isnull=False).order_by('-created_at').first().sent_at.strftime('%Y%m%d') if item_com.new_wallet.remittances and item_com.new_wallet.remittances.filter(sent_at__isnull=False).count() > 0 and payment_method_value == "02" else "",
                main_company_bank.iban if payment_method_value in ['02', '03'] else "",
                item_com.new_wallet.remittances.filter(sent_at__isnull=False).order_by('-created_at').first().sent_at.strftime('%Y%m%d') if item_com.new_wallet.remittances and item_com.new_wallet.remittances.filter(sent_at__isnull=False).count() > 0 and item_com.new_wallet.status.token in [paid_status_token, piggy_status_token] else item_com.new_wallet.payment_date.strftime('%Y%m%d') if item_com.new_wallet.payment_date and item_com.new_wallet.status.token in [paid_status_token, piggy_status_token] else item_com.timestamp.strftime('%Y%m%d'),
                item_com.new_wallet.amount if item_com.previous_remaining < item_com.current_remaining else -item_com.new_wallet.amount, 
                item_com.new_wallet.payment_bank, item_com.object.commitment_deposit.token, "N"
            ]
            row = add_row(sheet, row, line_row)
            
        counter_payments_without_log = 0
        for item_log in payments_without_log:
            
            counter_payments_without_log += 1
            if item_log.status.token == payment_cancelled_token:
                continue

            used_date_str = _get_used_date_str_for_payment(item_log)
            if date_range and start_date_date and end_date_date:
                used_date_obj = datetime.datetime.strptime(used_date_str, '%Y%m%d').date()
                if not (start_date_date <= used_date_obj <= end_date_date):
                    continue

            if counter_payments_without_log % 1000 == 0:
                print(
                    f"total payments without log processed: {counter_payments_without_log} "
                    f"out of {payments_without_log.count()}"
                )
            bail = None
            try:
                bail = Bail.objects.filter(invoice=item_log.invoice).first() if item_log.invoice else None
            except Exception:
                pass
            try:
                payment_method_value = payment_method_values[item_log.payment_type_token]
            except Exception:
                payment_method_value = "01"
            line_row = [
                "04" if bail else "90" if item_log.commitment_deposit else "01",
                item_log.commitment_deposit.created_at.strftime('%Y') if item_log.commitment_deposit else "",
                item_log.commitment_deposit.token if item_log.commitment_deposit else "",
                payment_method_value,
                item_log.commitment_deposit.due_date.strftime('%Y%m%d')
                if item_log.commitment_deposit and item_log.commitment_deposit.due_date
                else "",
                item_log.remittances.filter(sent_at__isnull=False)
                .order_by('created_at')
                .first()
                .token
                if item_log.remittances and payment_method_value == "02"
                and item_log.remittances.filter(sent_at__isnull=False).count() > 0
                else "",
                item_log.remittances.filter(sent_at__isnull=False)
                .order_by('-created_at')
                .first()
                .sent_at.strftime('%Y%m%d')
                if item_log.remittances and payment_method_value == "02"
                and item_log.remittances.filter(sent_at__isnull=False).count() > 0
                else "",
                main_company_bank.iban if payment_method_value in ['02', '03'] else "",
                used_date_str,
                item_log.amount,
                item_log.payment_bank,
                item_log.invoice.serie_final if item_log.invoice else "", "N"
            ]
            row = add_row(sheet, row, line_row)

        if row >= 2:
            data_rows = []
            for r in range(2, row + 1):
                data_rows.append([sheet.cell(row=r, column=c).value for c in range(1, 14)])
            data_rows.sort(key=lambda r: r[8] or "")
            for i, data_row in enumerate(data_rows):
                for c, val in enumerate(data_row, start=1):
                    sheet.cell(row=2 + i, column=c).value = val
    elif accounting_type == "INVOICE":
        sheet.title = _("Factures")
        titles = [
            _("Núm. Factura"), _("Data emissió"), _("Titular"), 
            _("NIF"), _("Base"), _("IVA"), _("Total"), _("Estat")
        ]
        row = add_row(sheet, row, titles, fill=black_fill, font=white_bold_font)
        
        # Sort invoices by issue_date
        sorted_invoices = invoices.select_related('status').order_by('issue_date')
        total_invoices = invoices.count()

        for idx, inv in enumerate(sorted_invoices, start=1):
            if task and idx % 50 == 0:
                task.update_state(
                    state='PROGRESS',
                    meta={
                        'current': idx,
                        'total': total_invoices,
                        'percent': round((idx / total_invoices) * 100, 2) if total_invoices else 0.0
                    }
                )
            row_data = [
                inv.serie_final or inv.number or inv.token,
                inv.issue_date.strftime('%d/%m/%Y') if inv.issue_date else "",
                inv.customer_final or "",
                inv.customer_token_final or "",
                round(inv.subtotal_final or 0, 2),
                round((inv.total_final or 0) - (inv.subtotal_final or 0), 2),
                round(inv.total_final or 0, 2),
                inv.status.name if inv.status else ""
            ]
            row = add_row(sheet, row, row_data)

    elif accounting_type == "COMMITMENT":
        sheet.title = _("Compromisos")
        titles = [
            _("Token"), _("Data creació"), _("Titular"), 
            _("NIF"), _("Import"), _("Pendent"), _("Estat")
        ]
        row = add_row(sheet, row, titles, fill=black_fill, font=white_bold_font)
        
        # commitments is already fetched earlier
        sorted_commitments = commitments.select_related('status', 'contract__holder').order_by('created_at')
        total_commitments = commitments.count()

        for idx, com in enumerate(sorted_commitments, start=1):
            if task and idx % 50 == 0:
                task.update_state(
                    state='PROGRESS',
                    meta={
                        'current': idx,
                        'total': total_commitments,
                        'percent': round((idx / total_commitments) * 100, 2) if total_commitments else 0.0
                    }
                )
            row_data = [
                com.token,
                com.created_at.strftime('%d/%m/%Y') if com.created_at else "",
                str(com.contract.holder) if com.contract and com.contract.holder else "",
                com.contract.holder.token if com.contract and com.contract.holder else "",
                round(com.total or 0, 2),
                round(com.remaining or 0, 2),
                com.status.name if com.status else ""
            ]
            row = add_row(sheet, row, row_data)
    else:
        raise Exception("Report not implemented")
    
    adjust_column_widths(sheet)

    filename = f"accounting_values_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb.save(filename)

    with open(filename, "rb") as f:
        response = HttpResponse(
            f.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f"attachment; filename={filename}"

    document_id = save_report(
        content=wb,
        filename=filename,
        name=name,
        type_id=type_id,
        start_date=start_date,
        end_date=end_date,
    )
    os.remove(filename)
    return response, document_id, None

def generate_aqua_remittance_report(request, black_fill=None, white_bold_font=None, task=None):
    
    date_range = request.data.get('date_range', None)
    name = request.data.get('name', '')
    type_id = request.data.get('type_id', None)
    remittance_id = request.data.get('remittance_id', None)
    
    if not remittance_id:
        return Response({"error": "remittance id is required."}, status=status.HTTP_400_BAD_REQUEST), None, None
    
    start_date = None
    end_date = None
    
    if remittance_id:
        remittance = PaymentRemittance.objects.get(id=remittance_id)
        payments = remittance.payments.all()
    else:
        raise Exception("Error obtaining invoices for report billing summary")
    
    print("payments: ", payments.count())
    
    payments_total = sum(payment.amount for payment in payments)
    bank = remittance.company_bank
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "REMESA AQUA"
    
    row = 0
    titles = [ 
              "Nº remesa", "Fecha registro", "Nº Banco", "CIF Cliente", 
              "Nº Factura", "Fecha vencimiento", "Importe factura", "Importe total remesa", 
              ]

    row = add_row(sheet, row, titles, fill=PatternFill(start_color="b8cce4", end_color="b8cce4", fill_type="solid"))
    for item in payments:
        parent_token = item.invoice.serie_final if item.invoice else item.commitment_deposit.token if item.commitment_deposit else item.token
        used_date = remittance.sent_at.strftime('%d/%m/%Y') if remittance.sent_at else item.payment_date.strftime('%d/%m/%Y') if item.payment_date else item.due_date.strftime('%d/%m/%Y')
        line_row = [
            remittance.token, used_date, bank.bank.token if bank and bank.bank else "", item.customer_token_final, 
            parent_token, item.due_date.strftime('%d/%m/%Y'), item.amount, payments_total
        ]
        row = add_row(sheet, row, line_row)
        
    adjust_column_widths(sheet)
        
    filename = f"remesa_aqua_{remittance.token}_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    
    wb.save(filename)

    # Send the file as a response
    with open(filename, "rb") as f:
        response = HttpResponse(f.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f"attachment; filename={filename}"
    
    document_id = save_report(content=wb, filename=filename, name=name, type_id=type_id, start_date=start_date, end_date=end_date)
    print("report saved")
    os.remove(filename)
    return response, document_id, None

def get_translated_alert_notes(alert_text):
    if not alert_text:
        return ""
    
    alert_translations = {
        'reading_alert_turn': {
            'ca': 'Volta de comptador',
            'es': 'Vuelta de contador',
            'en': 'Meter turn'
        },
        'reading_alert_zero': {
            'ca': 'Consum zero',
            'es': 'Consumo cero',
            'en': 'Zero consumption'
        },
        'reading_alert_negative': {
            'ca': 'Consum negatiu',
            'es': 'Consumo negativo',
            'en': 'Negative consumption'
        },
        'reading_alert_estimated': {
            'ca': 'Consum estimat',
            'es': 'Consumo estimado',
            'en': 'Estimated consumption'
        },
        'reading_alert_unusual': {
            'ca': 'Lectura inusual',
            'es': 'Lectura inusual',
            'en': 'Unusual reading'
        },
        'reading_alert_unusual_consumption': {
            'ca': 'Lectura inusual',
            'es': 'Lectura inusual',
            'en': 'Unusual reading'
        },
        'reading_alert_low_consumption': {
            'ca': 'Consum baix',
            'es': 'Consumo bajo',
            'en': 'Low consumption'
        }
    }
    
    from django.utils.translation import get_language
    lang = get_language() or 'ca'
    if '-' in lang:
        lang = lang.split('-')[0]
        
    stripped_text = alert_text.strip()
    if stripped_text in alert_translations:
        return alert_translations[stripped_text].get(lang, alert_translations[stripped_text].get('ca', alert_text))
        
    return alert_text


def _address_display_relations(prefix):
    """Relacions que llegeix str(Address) / get_address_complete_without_city."""
    return [
        f'{prefix}__street__type',
        f'{prefix}__street_number__number_type',
        f'{prefix}__city',
        f'{prefix}__country',
    ]


def generate_register_billing_summary(request, black_fill, white_bold_font, task=None):
    
    #PADRO DE FACTURACIÓ
    
    print("request: ", request.data)
    id = request.data.get('id', None)
    date_range = request.data.get('date_range', None)
    billing_ids = get_billing_ids(request.data)
    serie_final_from = request.data.get('serie_final_from', None)
    serie_final_to = request.data.get('serie_final_to', None)
    prefix = request.data.get('prefix', None)
    name = request.data.get('name', '')
    ids = request.data.get('product_ids', '')
    type_id = request.data.get('type_id', None)
    exploitation_id = request.data.get('exploitation_id', None)
    generate_object = request.data.get('generate_object', True)

    print("id: ", id)
    print("billing_ids: ", billing_ids)
    print("date_range: ", date_range)
    print("serie_final_from: ", serie_final_from)
    print("serie_final_to: ", serie_final_to)
    print("prefix: ", prefix)
    print("name: ", name)
    print("ids: ", ids)
    print("type_id: ", type_id)
    print("exploitation_id: ", exploitation_id)
    print("generate_object: ", generate_object)
    if not generate_object:
        request.data['include_preinvoices'] = True

    has_serie_range = has_serie_final_range(request.data)
    if not billing_ids and not date_range and not has_serie_range:
        return Response({"error": "billing id, date range or serie_final range is required."}, status=status.HTTP_400_BAD_REQUEST), None, None

    start_date = None
    end_date = None

    product_ids = None
    if ids:
        product_ids = [int(i) for i in ids.split(',')]

    exploitation = None
    if exploitation_id:
        exploitation = Exploitation.objects.get(id=exploitation_id)

    if billing_ids:
        invoices = Invoice.objects.filter(billing_id__in=billing_ids)
        invoices = filter_pending_invoices(invoices, request)
    elif date_range:
        start_date = datetime.datetime.strptime(date_range[0], '%Y-%m-%dT%H:%M:%S.%fZ')
        end_date = datetime.datetime.strptime(date_range[1], '%Y-%m-%dT%H:%M:%S.%fZ')
        invoice_type = ConfigProject.objects.get(token="invoice_type_invoice_token").value
        invoices = Invoice.objects.filter(
            issue_date__range=(start_date.date(), end_date.date()),
            type_final=invoice_type,
        ).distinct()
        invoices = filter_pending_invoices(invoices, request)
    elif has_serie_range:
        invoice_type = ConfigProject.objects.get(token="invoice_type_invoice_token").value
        invoices = filter_invoices_by_serie_final_range(
            Invoice.objects.filter(type_final=invoice_type),
            serie_final_from=serie_final_from,
            serie_final_to=serie_final_to,
            prefix=prefix,
        ).distinct()
        # Filtre d'usuari del flux de rangs (general-billing-summary.vue, que sempre
        # l'envia): treu les factures ja marcades com a revisades perquè no es tornin
        # a incloure en una nova generació del mateix rang. Nomes s'aplica si es
        # demana: la pantalla d'informes (billing/reports/add) no porta aquesta casella
        # i no ha de canviar en silenci les factures que entren a l'informe.
        if request.data.get('exclude_reviewed', False):
            invoices = invoices.filter(reviewed=False)
        invoices = filter_pending_invoices(invoices, request)
        # El periode surt de les factures del rang: sense date_range no n'hi ha cap
        # altre per guardar amb l'informe.
        start_date, end_date = serie_final_range_period(invoices)
    else:
        raise Exception("Error obtaining invoices for report billing summary")
    
    if exploitation:
        invoices = invoices.filter(exploitation=exploitation)

    multi_q = invoice_multi_filter_q(**get_multi_ids(request.data))
    if multi_q:
        invoices = invoices.filter(multi_q).distinct()

    line_items_qs = (
        InvoiceLineItem.objects.filter(is_active=True)
        .select_related('company')
        .prefetch_related('adjustments')
    )

    readings_qs = Reading.objects.select_related(
        'meter__caliber',
        'supply_point__source',
        'supply_point__type',
        'supply_point__supply_type',
        'previous_reading',
        'alert',
        'remote_alert',
    )

    invoices = invoices.select_related(
        'exploitation', 'billing', 'billing__biller',
        'contract', 'contract__holder', 'contract__address_billing__address',
        'contract__supply_point_default__address',
        'contract__supply_point_default__connection__dma',
        'contract__supply_point_default__property__route_position__route__route_zone',
        'contract_termination', 'contract_termination__contract',
        'contract_termination__contract__supply_point_default__address',
        'contract_termination__contract__supply_point_default__connection__dma',
        'contract_termination__contract__supply_point_default__property__route_position__route__route_zone',
        'contract_termination__contract__address_contact__address',
        'contract_termination__contract__use_type',
        'contract_request',
        'contract__address_contact__address', 'serie',
        'contract__use_type',
        # str(Address) recorre carrer, tipus de via, número, població i país: sense
        # això cada factura llançava unes dotze consultes (N+1).
        *_address_display_relations('contract__address_contact__address'),
        *_address_display_relations('contract__supply_point_default__address'),
        *_address_display_relations('contract_termination__contract__address_contact__address'),
        *_address_display_relations('contract_termination__contract__supply_point_default__address'),
    ).prefetch_related(
        Prefetch('line_items', queryset=line_items_qs, to_attr='prefetched_active_line_items'),
        Prefetch('readings', queryset=readings_qs),
        'contract__variables',
        'contract_termination__contract__variables',
    ).order_by('exploitation')

    # Evaluate once — avoid repeated COUNT / re-iteration of the same queryset
    invoice_list = list(invoices)
    total_invoices = len(invoice_list)
    print("invoices: ", total_invoices)

    # Bulk-resolve contracts linked only via contract_request (avoids N+1)
    request_ids = [
        inv.contract_request_id
        for inv in invoice_list
        if not inv.contract_id
        and not (inv.contract_termination_id and inv.contract_termination and inv.contract_termination.contract_id)
        and inv.contract_request_id
    ]
    contracts_by_request = {}
    if request_ids:
        for contract in (
            Contract.objects.filter(contract_request_id__in=request_ids)
            .select_related(
                'use_type',
                'address_contact__address',
                'supply_point_default__address',
                'supply_point_default__connection__dma',
                'supply_point_default__property__route_position__route__route_zone',
                *_address_display_relations('address_contact__address'),
                *_address_display_relations('supply_point_default__address'),
            )
            .prefetch_related('variables')
        ):
            contracts_by_request[contract.contract_request_id] = contract

    def _normalize_tax_key(value):
        try:
            return float(value)
        except (TypeError, ValueError):
            return 0.0

    def _get_line_items(invoice, for_products=False):
        if hasattr(invoice, 'prefetched_active_line_items'):
            items = invoice.prefetched_active_line_items
        else:
            items = [li for li in invoice.line_items.all() if li.is_active]
        if for_products and product_ids:
            items = [li for li in items if li.product_id in product_ids]
        return items

    def _totals_from_line_items(line_items):
        """Same logic as get_totals(), but without hitting the DB."""
        subtotal = 0
        taxes_map = {}
        taxes_base = {}
        for line_item in line_items:
            subtotal += round_ceil(line_item.price)
            if line_item.tax_percent is None:
                continue
            if int(line_item.tax_percent) == 0:
                continue
            percent = str(line_item.tax_percent)
            taxes_map[percent] = taxes_map.get(percent, 0) + round_ceil(line_item.tax_price)
            taxes_base[percent] = taxes_base.get(percent, 0) + round_ceil(line_item.price)
        total_taxes = sum(round_ceil(t) for t in taxes_map.values())
        return round_ceil(subtotal), taxes_map, taxes_base, round_ceil(subtotal) + round_ceil(total_taxes)

    def _resolve_contract(invoice):
        contract = invoice.contract
        if contract:
            return contract
        if invoice.contract_termination and invoice.contract_termination.contract:
            return invoice.contract_termination.contract
        if invoice.contract_request_id:
            return contracts_by_request.get(invoice.contract_request_id)
        return None

    # --- Dimensionament estable de les columnes -----------------------------
    # Abans, l'amplada del full es derivava de les factures del rang exportat
    # (màxim de lectures d'una factura, productes que hi apareixien, línies per
    # producte i correctors vistos), de manera que dues exportacions del mateix
    # informe amb rangs diferents tenien columnes diferents i desplaçades.
    # Ara es deriva del catàleg, del filtre de productes i de l'històric de la
    # BD; les dades del rang només poden AMPLIAR la reserva, mai reduir-la, per
    # no perdre cap import si una factura porta més línies de les reservades.
    #
    # Tot el que no depèn de l'explotació es calcula un sol cop aquí, i el que
    # sí que en depèn arriba agrupat per explotació en una única consulta,
    # perquè el nombre de consultes no creixi amb el nombre de pestanyes.
    def _config_int(token, default):
        try:
            value = int(ConfigProject.objects.get(token=token).value)
        except (ConfigProject.DoesNotExist, TypeError, ValueError):
            return default
        return value if value >= 0 else default

    layout_readings = _config_int('report_register_billing_readings', 2)
    adjustments_cap = _config_int('report_register_billing_adjustments', 4)
    other_product_blocks = _config_int('report_register_billing_other_blocks', 3)
    variable_values = _config_int('report_register_billing_variable_values', 3)
    # Pes mínim (en tant per mil de les línies de l'explotació) per donar bloc
    # propi a un producte d'una altra explotació. A 0 no es filtra res.
    product_min_permille = _config_int('report_register_billing_product_min_permille', 1)

    line_item_table = InvoiceLineItem._meta.db_table
    invoice_table = Invoice._meta.db_table
    adjustment_link_table = InvoiceLineItem._meta.get_field('adjustments').remote_field.through._meta.db_table
    reading_link_table = Invoice._meta.get_field('readings').remote_field.through._meta.db_table

    # Màxims històrics per explotació i producte (línies d'un mateix producte en
    # una factura, correctors d'una línia, lectures d'una factura). Van en SQL
    # perquè són agregacions de dos nivells —el màxim d'un recompte per
    # factura— i l'ORM no les pot expressar: `Max('total')` sobre un `Count`
    # dona FieldError. Són tres consultes per a totes les pestanyes.
    #
    # Només compten les factures emeses des de `ov_pdf_from` (o 2026-01-01 si
    # no està configurat): les importades d'anys enrere porten fins a 120
    # línies de cànon en una sola factura i eixamplaven cada exportació amb
    # milers de columnes buides. Les factures del rang encara poden ampliar la
    # reserva, de manera que cap import es perd.
    history_from_config = ConfigProject.objects.filter(token='ov_pdf_from').first()
    history_from = (
        parse_date(history_from_config.value.strip())
        if history_from_config and history_from_config.value else None
    ) or datetime.date(2026, 1, 1)
    history_blocks = {}
    history_lines = {}
    history_adjustments = {}
    with connection.cursor() as cursor:
        cursor.execute(f"""
            SELECT exploitation_id, product_id, max(total), sum(total) FROM (
                SELECT inv.exploitation_id, item.invoice_id, item.product_id, count(*) AS total
                FROM {line_item_table} item
                JOIN {invoice_table} inv ON inv.id = item.invoice_id
                WHERE item.is_active AND item.product_id IS NOT NULL
                  AND inv.issue_date >= %s
                GROUP BY 1, 2, 3
            ) per_invoice GROUP BY 1, 2
        """, [history_from])
        for exploitation_key, product_key, peak, total in cursor.fetchall():
            history_blocks.setdefault(exploitation_key, {})[product_key] = peak
            history_lines.setdefault(exploitation_key, {})[product_key] = total
        cursor.execute(f"""
            SELECT inv.exploitation_id, item.product_id, max(per_item.total) FROM (
                SELECT invoicelineitem_id, count(*) AS total
                FROM {adjustment_link_table} GROUP BY invoicelineitem_id
            ) per_item
            JOIN {line_item_table} item ON item.id = per_item.invoicelineitem_id
            JOIN {invoice_table} inv ON inv.id = item.invoice_id
            WHERE item.is_active AND item.product_id IS NOT NULL
              AND inv.issue_date >= %s
            GROUP BY 1, 2
        """, [history_from])
        for exploitation_key, product_key, total in cursor.fetchall():
            history_adjustments.setdefault(exploitation_key, {})[product_key] = total
        cursor.execute(f"""
            SELECT exploitation_id, max(total) FROM (
                SELECT inv.exploitation_id, link.invoice_id, count(*) AS total
                FROM {reading_link_table} link
                JOIN {invoice_table} inv ON inv.id = link.invoice_id
                WHERE inv.issue_date >= %s
                GROUP BY 1, 2
            ) per_invoice GROUP BY 1
        """, [history_from])
        history_readings = dict(cursor.fetchall())

    # Blocs "Concepte" per producte = trams (BillingRange) actius del producte;
    # correctors = ajustos configurats als seus LineItemType, limitats per
    # `adjustments_cap`. Cap de les dues coses depèn de l'explotació.
    catalog_blocks = {}
    for row_data in (
        BillingRange.objects.filter(price_rate__is_active=True)
        .values('price_rate__product_id').annotate(total=Count('id', distinct=True))
    ):
        catalog_blocks[row_data['price_rate__product_id']] = row_data['total']

    catalog_adjustments = {}
    for row_data in (
        Adjustment.objects
        .filter(is_active=True, line_item_type__billing_range__price_rate__is_active=True)
        .values('line_item_type__billing_range__price_rate__product_id', 'line_item_type_id')
        .annotate(total=Count('id'))
    ):
        product_key = row_data['line_item_type__billing_range__price_rate__product_id']
        if row_data['total'] > catalog_adjustments.get(product_key, 0):
            catalog_adjustments[product_key] = row_data['total']

    # Variables: el catàleg de noms actius, no els que surtin al rang.
    catalog_variable_names = sorted(
        name for name in (
            Variable.objects.filter(is_active=True)
            .exclude(name__isnull=True).exclude(name='')
            .values_list('name', flat=True).distinct()
        ) if name
    )

    # Productes que poden entrar a qualsevol pestanya, carregats un sol cop.
    candidate_product_ids = set(product_ids or [])
    if not product_ids:
        for products_of_exploitation in history_blocks.values():
            candidate_product_ids.update(products_of_exploitation)
    products_by_id = {
        product.id: product
        for product in Product.objects.filter(id__in=candidate_product_ids).select_related('exploitation')
    }

    taxes = list(Tax.objects.exclude(percent=0).order_by("percent").distinct())

    # Mode només escriptura i sense estils: el full detallat té milers de
    # columnes i milions de cel·les, i així les files s'escriuen en streaming
    # en lloc de construir un objecte Cell per cadascuna. Les amplades s'han de
    # fixar abans de la primera fila i cada full només es pot desar un cop.
    wb = openpyxl.Workbook(write_only=True)
    company_vat = ConfigProject.objects.get(token="main_company_token").value
    company_main = Company.objects.get(vat=company_vat)

    biller_types = {
        0: "Puntual",
        360: "Anual",
        180: "Semestral",
        90: "Trimestral",
        60: "Quadrimestral",
        30: "Mensual",
    }

    # Els noms de les columnes repetides es qualifiquen amb el bloc a què
    # pertanyen (lectura, producte + concepte, corrector, variable), perquè es
    # puguin identificar pel nom i no per la posició: abans hi havia desenes de
    # columnes que es deien igual ("Concepte", "Total", "Corrector", "Valor").
    def _qualified(prefix, label):
        return f"{prefix} · {label}" if prefix else label

    def _sheet_layout(exploitation_keys, sheet_exploitation, sheet_invoices):
        """Disposició de columnes d'una pestanya.

        Els productes són els demanats al filtre o, si no n'hi ha, els que
        corresponen a aquella explotació. Per decidir-ho no es pot fer servir el
        camp `exploitation` del producte tot sol: hi ha productes assignats a un
        municipi que es facturen a d'altres (p. ex. el cànon de l'ACA, encara
        vigent). Ni tampoc n'hi ha prou amb "s'ha facturat aquí alguna vegada",
        perquè hi ha usos residuals per error de tarifa (productes d'altres
        municipis amb poques línies que reservaven centenars de columnes buides,
        els seus trams de cànon). El criteri és, doncs: el producte és de
        l'explotació o compartit, o bé s'hi factura amb un pes significatiu.
        """
        exploitation_blocks = {}
        exploitation_adjustments = {}
        exploitation_lines = {}
        for exploitation_key in exploitation_keys:
            for product_key, value in history_blocks.get(exploitation_key, {}).items():
                if value > exploitation_blocks.get(product_key, 0):
                    exploitation_blocks[product_key] = value
            for product_key, value in history_adjustments.get(exploitation_key, {}).items():
                if value > exploitation_adjustments.get(product_key, 0):
                    exploitation_adjustments[product_key] = value
            for product_key, value in history_lines.get(exploitation_key, {}).items():
                exploitation_lines[product_key] = exploitation_lines.get(product_key, 0) + value
        total_lines = sum(exploitation_lines.values())

        def _belongs_to_sheet(product):
            if product.exploitation_id is None or product.exploitation_id in exploitation_keys:
                return True
            if not product_min_permille:
                return True
            return exploitation_lines.get(product.id, 0) * 1000 >= total_lines * product_min_permille

        if product_ids:
            sheet_product_ids = [pid for pid in product_ids if pid in products_by_id]
        else:
            sheet_product_ids = [
                pid for pid in exploitation_blocks
                if pid in products_by_id and _belongs_to_sheet(products_by_id[pid])
            ]
        sheet_products = sorted(
            (products_by_id[pid] for pid in sheet_product_ids),
            key=lambda item: (item.position if item.position is not None else 0, item.name or '', item.id),
        )
        sheet_product_id_set = {product.id for product in sheet_products}

        # Un mateix nom de producte pot sortir més d'una vegada a la pestanya
        # (hi pot haver un "AIGUA" per explotació i les factures d'un municipi en
        # poden portar de diversos): dues columnes amb el mateix nom tornarien a
        # obligar a mapejar per posició, així que s'hi afegeix l'explotació.
        label_counts = {}
        for product in sheet_products:
            label_counts[product.name or ''] = label_counts.get(product.name or '', 0) + 1
        labels = {}
        for product in sheet_products:
            label = product.name or ''
            if label_counts.get(label, 0) > 1:
                qualifier = (
                    product.exploitation.name if product.exploitation and product.exploitation.name
                    else product.token or str(product.id)
                )
                label = f"{label} ({qualifier})"
            labels[product.id] = label

        # Passada per les factures de la pestanya: només per ampliar la reserva.
        observed_readings = 0
        observed_blocks = {}
        observed_adjustments = {}
        observed_variables = {}
        for invoice in sheet_invoices:
            invoice_readings = invoice.readings.all()
            reading_count = len(invoice_readings) if hasattr(invoice_readings, '__len__') else invoice_readings.count()
            if reading_count > observed_readings:
                observed_readings = reading_count
            counts_in_invoice = {}
            for line_item in _get_line_items(invoice, for_products=True):
                product_key = line_item.product_id if line_item.product_id in sheet_product_id_set else None
                counts_in_invoice[product_key] = counts_in_invoice.get(product_key, 0) + 1
                if counts_in_invoice[product_key] > observed_blocks.get(product_key, 0):
                    observed_blocks[product_key] = counts_in_invoice[product_key]
                # Use prefetched cache — never .count() (that hits DB)
                adjustment_count = len(line_item.adjustments.all())
                if adjustment_count > observed_adjustments.get(product_key, 0):
                    observed_adjustments[product_key] = adjustment_count
            contract_for_variables = _resolve_contract(invoice)
            if contract_for_variables:
                variables_in_invoice = {}
                for variable in contract_for_variables.variables.all():
                    if not variable.is_active:
                        continue
                    variable_name = variable.name or ''
                    if not variable_name:
                        continue
                    variables_in_invoice[variable_name] = variables_in_invoice.get(variable_name, 0) + 1
                    if variables_in_invoice[variable_name] > observed_variables.get(variable_name, 0):
                        observed_variables[variable_name] = variables_in_invoice[variable_name]

        expanded = []
        reserved_readings = max(
            [layout_readings] + [history_readings.get(key, 0) for key in exploitation_keys]
        )
        if observed_readings > reserved_readings:
            expanded.append(f"lectures: {observed_readings}")

        layout = []
        for product in sheet_products:
            reserved_blocks = max(
                catalog_blocks.get(product.id, 1), 1, exploitation_blocks.get(product.id, 0)
            )
            reserved_adjustments = max(
                min(catalog_adjustments.get(product.id, 0), adjustments_cap),
                exploitation_adjustments.get(product.id, 0),
            )
            if observed_blocks.get(product.id, 0) > reserved_blocks:
                expanded.append(f"conceptes de {labels[product.id]}: {observed_blocks[product.id]}")
            if observed_adjustments.get(product.id, 0) > reserved_adjustments:
                expanded.append(f"correctors de {labels[product.id]}: {observed_adjustments[product.id]}")
            layout.append((
                product.id,
                labels[product.id],
                max(reserved_blocks, observed_blocks.get(product.id, 0)),
                max(reserved_adjustments, observed_adjustments.get(product.id, 0)),
            ))
        if not product_ids:
            # Bloc "Altres" per a les línies sense producte del catàleg (línies
            # manuals amb nom lliure). Amb filtre de productes ja queden fora.
            layout.append((
                None,
                _("Altres"),
                max(other_product_blocks, observed_blocks.get(None, 0)),
                observed_adjustments.get(None, 0),
            ))

        sheet_variables = [
            (name, max(variable_values, observed_variables.get(name, 0)))
            for name in catalog_variable_names
        ]
        for name in sorted(observed_variables):
            if name not in catalog_variable_names:
                sheet_variables.append((name, observed_variables[name]))
        for name, reserved in sheet_variables:
            if observed_variables.get(name, 0) > variable_values:
                expanded.append(f"valors de {name}: {observed_variables[name]}")

        if expanded:
            print("[padro] disposició ampliada per les dades del rang: " + "; ".join(expanded))
        return max(reserved_readings, observed_readings), layout, sheet_variables

    def _render_sheet(sheet, exploitation_keys, sheet_exploitation, sheet_invoices, progress_offset, layout=None):
        max_invoice_readings, layout_products, layout_variables = layout or _sheet_layout(
            exploitation_keys, sheet_exploitation, sheet_invoices
        )
        # Productes amb bloc propi en aquesta pestanya: la resta de línies (les
        # manuals, sense producte del catàleg) van al bloc "Altres" (clau None).
        sheet_product_id_set = {key for key, _label, _blocks, _adjustments in layout_products if key is not None}

        titles = [
                    _("Abonat"), _("NIF/CIF"), _("Dom. del subministre"),
                    _("Adreça fiscal"), _("Sector"), _("Població"),
                    _("Identificador"), _("Any factura"), _("Data emissió"),
                    _("Tipus pagament"), _("Núm. factura"), _("Tipus factura"),
                    _("Periodicitat"),
                    _("Tipus d'ús contracte"),
                    _("Ús ACA"),
                    ]
        for i in range(max_invoice_readings):
            reading_prefix = "%s %s" % (_("Lectura"), i + 1)
            titles.extend([
                    _qualified(reading_prefix, _("Lect. anterior")),
                    _qualified(reading_prefix, _("Data lec. anterior")),
                    _qualified(reading_prefix, _("Lect. actual")),
                    _qualified(reading_prefix, _("Data lec. actual")),
                    _qualified(reading_prefix, _("Dies periode")),
                    _qualified(reading_prefix, _("Estimada")),
                    _qualified(reading_prefix, _("Observ. actual")),
                    ])
        titles.extend([
                    _("Consum (m3)")
                    ])
        for tax in taxes:
            titles.append(_("Import IVA %(percent)s") % {"percent": str(int(tax.percent))})
        titles.extend([
                    _("Import IVA Total"),
                    _("Total sense IVA"),
                    _("Total")
                    ])
    
        for product_key, product_label, max_count, max_adjustments in layout_products:
            titles.append(_qualified(product_label, _("Producte")))
            for idx in range(max_count):
                item_prefix = "%s · %s %s" % (product_label, _("Concepte"), idx + 1)
                titles.append(item_prefix)
                titles.append(_qualified(item_prefix, _("Quantitat")))
                titles.append(_qualified(item_prefix, _("Preu Ut.")))
                titles.append(_qualified(item_prefix, _("IVA")))
                titles.append(_qualified(item_prefix, _("Total IVA")))
                titles.append(_qualified(item_prefix, _("Subtotal")))
                titles.append(_qualified(item_prefix, _("Total")))
                titles.append(_qualified(item_prefix, _("Limit")))
                for adj_idx in range(max_adjustments):
                    adjustment_prefix = "%s · %s %s" % (item_prefix, _("Corrector"), adj_idx + 1)
                    titles.append(adjustment_prefix)
                    titles.append(_qualified(adjustment_prefix, _("Cantidad corrector")))
            titles.append(_qualified(product_label, _("Propietari")))
        for i in range(max_invoice_readings):
            meter_prefix = "%s %s" % (_("Comptador"), i + 1)
            titles.extend([
                meter_prefix,
                _qualified(meter_prefix, _("Calibre")),
                _qualified(meter_prefix, _("Ubicació")),
                _qualified(meter_prefix, _("Ús submministre")),
                _qualified(meter_prefix, _("Descripció tipus subm.")),
            ])
        titles.extend([
            _("Núm. persones"), _("Domicili alternatiu"), _("Codi postal alternatiu"),
            _("A l'atenció de"), _("Codi ruta"), _("Zona")
        ])

        for variable_name, max_count in layout_variables:
            titles.append(_qualified(variable_name, _("Variable")))
            for idx in range(max_count):
                titles.append("%s · %s %s" % (variable_name, _("Valor"), idx + 1))
    
        # Amplades des de les capçaleres (no recórrer el full), abans de cap fila.
        for col_idx, title in enumerate(titles, start=1):
            width = min(max(len(str(title)) + 2, 10), 40)
            sheet.column_dimensions[get_column_letter(col_idx)].width = width
        sheet.append(titles)

        for idx, i in enumerate(sheet_invoices, start=1):
            if task and (progress_offset + idx) % 200 == 0:
                task.update_state(
                    state='PROGRESS',
                    meta={
                        'current': progress_offset + idx,
                        'total': total_invoices,
                        'percent': round(((progress_offset + idx) / total_invoices) * 100, 2) if total_invoices else 100,
                    }
                )
            contract = _resolve_contract(i)
            readings = list(i.readings.all())
            all_line_items = _get_line_items(i, for_products=False)
            invoice_line_items = (
                [li for li in all_line_items if li.product_id in product_ids]
                if product_ids else all_line_items
            )

            _subtotal, taxes_invoice, taxes_base, _total = _totals_from_line_items(all_line_items)
            final_taxes = {}
            for tax_rate in taxes_invoice:
                normalized_key = _normalize_tax_key(tax_rate)
                final_taxes[normalized_key] = {
                    "tax_value": taxes_invoice[tax_rate],
                    "tax_base": taxes_base[tax_rate],
                }

            different_address = False
            if contract:
                if (
                    contract.address_contact and contract.address_contact.address
                    and contract.supply_point_default and contract.supply_point_default.address
                ):
                    if (
                        contract.address_contact.address.id != contract.supply_point_default.address.id
                        and str(contract.address_contact.address) != str(contract.supply_point_default.address)
                    ):
                        different_address = True
        
            i_data = [
                i.customer_final, i.customer_token_final, 
                get_address_complete_without_city(contract.supply_point_default.address) if contract and contract.supply_point_default and contract.supply_point_default.address else i.address_final ,
                i.address_final, 
                (
                    i.contract.supply_point_default.connection.dma.name
                    if i.contract and i.contract.supply_point_default
                    and i.contract.supply_point_default.connection
                    and i.contract.supply_point_default.connection.dma
                    else ""
                ),
                i.exploitation.name if i.exploitation else "",
                contract.token if contract else i.customer_token_final,
                i.issue_date.year if i.issue_date else "",
                i.issue_date.strftime('%d/%m/%Y') if i.issue_date else "",
                i.payment_type_final, i.serie_final, i.serie.name if i.serie else "Factura Original", 
                biller_types.get(i.billing_period_days) if i.billing_period_days else (
                    i.billing.biller.period_type.capitalize() if i.billing and i.billing.biller else "Puntual"
                ),
                contract.use_type.name if contract and contract.use_type else "",
                contract.get_use_aca_display() if contract and contract.use_aca else "",
                ]
            total_estimated_used = 0
            for j in range(max_invoice_readings):
                if j < len(readings):
                    reading = readings[j]
                    total_estimated_used += reading.estimated_used if reading.estimated_used else 0
                    previous_date = None
                    if reading.previous_reading:
                        previous_date = reading.previous_reading.reading_date
                    elif reading.reading_date and i.contract and i.contract.created_at:
                        previous_date = i.contract.created_at.date()

                    i_data.extend([
                        reading.previous_reading.reading_value if reading.previous_reading else 0,
                        previous_date if previous_date else "",
                        reading.reading_value if reading.reading_value is not None else 0,
                        reading.reading_date if reading.reading_date else "",
                        reading.consumption_days if reading.consumption_days is not None else 0,
                        "S" if reading.is_estimated else "N",
                        get_translated_alert_notes(
                            reading.alert_notes if reading.alert_notes
                            else reading.remote_alert.name if reading.remote_alert
                            else reading.alert.name if reading.alert else ""
                        )
                    ])
                else:
                    i_data.extend([
                        0, "", 0, "", 0, "", ""
                    ])
        
            i_data.extend([
                            float(i.consumption) - float(total_estimated_used) if i.consumption else 0
                        ])
        
            for tax in taxes:
                tax_key = _normalize_tax_key(tax.percent)
                tax_value = final_taxes.get(tax_key, {}).get("tax_value", 0)
                i_data.append(0 if tax_value is None else tax_value)
        
            i_data.extend([
                            i.total_final - i.subtotal_final,
                            i.subtotal_final,
                            i.total_final
                        ])
        
            invoice_line_items = sorted(invoice_line_items, key=lambda x: x.id)
        
            # S'agrupa per producte (product_id) i no pel nom desat a la línia: el
            # nom és una còpia de text i les línies manuals en porten un de lliure,
            # cosa que abans creava un bloc de columnes nou per cada nom vist.
            line_items_by_product = {}
            for line_item in invoice_line_items:
                product_key = line_item.product_id if line_item.product_id in sheet_product_id_set else None
                line_items_by_product.setdefault(product_key, []).append(line_item)

            for product_key, product_label, max_count, max_adjustments in layout_products:
                line_items_for_product = line_items_by_product.get(product_key, [])

                i_data.append(product_label)
            
                company_name = company_main.name 
            
                for item_idx in range(max_count):
                    if item_idx < len(line_items_for_product):
                        line_item = line_items_for_product[item_idx]
                        if item_idx == 0:
                            company_name = line_item.company.name if line_item.company else company_main.name
                        price_value = round(line_item.price or 0, 4)
                        tax_price_value = round(line_item.tax_price or 0, 4)
                    
                        line_name = f"{line_item.price_rate_name} {line_item.name.split('-')[0]}"
                        if line_item.interval and line_item.interval > 0:
                            line_name += f" TRAM {line_item.interval}"
                    
                        i_data.append(line_name)
                        i_data.append(round(line_item.units, 4))
                        i_data.append(round(line_item.price_unit, 4))
                        i_data.append(line_item.tax_percent)
                        i_data.append(tax_price_value)
                        i_data.append(price_value)
                        i_data.append(price_value + tax_price_value)
                        i_data.append(line_item.end_stretch if line_item.end_stretch else "")
                    
                        line_item_adjustments = list(line_item.adjustments.all())
                        for adj_idx in range(max_adjustments):
                            if adj_idx < len(line_item_adjustments):
                                adjustment = line_item_adjustments[adj_idx]
                                adjustment_name = adjustment.name or adjustment.token or ""
                                adjustment_amount = round(float(adjustment.total_applied or 0), 2)
                                i_data.append(adjustment_name)
                                i_data.append(adjustment_amount)
                            else:
                                i_data.append("")  
                                i_data.append(0)   
                    else:
                        i_data.append("")  
                        i_data.append(0)   
                        i_data.append(0)   
                        i_data.append(0)   
                        i_data.append(0)   
                        i_data.append(0)   
                        i_data.append(0)   
                        i_data.append("")   
                    
                        for adj_idx in range(max_adjustments):
                            i_data.append("")  
                            i_data.append(0)   
            
                i_data.append(company_name)
        
            for j in range(max_invoice_readings):
                if j < len(readings):
                    reading = readings[j]
                    i_data.extend([
                        reading.meter.code if reading.meter else "",
                        reading.meter.caliber.token if reading.meter and reading.meter.caliber else "",
                        reading.supply_point.source.name if reading.supply_point and reading.supply_point.source else "",
                        reading.supply_point.type.name if reading.supply_point and reading.supply_point.type else "",
                        reading.supply_point.supply_type.name if reading.supply_point and reading.supply_point.supply_type else "",
                    ])
                else:
                    i_data.extend([
                        "", "", "", "", "",
                    ])
        
            i_data.extend([
                i.persons_final if not i.customer_is_juridic else 0, 
                f"{str(contract.address_contact.address.street)}, {str(contract.address_contact.address.street_number)}" if contract and contract.address_contact and contract.address_contact.address and different_address else "",
                contract.address_contact.address.postal_code if contract and contract.address_contact and contract.address_contact.address and different_address else "",
                contract.address_contact.attention_to if contract and contract.address_contact and different_address else "",
                contract.supply_point_default.property.route_position.route.token if contract and contract.supply_point_default and contract.supply_point_default.property and contract.supply_point_default.property.route_position and contract.supply_point_default.property.route_position.route else "",
                contract.supply_point_default.property.route_position.route.route_zone.name if contract and contract.supply_point_default and contract.supply_point_default.property and contract.supply_point_default.property.route_position and contract.supply_point_default.property.route_position.route and contract.supply_point_default.property.route_position.route.route_zone else "",
            ])
        
            contract_variables_by_name = {}
            if contract:
                for variable in contract.variables.all():
                    if not variable.is_active:
                        continue
                    variable_name = variable.name or ''
                    if not variable_name:
                        continue
                    contract_variables_by_name.setdefault(variable_name, []).append(variable)
        
            for variable_name, max_count in layout_variables:
                variables_for_name = contract_variables_by_name.get(variable_name, [])
                i_data.append(variable_name if variables_for_name else "")
                for var_idx in range(max_count):
                    if var_idx < len(variables_for_name):
                        variable_value = variables_for_name[var_idx].value
                        i_data.append(variable_value if variable_value is not None else "")
                    else:
                        i_data.append("")
        
            sheet.append(i_data)

    def _render_summary_sheet(sheet, layout_products, sheet_invoices):
        """Full resum: una fila per factura i tres columnes per producte.

        El full detallat obre un bloc de columnes per cada concepte, tram i
        corrector de cada producte, i acaba tenint centenars o milers de
        columnes. Aquí només s'hi agrega l'import per producte, que és el que
        es fa servir per quadrar la facturació.
        """
        sheet_product_id_set = {key for key, _label, _blocks, _adjustments in layout_products if key is not None}

        titles = [
            _("Abonat"), _("NIF/CIF"), _("Dom. del subministre"),
            _("Població"), _("Identificador"), _("Data emissió"),
            _("Núm. factura"), _("Tipus factura"), _("Consum (m3)"),
            _("Total sense IVA"), _("Import IVA Total"), _("Total"),
        ]
        for _product_key, product_label, _max_count, _max_adjustments in layout_products:
            titles.append(_qualified(product_label, _("Subtotal")))
            titles.append(_qualified(product_label, _("IVA")))
            titles.append(_qualified(product_label, _("Total")))

        for col_idx, title in enumerate(titles, start=1):
            width = min(max(len(str(title)) + 2, 10), 40)
            sheet.column_dimensions[get_column_letter(col_idx)].width = width
        sheet.append(titles)

        # Fila final de totals: mateixes columnes numèriques que les files de dades.
        totals_by_column = {}

        for i in sheet_invoices:
            contract = _resolve_contract(i)
            all_line_items = _get_line_items(i, for_products=False)
            invoice_line_items = (
                [li for li in all_line_items if li.product_id in product_ids]
                if product_ids else all_line_items
            )

            total_estimated_used = sum(
                reading.estimated_used or 0 for reading in i.readings.all()
            )

            i_data = [
                i.customer_final, i.customer_token_final,
                get_address_complete_without_city(contract.supply_point_default.address)
                if contract and contract.supply_point_default and contract.supply_point_default.address
                else i.address_final,
                i.exploitation.name if i.exploitation else "",
                contract.token if contract else i.customer_token_final,
                i.issue_date.strftime('%d/%m/%Y') if i.issue_date else "",
                i.serie_final, i.serie.name if i.serie else "Factura Original",
                float(i.consumption) - float(total_estimated_used) if i.consumption else 0,
                i.subtotal_final,
                i.total_final - i.subtotal_final,
                i.total_final,
            ]

            amounts_by_product = {}
            for line_item in invoice_line_items:
                product_key = line_item.product_id if line_item.product_id in sheet_product_id_set else None
                subtotal, tax_value = amounts_by_product.get(product_key, (0, 0))
                amounts_by_product[product_key] = (
                    subtotal + round_ceil(line_item.price),
                    tax_value + round_ceil(line_item.tax_price),
                )

            for product_key, _product_label, _max_count, _max_adjustments in layout_products:
                subtotal, tax_value = amounts_by_product.get(product_key, (0, 0))
                subtotal = round_ceil(subtotal)
                tax_value = round_ceil(tax_value)
                i_data.extend([subtotal, tax_value, round_ceil(subtotal + tax_value)])

            # Els imports de la factura són Decimal i els agregats per producte
            # float: es normalitzen per poder-los sumar a la fila de totals.
            for column_idx in range(8, len(i_data)):
                try:
                    value = float(i_data[column_idx])
                except (TypeError, ValueError):
                    continue
                totals_by_column[column_idx] = totals_by_column.get(column_idx, 0) + value

            sheet.append(i_data)

        if sheet_invoices:
            totals_row = [""] * len(titles)
            totals_row[0] = _("TOTAL")
            for column_idx, value in totals_by_column.items():
                totals_row[column_idx] = round_ceil(value)
            sheet.append(totals_row)

    # Una pestanya per explotació quan el client en té més d'una i no se n'ha
    # demanat cap: cada explotació té els seus productes i, per tant, la seva
    # disposició. Ficar-ho tot en un full obligava a arrossegar les columnes de
    # totes (amb diverses explotacions, milers de columnes) i deixava la
    # major part del full buida a cada fila.
    invoices_by_exploitation = {}
    for invoice in invoice_list:
        invoices_by_exploitation.setdefault(invoice.exploitation_id, []).append(invoice)

    # Cada entrada del pla és (claus d'explotació que cobreix el full, explotació
    # del full, factures). Les claus són les que fan servir els màxims històrics
    # i el criteri de pertinença dels productes.
    if exploitation:
        sheet_plan = [({exploitation.id}, exploitation, invoice_list)]
    else:
        client_exploitations = list(Exploitation.objects.all().order_by('name', 'id'))
        if len(client_exploitations) <= 1:
            single_exploitation = client_exploitations[0] if client_exploitations else None
            # Un sol full per a tot: hi caben també les factures sense explotació.
            keys = {single_exploitation.id} if single_exploitation else set()
            keys.add(None)
            sheet_plan = [(keys, single_exploitation, invoice_list)]
        else:
            sheet_plan = [
                (
                    {client_exploitation.id},
                    client_exploitation,
                    invoices_by_exploitation.get(client_exploitation.id, []),
                )
                for client_exploitation in client_exploitations
            ]
            if Invoice.objects.filter(exploitation__isnull=True).exists():
                sheet_plan.append(({None}, None, invoices_by_exploitation.get(None, [])))

    def _sheet_title(sheet_exploitation, used_titles):
        base = (
            sheet_exploitation.name if sheet_exploitation and sheet_exploitation.name
            else _("Sense explotació")
        )
        # Excel no accepta : \ / ? * [ ] i talla els títols a 31 caràcters.
        for forbidden in ':\\/?*[]':
            base = base.replace(forbidden, ' ')
        base = ' '.join(base.split())[:31] or _('Full')
        title = base
        suffix = 2
        while title in used_titles:
            title = f"{base[:28]} {suffix}"
            suffix += 1
        used_titles.add(title)
        return title

    def _summary_sheet_title(sheet_exploitation, used_titles, sheet_count):
        if sheet_count == 1:
            base = str(_("Resum per producte"))
        else:
            base = "%s %s" % (_("Resum"), (
                sheet_exploitation.name if sheet_exploitation and sheet_exploitation.name
                else _("Sense explotació")
            ))
        for forbidden in ':\\/?*[]':
            base = base.replace(forbidden, ' ')
        base = ' '.join(base.split())[:31] or _('Resum')
        title = base
        suffix = 2
        while title in used_titles:
            title = f"{base[:28]} {suffix}"
            suffix += 1
        used_titles.add(title)
        return title

    used_sheet_titles = set()
    rendered_invoices = 0
    for position, (exploitation_keys, sheet_exploitation, sheet_invoices) in enumerate(sheet_plan):
        if len(sheet_plan) == 1:
            target_sheet = wb.create_sheet(title=_("Facturació detallada"))
        else:
            target_sheet = wb.create_sheet(title=_sheet_title(sheet_exploitation, used_sheet_titles))
        sheet_layout = _sheet_layout(exploitation_keys, sheet_exploitation, sheet_invoices)
        _render_sheet(
            target_sheet, exploitation_keys, sheet_exploitation, sheet_invoices,
            rendered_invoices, layout=sheet_layout,
        )
        # Segon full amb els imports agregats per producte, sense el detall de
        # conceptes, trams i correctors del full anterior.
        summary_sheet = wb.create_sheet(
            title=_summary_sheet_title(sheet_exploitation, used_sheet_titles, len(sheet_plan))
        )
        _render_summary_sheet(summary_sheet, sheet_layout[1], sheet_invoices)
        rendered_invoices += len(sheet_invoices)


    filename = f"billing_detailed_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

    # Es serialitza un sol cop: amb milers de columnes cada wb.save() triga
    # minuts, i abans es feia dues vegades (a disc i dins de save_report).
    buffer = BytesIO()
    wb.save(buffer)
    content = buffer.getvalue()

    response = HttpResponse(content, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f"attachment; filename={filename}"

    document_id = save_report(content=content, filename=filename, name=name, type_id=type_id, start_date=start_date, end_date=end_date)
    return response, document_id, None


def generate_mini_register_billing_summary(request, black_fill, white_bold_font, task=None):
    
    #PADRO DE FACTURACIÓ RESUMIT
    
    id = request.data.get('id', None)
    date_range = request.data.get('date_range', None)
    billing_ids = get_billing_ids(request.data)
    name = request.data.get('name', '')
    ids = request.data.get('product_ids', '')
    type_id = request.data.get('type_id', None)
    exploitation_id = request.data.get('exploitation_id', None)
    generate_object = request.data.get('generate_object', True)
    
    if not generate_object:
        request.data['include_preinvoices'] = True
    
    has_serie_range = has_serie_final_range(request.data)
    if not billing_ids and not date_range and not has_serie_range:
        return Response({"error": "billing id, date range or serie_final range is required."}, status=status.HTTP_400_BAD_REQUEST), None, None
    
    start_date = None
    end_date = None

    if ids:
        product_ids = [int(i) for i in ids.split(',')]
    else:
        product_ids = None

    exploitation = None
    if exploitation_id:
        exploitation = Exploitation.objects.get(id=exploitation_id)
        
    if billing_ids:
        invoices = Invoice.objects.filter(billing_id__in=billing_ids)
        invoices = filter_pending_invoices(invoices, request)
    elif date_range:
        start_date = datetime.datetime.strptime(date_range[0], '%Y-%m-%dT%H:%M:%S.%fZ')
        end_date = datetime.datetime.strptime(date_range[1], '%Y-%m-%dT%H:%M:%S.%fZ')
        invoice_type = ConfigProject.objects.get(token="invoice_type_invoice_token").value
        invoices = Invoice.objects.filter(issue_date__range=(start_date.date(), end_date.date()), type_final=invoice_type).distinct()
        invoices = filter_pending_invoices(invoices, request)
    elif has_serie_range:
        # Seleccio nomes pel rang de numero de factura: el rang ja identifica les
        # factures, sense lot ni periode. filter_pending_invoices hi aplica el rang
        # (apply_serie_final_range) i el periode surt de les factures resultants,
        # que es el que fan servir els titols i el nom del fitxer.
        invoices = filter_pending_invoices(serie_final_range_invoices(), request)
        start_date, end_date = serie_final_range_period(invoices)
        if start_date is None:
            return Response({"error": "the serie_final range has no invoices."}, status=status.HTTP_400_BAD_REQUEST), None, None
    else:
        raise Exception("Error obtaining invoices for report billing summary")
    
    if exploitation:
        invoices = invoices.filter(exploitation=exploitation)

    multi_q = invoice_multi_filter_q(**get_multi_ids(request.data))
    if multi_q:
        invoices = invoices.filter(multi_q).distinct()

    invoices = invoices.order_by('exploitation')

    # Line items for product columns (optionally filter by product_ids)
    line_items_qs = InvoiceLineItem.objects.filter(
        invoice__in=invoices,
        is_active=True,
    )
    if product_ids:
        line_items_qs = line_items_qs.filter(product__id__in=product_ids)

    # Discover all (product_name, interval) pairs and sort for stable column order
    def _product_column_label(product_name, interval):
        pn = (product_name or "").strip()
        iv = interval
        if iv is not None and iv != 0:
            return _("%(product)s TRAM %(interval)s") % {"product": pn, "interval": iv}
        return pn

    all_pairs = list(
        set(line_items_qs.values_list("product_name", "interval"))
    )
    all_pairs_sorted = sorted(
        all_pairs,
        key=lambda x: (str(x[0] or "").upper(), x[1] or 0),
    )
    product_interval_pairs = [
        (pn, iv) for pn, iv in all_pairs_sorted if _product_column_label(pn, iv)
    ]

    # Per-exploitation, per (product_name, interval): sum units and sum total
    line_aggregates = (
        line_items_qs.values(
            "invoice__exploitation_id",
            "product_name",
            "interval",
        )
        .annotate(
            units_sum=Sum("units"),
            total_sum=Sum("total"),
            price_sum=Sum("price"),
        )
    )

    # Build lookup: (exploitation_id, (product_name, interval)) -> {units, total}
    exp_product_sums = {}
    for agg in line_aggregates:
        eid = agg["invoice__exploitation_id"]
        pn = agg["product_name"] or ""
        iv = agg["interval"]
        units = float(agg["units_sum"] or 0.0)
        total = float(agg["total_sum"] or 0.0)
        price = float(agg["price_sum"] or 0.0)

        key = (eid, (pn, iv))
        if key not in exp_product_sums:
            exp_product_sums[key] = {
                "units": 0.0,
                "total": 0.0,
                "price": 0.0,
            }

        exp_product_sums[key]["units"] += units
        exp_product_sums[key]["total"] += total
        exp_product_sums[key]["price"] += price

    # Group and aggregate invoices by exploitation
    invoices_grouped_by_exploitation = list(
        invoices
        .values("exploitation_id", "exploitation__name")
        .annotate(
            total_m3=Sum("consumption"),
            total_final_sum=Sum("total_final"),
            subtotal_sum=Sum("subtotal_final"),
        )
        .order_by("exploitation__name")
    )

    # Compute estimated used totals per exploitation
    estimated_used_by_exp = {}
    # Pre-càlcul del Pou Propi per a cada explotació
    pou_propi_data = {}
    for exp in invoices_grouped_by_exploitation:
        eid = exp["exploitation_id"]
        invs_exp = invoices.filter(exploitation_id=eid)
        est_sum = Reading.objects.filter(invoices__in=invs_exp).aggregate(sum_est=Sum('estimated_used'))['sum_est'] or 0.0
        estimated_used_by_exp[eid] = float(est_sum)
        invs_without_aigua = invs_exp.exclude(
            id__in=InvoiceLineItem.objects.filter(
                invoice__in=invs_exp,
                product_name__in=['AIGUA', 'AIGUA C'],
                is_active=True
            ).values_list('invoice_id', flat=True)
        )
        pp_brut = float(invs_without_aigua.aggregate(Sum('consumption'))['consumption__sum'] or 0.0)
        pp_est = float(Reading.objects.filter(invoices__in=invs_without_aigua).aggregate(sum_est=Sum('estimated_used'))['sum_est'] or 0.0)
        pou_propi_data[eid] = {
            'brut': pp_brut,
            'estimat': pp_est,
            'net': pp_brut - pp_est
        }

    # --- Desglosament per USE TYPE de contracte (a nivell de factura) ---
    use_type_agg = list(
        invoices.values("exploitation_id", "contract__use_type__name")
        .annotate(
            total_m3=Sum("consumption"),
            total_final_sum=Sum("total_final"),
            subtotal_sum=Sum("subtotal_final"),
        )
        .order_by("exploitation_id", "contract__use_type__name")
    )
    all_use_types = sorted(set(r["contract__use_type__name"] or _("(Sense tipus)") for r in use_type_agg))
    # (eid, ut) -> {total_m3, total_final_sum, subtotal_sum}
    use_type_inv = {}
    for r in use_type_agg:
        ut = r["contract__use_type__name"] or _("(Sense tipus)")
        use_type_inv[(r["exploitation_id"], ut)] = {
            "total_m3": float(r["total_m3"] or 0),
            "total_final_sum": float(r["total_final_sum"] or 0),
            "subtotal_sum": float(r["subtotal_sum"] or 0),
        }

    # --- Desglosament per USE TYPE a nivell de línia de producte ---
    ut_line_agg = (
        line_items_qs
        .values("invoice__exploitation_id", "contract__use_type__name", "product_name", "interval")
        .annotate(units_sum=Sum("units"), total_sum=Sum("total"), price_sum=Sum("price"))
    )
    # (eid, ut, (pn, iv)) -> {units, total, price}
    ut_product_sums = {}
    for agg in ut_line_agg:
        eid = agg["invoice__exploitation_id"]
        ut = agg["contract__use_type__name"] or _("(Sense tipus)")
        pn = agg["product_name"] or ""
        iv = agg["interval"]
        key = (eid, ut, (pn, iv))
        if key not in ut_product_sums:
            ut_product_sums[key] = {"units": 0.0, "total": 0.0, "price": 0.0}
        ut_product_sums[key]["units"] += float(agg["units_sum"] or 0)
        ut_product_sums[key]["total"] += float(agg["total_sum"] or 0)
        ut_product_sums[key]["price"] += float(agg["price_sum"] or 0)

    # Lookup per eid per als totals de factura
    exp_totals = {e["exploitation_id"]: e for e in invoices_grouped_by_exploitation}

    multiple_exploitations = len(invoices_grouped_by_exploitation) > 1

    # Nombre de columnes per bloc d'explotació: ut_cols + ut_total
    cols_per_exp = len(all_use_types) + 1

    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = _("Facturació resumida")
    row = 0
    print("invoices: ", invoices.count())

    # --- Resum superior (TOTAL M3 i desglossament per tipus d'ús) ---
    import unicodedata
    from collections import defaultdict

    def _normalize_use_type(name):
        name = unicodedata.normalize('NFKD', name or "").encode('ascii', 'ignore').decode('ascii')
        return name.strip().upper()

    use_type_m3_totals = defaultdict(float)
    for r in use_type_agg:
        ut = r["contract__use_type__name"] or _("(Sense tipus)")
        use_type_m3_totals[_normalize_use_type(ut)] += float(r["total_m3"] or 0)

    total_m3_grand = sum(float(exp["total_m3"] or 0) for exp in invoices_grouped_by_exploitation)

    top_summary_categories = [
        (_("DOMESTIC"), "DOMESTIC"),
        (_("INDUSTRIAL"), "INDUSTRIAL"),
        (_("MUNICIPAL"), "MUNICIPAL"),
        (_("AGRICOLA"), "AGRICOLA"),
    ]

    row = add_row(sheet, row, [_("TOTAL M3"), f"{round(total_m3_grand, 4)}M3"], fill=black_fill, font=white_bold_font)
    categorized_m3 = 0.0
    for label, ut_key in top_summary_categories:
        value = use_type_m3_totals.get(ut_key, 0.0)
        categorized_m3 += value
        row = add_row(sheet, row, [label, f"{round(value, 4)}M3"])
    row = add_row(sheet, row, [_("ALTRES"), f"{round(total_m3_grand - categorized_m3, 4)}M3"])
    row = add_row(sheet, row, [""])

    # --- Capçaleres ---
    # Si múltiples explotacions: 3 files
    #   Fila 1: [""] + [nom_exp fusionat per cols_per_exp] * N + ["TOTAL GENERAL"]
    #   Fila 2: [""] + ["USE TYPE" fusionat] * N + [""]
    #   Fila 3: [""] + [ut1..utN, "TOTAL"] * N + [""]
    # Si una sola explotació: 2 files
    #   Fila 1: [""] + ["USE TYPE" fusionat]
    #   Fila 2: [""] + [ut1..utN, "TOTAL"]

    if multiple_exploitations:
        hr1 = [""]
        for exp in invoices_grouped_by_exploitation:
            hr1.append(exp["exploitation__name"] or "")
            hr1.extend([""] * (cols_per_exp - 1))
        hr1.append(_("TOTAL GENERAL"))
        row = add_row(sheet, row, hr1, fill=black_fill, font=white_bold_font)
        exp_col_start = 2
        for _exp in invoices_grouped_by_exploitation:
            if cols_per_exp > 1:
                sheet.merge_cells(
                    start_row=row, start_column=exp_col_start,
                    end_row=row, end_column=exp_col_start + cols_per_exp - 1
                )
            exp_col_start += cols_per_exp

    ut_span = len(all_use_types) + 1
    hr2 = [""]
    for _exp in invoices_grouped_by_exploitation:
        hr2.append(_("USE TYPE"))
        hr2.extend([""] * (ut_span - 1))
    hr2.append("" if multiple_exploitations else "")
    row = add_row(sheet, row, hr2, fill=black_fill, font=white_bold_font)
    exp_col_start = 2
    for _exp in invoices_grouped_by_exploitation:
        if ut_span > 1:
            sheet.merge_cells(
                start_row=row, start_column=exp_col_start,
                end_row=row, end_column=exp_col_start + ut_span - 1
            )
        exp_col_start += cols_per_exp

    hr3 = [""]
    for _exp in invoices_grouped_by_exploitation:
        hr3.extend(all_use_types)
        hr3.append(_("TOTAL"))
    hr3.append("" if multiple_exploitations else "")
    row = add_row(sheet, row, hr3, fill=black_fill, font=white_bold_font)

    # --- Funció per construir una fila de dades ---
    # get_ut_val(eid, ut) -> float|None   (None = sense desglossament per aquest camp)
    # get_total_val(eid) -> float|""
    def _data_row(label, get_ut_val, get_total_val, decimals=4):
        cells = [label]
        grand_total = 0.0
        is_numeric = False
        for exp in invoices_grouped_by_exploitation:
            eid = exp["exploitation_id"]
            total_val = get_total_val(eid)
            is_num = isinstance(total_val, (int, float))

            if get_ut_val is not None and is_num:
                ut_vals = [get_ut_val(eid, ut) for ut in all_use_types]
                ut_cells = [round(v, decimals) if isinstance(v, (int, float)) else "" for v in ut_vals]
                ut_total = round(sum(v for v in ut_vals if isinstance(v, (int, float))), decimals)
            else:
                ut_cells = [""] * len(all_use_types)
                ut_total = round(total_val, decimals) if is_num else ""

            cells.extend(ut_cells)
            cells.append(ut_total)

            if is_num:
                grand_total += total_val
                is_numeric = True

        if multiple_exploitations:
            cells.append(round(grand_total, decimals) if is_numeric else "")
        return cells

    # --- Files de mètriques de resum ---
    from collections import defaultdict

    def _sum_agua_trams(eid):
        return sum(
            data["units"] for key, data in exp_product_sums.items()
            if key[0] == eid and key[1][0] in ["AIGUA", "AIGUA C"]
            and key[1][1] is not None and key[1][1] > 0
        )

    def _sum_agua_notram(eid):
        return sum(
            data["units"] for key, data in exp_product_sums.items()
            if key[0] == eid and key[1][0] in ["AIGUA", "AIGUA C"]
            and (key[1][1] is None or key[1][1] <= 0)
        )

    summary_metrics = [
        (
            _("TOTAL M3 (Consum Brut)"),
            lambda eid, ut: use_type_inv.get((eid, ut), {}).get("total_m3"),
            lambda eid: float(exp_totals.get(eid, {}).get("total_m3") or 0),
            4,
        ),
        (
            _("TOTAL M3 FACTURAT ALS TRAMS (AIGUA)"),
            None,
            lambda eid: round(_sum_agua_trams(eid), 4),
            4,
        ),
        (
            _("TOTAL ESTIMAT RESTAT M3"),
            None,
            lambda eid: round(estimated_used_by_exp.get(eid, 0.0), 4),
            4,
        ),
        (
            _("TOTAL M3 POU PROPI / SENSE PRODUCTE AIGUA"),
            None,
            lambda eid: round(max(0.0, float(exp_totals.get(eid, {}).get("total_m3") or 0) - (
                _sum_agua_trams(eid) + estimated_used_by_exp.get(eid, 0.0)
            )), 4),
            4,
        ),
        (
            "  -> " + _("M3 POU PROPI (Net)"),
            None,
            lambda eid: round(pou_propi_data.get(eid, {}).get('net', 0.0), 4),
            4,
        ),
        (
            "  -> " + _("M3 POU PROPI (Estimat)"),
            None,
            lambda eid: round(pou_propi_data.get(eid, {}).get('estimat', 0.0), 4),
            4,
        ),
        (
            "  -> " + _("M3 AIGUA FACTURADA SENSE TRAM (Interval 0)"),
            None,
            lambda eid: round(_sum_agua_notram(eid), 4),
            4,
        ),
        (
            "  -> " + _("M3 DESVIAMENTS / ARRODONIMENTS"),
            None,
            lambda eid: round(
                max(0.0, float(exp_totals.get(eid, {}).get("total_m3") or 0) - (
                    _sum_agua_trams(eid) + estimated_used_by_exp.get(eid, 0.0)
                ))
                - pou_propi_data.get(eid, {}).get('net', 0.0)
                - _sum_agua_notram(eid),
                4
            ),
            4,
        ),
        (
            _("TOTAL"),
            lambda eid, ut: use_type_inv.get((eid, ut), {}).get("total_final_sum"),
            lambda eid: float(exp_totals.get(eid, {}).get("total_final_sum") or 0),
            2,
        ),
        (
            _("TOTAL SENSE IVA"),
            lambda eid, ut: use_type_inv.get((eid, ut), {}).get("subtotal_sum"),
            lambda eid: float(exp_totals.get(eid, {}).get("subtotal_sum") or 0),
            2,
        ),
    ]

    for label, get_ut, get_total, dec in summary_metrics:
        row = add_row(sheet, row, _data_row(label, get_ut, get_total, dec))

    row = add_row(sheet, row, [""])

    # --- Files per producte ---
    product_to_intervals = defaultdict(list)
    for pn, iv in product_interval_pairs:
        product_to_intervals[pn].append(iv)

    def _is_m3_product(product_name):
        p_upper = (product_name or "").upper()
        if "QUOTA FIXA" in p_upper or "CONSERVACIÓ" in p_upper or "CONSERVACIO" in p_upper or "MANTENIMENT" in p_upper:
            return False
        if any(k in p_upper for k in ["AIGUA", "CANON", "ACA", "CLAVEGUERAM", "CONSUM"]):
            return True
        return False

    sorted_products = sorted(product_to_intervals.keys(), key=lambda x: x.upper())
    for pn in sorted_products:
        intervals = product_to_intervals[pn]
        for iv in intervals:
            label_base = _product_column_label(pn, iv)
            if iv != 0 or _is_m3_product(pn):
                label_units = _("%(product)s M3") % {"product": label_base}
            else:
                label_units = _("%(product)s QUOTES/UNITATS") % {"product": label_base}
            label_eur = _("%(product)s TOTAL") % {"product": label_base}

            label_eur_no_tax = _("%(product)s TOTAL SENSE IVA") % {"product": label_base}

            row = add_row(sheet, row, _data_row(
                label_units,
                lambda eid, ut, p=pn, i=iv: ut_product_sums.get((eid, ut, (p, i)), {}).get("units"),
                lambda eid, p=pn, i=iv: round(exp_product_sums.get((eid, (p, i)), {}).get("units", 0.0), 4),
                4,
            ))
            row = add_row(sheet, row, _data_row(
                label_eur_no_tax,
                lambda eid, ut, p=pn, i=iv: ut_product_sums.get((eid, ut, (p, i)), {}).get("price"),
                lambda eid, p=pn, i=iv: round(exp_product_sums.get((eid, (p, i)), {}).get("price", 0.0), 2),
                2,
            ))
            row = add_row(sheet, row, _data_row(
                label_eur,
                lambda eid, ut, p=pn, i=iv: ut_product_sums.get((eid, ut, (p, i)), {}).get("total"),
                lambda eid, p=pn, i=iv: round(exp_product_sums.get((eid, (p, i)), {}).get("total", 0.0), 2),
                2,
            ))
            row = add_row(sheet, row, [""])

        if pn in ["AIGUA", "AIGUA C"]:
            row = add_row(sheet, row, _data_row(
                _("%(product)s ESTIMAT RESTAT M3") % {"product": pn},
                None,
                lambda eid: round(estimated_used_by_exp.get(eid, 0.0), 4),
                4,
            ))
            row = add_row(sheet, row, [""])

    adjust_column_widths(sheet)
    
    filename = f"billing_detailed_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    
    wb.save(filename)

    # Send the file as a response
    with open(filename, "rb") as f:
        response = HttpResponse(f.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f"attachment; filename={filename}"
    
    
    document_id = save_report(content=wb, filename=filename, name=name, type_id=type_id, start_date=start_date, end_date=end_date)
    
    
    os.remove(filename)
    return response, document_id, None


    
def generate_order_time_summary(request, black_fill=None, white_bold_font=None, task=None):
    try:
        print("getting order time summary")
        date_range = request.data.get('date_range', None)
        name = request.data.get('name', '')
        type_id = request.data.get('type_id', None)
        id = request.data.get('id', None)
        exploitation_id = request.data.get('exploitation_id', None)
        
        if not id and not date_range:
            return Response({"error": "id or date range is required."}, status=status.HTTP_400_BAD_REQUEST), None, None
        print("got id or date range")
        start_date = None
        end_date = None
        
        if date_range:
            start_date = datetime.datetime.strptime(date_range[0], '%Y-%m-%dT%H:%M:%S.%fZ')
            end_date = datetime.datetime.strptime(date_range[1], '%Y-%m-%dT%H:%M:%S.%fZ')
            orders = Order.objects.filter(created_at__range=(start_date.date(), end_date.date()))
            if exploitation_id:
                orders = orders.filter(
                    Q(contract__supply_point_default__connection__exploitation__id=exploitation_id) |
                    Q(contract_request__supply_point_default__connection__exploitation__id=exploitation_id) |
                    Q(supply_point__connection__exploitation__id=exploitation_id) |
                    Q(connection__exploitation__id=exploitation_id) |
                    Q(connection_request__exploitation__id=exploitation_id) |
                    Q(address__city__exploitations__id=exploitation_id) |
                    Q(claim_request__payments__contract__supply_point_default__connection__exploitation__id=exploitation_id)
                    ).distinct()
            reports_qs = OrderReport.objects.filter(order__in=orders)
            # Use a separate, sorted list for iteration so we can still
            # use the queryset (`reports_qs`) for ORM operations below. 
            reports = sorted(reports_qs, key=lambda x: x.order.created_at)
        else:
            raise Exception("Error obtaining orders for report time summary")
        orders_without_reports = Order.objects.filter(
            id__in=orders.values_list('id', flat=True),
            created_at__range=(start_date.date(), end_date.date())
        ).exclude(
            id__in=reports_qs.values_list('order__id', flat=True)
        )
        
        if exploitation_id:
            orders_without_reports = orders_without_reports.filter(
                Q(contract__supply_point_default__connection__exploitation__id=exploitation_id) |
                Q(contract_request__supply_point_default__connection__exploitation__id=exploitation_id) |
                Q(supply_point__connection__exploitation__id=exploitation_id) |
                Q(connection__exploitation__id=exploitation_id) |
                Q(connection_request__exploitation__id=exploitation_id) |
                Q(address__city__exploitations__id=exploitation_id) |
                Q(claim_request__payments__contract__supply_point_default__connection__exploitation__id=exploitation_id)
                ).distinct()
        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.title = _("Resum temps ordres")
        
        row = 0
        
        title = _("Ordres de treball %(start)s - %(end)s") % {"start": start_date.strftime('%d/%m/%Y'), "end": end_date.strftime('%d/%m/%Y')} if date_range else _("Ordre de treball %(token)s") % {"token": order.token} if id else ""
        row = add_row(sheet, row, [title])
        sheet.merge_cells(f'A{row}:D{row}')
        
        row = jump_row(row)
        
        titles = [_("Ordre"), _("Data creació"), _("Data finalització"), _("Temps dedicat (min)"), _("Tècnic"), _("Adreça"), _("Tipus"), _("Motiu"), _("Observacions")]
        row = add_row(sheet, row, titles, fill=black_fill, font=white_bold_font)
        
        current_order = None
        order_time = 0
        total_time = 0
        operator_times = {}
        for order in orders_without_reports:
            order_address = get_order_address(order)
            order_row = [
                order.token,
                order.created_at.strftime('%d/%m/%Y'),
                order.completed_at.strftime('%d/%m/%Y') if order.completed_at else "",
                0,
                f"{order.operators.first().name} {order.operators.first().surname} ({order.operators.first().token})" if len(order.operators.all()) > 0 else "",
                order_address,
                order.type.name if order.type else "",
                order.reason.name if order.reason else "",
                "",
            ]
            row = add_row(sheet, row, order_row)
        
        total_reports = len(reports)
        for idx_report, report in enumerate(reports, start=1):
            if task and idx_report % 50 == 0:
                task.update_state(
                    state='PROGRESS',
                    meta={
                        'current': idx_report,
                        'total': total_reports,
                        'percent': round((idx_report / total_reports) * 100, 2) if total_reports else 0.0
                    }
                )
            if current_order and current_order != report.order.id:
                order_address = get_order_address(report.order)
                
                for operator_id, time in operator_times.items():
                    operator = Operator.objects.get(id=operator_id)
                    operator_row = [
                        report.order.token,
                        report.order.created_at.strftime('%d/%m/%Y'),
                        report.order.completed_at.strftime('%d/%m/%Y') if report.order.completed_at else "",
                        time,
                        f"{operator.name} {operator.surname} ({operator.token})" if operator else "",
                        order_address,
                        report.order.type.name if report.order.type else "",
                        report.order.reason.name if report.order.reason else "",
                        report.observation if report.observation else ""
                    ]
                    row = add_row(sheet, row, operator_row)
                
                order_time = 0
                operator_times = {}
            
            current_order = report.order.id
            order_time += report.time_dedicated if report.time_dedicated else 0
            total_time += report.time_dedicated if report.time_dedicated else 0
            
            if report.operator:
                if report.operator.id not in operator_times:
                    operator_times[report.operator.id] = 0
                operator_times[report.operator.id] += report.time_dedicated if report.time_dedicated else 0
        
        
        if reports:
            order_address = get_order_address(reports[-1].order)
            for operator_id, time in operator_times.items():
                operator = Operator.objects.get(id=operator_id)
                operator_row = [
                    reports[-1].order.token,
                    reports[-1].order.created_at.strftime('%d/%m/%Y'),
                    reports[-1].order.completed_at.strftime('%d/%m/%Y') if reports[-1].order.completed_at else "",
                    time,
                    f"{operator.name} {operator.surname} ({operator.token})" if operator else "",
                    order_address,
                    reports[-1].order.type.name if reports[-1].order.type else "",
                    reports[-1].order.reason.name if reports[-1].order.reason else "",
                    reports[-1].observation if reports[-1].observation else ""
                ]
                row = add_row(sheet, row, operator_row)
            
            # order_total = ["Total ordre", "", "", f"{order_time} min", ""]
            # row = add_row(sheet, row, order_total, fill=black_fill, font=white_bold_font)
            # row = jump_row(row)
        
        
        grand_total = [_("TOTAL GENERAL"), "", "", f"{total_time} min"]
        row = add_row(sheet, row, grand_total, fill=black_fill, font=white_bold_font)
        
        adjust_column_widths(sheet)
        
        filename = f"orders_time_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        
        wb.save(filename)

        
        with open(filename, "rb") as f:
            response = HttpResponse(f.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            response["Content-Disposition"] = f"attachment; filename={filename}"
        
        document_id = save_report(content=wb, filename=filename, name=name, type_id=type_id, start_date=start_date, end_date=end_date)
        print("report saved")
        
        os.remove(filename)
        return response, document_id, None
    except Exception as e:
        return Response({"error": f"Ha ocorregut un error inesperat: {e}"}, status=status.HTTP_500_INTERNAL_SERVER_ERROR), None, None

def get_order_address(order):
    order_address = ""
    if order.address:
        order_address = str(order.address)
    if not order_address or order_address == "":
        if order.supply_point and order.supply_point.address:
            order_address = str(order.supply_point.address)
    if not order_address or order_address == "":
        if order.contract_request and order.contract_request.supply_point_default and order.contract_request.supply_point_default.address:
            order_address = str(order.contract_request.supply_point_default.address)
        elif order.contract and order.contract.supply_point_default and order.contract.supply_point_default.address:
            order_address = str(order.contract.supply_point_default.address)
        elif order.contract_termination_request and order.contract_termination_request.contract.supply_point_default and order.contract_termination_request.contract.supply_point_default.address:
            order_address = str(order.contract_termination_request.contract.supply_point_default.address)
        elif order.connection or order.connection_request:
            conn = order.connection if order.connection else order.connection_request.connection
            order_address = f"{str(conn.address_street)}, {str(conn.address_street_number)} - {str(conn.address_postal_code)} {str(conn.address_city)}"
    return order_address

def generate_bails_report(request, black_fill=None, white_bold_font=None, task=None):
    try:
        print("getting bails report txt")
        date_range = request.data.get('date_range', None)
        name = request.data.get('name', '')
        type_id = request.data.get('type_id', None)
        exploitation_id = request.data.get('exploitation_id', None)
        
        if not date_range:
            return Response({"error": "date range is required."}, status=status.HTTP_400_BAD_REQUEST), None, None
        
        bail_status_returned_token = ConfigProject.objects.get(token='bail_status_returned_token').value
        bail_status_pending_token = ConfigProject.objects.get(token='bail_status_pending_token').value
        
        try:
            start_date = datetime.datetime.strptime(date_range[0], '%Y-%m-%dT%H:%M:%S.%fZ')
            end_date = datetime.datetime.strptime(date_range[1], '%Y-%m-%dT%H:%M:%S.%fZ')
            contracts = Contract.objects.filter(
                created_at__range=(start_date, end_date)
            )
            multi_ids = get_multi_ids(request.data)
            multi_q = contract_multi_filter_q(person_ids=multi_ids['person_ids'], contract_ids=multi_ids['contract_ids'])
            if multi_q:
                contracts = contracts.filter(multi_q).distinct()
            bails_terminations = Bail.objects.filter(
                return_date__range=(start_date, end_date)
            )
            bails_contracts = Bail.objects.filter(
                contract__id__in=contracts.values_list('id', flat=True),
            )
            
        except Exception as e:
            return Response({"error": f"Ha ocorregut un error inesperat: {e}"}, status=status.HTTP_500_INTERNAL_SERVER_ERROR), None, None
        
        wb = openpyxl.Workbook()
        sheet = wb.active
        incasol_num = resolve_incasol_num(exploitation_id)
        n_con = f"S{incasol_num}"
        trimestre = "01" if end_date.month in [1, 2, 3] else "02" if end_date.month in [4, 5, 6] else "03" if end_date.month in [7, 8, 9] else "04"
        n_liq = f"77{trimestre}{end_date.strftime('%y')}{incasol_num}"
        date = end_date.strftime('%d/%m/%Y')
        total_contracts = bails_contracts.count()
        total_terminations = bails_terminations.count()
        price_contracts = round(sum(bail_contract.amount for bail_contract in bails_contracts), 2)
        price_terminations = -round(sum(bail_ter.amount for bail_ter in bails_terminations), 2)
        filename = f"INCASOL_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        
        sheet.title = "Capçalera"
        print("before sheet titles")
        sheet_titles = [
            "Núm. Concert", "Núm. Liquidació", "Data Liquidació", 
            "Núm. Altes", "Núm. Baixes", 
            "Import Altes", "Import Baixes", "Fitxer Errors"
            ]
        sheet_title_values = [
            n_con, n_liq, date, 
            total_contracts, total_terminations, 
            price_contracts, price_terminations,
            0,
        ]
        row = 0
        row = add_row(sheet, row, sheet_titles)
        row = add_row(sheet, row, sheet_title_values)
        adjust_column_widths(sheet)
        print("after sheet titles")
        
        print("before sheet lines")
        sheet_lines = wb.create_sheet(title="Moviments")
        sheet_lines_titles = [
            "tipus", "número referència propi contracte", "data fiança", 
            "import fiança", "tipus carrer", "nom carrer", 
            
            "núm. carrer", "escala", "pis",
            "porta", "codi postal", "municipi",
            
            "altre identificador adreça", "referència cadastral",
            "NIF abonat", "cognom/raó social abonat",
            
            "nom abonat", "Id. Error Fiança", "Descripció Error"
        ]
        row_lines = 0
        row_lines = add_row(sheet_lines, row_lines, sheet_lines_titles)
        
        street_types_by_id = {st.id: st for st in StreetType.objects.all()}
        for idx_bail_contract, bail_contract in enumerate(bails_contracts, start=1):
            if task and idx_bail_contract % 50 == 0:
                task.update_state(
                    state='PROGRESS',
                    meta={
                        'current': idx_bail_contract,
                        'total': total_contracts,
                        'percent': round((idx_bail_contract / total_contracts) * 100, 2) if total_contracts else 0.0
                    }
                )
            contract = bail_contract.contract
            bail_total = round(bail_contract.amount, 2)
            address = contract.supply_point_default.address if contract.supply_point_default else None
            street_number = address.street_number if address and address.street_number else None
            street = address.street if address else None
            st_type = street_types_by_id.get(street.type_id) if street and street.type_id else None
            street_type = (
                (st_type.aca_abbreviation or st_type.abbreviation or "C")
                if st_type else ""
            )
            
            exploitation_name = ""
            if (contract.supply_point_default and 
                contract.supply_point_default.connection and 
                contract.supply_point_default.connection.exploitation and 
                contract.supply_point_default.connection.exploitation.name):
                exploitation_name = contract.supply_point_default.connection.exploitation.name.upper()
            
            lines_content = [
                "FI", contract.token, bail_contract.created_at.strftime('%d/%m/%Y'),
                bail_total, street_type, 
                (f"{street.name if street else ''}{', '+str(street_number) if street_number else ''}{', ' +address.stair if address and address.stair else ''}{', ' +address.floor if address and address.floor else ''}{', ' +address.door if address and address.door else ''}").upper()[:50],
                
                "", "", "",
                "", address.postal_code if address else "", exploitation_name,
                
                address.building if address else "", "",
                contract.holder.token, 
                (f"{contract.holder.surname + ', ' if contract.holder.surname else ''}{contract.holder.name if contract.holder.name else ''}").upper()[:50],
                
                "", "", ""
            ]
            row_lines = add_row(sheet_lines, row_lines, lines_content)
        
        print("after sheet lines new")
        for idx_bail_ter, bail_ter in enumerate(bails_terminations, start=1):
            if task and idx_bail_ter % 50 == 0:
                task.update_state(
                    state='PROGRESS',
                    meta={
                        'current': idx_bail_ter,
                        'total': total_terminations,
                        'percent': round((idx_bail_ter / total_terminations) * 100, 2) if total_terminations else 0.0
                    }
                )
            contract = bail_ter.contract
            bail_total = round(bail_ter.amount, 2)
            address = contract.supply_point_default.address if contract.supply_point_default else None
            street_number = address.street_number if address and address.street_number else None
            street = address.street if address else None
            st_type = street_types_by_id.get(street.type_id) if street and street.type_id else None
            street_type = (
                (st_type.abbreviation or "")
                if st_type else ""
            )
            
            exploitation_name = ""
            if (contract.supply_point_default and 
                contract.supply_point_default.connection and 
                contract.supply_point_default.connection.exploitation and 
                contract.supply_point_default.connection.exploitation.name):
                exploitation_name = contract.supply_point_default.connection.exploitation.name.upper()

            lines_content = [
                "CA", contract.token, bail_ter.return_date.strftime('%d/%m/%Y'),
                -bail_total, street_type,
                (f"{street.name if street else ''}{', '+str(street_number) if street_number else ''}{', ' +address.stair if address and address.stair else ''}{', ' +address.floor if address and address.floor else ''}{', ' +address.door if address and address.door else ''}").upper()[:50],
                
                "", "", "",
                "", address.postal_code if address else "", exploitation_name,
                
                address.building if address else "", "",
                contract.holder.token, 
                (f"{contract.holder.surname + ', ' if contract.holder.surname else ''}{contract.holder.name if contract.holder.name else ''}").upper()[:50],
                
                "", "", ""
            ]
            row_lines = add_row(sheet_lines, row_lines, lines_content)
            
        print("after sheet lines return")
        adjust_column_widths(sheet_lines)

        data_row_count = sheet_lines.max_row - 1 
        if data_row_count > 0:
            data_rows = list(
                sheet_lines.iter_rows(min_row=2, max_row=sheet_lines.max_row, values_only=True)
            )

            def _data_fianca_sort_key(row):
                val = row[2] 
                if not val:
                    return datetime.datetime.min
                try:
                    return datetime.datetime.strptime(val, "%d/%m/%Y")
                except (ValueError, TypeError):
                    return datetime.datetime.min

            data_rows.sort(key=_data_fianca_sort_key)
            for i, row_data in enumerate(data_rows, start=2):
                for j, value in enumerate(row_data, start=1):
                    sheet_lines.cell(row=i, column=j, value=value)

        wb.save(filename)
        
        
        with open(filename, "rb") as f:
            response = HttpResponse(f.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            response["Content-Disposition"] = f"attachment; filename={filename}"
        
        document_id = save_report(content=wb, filename=filename, name=name, type_id=type_id, start_date=start_date, end_date=end_date)
        
        os.remove(filename)
        return response, document_id, None
    except Exception as e:
        return Response({"error": f"Ha ocorregut un error inesperat: {e}"}, status=status.HTTP_500_INTERNAL_SERVER_ERROR), None, None
    
    
def generate_sii_report(request, black_fill=None, white_bold_font=None, task=None):
    try:
        """ 
        INFO
        https://sede.agenciatributaria.gob.es/static_files/Sede/Procedimiento_ayuda/G417/FicherosSuministros/V_1_1/Validaciones_ErroresSII_v1.1.pdf
        https://sede.agenciatributaria.gob.es/static_files/Sede/Procedimiento_ayuda/G417/FicherosSuministros/V_1_1/FaqGral/FAQs_SII_23_04_2026.pdf
        https://www.promer-asesores.com/DOCUMENTOS/Ref.Web.18.10.SII.pdf
        """
        
        id = request.data.get('id', None)
        date_range = request.data.get('date_range', None)
        billing_ids = get_billing_ids(request.data)
        title_name = request.data.get('name', '')
        type_id = request.data.get('type_id', None)
        exploitation_id = request.data.get('exploitation_id', None)
        
        start_date = None
        end_date = None
        
        has_serie_range = has_serie_final_range(request.data)
        if not billing_ids and not date_range and not has_serie_range:
            return Response({"error": "billing id, date range or serie_final range is required."}, status=status.HTTP_400_BAD_REQUEST), None, None
            
        if billing_ids:
            invoices = Invoice.objects.filter(billing_id__in=billing_ids)
            invoices = filter_pending_invoices(invoices, request)
        elif date_range:
            start_date = datetime.datetime.strptime(date_range[0], '%Y-%m-%dT%H:%M:%S.%fZ')
            end_date = datetime.datetime.strptime(date_range[1], '%Y-%m-%dT%H:%M:%S.%fZ')
            invoice_type = ConfigProject.objects.get(token="invoice_type_invoice_token").value
            invoices = Invoice.objects.filter(issue_date__range=(start_date.date(), end_date.date()), type_final=invoice_type).distinct()
            invoices = filter_pending_invoices(invoices, request)
        elif has_serie_range:
            # Seleccio nomes pel rang de numero de factura: el rang ja identifica les
            # factures, sense lot ni periode. filter_pending_invoices hi aplica el rang
            # (apply_serie_final_range) i el periode surt de les factures resultants,
            # que es el que fan servir els titols i el nom del fitxer.
            invoices = filter_pending_invoices(serie_final_range_invoices(), request)
            start_date, end_date = serie_final_range_period(invoices)
            if start_date is None:
                return Response({"error": "the serie_final range has no invoices."}, status=status.HTTP_400_BAD_REQUEST), None, None
        else:
            raise Exception("Error obtaining invoices for report billing summary")
        if exploitation_id:
            exploitation = Exploitation.objects.get(id=exploitation_id)
            invoices = invoices.filter(exploitation__id=exploitation_id)
            company = exploitation.company
        else:
            main_company_token = ConfigProject.objects.get(token="main_company_token").value
            company = Company.objects.get(vat=main_company_token)

        filename = (
            f'SII_{start_date.strftime("%Y%m%d")}_{end_date.strftime("%Y%m%d")}.xml'
            if start_date and end_date
            else f'SII_{invoices[0].billing.token}.xml'
        )

        root = ET.Element("soapenv:Envelope", {
            "xmlns:soapenv": "http://schemas.xmlsoap.org/soap/envelope/",
            "xmlns:sum": "https://www2.agenciatributaria.gob.es/static_files/common/internet/dep/aplicaciones/es/aeat/ssii/fact/ws/SuministroLR.xsd",
            "xmlns:sum1": "https://www2.agenciatributaria.gob.es/static_files/common/internet/dep/aplicaciones/es/aeat/ssii/fact/ws/SuministroInformacion.xsd"
        })
        
        ET.SubElement(root, "soapenv:Header")
        
        body = ET.SubElement(root, "soapenv:Body")
        sumSumLRFacEmit = ET.SubElement(body, "sum:SuministroLRFacturasEmitidas")
        sumCab = ET.SubElement(sumSumLRFacEmit, "sum1:Cabecera")
        ET.SubElement(sumCab, "sum1:IDVersionSii").text = "1.1"
        sumTitl = ET.SubElement(sumCab, "sum1:Titular")
        ET.SubElement(sumTitl, "sum1:NombreRazon").text = str(company.name) if company.name else str(company.alias) 
        ET.SubElement(sumTitl, "sum1:NIFRepresentante").text = str(company.vat) 
        ET.SubElement(sumTitl, "sum1:NIF").text = str(company.vat) 
        ET.SubElement(sumCab, "sum1:TipoComunicacion").text = 'A0'
        
        for invoice in invoices:
            line_items = invoice.line_items.filter(product__company=company, is_active=True)
            subtotal, taxes, taxes_base, total = get_totals(invoice, None, line_items)
            sumRegLRFactEmt = ET.SubElement(sumSumLRFacEmit, "sum:RegistroLRFacturasEmitidas")
            
            sumPerLiq = ET.SubElement(sumRegLRFactEmt, "sum1:PeriodoLiquidacion")
            ET.SubElement(sumPerLiq, "sum1:Ejercicio").text = str(invoice.issue_date.year) 
            ET.SubElement(sumPerLiq, "sum1:Periodo").text = str(invoice.issue_date.month).zfill(2) 
            
            sumIDFact = ET.SubElement(sumRegLRFactEmt, "sum:IDFactura")
            sumIDEmsFact = ET.SubElement(sumIDFact, "sum1:IDEmisorFactura")
            ET.SubElement(sumIDEmsFact, "sum1:NIF").text = str(company.vat) 
            ET.SubElement(sumIDFact, "sum1:NumSerieFacturaEmisor").text = str(invoice.serie_final) 
            ET.SubElement(sumIDFact, "sum1:FechaExpedicionFacturaEmisor").text = invoice.issue_date.strftime('%d-%m-%Y') 
            
            sumFactExp = ET.SubElement(sumRegLRFactEmt, "sum:FacturaExpedida")
            invoice_type = 'F1' if invoice.total_final >= 3000 and not invoice.refactored_token else 'R1' if invoice.total_final >= 3000 and invoice.refactored_token else 'R5' if invoice.refactored_token else 'F2'
            ET.SubElement(sumFactExp, "sum1:TipoFactura").text = invoice_type
            if invoice_type == 'R1' or invoice_type == 'R5':
                ET.SubElement(sumFactExp, "sum1:TipoRectificativa").text = "S"
                importeRectificacion = ET.SubElement(sumFactExp, "sum1:ImporteRectificacion")
                ET.SubElement(importeRectificacion, "sum1:BaseRectificada").text = "0.00"
                ET.SubElement(importeRectificacion, "sum1:CuotaRectificada").text = "0.00"
            
            ET.SubElement(sumFactExp, "sum1:ClaveRegimenEspecialOTrascendencia").text = '01' 
            ET.SubElement(sumFactExp, "sum1:ImporteTotal").text = str(total) 
            
            if invoice_type == 'R1' or invoice_type == 'F1':
                contraparte = ET.SubElement(sumFactExp, "sum1:Contraparte")
                ET.SubElement(contraparte, "sum1:NIF").text = str(invoice.customer_token_final)
                ET.SubElement(contraparte, "sum1:NombreRazonSocial").text = str(invoice.customer_final)
            
            ET.SubElement(sumFactExp, "sum1:DescripcionOperacion").text = invoice.title_final 
            ET.SubElement(sumFactExp, "sum1:Macrodato").text = 'N' if invoice.total_final < 100000000 else 'S' 
            
            sumTpDesgl = ET.SubElement(sumFactExp, "sum1:TipoDesglose")
            sumDesglTpOp = ET.SubElement(sumTpDesgl, "sum1:DesgloseTipoOperacion") 
            sumPrsServ = ET.SubElement(sumDesglTpOp, "sum1:PrestacionServicios")
            sumSuj = ET.SubElement(sumPrsServ, "sum1:Sujeta")  
            
            # get_totals drops 0% lines from taxes_base, so the exempt base is the
            # subtotal minus the 10%/21% (and any other taxed) bases.
            taxed_base = sum(round_ceil(float(base or 0)) for base in taxes_base.values())
            exempt_base = round_ceil(round_ceil(float(subtotal or 0)) - taxed_base)
            has_non_exempt = any(tax != 0 for tax in taxes.values())

            if exempt_base != 0 or not has_non_exempt:
                sumEx = ET.SubElement(sumSuj, "sum1:Exenta")
                sumDetEx = ET.SubElement(sumEx, "sum1:DetalleExenta")
                ET.SubElement(sumDetEx, "sum1:CausaExencion").text = 'E1'
                ET.SubElement(sumDetEx, "sum1:BaseImponible").text = "{:.2f}".format(exempt_base)

            if not has_non_exempt:
                continue
            sumNoEx = ET.SubElement(sumSuj, "sum1:NoExenta") 
            ET.SubElement(sumNoEx, "sum1:TipoNoExenta").text = 'S1' 
            sumDesglIVA = ET.SubElement(sumNoEx, "sum1:DesgloseIVA") 
            try:
                for tax in taxes_base:
                    cuota = round_ceil(float(taxes.get(tax, 0)))
                    base = round_ceil(float(taxes_base.get(tax, 0)))
                    sumDetIVA = ET.SubElement(sumDesglIVA, "sum1:DetalleIVA") 
                    ET.SubElement(sumDetIVA, "sum1:TipoImpositivo").text = str(int(float(tax))) 
                    ET.SubElement(sumDetIVA, "sum1:BaseImponible").text = "{:.2f}".format(base) 
                    ET.SubElement(sumDetIVA, "sum1:CuotaRepercutida").text = "{:.2f}".format(cuota) 
            except Exception as e:
                print("Error in sum1:DetalleIVA: ", e)
                raise e
        
        xml_string = ET.tostring(root, encoding='utf-8', xml_declaration=True).decode('utf-8')

        response = HttpResponse(xml_string, content_type="application/xml")
        response["Content-Disposition"] = f'attachment; filename="{filename}"'

        document_id = save_report(
            content=xml_string,
            filename=filename,
            name=title_name,
            type_id=type_id,
            start_date=start_date,
            end_date=end_date,
        )

        # os.remove(filename)
        return response, document_id, filename
    except Exception as e:
        print("error: ", e)
        return Response({"error": f"Ha ocorregut un error inesperat: {e}"}, status=status.HTTP_500_INTERNAL_SERVER_ERROR), None, None


    
""" 
def generate_pending_invoices_values_report(request, black_fill=None, white_bold_font=None):
    
    id = request.data.get('id', None)
    date_range = request.data.get('date_range', None)
    name = request.data.get('name', '')
    type_id = request.data.get('type_id', None)
    exploitation_id = request.data.get('exploitation_id', None)
    accounting_type = request.data.get('accounting_type', None)
    
        
    cutoff_date = datetime.date(2025, 12, 31)
    cancelled_issue_date_cutoff = datetime.date(2025, 10, 1)
    post_2025_cutoff = datetime.date(2025, 12, 31)
    paid_status_token = ConfigProject.objects.get(token="payment_status_paid_token").value

    # Invoices issued on/before cutoff that were NOT paid during the period:
    # - include if no "paid" payments exist yet
    # - include if latest "paid" payment_date is AFTER the cutoff (paid later)
    # - exclude if latest "paid" payment_date is on/before cutoff (paid within period)
    payment_paid_token = ConfigProject.objects.get(token="payment_status_paid_token").value
    invoice_payoff_token = ConfigProject.objects.get(token="invoice_status_payoff_token").value
    invoice_pending_token = ConfigProject.objects.get(token="invoice_status_pending_token").value
    invoice_cancelled_token = ConfigProject.objects.get(token="invoice_status_cancelled_token").value
    payment_balance_token = ConfigProject.objects.get(token="payment_status_piggy_token").value
    invoices = (
        Invoice.objects.filter(
            type_final="F",
            issue_date__lte=cutoff_date,
        ).exclude(
            status__token__in=[invoice_payoff_token,invoice_pending_token]
        ).exclude(
            payments__status__token=payment_balance_token
        ).filter(
            Q(payments__payment_date__gt=cutoff_date, payments__status__token=payment_paid_token) | 
            Q(payments__status__token__in=["-1", "1", "-2", "-3", "-5"])
        )
    )

    # Cancelled invoices handling (return invoices):
    # - Cancelled with issue_date < 2025-10-01 => ignore
    # - Cancelled with no return_token => ignore
    # - Cancelled with return_token => keep ONLY if the referenced invoice issue_date is post-2025 (>= 2026-01-01)
    print("invoices: ", invoices.count())
    invoices = invoices.annotate(
        return_invoice_issue_date=Subquery(
            Invoice.objects.filter(token=OuterRef("return_token"))
            .values("issue_date")[:1]
        )
    ).exclude(
        Q(status__token=invoice_cancelled_token)
        & (
            Q(issue_date__lt=cancelled_issue_date_cutoff)
            | Q(return_token__isnull=True)
            | Q(return_token__exact="")
            | Q(return_invoice_issue_date__isnull=True)
            | Q(return_invoice_issue_date__lte=post_2025_cutoff)
        )
    )

    # Collect all related contract IDs from different invoice relationships.
    contract_ids = set(
        invoices.values_list("contract__id", flat=True).distinct()
    )
    contract_request_ids = set(
        Contract.objects.filter(
            contract_request__id__in=invoices.values_list(
                "contract_request__id", flat=True
            ).distinct()
        )
        .values_list("id", flat=True)
        .distinct()
    )
    contract_termination_ids = set(
        invoices.values_list(
            "contract_termination__contract__id", flat=True
        ).distinct()
    )

    all_contract_ids = (
        contract_ids | contract_request_ids | contract_termination_ids
    )

    contracts = Contract.objects.filter(id__in=all_contract_ids).distinct()

    # Accumulate debt per contract (keyed by token); in-memory dict so it persists.
    contract_debt = {}
    
    single_invoices = {}

    wb = openpyxl.Workbook()
    contract_debt_sheet = wb.create_sheet(title=_("Contractes amb deute"))
    pending_invoices_sheet = wb.create_sheet(title=_("Factures pendents"))
    
    row = 0
    inv_row = 0
    titles = [_("Explotació"), _("Contracte"), _("Abonat"), _("NIF Abonat"), _("Deute Pendent")]
    row = add_row(contract_debt_sheet, row, titles, fill=black_fill, font=white_bold_font)
      
    inv_titles = [_("Contracte"), _("NIF Abonat"), _("Núm Factura"), _("Data Factura"), _("Estat"), _("Data pagament"), _("Total Factura"), _("Total Pendent")]
    inv_row = add_row(pending_invoices_sheet, inv_row, inv_titles, fill=black_fill, font=white_bold_font)
    count = 0
    for inv in invoices:
        if count % 1000 == 0:
            print(f"processing invoice {count} out of {invoices.count()}")
        count += 1
        contract_token = None
        if inv.contract:
            contract_token = inv.contract.token
        elif inv.contract_request:
            contract_token = inv.contract_request.token
            try:
                contract_token = Contract.objects.get(token=contract_token).token
            except Exception as e:
                print(f"error getting contract {contract_token}: {e}")
                continue
        elif inv.contract_termination:
            contract_token = inv.contract_termination.contract.token
        if contract_token is not None:
            contract_debt[contract_token] = contract_debt.get(contract_token, 0) + inv.left_to_pay
        else:
            # Aggregate debt for customers whose invoices are not linked to a contract.
            prev_data = single_invoices.get(inv.customer_token_final, {})
            customer_debt = prev_data.get("left_to_pay", 0) + inv.left_to_pay
            single_invoices[inv.customer_token_final] = {
                "exploitation": inv.exploitation.name if inv.exploitation else "",
                "customer_token_final": inv.customer_token_final,
                "customer_final": inv.customer_final,
                "left_to_pay": customer_debt,
            }
        payment = inv.payments.all().order_by('-payment_date').first()
        inv_row_data = [
            inv.contract.token if inv.contract else inv.contract_request.token if inv.contract_request else inv.contract_termination.contract.token if inv.contract_termination else "",
            inv.customer_token_final,
            inv.serie_final,
            inv.issue_date.strftime('%d-%m-%Y') if inv.issue_date else "",
            inv.status.name,
            payment.payment_date.strftime('%d-%m-%Y') if payment and payment.payment_date and payment.status.token == payment_paid_token else "",
            inv.total_final,
            inv.left_to_pay,
        ]
        inv_row = add_row(pending_invoices_sheet, inv_row, inv_row_data)
    
    adjust_column_widths(pending_invoices_sheet)    
    contract_count = 0
    for contract in contracts:
        if contract_count % 1000 == 0:
            print(f"processing contract {contract_count} out of {contracts.count()}")
        contract_count += 1
        total_debt = contract_debt.get(contract.token, 0)
        contract_row_data = [
            contract.supply_point_default.connection.exploitation.name if contract.supply_point_default and contract.supply_point_default.connection and contract.supply_point_default.connection.exploitation else "",
            contract.token,
            str(contract.holder),
            contract.holder.token,
            total_debt,
        ]
        row = add_row(contract_debt_sheet, row, contract_row_data)
    
    # Add rows for "single" invoices (customers without a linked contract) to the same debt sheet.
    for data in single_invoices.values():
        contract_row_data = [
            data["exploitation"],
            "",
            data["customer_final"],
            data["customer_token_final"],
            data["left_to_pay"],
        ]
        row = add_row(contract_debt_sheet, row, contract_row_data)
    
    adjust_column_widths(contract_debt_sheet)  
    
    filename = f"puntual_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    
    wb.save(filename)

    # Send the file as a response
    with open(filename, "rb") as f:
        response = HttpResponse(f.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f"attachment; filename={filename}"
    
    document_id = save_report(content=wb, filename=filename, name=name, type_id=type_id)
    os.remove(filename)
    return response, document_id, None


"""


# Permet personalitzar els reports per client: si existeix report_service_personalized,
# s'usen les seves funcions en lloc de les d'aquest mòdul.
try:
    from statistics.utils import report_service_personalized
    for _attr in dir(report_service_personalized):
        if not _attr.startswith("_"):
            globals()[_attr] = getattr(report_service_personalized, _attr)
except ImportError:
    pass