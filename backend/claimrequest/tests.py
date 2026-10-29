from types import SimpleNamespace

from django.contrib.auth.models import User
from django.db.models import Q
from django.test import TestCase
from rest_framework.test import APIClient

from billing.models import InvoiceStatus, ReadingBatchTemplate
from claimrequest.tasks import _build_claim_request_filters
from claimrequest.utils.claim_request_service import (
    contract_from_finished_request,
    resolve_invoice_contract,
)
from contract.models import Contract, ContractRequest, ContractRequestStatus
from coredata.models import ConfigProject, Person
from service.models import Route, RouteZone, SupplyPoint


def q_lookups(q):
    """Totes les condicions (lookup, valor) d'un Q, incloses les dels Q niats."""
    lookups = []
    for child in q.children:
        if isinstance(child, Q):
            lookups.extend(q_lookups(child))
        else:
            lookups.append(child)
    return lookups


class ClaimRequestReadingBatchTemplateFilterTests(TestCase):
    """El filtre de reclamació d'impagats per plantilla de lot de lectura
    restringeix els contractes per la ruta del punt de subministrament."""

    def setUp(self):
        for token in (
            'payment_status_expired_token',
            'payment_status_returned_token',
            'payment_status_endowment_token',
            'payment_status_irrecoverable_token',
            'claim_request_status_accepted_token',
            'claim_request_status_pending_token',
        ):
            ConfigProject.objects.get_or_create(token=token, defaults={'value': token.upper()})
        expired = ConfigProject.objects.get_or_create(
            token='invoice_status_expired_token', defaults={'value': 'EXPIRED'}
        )[0].value
        InvoiceStatus.objects.get_or_create(token=expired, defaults={'name': expired})

    def test_filter_by_reading_batch_template(self):
        filters, *_ = _build_claim_request_filters({'selectedReadingBatchTemplatesIds': [3, 5]})
        self.assertIn(
            ('supply_point_default__property__route_position__route__reading_batch_template__id__in', [3, 5]),
            q_lookups(filters),
        )

    def test_no_filter_without_reading_batch_template(self):
        filters, *_ = _build_claim_request_filters({})
        self.assertFalse(any('reading_batch_template' in lookup for lookup, _ in q_lookups(filters)))

    def test_closed_termination_invoices_stay_eligible_by_positive_id(self):
        ConfigProject.objects.update_or_create(
            token='contract_termination_cancelled_token', defaults={'value': 'cancelled'}
        )
        ConfigProject.objects.update_or_create(
            token='contract_termination_completed_token', defaults={'value': 'completed'}
        )
        filters, payment_filters, annotation_filters, *_ = _build_claim_request_filters({})
        self.assertIn('invoices__id__in', [lookup for lookup, _ in q_lookups(filters)])
        self.assertIn('invoice_id__in', [lookup for lookup, _ in q_lookups(payment_filters)])
        self.assertIn('invoices__id__in', [lookup for lookup, _ in q_lookups(annotation_filters)])
        self.assertFalse(any('contract_termination' in lookup for lookup, _ in q_lookups(filters)))


class ContractFromFinishedRequestIdentityTests(TestCase):
    """El codi compartit no n'hi ha prou: el titular i el punt també han de coincidir."""

    def setUp(self):
        ConfigProject.objects.update_or_create(
            token='contract_request_finalize_token', defaults={'value': 'finalitzada'}
        )
        ConfigProject.objects.update_or_create(
            token='contract_terminated_status', defaults={'value': 'baixa'}
        )
        self.finalized = ContractRequestStatus.objects.create(token='finalitzada', name='Finalitzada')
        self.holder = Person.objects.create(token='H1', name='Titular')
        self.other_holder = Person.objects.create(token='H2', name='Altre')
        self.supply_point = SupplyPoint.objects.create(token='SP1')
        self.other_point = SupplyPoint.objects.create(token='SP2')

    def request(self, **kwargs):
        return ContractRequest.objects.create(
            status=self.finalized,
            holder=self.holder,
            supply_point_default=self.supply_point,
            **kwargs,
        )

    def test_accepts_token_when_holder_and_supply_point_match(self):
        request = self.request(token='C-1')
        contract = Contract.objects.create(
            token='C-1', holder=self.holder, supply_point_default=self.supply_point, is_active=True,
        )
        self.assertEqual(contract_from_finished_request(request).id, contract.id)

    def test_rejects_token_when_holder_differs_and_uses_the_matching_contract(self):
        request = self.request(token='C-1')
        Contract.objects.create(
            token='C-1', holder=self.other_holder, supply_point_default=self.supply_point, is_active=True,
        )
        matching = Contract.objects.create(
            token='C-ALT', holder=self.holder, supply_point_default=self.supply_point, is_active=True,
        )
        self.assertEqual(contract_from_finished_request(request).id, matching.id)

    def test_rejects_token_when_supply_point_differs(self):
        request = self.request(token='C-1')
        Contract.objects.create(
            token='C-1', holder=self.holder, supply_point_default=self.other_point, is_active=True,
        )
        self.assertIsNone(contract_from_finished_request(request))


class ResolveInvoiceContractTerminationTests(TestCase):
    """Una baixa tancada deixa la factura sense contracte reclamable, encara que
    la factura encara apunti al contracte. Sense baixa, o amb la baixa oberta,
    es conserva el contracte."""

    def setUp(self):
        ConfigProject.objects.update_or_create(
            token='contract_termination_cancelled_token', defaults={'value': 'cancelled'}
        )
        ConfigProject.objects.update_or_create(
            token='contract_termination_completed_token', defaults={'value': 'completed'}
        )
        self.contract = SimpleNamespace(id=1)

    def invoice(self, *, contract_id=1, termination_id=None, status_token=None, termination_active=True):
        termination = None
        if termination_id:
            termination = SimpleNamespace(
                is_active=termination_active,
                status=SimpleNamespace(token=status_token) if status_token else None,
                contract=self.contract,
            )
        return SimpleNamespace(
            contract_id=contract_id,
            contract=self.contract if contract_id else None,
            contract_request_id=None,
            contract_termination_id=termination_id,
            contract_termination=termination,
        )

    def test_closed_termination_hides_the_invoice_contract(self):
        invoice = self.invoice(termination_id=5, status_token='completed')
        self.assertIsNone(resolve_invoice_contract(invoice))

    def test_invoice_without_termination_keeps_its_contract(self):
        invoice = self.invoice()
        self.assertIs(resolve_invoice_contract(invoice), self.contract)

    def test_open_termination_keeps_the_invoice_contract(self):
        invoice = self.invoice(termination_id=5, status_token='pending')
        self.assertIs(resolve_invoice_contract(invoice), self.contract)

    def test_open_termination_without_contract_uses_the_termination_contract(self):
        invoice = self.invoice(contract_id=None, termination_id=5, status_token='pending')
        self.assertIs(resolve_invoice_contract(invoice), self.contract)


class ReadingBatchTemplateSelectorFilterTests(TestCase):
    """Filtres `route_zone` i `has_routes` del llistat de plantilles de lot de lectura."""

    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(User.objects.create_superuser('admin-test', 'admin@test.local', 'x'))

        self.zone_a = RouteZone.objects.create(name='Zona A')
        self.zone_b = RouteZone.objects.create(name='Zona B')
        self.template_a1 = ReadingBatchTemplate.objects.create(name='A/1')
        self.template_a2 = ReadingBatchTemplate.objects.create(name='A/2')
        self.template_b1 = ReadingBatchTemplate.objects.create(name='B/1')
        self.template_empty = ReadingBatchTemplate.objects.create(name='Sense rutes')

        Route.objects.create(name='R1', route_zone=self.zone_a, reading_batch_template=self.template_a1)
        Route.objects.create(name='R2', route_zone=self.zone_a, reading_batch_template=self.template_a1)
        Route.objects.create(name='R3', route_zone=self.zone_a, reading_batch_template=self.template_a2)
        Route.objects.create(name='R4', route_zone=self.zone_b, reading_batch_template=self.template_b1)

    def get_ids(self, query=''):
        response = self.client.get(f'/billing/reading-batch-template/?{query}')
        self.assertEqual(response.status_code, 200)
        return [item['id'] for item in response.json()['results']]

    def test_has_routes_excludes_templates_without_routes(self):
        ids = self.get_ids('has_routes=true')
        self.assertCountEqual(ids, [self.template_a1.id, self.template_a2.id, self.template_b1.id])

    def test_route_zone_filter_without_duplicates(self):
        ids = self.get_ids(f'has_routes=true&route_zone={self.zone_a.id}')
        self.assertCountEqual(ids, [self.template_a1.id, self.template_a2.id])

    def test_without_params_returns_all(self):
        self.assertEqual(len(self.get_ids()), 4)


class ClaimContractsExportTests(TestCase):
    """Columnes calculades de l'exportació de contractes dels llistats de reclamació."""

    def setUp(self):
        from billing.models import Invoice, InvoiceLineItem, Payment, PaymentStatus
        from contract.models import Contract
        from pricing.models import PriceRate
        from service.models import SupplyPoint

        # Els signals de contracte, factura i pagament llegeixen aquests valors en desar.
        for token, value in (
            ('contract_terminated_status', '99'),
            ('payment_status_paid_token', '0'),
            ('payment_status_expired_token', '-1'),
            ('payment_status_returned_token', '-2'),
            ('payment_status_cancelled_token', '-4'),
            ('invoice_status_paid_token', 'PAID'),
            ('invoice_status_cancelled_token', 'CANCELLED'),
        ):
            ConfigProject.objects.get_or_create(token=token, defaults={'value': value})

        self.supply_point = SupplyPoint.objects.create(token='SP-TEST')
        self.contract = Contract.objects.create(token='C-TEST', supply_point_default=self.supply_point)
        expired = PaymentStatus.objects.get_or_create(token='-1', defaults={'name': 'Vençut'})[0]
        paid = PaymentStatus.objects.get_or_create(token='0', defaults={'name': 'Pagat'})[0]

        return_fee_rate = PriceRate.objects.create(token='RETURN-FEE')
        ConfigProject.objects.update_or_create(
            token='invoice_return_price_rate_token', defaults={'value': return_fee_rate.token}
        )

        def invoice_with_payments(*payments, return_fee=False):
            invoice = Invoice.objects.create(contract=self.contract)
            for payment_status, amount in payments:
                Payment.objects.create(invoice=invoice, status=payment_status, amount=amount)
            if return_fee:
                InvoiceLineItem.objects.create(invoice=invoice, price_rate=return_fee_rate)
            return invoice

        # Dues factures impagades (una amb dos pagaments), una de pagada i una de recàrrec per devolució.
        invoice_with_payments((expired, 10), (expired, 5))
        invoice_with_payments((expired, 20))
        invoice_with_payments((paid, 30))
        invoice_with_payments((expired, 3), return_fee=True)

    def test_unpaid_invoices_exclude_paid_and_return_fee_invoices(self):
        from claimrequest.utils.claim_contract_export_service import unpaid_invoices_count_by_contract_id

        self.assertEqual(unpaid_invoices_count_by_contract_id([self.contract.id]), {self.contract.id: 2})

    def test_unpaid_invoice_with_expired_and_paid_payments_is_counted_once(self):
        from billing.models import Invoice, Payment, PaymentStatus
        from claimrequest.utils.claim_contract_export_service import unpaid_invoices_count_by_contract_id

        invoice = Invoice.objects.create(contract=self.contract)
        for token, amount in (('-1', 10), ('-1', 5), ('0', 20)):
            Payment.objects.create(
                invoice=invoice, status=PaymentStatus.objects.get(token=token), amount=amount,
            )

        self.assertEqual(unpaid_invoices_count_by_contract_id([self.contract.id]), {self.contract.id: 3})

    def test_exports_in_the_same_second_keep_separate_files_and_contents(self):
        from io import BytesIO
        from pathlib import Path
        from tempfile import TemporaryDirectory
        from unittest.mock import patch

        from django.utils import timezone
        from openpyxl import load_workbook

        from claimrequest.tasks import export_claim_contracts_task
        from documentmanager.models import Document

        now = timezone.now()
        with TemporaryDirectory() as storage_path, self.settings(
            DOCUMENT_STORAGE_PATH=storage_path, DOCUMENT_MANAGER_SERVICES={'claimrequest': 'hdd'},
        ), patch('claimrequest.tasks.timezone.now', return_value=now), patch(
            # Dues tasques poden llegir la mateixa versió abans de desar cap document.
            'documentmanager.utils.main_utils.get_next_version', return_value=1,
        ):
            first = export_claim_contracts_task([self.contract.id])
            second = export_claim_contracts_task([])
            self.assertEqual(first['status'], 'success', first)
            self.assertEqual(second['status'], 'success', second)

            first_document = Document.objects.get(pk=first['document_id'])
            second_document = Document.objects.get(pk=second['document_id'])
            self.assertNotEqual(first['document_name'], second['document_name'])
            self.assertNotEqual(first_document.location, second_document.location)
            first_rows = list(load_workbook(BytesIO(Path(first_document.location).read_bytes())).active.values)
            second_rows = list(load_workbook(BytesIO(Path(second_document.location).read_bytes())).active.values)
            self.assertEqual(first_rows[1][0], self.contract.token)
            self.assertEqual(len(first_rows), 2)
            self.assertEqual(len(second_rows), 1)

    def test_supply_cut_date_only_for_open_cuts(self):
        import datetime
        from django.utils import timezone
        from claimrequest.utils.claim_contract_export_service import claim_contracts_export_queryset, _supply_cut_date
        from service.models import SupplyCut, SupplyCutStatus

        # Catàlegs sembrats per migracions: reutilitzar, no duplicar.
        finished = SupplyCutStatus.objects.get_or_create(token='2', defaults={'name': 'Acabat'})[0]
        planned = SupplyCutStatus.objects.get_or_create(token='0', defaults={'name': 'Planificat'})[0]
        closed_cut = SupplyCut.objects.create(
            status=finished, date_start=timezone.make_aware(datetime.datetime(2026, 1, 10, 9)),
        )
        closed_cut.supply_points.add(self.supply_point)

        contract = claim_contracts_export_queryset().get(id=self.contract.id)
        self.assertEqual(_supply_cut_date(contract), '')

        planned_cut = SupplyCut.objects.create(
            status=planned, date_start=timezone.make_aware(datetime.datetime(2026, 11, 5, 9)),
        )
        planned_cut.supply_points.add(self.supply_point)

        contract = claim_contracts_export_queryset().get(id=self.contract.id)
        self.assertEqual(_supply_cut_date(contract), '05/11/2026')

    def test_export_endpoint_queues_download_job(self):
        from documentmanager.models import ExportJob
        from claimrequest.models import ClaimRequest, ClaimRequestPayment

        claim_request = ClaimRequest.objects.create(token='CR-TEST')
        ClaimRequestPayment.objects.create(claim_request=claim_request, contract=self.contract)
        scope = {'claim_request_id': claim_request.id}

        client = APIClient()
        client.force_authenticate(User.objects.create_user('no-perms', 'no-perms@test.local', 'x'))
        url = '/claimrequest/claim-request-manage/contracts-export/'
        self.assertEqual(client.post(url, {'contract_ids': [self.contract.id], **scope}, format='json').status_code, 403)

        client.force_authenticate(User.objects.create_superuser('admin-export', 'admin-export@test.local', 'x'))
        self.assertEqual(client.post(url, {'contract_ids': [], **scope}, format='json').status_code, 400)
        response = client.post(url, {'contract_ids': [self.contract.id, self.contract.id], **scope}, format='json')
        self.assertEqual(response.status_code, 202)
        job = ExportJob.objects.get(id=response.data['export_job_id'])
        self.assertEqual(job.kind, 'claim_contracts')
        self.assertEqual(job.params, {'total': 1})

    def test_export_endpoint_rejects_contracts_outside_scope(self):
        import json
        from django_celery_results.models import TaskResult
        from claimrequest.models import ClaimRequest
        from contract.models import Contract

        other_contract = Contract.objects.create(token='C-OTHER')
        client = APIClient()
        client.force_authenticate(User.objects.create_superuser('admin-scope', 'admin-scope@test.local', 'x'))
        url = '/claimrequest/claim-request-manage/contracts-export/'

        # Sense abast, o amb tots dos alhora, no s'accepta.
        self.assertEqual(client.post(url, {'contract_ids': [self.contract.id]}, format='json').status_code, 400)

        # Gestió creada: només els seus contractes.
        claim_request = ClaimRequest.objects.create(token='CR-SCOPE')
        response = client.post(url, {'contract_ids': [other_contract.id], 'claim_request_id': claim_request.id}, format='json')
        self.assertEqual(response.status_code, 403)

        # En crear la gestió: només els contractes que ha retornat la tasca de filtratge.
        TaskResult.objects.create(
            task_id='filter-task-test', status='SUCCESS', content_type='application/json', content_encoding='utf-8',
            result=json.dumps({'contracts': [{'id': self.contract.id}]}),
        )
        scope = {'filter_task_id': 'filter-task-test'}
        self.assertEqual(client.post(url, {'contract_ids': [self.contract.id, other_contract.id], **scope}, format='json').status_code, 403)
        self.assertEqual(client.post(url, {'contract_ids': [self.contract.id], **scope}, format='json').status_code, 202)
        self.assertEqual(client.post(url, {'contract_ids': [self.contract.id], 'filter_task_id': 'unknown'}, format='json').status_code, 400)

    # Full de factures: abast calculat al servidor i contingut.

    def expired_payments(self):
        from billing.models import Payment

        return list(Payment.objects.filter(invoice__contract=self.contract, status__token='-1').order_by('id'))

    def test_task_writes_invoice_sheet_ordered_by_invoice_date(self):
        import datetime
        from io import BytesIO
        from pathlib import Path
        from tempfile import TemporaryDirectory

        from openpyxl import load_workbook

        from claimrequest.tasks import export_claim_contracts_task
        from documentmanager.models import Document

        self.contract.holder = Person.objects.create(token='00000000T', name='Titular', surname='Prova')
        self.contract.save()
        # El primer i el tercer pagament vençuts són de factures diferents (10 i 20).
        payments = self.expired_payments()
        first, second = payments[0], payments[2]
        first.invoice.issue_date = datetime.date(2026, 3, 1)
        first.invoice.serie_final = 'F-2'
        first.invoice.save()
        second.invoice.issue_date = datetime.date(2026, 1, 1)
        second.invoice.serie_final = 'F-1'
        second.invoice.save()

        with TemporaryDirectory() as storage_path, self.settings(
            DOCUMENT_STORAGE_PATH=storage_path, DOCUMENT_MANAGER_SERVICES={'claimrequest': 'hdd'},
        ):
            result = export_claim_contracts_task(
                [self.contract.id], [[self.contract.id, first.id, False], [self.contract.id, second.id, True]],
            )
            self.assertEqual(result['status'], 'success', result)
            workbook = load_workbook(BytesIO(Path(Document.objects.get(pk=result['document_id']).location).read_bytes()))

        contract_rows = list(workbook.worksheets[0].values)
        self.assertEqual(contract_rows[1][:3], ('C-TEST', 'Titular Prova', '00000000T'))

        invoice_rows = list(workbook.worksheets[1].values)
        header = invoice_rows[0]
        self.assertEqual(len(invoice_rows), 3)
        self.assertEqual([row[header.index('Factura')] for row in invoice_rows[1:]], ['F-1', 'F-2'])
        self.assertEqual(invoice_rows[1][header.index('NIF')], '00000000T')
        self.assertEqual(invoice_rows[1][header.index('Data factura')], '01/01/2026')
        self.assertEqual(invoice_rows[1][header.index('Import pendent')], 20)
        self.assertEqual([row[header.index('Exclosa')] for row in invoice_rows[1:]], ['Sí', 'No'])

    def post_export(self, data):
        from unittest.mock import patch

        client = APIClient()
        client.force_authenticate(User.objects.create_superuser('admin-invoices', 'admin-invoices@test.local', 'x'))
        with patch('claimrequest.views.claim_contracts_export_view.enqueue_export') as enqueue:
            enqueue.return_value = SimpleNamespace(task_id='t', id=1, name='export')
            response = client.post('/claimrequest/claim-request-manage/contracts-export/', data, format='json')
        self.assertEqual(response.status_code, 202, response.data)
        return enqueue.call_args.kwargs['args']

    def test_claim_scope_marks_excluded_payments(self):
        from claimrequest.models import ClaimRequest, ClaimRequestPayment

        claim_request = ClaimRequest.objects.create(token='CR-INVOICES')
        included, excluded = self.expired_payments()[:2]
        ClaimRequestPayment.objects.create(claim_request=claim_request, contract=self.contract, payment=included)
        ClaimRequestPayment.objects.create(
            claim_request=claim_request, contract=self.contract, payment=excluded, is_excluded=True,
        )

        args = self.post_export({'contract_ids': [self.contract.id], 'claim_request_id': claim_request.id})
        self.assertEqual(args, [[self.contract.id], [
            [self.contract.id, included.id, False], [self.contract.id, excluded.id, True],
        ]])

    def test_filter_scope_uses_task_payments_and_marks_removed_ones(self):
        import json
        from django_celery_results.models import TaskResult

        kept, removed = self.expired_payments()[:2]
        TaskResult.objects.create(
            task_id='filter-task-invoices', status='SUCCESS', content_type='application/json',
            content_encoding='utf-8',
            result=json.dumps({'contracts': [{'id': self.contract.id, 'payments': [{'id': kept.id}, {'id': removed.id}]}]}),
        )

        args = self.post_export({
            'contract_ids': [self.contract.id], 'filter_task_id': 'filter-task-invoices',
            'exclude_payment_ids': [removed.id],
        })
        self.assertEqual(args, [[self.contract.id], [
            [self.contract.id, kept.id, False], [self.contract.id, removed.id, True],
        ]])
