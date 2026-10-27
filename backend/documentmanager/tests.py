import io
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier, Lock
from unittest.mock import Mock, patch

from botocore.exceptions import ClientError
from django.contrib.auth.models import User
from django.core.exceptions import ValidationError
from django.http import HttpResponse
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from rest_framework import status
from rest_framework.test import APIClient

from contract.models import Contract, ContractRequest
from coredata.models import ConfigProject
from documentmanager.models import Document, DocumentSign
from integrations.outbound.signing.exceptions import SigningApiError


@override_settings(AWS_STORAGE_BUCKET_NAME="test-documents")
class AwsUploadTests(TestCase):
    def _upload(self, name, reservations=None):
        from documentmanager.utils.aws_service import upload_to_aws

        document = Document(entity_id=1, service="aws", document_name=name)
        uploaded_file = io.BytesIO(b"mandate")
        with patch("documentmanager.utils.aws_service.get_current_user", return_value=None), patch(
            "documentmanager.utils.aws_service.boto3.client"
        ) as client, patch(
            "documentmanager.utils.aws_service.file_already_exists_check",
        ) as exists:
            client.return_value.put_object.side_effect = reservations
            upload_to_aws(uploaded_file, document, name, "CONTRACT", "SEPA")

        document.refresh_from_db()
        exists.assert_not_called()
        client.return_value.get_object.assert_not_called()
        checked_keys = [call.kwargs["Key"] for call in client.return_value.put_object.call_args_list]
        self.assertEqual(document.folder, checked_keys[-1])
        self.assertEqual(document.document_name, document.folder.rsplit("/", 1)[-1])
        client.return_value.upload_fileobj.assert_called_once_with(
            uploaded_file,
            "test-documents",
            document.folder,
            ExtraArgs={"Tagging": "user=Admin&category=SEPA"},
        )
        self.assertEqual(client.return_value.put_object.call_args.kwargs, {
            "Bucket": "test-documents", "Key": document.folder,
            "Body": b"", "IfNoneMatch": "*",
        })
        return document, checked_keys

    def test_collisions_reserve_new_keys_without_accumulating_suffixes(self):
        for name, expected_names in (
            ("mandate.pdf", ["mandate.pdf", "mandate(1).pdf", "mandate(2).pdf"]),
            ("mandate", ["mandate", "mandate(1)", "mandate(2)"]),
            ("mandate.signed.pdf", ["mandate.signed.pdf", "mandate(1).signed.pdf", "mandate(2).signed.pdf"]),
            ("mandate(2).pdf", ["mandate(2).pdf", "mandate(2)(1).pdf", "mandate(2)(2).pdf"]),
        ):
            with self.subTest(name=name):
                error = ClientError({"Error": {"Code": "PreconditionFailed"}}, "PutObject")
                document, checked_keys = self._upload(name, [error, error, {}])
                self.assertEqual([key.rsplit("/", 1)[-1] for key in checked_keys], expected_names)
                self.assertEqual(document.document_name, expected_names[-1])

    def test_no_collision_keeps_original_name_and_key(self):
        document, checked_keys = self._upload("mandate.pdf")

        self.assertEqual(document.document_name, "mandate.pdf")
        self.assertEqual(len(checked_keys), 1)

    def test_reservation_conflicts_advance_to_next_collision_name(self):
        for code in ("PreconditionFailed", "ConditionalRequestConflict"):
            with self.subTest(code=code):
                error = ClientError({"Error": {"Code": code}}, "PutObject")
                document, keys = self._upload("mandate.pdf", [error, {}])
                self.assertEqual(document.document_name, "mandate(1).pdf")
                self.assertNotEqual(keys[0], keys[1])

    def test_long_names_fit_folder_limit_and_preserve_extension_and_version(self):
        for extension in ("", ".pdf", ".signed.pdf"):
            with self.subTest(extension=extension):
                error = ClientError({"Error": {"Code": "PreconditionFailed"}}, "PutObject")
                document, keys = self._upload("m" * 300 + extension, [error] * 10 + [{}])
                prefix = keys[0].rsplit("/", 1)[0] + "/"
                for version, key in enumerate(keys):
                    suffix = (f"({version})" if version else "") + extension
                    self.assertEqual(len(key), 255)
                    self.assertTrue(key.startswith(prefix))
                    self.assertTrue(key.endswith(suffix))
                self.assertTrue(document.document_name.endswith("(10)" + extension))

    def test_unrepresentable_key_is_rejected_before_upload(self):
        from documentmanager.utils.aws_service import upload_to_aws

        for name, field in (("m." + "x" * 255, "CONTRACT"), ("m.pdf", "x" * 255)):
            with self.subTest(name=name, field=field), patch(
                "documentmanager.utils.aws_service.boto3.client"
            ) as client, self.assertRaises(ValidationError):
                upload_to_aws(io.BytesIO(b"mandate"), Mock(), name, field, "SEPA")
            client.return_value.get_object.assert_not_called()
            client.return_value.put_object.assert_not_called()
            client.return_value.upload_fileobj.assert_not_called()

    def test_only_missing_key_is_treated_as_absent(self):
        from documentmanager.utils.aws_service import file_already_exists_check

        for code in ("NoSuchKey", "AccessDenied", "NoSuchBucket", "SlowDown", "InternalError"):
            with self.subTest(code=code), patch(
                "documentmanager.utils.aws_service.boto3.client"
            ) as client:
                error = ClientError({"Error": {"Code": code}}, "GetObject")
                client.return_value.get_object.side_effect = error
                if code == "NoSuchKey":
                    self.assertFalse(file_already_exists_check("bucket", "key"))
                else:
                    with self.assertRaises(ClientError) as raised:
                        file_already_exists_check("bucket", "key")
                    self.assertIs(raised.exception, error)

    def test_non_collision_reservation_errors_prevent_upload(self):
        from documentmanager.utils.aws_service import upload_to_aws

        errors = [
            ClientError({"Error": {"Code": code}}, "PutObject")
            for code in ("AccessDenied", "NoSuchBucket", "SlowDown", "InternalError")
        ] + [RuntimeError("connection failed")]
        for error in errors:
            with self.subTest(error=error), patch(
                "documentmanager.utils.aws_service.boto3.client"
            ) as client, patch(
                "documentmanager.utils.aws_service.get_current_user", return_value=None
            ):
                document = Mock()
                client.return_value.put_object.side_effect = error
                with self.assertRaises(ValidationError) as raised:
                    upload_to_aws(io.BytesIO(b"mandate"), document, "m.pdf", "CONTRACT", "SEPA")
                self.assertIs(raised.exception.__context__, error)
                client.return_value.get_object.assert_not_called()
                client.return_value.put_object.assert_called_once()
                client.return_value.upload_fileobj.assert_not_called()
                client.return_value.delete_object.assert_not_called()
                document.save.assert_not_called()

    def test_failed_transfer_releases_reservation(self):
        from documentmanager.utils.aws_service import upload_to_aws

        with patch("documentmanager.utils.aws_service.boto3.client") as client, patch(
            "documentmanager.utils.aws_service.get_current_user", return_value=None
        ):
            client.return_value.upload_fileobj.side_effect = RuntimeError("transfer failed")
            document = Mock()
            with self.assertRaisesMessage(RuntimeError, "transfer failed"):
                upload_to_aws(io.BytesIO(b"mandate"), document, "m.pdf", "CONTRACT", "SEPA")
            client.return_value.delete_object.assert_called_once_with(
                Bucket="test-documents", Key=client.return_value.put_object.call_args.kwargs["Key"]
            )
            document.save.assert_not_called()

    def test_simultaneous_uploads_preserve_both_contents(self):
        from documentmanager.utils.aws_service import upload_to_aws

        objects = {}
        lock = Lock()
        reserving = Barrier(2, timeout=5)

        def put_object(Bucket, Key, Body, IfNoneMatch):
            self.assertEqual(IfNoneMatch, "*")
            if Key.endswith("/mandate.pdf"):
                reserving.wait()
            with lock:
                if Key in objects:
                    raise ClientError({"Error": {"Code": "PreconditionFailed"}}, "PutObject")
                objects[Key] = Body

        def transfer(file, bucket, key, ExtraArgs):
            with lock:
                objects[key] = file.read()

        with patch("documentmanager.utils.aws_service.boto3.client") as client, patch(
            "documentmanager.utils.aws_service.get_current_user", return_value=None
        ):
            client.return_value.put_object.side_effect = put_object
            client.return_value.upload_fileobj.side_effect = transfer
            documents = [Mock(), Mock()]
            with ThreadPoolExecutor(max_workers=2) as pool:
                futures = [pool.submit(
                    upload_to_aws, io.BytesIO(content), document, "mandate.pdf", "CONTRACT", "SEPA"
                ) for document, content in zip(documents, (b"first", b"second"))]
                for future in futures:
                    future.result(timeout=10)
            client.return_value.get_object.assert_not_called()

        self.assertEqual({document.document_name for document in documents}, {"mandate.pdf", "mandate(1).pdf"})
        for document, content in zip(documents, (b"first", b"second")):
            self.assertEqual(objects[document.folder], content)


class DocumentSignApiTest(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.user = User.objects.create_superuser(
            username="docsign", email="docsign@example.com", password="secret"
        )
        self.client.force_authenticate(user=self.user)
        ConfigProject.objects.filter(token="DOCUMENT_SIGN_ENABLED").update(value="true")
        if not ConfigProject.objects.filter(token="DOCUMENT_SIGN_ENABLED").exists():
            ConfigProject.objects.create(token="DOCUMENT_SIGN_ENABLED", value="true")
        ConfigProject.objects.get_or_create(
            token="contract_terminated_status", defaults={"value": "-1"}
        )
        self.contract_request = ContractRequest.objects.create(token="SOL-1")
        self.contract = Contract.objects.create(token="260929006")
        self.original = Document.objects.create(
            entity_id=1, service="hdd", document_name="original.pdf"
        )
        self.signed = Document.objects.create(
            entity_id=1, service="hdd", document_name="signed.pdf"
        )

    def _document_sign(self, **overrides):
        data = {
            "contract_request": self.contract_request,
            "contract_file": self.original,
            "status": DocumentSign.STATUS_SENDED,
            "token": "550e8400-e29b-41d4-a716-446655440000",
            "otp_name": "Nom Cognom",
            "otp_email": "info@example.com",
            "otp_phone": "679000000",
        }
        data.update(overrides)
        return DocumentSign.objects.create(**data)

    def test_create_stays_pending_and_does_not_call_sign(self):
        with patch(
            "integrations.outbound.signing.services.create_document_sign_session"
        ) as mock_send:
            response = self.client.post(
                reverse("documentsign-list"),
                {
                    "contract_request": self.contract_request.id,
                    "contract_file": self.original.id,
                    "otp_name": "Nom Cognom",
                    "otp_email": "info@example.com",
                    "otp_phone": "679000000",
                },
                format="json",
            )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data["status"], DocumentSign.STATUS_PENDING)
        self.assertIsNone(response.data["contract_file_signed_url"])
        mock_send.assert_not_called()

    def test_create_requires_contract_or_request(self):
        response = self.client.post(
            reverse("documentsign-list"),
            {"otp_name": "Albert", "otp_email": "a@b.c"},
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_list_by_contract_finds_request_migrated_to_contract(self):
        document_sign = self._document_sign(contract=self.contract)
        response = self.client.get(
            reverse("document_signs_by_contract"),
            {"contract_id": self.contract.id},
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(len(response.data["document_signs"]), 1)
        row = response.data["document_signs"][0]
        self.assertEqual(row["id"], document_sign.id)
        self.assertEqual(row["contract"], self.contract.id)
        self.assertEqual(row["contract_request"], self.contract_request.id)
        self.assertEqual(row["contract_token"], "260929006")
        self.assertEqual(row["contract_request_token"], "SOL-1")
        self.assertIsNone(row["contract_file_signed_url"])

    @patch("documentmanager.views.download_document")
    def test_download_does_not_serve_original_before_signed(self, mock_download):
        document_sign = self._document_sign(contract_file_signed=self.signed)
        response = self.client.get(reverse("document_sign_download", args=[document_sign.id]))
        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)
        mock_download.assert_not_called()

    @patch("documentmanager.views.download_document")
    def test_download_serves_signed_pdf_only(self, mock_download):
        mock_download.return_value = HttpResponse(b"%PDF-signed", content_type="application/octet-stream")
        document_sign = self._document_sign(
            status=DocumentSign.STATUS_SIGNED,
            contract_file_signed=self.signed,
            signed_at=timezone.now(),
        )
        response = self.client.get(reverse("document_sign_download", args=[document_sign.id]))
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertIn('filename="signed.pdf"', response["Content-Disposition"])
        self.assertEqual(response.content, b"%PDF-signed")
        self.assertEqual(mock_download.call_args.args[0].id, self.signed.id)

    @patch("integrations.outbound.signing.polling.retrieve_signed_document")
    def test_retrieve_pending_returns_json_and_keeps_status(self, mock_retrieve):
        mock_retrieve.return_value = {"action": "pending"}
        document_sign = self._document_sign()
        response = self.client.post(reverse("documentsign-retrieve", args=[document_sign.id]))
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response["Content-Type"], "application/json")
        self.assertEqual(response.data["status"], DocumentSign.STATUS_SENDED)
        self.assertIsNone(response.data["contract_file_signed_url"])
        self.assertIn(
            f"/documentmanager/view-document/{self.original.id}/",
            response.data["contract_file_url"],
        )

    @patch("integrations.outbound.signing.polling.retrieve_signed_document")
    def test_retrieve_signed_returns_updated_object(self, mock_retrieve):
        document_sign = self._document_sign()

        def _sign(obj):
            obj.status = DocumentSign.STATUS_SIGNED
            obj.contract_file_signed = self.signed
            obj.signed_at = timezone.now()
            obj.save()
            return {"action": "signed"}

        mock_retrieve.side_effect = _sign
        response = self.client.post(reverse("documentsign-retrieve", args=[document_sign.id]))
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["status"], DocumentSign.STATUS_SIGNED)
        self.assertIn(
            f"/document-sign-download/{document_sign.id}/",
            response.data["contract_file_signed_url"],
        )
        self.assertIsNotNone(response.data["signed_at"])

    @patch("integrations.outbound.signing.polling.retrieve_signed_document")
    def test_retrieve_remote_failure_marks_error(self, mock_retrieve):
        document_sign = self._document_sign()

        def _fail(obj):
            obj.status = DocumentSign.STATUS_ERROR
            obj.error_report = "timeout"
            obj.save()
            raise SigningApiError("timeout")

        mock_retrieve.side_effect = _fail
        response = self.client.post(reverse("documentsign-retrieve", args=[document_sign.id]))
        self.assertEqual(response.status_code, status.HTTP_502_BAD_GATEWAY)
        self.assertEqual(response.data["status"], DocumentSign.STATUS_ERROR)
        self.assertEqual(response.data["error_report"], "timeout")

    @patch("integrations.outbound.signing.polling.retrieve_signed_document")
    def test_retrieve_rejects_pending_local(self, mock_retrieve):
        document_sign = self._document_sign(status=DocumentSign.STATUS_PENDING, token=None)
        response = self.client.post(reverse("documentsign-retrieve", args=[document_sign.id]))
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        mock_retrieve.assert_not_called()
        document_sign.refresh_from_db()
        self.assertEqual(document_sign.status, DocumentSign.STATUS_PENDING)

    def test_delete_removes_the_request(self):
        document_sign = self._document_sign(status=DocumentSign.STATUS_SIGNED)
        response = self.client.delete(reverse("documentsign-detail", args=[document_sign.id]))
        self.assertEqual(response.status_code, status.HTTP_204_NO_CONTENT)
        self.assertFalse(DocumentSign.objects.filter(id=document_sign.id).exists())

    def test_retrieve_requires_authentication(self):
        document_sign = self._document_sign()
        self.client.force_authenticate(user=None)
        response = self.client.post(reverse("documentsign-retrieve", args=[document_sign.id]))
        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)


def _blank_pdf():
    from PyPDF2 import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    buffer = io.BytesIO()
    writer.write(buffer)
    buffer.seek(0)
    return buffer


class DocumentStorageLocationTest(TestCase):
    """La ubicació del fitxer no es pot canviar per API ni sortir de la carpeta de documents."""

    def setUp(self):
        import os
        import tempfile

        from django.test import override_settings

        self.storage = tempfile.TemporaryDirectory()
        self.addCleanup(self.storage.cleanup)
        self.outside = tempfile.NamedTemporaryFile(delete=False)
        self.outside.write(b"secret")
        self.outside.close()
        self.addCleanup(os.remove, self.outside.name)
        settings_override = override_settings(DOCUMENT_STORAGE_PATH=self.storage.name)
        settings_override.enable()
        self.addCleanup(settings_override.disable)

        self.inside_path = os.path.join(self.storage.name, "doc.pdf")
        with open(self.inside_path, "wb") as f:
            f.write(b"document")
        self.document = Document.objects.create(
            entity="CONTRACT", entity_id=1, service="hdd", document_name="doc.pdf", location=self.inside_path
        )
        self.client = APIClient()
        self.client.force_authenticate(User.objects.create_superuser("docs", "docs@example.com", "pwd"))

    def test_storage_fields_are_read_only(self):
        response = self.client.patch(
            f"/documentmanager/document/{self.document.id}/",
            {"service": "aws", "location": self.outside.name, "document_name": "nou.pdf"},
            format="json",
        )
        self.assertEqual(response.status_code, 200)
        self.document.refresh_from_db()
        self.assertEqual(self.document.location, self.inside_path)
        self.assertEqual(self.document.service, "hdd")
        self.assertEqual(self.document.document_name, "nou.pdf")

    def test_download_inside_storage(self):
        response = self.client.get(f"/documentmanager/view-document/{self.document.id}/")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.content, b"document")

    def test_download_outside_storage_is_not_served(self):
        Document.objects.filter(id=self.document.id).update(location=self.outside.name)
        response = self.client.get(f"/documentmanager/view-document/{self.document.id}/")
        self.assertEqual(response.status_code, 404)

    def test_traversal_outside_storage_is_not_served(self):
        import os

        traversal = os.path.join(self.storage.name, "..", os.path.basename(self.outside.name))
        Document.objects.filter(id=self.document.id).update(location=traversal)
        response = self.client.get(f"/documentmanager/view-document/{self.document.id}/")
        self.assertEqual(response.status_code, 404)

    def test_delete_outside_storage_keeps_file(self):
        import os

        from documentmanager.utils.hdd_service import delete_document_from_hdd

        self.document.location = self.outside.name
        delete_document_from_hdd(self.document)
        self.assertTrue(os.path.exists(self.outside.name))


class DocumentOrderTest(TestCase):
    """El PDF únic respecta l'ordre dels ids rebuts, no l'ordre de creació."""

    def setUp(self):
        self.client = APIClient()
        self.user = User.objects.create_superuser(
            username="docorder", email="docorder@example.com", password="secret"
        )
        self.client.force_authenticate(user=self.user)
        self.first = Document.objects.create(entity_id=1, service="hdd", document_name="first.pdf")
        self.second = Document.objects.create(entity_id=1, service="hdd", document_name="second.pdf")
        self.third = Document.objects.create(entity_id=1, service="hdd", document_name="third.pdf")

    def test_in_order_keeps_requested_order_without_duplicates(self):
        ids = [self.third.id, str(self.first.id), self.third.id, self.second.id, 999999]
        self.assertEqual(Document.in_order(ids), [self.third, self.first, self.second])

    @patch("documentmanager.views.sign_pdf", side_effect=lambda buffer, *args: buffer)
    @patch("documentmanager.views.CompanySerializer")
    @patch("documentmanager.views.Exploitation")
    @patch("documentmanager.views.download_document")
    def test_single_pdf_merges_in_requested_order(self, download_document, *mocks):
        download_document.side_effect = lambda document: _blank_pdf()

        response = self.client.post(
            reverse("download_single_pdf_document"),
            {"doc_ids": [self.third.id, self.first.id, self.second.id]},
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        merged = [call.args[0] for call in download_document.call_args_list]
        self.assertEqual(merged, [self.third, self.first, self.second])


class GenericExportXlsxTest(TestCase):
    """Les cadenes de les exportacions XLSX es desen sempre com a text."""

    def test_text_que_comenca_per_igual_no_es_formula(self):
        import openpyxl
        from documentmanager.utils.generic_export_service import build_export_xlsx_bytes_from_objects

        columns = [("text", "Text", lambda o: o["text"]), ("number", "Number", lambda o: o["number"])]
        file_bytes = build_export_xlsx_bytes_from_objects([{"text": "=1+1", "number": 5}], columns)

        row = list(openpyxl.load_workbook(io.BytesIO(file_bytes)).active.iter_rows(min_row=2))[0]
        self.assertEqual(row[0].value, "=1+1")
        self.assertEqual(row[0].data_type, "s")
        self.assertEqual(row[1].value, 5)
        self.assertEqual(row[1].data_type, "n")


class XlsxStringFormulaSafetyTest(TestCase):
    """Cap text escrit des del codi es desa com a fórmula (documentmanager.utils.xlsx_safety)."""

    def _reload(self, workbook):
        import openpyxl
        buffer = io.BytesIO()
        workbook.save(buffer)
        return openpyxl.load_workbook(io.BytesIO(buffer.getvalue())).active

    def test_text_amb_igual_es_text_en_mode_normal(self):
        import openpyxl
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet["A1"] = "=HYPERLINK(\"http://example.com\")"
        sheet.append(["=1+1", 7, "text"])
        sheet = self._reload(workbook)
        self.assertEqual(sheet["A1"].data_type, "s")
        self.assertEqual(sheet["A2"].data_type, "s")
        self.assertEqual(sheet["A2"].value, "=1+1")
        self.assertEqual(sheet["B2"].data_type, "n")

    def test_text_amb_igual_es_text_en_mode_escriptura(self):
        import openpyxl
        workbook = openpyxl.Workbook(write_only=True)
        workbook.create_sheet().append(["=1+1"])
        self.assertEqual(self._reload(workbook)["A1"].data_type, "s")

    def test_les_formules_dels_fitxers_llegits_es_conserven(self):
        import openpyxl
        workbook = openpyxl.Workbook()
        cell = workbook.active["A1"]
        cell.value = "=1+1"
        cell.data_type = "f"  # fórmula d'un fitxer pujat per l'usuari
        self.assertEqual(self._reload(workbook)["A1"].data_type, "f")


class CsvSafetyTest(TestCase):
    """Els textos que Excel executaria com a fórmula es desen amb una `'` al davant."""

    def test_textos_amb_prefix_de_formula_s_escapen(self):
        from documentmanager.utils.csv_safety import csv_safe_value
        for value in ("=1+1", "+cmd", "-1+1", "@SUM(A1)", "\tx", "-"):
            self.assertEqual(csv_safe_value(value), "'" + value)

    def test_numeros_i_textos_normals_no_es_toquen(self):
        from documentmanager.utils.csv_safety import csv_safe_value
        for value in ("-12,50", "+34 600 000 000", "1.234,5", "Carrer Major", "", None, -5, 3.2):
            self.assertEqual(csv_safe_value(value), value)

    def test_el_writer_escapa_cada_fila(self):
        from documentmanager.utils.csv_safety import SafeCsvWriter
        buffer = io.StringIO()
        writer = SafeCsvWriter(buffer, delimiter=";")
        writer.writerow(["=1+1", 5])
        writer.writerows([["-3", "@x"]])
        self.assertEqual(buffer.getvalue().splitlines(), ["'=1+1;5", "-3;'@x"])

    def test_la_importacio_recupera_el_valor_original(self):
        from documentmanager.utils.csv_safety import csv_safe_value, csv_unescape_value
        for value in ("=1+1", "-DN15", "@x", "DN15", "'text normal", "-12,50", None, 7):
            self.assertEqual(csv_unescape_value(csv_safe_value(value)), value)
        # Una `'` que no protegeix cap fórmula es conserva.
        self.assertEqual(csv_unescape_value("'text normal"), "'text normal")
